from __future__ import annotations

from io import BytesIO
import json
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault(
    "AI_TUTOR_DB_PATH",
    str(Path(__file__).resolve().with_name("app.test-runtime-api-contract.db")),
)

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.main import app
from backend.seed_db import seed_database
from backend import public_seed, services
from backend.settings import DB_PATH
from backend.settings import CURRENT_DATE


def fetch_one(query: str, params: tuple = {}) -> dict:
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        return dict(conn.execute(query, params).fetchone())


def fetch_optional(query: str, params: tuple = {}) -> dict | None:
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        row = conn.execute(query, params).fetchone()
        return dict(row) if row else None


def execute_sql(query: str, params: tuple = {}) -> None:
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(query, params)
        conn.commit()


def fetch_all(query: str, params: tuple = {}) -> list[dict]:
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        return [dict(row) for row in conn.execute(query, params).fetchall()]


def insert_daily_activity(user_id: str, activity_date: str, completed_minutes: int) -> None:
    execute_sql(
        """
        INSERT INTO student_daily_activity (
            activity_id, user_id, activity_date, is_studied, intensity_level,
            target_minutes, completed_minutes, completed_questions,
            submitted_type_count, checkin_status, created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            f"ACT-{user_id}-{activity_date.replace('-', '')}",
            user_id,
            activity_date,
            "1" if completed_minutes > 0 else "0",
            "low" if completed_minutes > 0 else "none",
            "180",
            str(completed_minutes),
            "0",
            "0",
            "completed" if completed_minutes > 0 else "pending",
            f"{activity_date} 08:00:00",
            f"{activity_date} 20:00:00",
        ),
    )


def reset_test_clock(date_text: str = CURRENT_DATE) -> None:
    execute_sql(
        """
        INSERT OR REPLACE INTO test_runtime_state (state_key, state_value, updated_at)
        VALUES ('current_test_date', ?, '2026-06-10 00:00:00')
        """,
        (date_text,),
    )


def auth_headers(payload: dict, device_id: str = "device-a") -> dict[str, str]:
    return {
        "Authorization": f"Bearer {payload['token']}",
        "X-Session-Id": payload["sessionId"],
        "X-Device-Id": device_id,
    }


INTERNAL_TEST_CREDENTIALS = {
    "10001": "98106374",
    "10002": "53087197",
    "10003": "65674025",
    "10004": "25068969",
    "10005": "87487965",
    "10006": "30253556",
    "10007": "74363128",
    "10008": "97738287",
    "10009": "26989080",
    "10010": "39843025",
}


class RuntimeApiContractTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if DB_PATH.exists():
            DB_PATH.unlink()
        seed_database()
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.client.close()

    def student_auth_payload(
        self,
        user_id: str,
        password: str | None = None,
        device_id: str | None = None,
    ) -> dict:
        user = fetch_one(
            "SELECT login_account, login_password FROM users WHERE user_id = ?",
            (user_id,),
        )
        device_id = device_id or f"test-device-{user_id}"
        response = self.client.post(
            "/auth/login",
            json={
                "loginAccount": user["login_account"],
                "password": password
                if password is not None
                else INTERNAL_TEST_CREDENTIALS.get(user_id, user["login_password"]),
                "deviceId": device_id,
                "deviceName": "Contract Test",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def student_auth_headers(
        self,
        user_id: str,
        password: str | None = None,
        device_id: str | None = None,
    ) -> dict[str, str]:
        device_id = device_id or f"test-device-{user_id}"
        return auth_headers(self.student_auth_payload(user_id, password, device_id), device_id)

    def teacher_auth_payload(
        self,
        login_account: str = "",
        password: str | None = None,
        device_id: str = "teacher-device-a",
    ) -> dict:
        if not login_account:
            row = fetch_one(
                "SELECT login_account, login_password FROM users WHERE role = 'teacher' ORDER BY user_id LIMIT 1"
            )
            login_account = row["login_account"]
            password = password if password is not None else row["login_password"]
        response = self.client.post(
            "/auth/login",
            json={
                "loginAccount": login_account,
                "password": password,
                "deviceId": device_id,
                "deviceName": "Teacher Contract Test",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["user"]["role"], "teacher")
        return payload

    def teacher_auth_headers(
        self,
        login_account: str = "",
        password: str | None = None,
        device_id: str = "teacher-device-a",
    ) -> dict[str, str]:
        return auth_headers(self.teacher_auth_payload(login_account, password, device_id), device_id)

    def admin_auth_payload(
        self,
        login_account: str = "",
        password: str | None = None,
        device_id: str = "admin-device-a",
    ) -> dict:
        if login_account:
            row = fetch_one(
                """
                SELECT user_id, login_account, login_password
                FROM users
                WHERE role = 'admin' AND (login_account = ? OR user_id = ?)
                """,
                (login_account, login_account),
            )
        else:
            row = fetch_one(
                """
                SELECT user_id, login_account, login_password
                FROM users
                WHERE role = 'admin'
                ORDER BY user_id
                LIMIT 1
                """
            )
        login_account = row["login_account"] or row["user_id"]
        password = password if password is not None else row["login_password"]
        first_body = {
            "loginAccount": login_account,
            "password": password,
            "deviceId": device_id,
            "deviceName": "Admin Contract Test",
        }
        first_response = self.client.post("/auth/login", json=first_body)
        self.assertEqual(first_response.status_code, 200, first_response.text)
        first_payload = first_response.json()
        if first_payload.get("requiresTotpSetup"):
            manual_secret = first_payload["manualSecret"]
            totp_code = services.totp_code(manual_secret)
            response = self.client.post(
                "/auth/login",
                json={
                    **first_body,
                    "setupToken": first_payload["setupToken"],
                    "totpCode": totp_code,
                },
            )
        else:
            self.assertTrue(first_payload.get("requiresTotp"), first_payload)
            state = fetch_one(
                "SELECT totp_secret FROM admin_totp_state WHERE admin_user_id = ?",
                (row["user_id"],),
            )
            response = self.client.post(
                "/auth/login",
                json={**first_body, "totpCode": services.totp_code(state["totp_secret"])},
            )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["user"]["role"], "admin")
        self.assertFalse(payload.get("needsOnboarding", False))
        return payload

    def admin_auth_headers(
        self,
        login_account: str = "",
        password: str | None = None,
        device_id: str = "admin-device-a",
    ) -> dict[str, str]:
        return auth_headers(self.admin_auth_payload(login_account, password, device_id), device_id)

    def verify_admin_risk_authorization(self, headers: dict[str, str], admin_user_id: str = "adm001") -> dict:
        state = fetch_one("SELECT totp_secret FROM admin_totp_state WHERE admin_user_id = ?", (admin_user_id,))
        response = self.client.post(
            "/admin/risk-authorizations/verify",
            headers=headers,
            json={"totpCode": services.totp_code(state["totp_secret"])},
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertTrue(payload["authorizedUntil"])
        return payload

    def test_admin_first_login_binds_totp_then_uses_single_device_session(self) -> None:
        admin = fetch_one(
            "SELECT user_id, login_account, login_password, current_session_id FROM users WHERE role = 'admin' ORDER BY user_id LIMIT 1"
        )
        original_totp = fetch_optional("SELECT * FROM admin_totp_state WHERE admin_user_id = ?", (admin["user_id"],))
        session_ids: list[str] = []
        try:
            execute_sql(
                """
                UPDATE admin_totp_state
                SET totp_secret = '', totp_bound_at = '', setup_token_hash = '',
                    setup_token_expires_at = '', last_verified_at = ''
                WHERE admin_user_id = ?
                """,
                (admin["user_id"],),
            )
            setup_response = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": admin["login_account"],
                    "password": admin["login_password"],
                    "deviceId": "admin-phone-a",
                    "deviceName": "Admin Missing TOTP",
                },
            )
            self.assertEqual(setup_response.status_code, 200, setup_response.text)
            setup_payload = setup_response.json()
            self.assertTrue(setup_payload["requiresTotpSetup"])
            self.assertIn("setupToken", setup_payload)
            self.assertIn("manualSecret", setup_payload)
            self.assertNotIn("token", setup_payload)

            first = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": admin["login_account"],
                    "password": admin["login_password"],
                    "deviceId": "admin-phone-a",
                    "deviceName": "Admin Setup TOTP",
                    "setupToken": setup_payload["setupToken"],
                    "totpCode": services.totp_code(setup_payload["manualSecret"]),
                },
            )
            self.assertEqual(first.status_code, 200, first.text)
            first_payload = first.json()
            first_state = fetch_one("SELECT totp_bound_at FROM admin_totp_state WHERE admin_user_id = ?", (admin["user_id"],))
            self.assertTrue(first_state["totp_bound_at"])
            self.assertEqual(first_payload["user"]["role"], "admin")
            session_ids.append(first_payload["sessionId"])

            challenge = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": admin["login_account"],
                    "password": admin["login_password"],
                    "deviceId": "admin-phone-b",
                    "deviceName": "Admin Bound TOTP",
                },
            )
            self.assertEqual(challenge.status_code, 200, challenge.text)
            self.assertTrue(challenge.json()["requiresTotp"])
            self.assertNotIn("token", challenge.json())

            state = fetch_one("SELECT totp_secret FROM admin_totp_state WHERE admin_user_id = ?", (admin["user_id"],))
            second = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": admin["login_account"],
                    "password": admin["login_password"],
                    "deviceId": "admin-phone-b",
                    "deviceName": "Admin Bound TOTP",
                    "totpCode": services.totp_code(state["totp_secret"]),
                },
            )
            self.assertEqual(second.status_code, 200, second.text)
            second_payload = second.json()
            session_ids.append(second_payload["sessionId"])

            old_request = self.client.get("/admin/accounts", headers=auth_headers(first_payload, "admin-phone-a"))
            self.assertEqual(old_request.status_code, 401)
            self.assertEqual(old_request.json()["detail"]["code"], "SESSION_REPLACED")

            new_request = self.client.get("/admin/accounts", headers=auth_headers(second_payload, "admin-phone-b"))
            self.assertEqual(new_request.status_code, 200, new_request.text)
        finally:
            for session_id in session_ids:
                execute_sql("DELETE FROM user_sessions WHERE session_id = ?", (session_id,))
            execute_sql(
                "UPDATE users SET current_session_id = ? WHERE user_id = ?",
                (admin["current_session_id"], admin["user_id"]),
            )
            if original_totp:
                execute_sql(
                    """
                    UPDATE admin_totp_state
                    SET formal_activations_since_totp = ?, totp_secret = ?, totp_bound_at = ?,
                        setup_token_hash = ?, setup_token_expires_at = ?, last_verified_at = ?, updated_at = ?
                    WHERE admin_user_id = ?
                    """,
                    (
                        original_totp.get("formal_activations_since_totp", "0"),
                        original_totp.get("totp_secret", ""),
                        original_totp.get("totp_bound_at", ""),
                        original_totp.get("setup_token_hash", ""),
                        original_totp.get("setup_token_expires_at", ""),
                        original_totp.get("last_verified_at", ""),
                        original_totp.get("updated_at", ""),
                        admin["user_id"],
                    ),
                )

    def test_admin_create_teacher_returns_initial_password_and_teacher_can_login(self) -> None:
        user_id = ""
        try:
            admin_payload = self.admin_auth_payload(device_id="admin-create-teacher")
            headers = auth_headers(admin_payload, "admin-create-teacher")
            self.verify_admin_risk_authorization(headers, admin_payload["user"]["userId"])
            response = self.client.post(
                "/admin/teachers",
                headers=headers,
                json={
                    "loginAccount": "tea_contract_create",
                    "nickname": "契约老师",
                    "contactQQ": "2001002001",
                    "status": "active",
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertIn("initialPassword", payload)
            self.assertRegex(payload["initialPassword"], r"^T[A-Z0-9]{7}$")
            teacher = payload["teacher"]
            user_id = teacher["userId"]
            self.assertEqual(teacher["loginAccount"], "tea_contract_create")
            self.assertEqual(teacher["status"], "active")

            login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": "tea_contract_create",
                    "password": payload["initialPassword"],
                    "deviceId": "teacher-create-device",
                    "deviceName": "Teacher Create Test",
                },
            )
            self.assertEqual(login.status_code, 200, login.text)
            self.assertEqual(login.json()["user"]["role"], "teacher")
        finally:
            if user_id:
                execute_sql("DELETE FROM user_sessions WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM users WHERE user_id = ?", (user_id,))

    def test_admin_teacher_management_requires_session_and_recent_totp(self) -> None:
        created_user_id = ""
        admin_payload = self.admin_auth_payload(device_id="admin-teacher-risk")
        admin_user_id = admin_payload["user"]["userId"]
        headers = auth_headers(admin_payload, "admin-teacher-risk")
        try:
            missing_session = self.client.post(
                "/admin/teachers",
                json={"loginAccount": "tea_requires_auth", "nickname": "No Auth"},
            )
            self.assertEqual(missing_session.status_code, 401)
            self.assertEqual(missing_session.json()["detail"]["code"], "SESSION_REQUIRED")

            execute_sql(
                """
                UPDATE admin_totp_state
                SET last_verified_at = ?
                WHERE admin_user_id = ?
                """,
                ("2026-06-01 00:00:00", admin_user_id),
            )
            stale_risk = self.client.post(
                "/admin/teachers",
                headers=headers,
                json={"loginAccount": "tea_requires_totp", "nickname": "Needs TOTP"},
            )
            self.assertEqual(stale_risk.status_code, 403)
            self.assertEqual(stale_risk.json()["detail"], "TOTP_REQUIRED")

            invalid_totp = self.client.post(
                "/admin/risk-authorizations/verify",
                headers=headers,
                json={"totpCode": "000000"},
            )
            self.assertEqual(invalid_totp.status_code, 400)

            self.verify_admin_risk_authorization(headers, admin_user_id)
            created = self.client.post(
                "/admin/teachers",
                headers=headers,
                json={
                    "loginAccount": "tea_requires_totp",
                    "nickname": "Needs TOTP",
                    "contactQQ": "2001002002",
                    "status": "active",
                },
            )
            self.assertEqual(created.status_code, 200, created.text)
            created_user_id = created.json()["teacher"]["userId"]
        finally:
            if created_user_id:
                execute_sql("DELETE FROM user_sessions WHERE user_id = ?", (created_user_id,))
                execute_sql("DELETE FROM users WHERE user_id = ?", (created_user_id,))

    def test_teacher_can_change_password_and_keep_current_session(self) -> None:
        teacher = fetch_one(
            """
            SELECT user_id, login_account, login_password, password_hash,
                   password_changed_at, current_session_id, account_status
            FROM users
            WHERE role = 'teacher'
            ORDER BY user_id
            LIMIT 1
            """
        )
        old_password = teacher["login_password"]
        new_password = "Teach888"
        payload = None
        replacement_payload = None
        try:
            payload = self.teacher_auth_payload(
                teacher["login_account"],
                old_password,
                "teacher-change-password-a",
            )
            headers = auth_headers(payload, "teacher-change-password-a")
            response = self.client.post(
                "/teacher/password",
                headers=headers,
                json={
                    "currentPassword": old_password,
                    "newPassword": new_password,
                    "confirmPassword": new_password,
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["sessionId"], payload["sessionId"])

            still_logged_in = self.client.get("/teacher/classes", headers=headers)
            self.assertEqual(still_logged_in.status_code, 200, still_logged_in.text)

            old_login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": teacher["login_account"],
                    "password": old_password,
                    "deviceId": "teacher-old-password",
                },
            )
            self.assertEqual(old_login.status_code, 400)

            new_login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": teacher["login_account"],
                    "password": new_password,
                    "deviceId": "teacher-change-password-b",
                },
            )
            self.assertEqual(new_login.status_code, 200, new_login.text)
            replacement_payload = new_login.json()
        finally:
            execute_sql("DELETE FROM user_sessions WHERE user_id = ?", (teacher["user_id"],))
            execute_sql(
                """
                UPDATE users
                SET login_password = ?, password_hash = ?, password_changed_at = ?,
                    current_session_id = ?, account_status = ?
                WHERE user_id = ?
                """,
                (
                    teacher["login_password"],
                    teacher["password_hash"],
                    teacher["password_changed_at"],
                    teacher["current_session_id"],
                    teacher["account_status"],
                    teacher["user_id"],
                ),
            )

    def test_teacher_can_update_own_contact_qq(self) -> None:
        teacher = fetch_one(
            """
            SELECT user_id, login_account, login_password, QQ
            FROM users
            WHERE role = 'teacher'
            ORDER BY user_id
            LIMIT 1
            """
        )
        payload = self.teacher_auth_payload(
            teacher["login_account"],
            teacher["login_password"],
            "teacher-profile-qq",
        )
        headers = auth_headers(payload, "teacher-profile-qq")
        try:
            response = self.client.post(
                "/teacher/profile",
                headers=headers,
                json={"contactQQ": "288888888"},
            )
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["teacher"]["contactQQ"], "288888888")

            classes = self.client.get("/teacher/classes", headers=headers)
            self.assertEqual(classes.status_code, 200, classes.text)
            self.assertEqual(classes.json()["teacher"]["contactQQ"], "288888888")

            stored = fetch_one("SELECT QQ FROM users WHERE user_id = ?", (teacher["user_id"],))
            self.assertEqual(stored["QQ"], "288888888")

            invalid = self.client.post(
                "/teacher/profile",
                headers=headers,
                json={"contactQQ": "qq-288"},
            )
            self.assertEqual(invalid.status_code, 400)
        finally:
            execute_sql("DELETE FROM user_sessions WHERE user_id = ?", (teacher["user_id"],))
            execute_sql("UPDATE users SET QQ = ? WHERE user_id = ?", (teacher["QQ"], teacher["user_id"]))

    def test_admin_teacher_reset_and_disable_revoke_teacher_sessions(self) -> None:
        teacher = fetch_one(
            """
            SELECT user_id, login_account, login_password, password_hash,
                   password_changed_at, current_session_id, account_status
            FROM users
            WHERE role = 'teacher'
            ORDER BY user_id
            LIMIT 1
            """
        )
        admin_payload = self.admin_auth_payload(device_id="admin-teacher-reset")
        headers = auth_headers(admin_payload, "admin-teacher-reset")
        teacher_payload = None
        try:
            self.verify_admin_risk_authorization(headers, admin_payload["user"]["userId"])
            teacher_payload = self.teacher_auth_payload(
                teacher["login_account"],
                teacher["login_password"],
                "teacher-before-reset",
            )
            reset = self.client.post(
                f"/admin/accounts/{teacher['user_id']}/reset-password",
                headers=headers,
                json={"actor": "adm001"},
            )
            self.assertEqual(reset.status_code, 200, reset.text)
            new_password = reset.json()["newPassword"]

            old_session = self.client.get(
                "/teacher/classes",
                headers=auth_headers(teacher_payload, "teacher-before-reset"),
            )
            self.assertEqual(old_session.status_code, 401)

            relogin = self.teacher_auth_payload(
                teacher["login_account"],
                new_password,
                "teacher-after-reset",
            )
            disable = self.client.patch(
                f"/admin/teachers/{teacher['user_id']}/status",
                headers=headers,
                json={"status": "disabled", "actor": "adm001"},
            )
            self.assertEqual(disable.status_code, 200, disable.text)
            disabled_session = self.client.get(
                "/teacher/classes",
                headers=auth_headers(relogin, "teacher-after-reset"),
            )
            self.assertEqual(disabled_session.status_code, 401)
            disabled_login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": teacher["login_account"],
                    "password": new_password,
                    "deviceId": "teacher-disabled-login",
                },
            )
            self.assertEqual(disabled_login.status_code, 400)
        finally:
            execute_sql("DELETE FROM user_sessions WHERE user_id = ?", (teacher["user_id"],))
            execute_sql(
                """
                UPDATE users
                SET login_password = ?, password_hash = ?, password_changed_at = ?,
                    current_session_id = ?, account_status = ?
                WHERE user_id = ?
                """,
                (
                    teacher["login_password"],
                    teacher["password_hash"],
                    teacher["password_changed_at"],
                    teacher["current_session_id"],
                    teacher["account_status"],
                    teacher["user_id"],
                ),
            )

    def test_admin_teacher_status_update_can_change_nickname(self) -> None:
        teacher = fetch_one(
            """
            SELECT user_id, nickname, account_status
            FROM users
            WHERE role = 'teacher' AND account_status = 'active'
            ORDER BY user_id
            LIMIT 1
            """
        )
        admin_payload = self.admin_auth_payload(device_id="admin-teacher-nickname")
        headers = auth_headers(admin_payload, "admin-teacher-nickname")
        new_nickname = "合约改名老师"
        try:
            self.verify_admin_risk_authorization(headers, admin_payload["user"]["userId"])
            response = self.client.patch(
                f"/admin/teachers/{teacher['user_id']}/status",
                headers=headers,
                json={
                    "status": teacher["account_status"],
                    "nickname": new_nickname,
                    "actor": "adm001",
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["teacher"]["nickname"], new_nickname)
            row = fetch_one("SELECT nickname FROM users WHERE user_id = ?", (teacher["user_id"],))
            self.assertEqual(row["nickname"], new_nickname)
        finally:
            execute_sql(
                "UPDATE users SET nickname = ?, account_status = ? WHERE user_id = ?",
                (teacher["nickname"], teacher["account_status"], teacher["user_id"]),
            )

    def test_teacher_apis_require_current_teacher_session(self) -> None:
        teacher = fetch_one(
            "SELECT user_id, login_account, login_password FROM users WHERE role = 'teacher' ORDER BY user_id LIMIT 1"
        )
        class_row = fetch_one(
            "SELECT class_id FROM classes WHERE teacher_user_id = ? ORDER BY class_id LIMIT 1",
            (teacher["user_id"],),
        )
        headers = self.teacher_auth_headers(
            teacher["login_account"],
            teacher["login_password"],
            "teacher-scope-a",
        )

        missing = self.client.get("/teacher/classes")
        self.assertEqual(missing.status_code, 401)
        self.assertEqual(missing.json()["detail"]["code"], "SESSION_REQUIRED")

        mine = self.client.get("/teacher/classes", headers=headers)
        self.assertEqual(mine.status_code, 200, mine.text)
        class_ids = {item["classId"] for item in mine.json()["items"]}
        self.assertIn(class_row["class_id"], class_ids)

        own_status = self.client.get(
            f"/teacher/classes/{class_row['class_id']}/daily-feedback-status",
            params={"date": "2026-06-10", "teacherUserId": "wrong-teacher-id"},
            headers=headers,
        )
        self.assertEqual(own_status.status_code, 200, own_status.text)

    def test_teacher_classes_serializes_before_runtime_connection_closes(self) -> None:
        teacher = fetch_one(
            "SELECT user_id FROM users WHERE role = 'teacher' ORDER BY user_id LIMIT 1"
        )

        class ClosingSqliteConnection:
            def __init__(self) -> None:
                self.conn = sqlite3.connect(DB_PATH)
                self.conn.row_factory = sqlite3.Row
                self.conn.execute("PRAGMA foreign_keys = ON")

            def execute(self, *args, **kwargs):
                return self.conn.execute(*args, **kwargs)

            def commit(self) -> None:
                self.conn.commit()

            def rollback(self) -> None:
                self.conn.rollback()

            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb) -> None:
                if exc_type:
                    self.rollback()
                else:
                    self.commit()
                self.conn.close()

        with patch.object(services, "get_connection", side_effect=ClosingSqliteConnection):
            payload = services.list_teacher_classes(teacher["user_id"])

        self.assertEqual(payload["teacher"]["userId"], teacher["user_id"])
        self.assertTrue(payload["items"])

    def open_plan_type(self, user_id: str, plan_date: str) -> tuple[dict, dict]:
        response = self.client.get(
            f"/students/{user_id}/daily-plan",
            params={"date": plan_date},
            headers=self.student_auth_headers(user_id),
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        open_type = next(
            row for row in payload["plan"]["types"] if row["feedbackOpenStatus"] == "open"
        )
        return payload, open_type

    def delete_temp_student(self, user_id: str) -> None:
        for table in (
            "question_feedbacks",
            "daily_plans",
            "student_stats",
            "student_daily_activity",
            "student_chapter_completions",
            "student_weekly_targets",
            "type_mastery_stats",
            "loss_reason_stats",
            "student_study_goals",
            "focus_sessions",
            "user_sessions",
            "student_progress",
            "student_profiles",
            "users",
        ):
            execute_sql(f"DELETE FROM {table} WHERE user_id = ?", (user_id,))

    def create_plan_override_student(
        self,
        user_id: str,
        current_type_id: str,
        primary_book_1: str = "m2_880",
        primary_book_2: str = "",
    ) -> None:
        self.delete_temp_student(user_id)
        type_row = fetch_one("SELECT * FROM question_types WHERE type_id = ?", (current_type_id,))
        created_at = "2026-06-10 06:00:00"
        execute_sql(
            """
            INSERT INTO users (
                user_id, login_account, password_hash, login_password, role, nickname,
                avatar_url, QQ, account_status, account_type, activated_at, created_at, updated_at
            ) VALUES (?, ?, '', '12345678', 'student', ?, '', '10000', 'active', 'formal', ?, ?, ?)
            """,
            (user_id, f"stu{user_id}", user_id, created_at, created_at, created_at),
        )
        execute_sql(
            """
            INSERT INTO student_profiles (
                user_id, math_type, target_exam_year, target_exam_date, target_score,
                class_id, level, plan_enabled, paid_until, primary_book_1, primary_book_2,
                backup_book, current_type_id, current_stage, last_plan_date, created_at,
                updated_at, intensive_days, review_order, plan_variant_id,
                diagnostic_json, primary_book_1_scope, primary_book_2_scope
            ) VALUES (?, '数学二', '2026', '2026-12-19', '120',
                'A201', 'A', '1', '2026-12-31', ?, ?, 'm2_660', ?,
                '强化', '', ?, ?, '45', 'mixed', 'math2_mixed',
                '{}', 'both', 'both')
            """,
            (user_id, primary_book_1, primary_book_2, current_type_id, created_at, created_at),
        )
        execute_sql(
            """
            INSERT INTO student_progress (
                progress_id, user_id, math_type, plan_stage, current_type_id,
                current_type_sort_code, progress_status, current_type_status,
                last_generated_plan_date, last_completed_type_id, unfinished_type_id,
                unfinished_question_ids, carryover_type_id, carryover_question_ids,
                created_at, updated_at
            ) VALUES (?, ?, '数学二', '强化', ?, ?, 'active', 'active',
                '', '', '', '[]', '', '[]', ?, ?)
            """,
            (
                f"PR-{user_id}",
                user_id,
                current_type_id,
                type_row["sort_code"],
                created_at,
                created_at,
            ),
        )

    def first_assignable_type(self) -> dict:
        return fetch_one(
            """
            SELECT *
            FROM question_types
            WHERE direct_assignable = '1'
            ORDER BY sort_code
            LIMIT 1
            """
        )

    def first_assignable_type_in_later_chapter(self, after_chapter_id: str) -> dict:
        return fetch_one(
            """
            SELECT qt.*
            FROM question_types qt
            WHERE qt.direct_assignable = '1'
              AND qt.chapter_id <> ?
            ORDER BY qt.sort_code
            LIMIT 1
            """,
            (after_chapter_id,),
        )

    def test_plan_variant_duration_uses_student_level_column(self) -> None:
        type_id = self.first_assignable_type()["type_id"]
        plan_variant_id = "test_duration_variant"
        execute_sql(
            """
            CREATE TABLE IF NOT EXISTS question_type_orders (
                plan_variant_id TEXT NOT NULL,
                type_id TEXT NOT NULL,
                type_sort_in_chapter TEXT NOT NULL DEFAULT '1',
                final_type_label TEXT NOT NULL DEFAULT '',
                sort_code TEXT NOT NULL DEFAULT '',
                direct_assignable TEXT NOT NULL DEFAULT '1',
                duration_130_days TEXT NOT NULL DEFAULT '0',
                duration_100_days TEXT NOT NULL DEFAULT '0',
                duration_80_days TEXT NOT NULL DEFAULT '0',
                is_active TEXT NOT NULL DEFAULT '1',
                source_type_id TEXT NOT NULL DEFAULT '',
                source_sheet TEXT NOT NULL DEFAULT '',
                duration_basis TEXT NOT NULL DEFAULT '',
                note TEXT NOT NULL DEFAULT '',
                PRIMARY KEY (plan_variant_id, type_id)
            )
            """
        )
        try:
            execute_sql(
                """
                INSERT OR REPLACE INTO question_type_orders (
                    plan_variant_id, type_id, type_sort_in_chapter, final_type_label,
                    sort_code, direct_assignable, duration_130_days, duration_100_days,
                    duration_80_days, is_active
                )
                VALUES (?, ?, '1', '题型1', 'TEST-DURATION-001', '1', '1.30', '1.00', '0.80', '1')
                """,
                (plan_variant_id, type_id),
            )
            with sqlite3.connect(DB_PATH) as conn:
                conn.row_factory = sqlite3.Row
                self.assertEqual(
                    services.question_type_for_plan_variant(conn, type_id, plan_variant_id, "A")["suggested_days"],
                    "1.30",
                )
                self.assertEqual(
                    services.question_type_for_plan_variant(conn, type_id, plan_variant_id, "B")["suggested_days"],
                    "1.00",
                )
                self.assertEqual(
                    services.question_type_for_plan_variant(conn, type_id, plan_variant_id, "C")["suggested_days"],
                    "1.00",
                )
                self.assertEqual(
                    services.question_type_for_plan_variant(conn, type_id, plan_variant_id, "D")["suggested_days"],
                    "0.80",
                )
        finally:
            execute_sql(
                """
                DELETE FROM question_type_orders
                WHERE plan_variant_id = ? AND type_id = ?
                """,
                (plan_variant_id, type_id),
            )

    def prepare_onboarding_start_chapter_fixture(self, seed_path: Path) -> dict:
        chapters = fetch_all(
            """
            SELECT chapter_id, MIN(sort_code) AS sort_code
            FROM question_types
            WHERE direct_assignable = '1'
              AND chapter_id <> ''
            GROUP BY chapter_id
            ORDER BY MIN(sort_code), chapter_id
            LIMIT 2
            """
        )
        self.assertGreaterEqual(len(chapters), 2)
        chapter_ids = [row["chapter_id"] for row in chapters]
        target_type = fetch_one(
            """
            SELECT *
            FROM question_types
            WHERE direct_assignable = '1'
              AND chapter_id = ?
            ORDER BY sort_code
            LIMIT 1
            """,
            (chapter_ids[1],),
        )

        created_chapter_table = fetch_optional(
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'chapters'"
        ) is None
        created_order_table = fetch_optional(
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'chapter_orders'"
        ) is None
        execute_sql(
            """
            CREATE TABLE IF NOT EXISTS chapters (
                chapter_id TEXT PRIMARY KEY,
                math_type TEXT NOT NULL,
                subject TEXT NOT NULL DEFAULT '',
                chapter_name TEXT NOT NULL,
                chapter_alias TEXT NOT NULL DEFAULT '',
                is_active TEXT NOT NULL DEFAULT '1',
                source_pdf TEXT NOT NULL DEFAULT '',
                source_pdf_page TEXT NOT NULL DEFAULT '',
                note TEXT NOT NULL DEFAULT ''
            )
            """
        )
        execute_sql(
            """
            CREATE TABLE IF NOT EXISTS chapter_orders (
                plan_variant_id TEXT NOT NULL,
                chapter_id TEXT NOT NULL,
                chapter_sort TEXT NOT NULL,
                is_active TEXT NOT NULL DEFAULT '1'
            )
            """
        )
        for index, chapter_id in enumerate(chapter_ids, start=1):
            execute_sql(
                """
                INSERT OR IGNORE INTO chapters (
                    chapter_id, math_type, subject, chapter_name, is_active
                ) VALUES (?, '数学二', '高数', ?, '1')
                """,
                (chapter_id, f"起始章节{index}"),
            )
            for variant_id in ("math2_mixed", "math2_gaoshu_first"):
                execute_sql(
                    """
                    INSERT OR IGNORE INTO chapter_orders (
                        plan_variant_id, chapter_id, chapter_sort, is_active
                    ) VALUES (?, ?, ?, '1')
                    """,
                    (variant_id, chapter_id, str(index)),
                )

        execute_sql("DELETE FROM questions WHERE question_id LIKE 'TEST-ONBOARDING-START-CHAPTER-%'")
        for question_id, plan_stage in (
            ("TEST-ONBOARDING-START-CHAPTER-ZH", "强化"),
            ("TEST-ONBOARDING-START-CHAPTER-LEGACY", "寮哄寲"),
        ):
            execute_sql(
                """
                INSERT INTO questions (
                    question_id, type_id, type_name_check, question_note, applicable_levels,
                    plan_stage, source, original_position_text, question_order, display_permission
                ) VALUES (?, ?, ?, '', 'A,B,C,D', ?, '880', '起始章节测试题', '1', '仅显示题号')
                """,
                (question_id, target_type["type_id"], target_type["type_name"], plan_stage),
            )

        with sqlite3.connect(seed_path) as conn:
            conn.executescript(
                """
                CREATE TABLE plan_variants (
                  plan_variant_id TEXT PRIMARY KEY,
                  math_type TEXT NOT NULL,
                  review_order TEXT NOT NULL,
                  variant_name TEXT NOT NULL,
                  student_label TEXT NOT NULL,
                  description TEXT,
                  is_active TEXT NOT NULL,
                  version TEXT NOT NULL,
                  sort_order TEXT NOT NULL
                );
                CREATE TABLE exercise_books (
                  book_id TEXT PRIMARY KEY,
                  book_name TEXT NOT NULL,
                  book_type TEXT NOT NULL,
                  math_type TEXT,
                  plan_stage TEXT,
                  is_registration_selectable TEXT NOT NULL,
                  is_active TEXT NOT NULL,
                  sort_order TEXT NOT NULL
                );
                CREATE TABLE classes (
                  class_id TEXT PRIMARY KEY,
                  class_name TEXT NOT NULL,
                  math_type TEXT NOT NULL,
                  level TEXT NOT NULL,
                  group_name TEXT,
                  qq_group_number TEXT,
                  capacity TEXT,
                  is_accepting_new_students TEXT NOT NULL
                );
                CREATE TABLE class_intake_settings (
                  setting_id TEXT PRIMARY KEY,
                  math_type TEXT NOT NULL,
                  level TEXT NOT NULL,
                  active_class_id TEXT NOT NULL,
                  is_active TEXT NOT NULL
                );
                CREATE TABLE chapters (
                  chapter_id TEXT PRIMARY KEY,
                  math_type TEXT NOT NULL,
                  subject TEXT NOT NULL,
                  chapter_name TEXT NOT NULL,
                  is_active TEXT NOT NULL
                );
                CREATE TABLE chapter_orders (
                  plan_variant_id TEXT NOT NULL,
                  chapter_id TEXT NOT NULL,
                  chapter_sort TEXT NOT NULL,
                  is_active TEXT NOT NULL
                );
                """
            )
            conn.executemany(
                """
                INSERT INTO plan_variants (
                  plan_variant_id, math_type, review_order, variant_name,
                  student_label, description, is_active, version, sort_order
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    ("math2_mixed", "数学二", "mixed", "数学二穿插版", "高数线代穿插", "", "1", "test", "1"),
                    (
                        "math2_gaoshu_first",
                        "数学二",
                        "gaoshu_first",
                        "数学二先高数版",
                        "先高数后线代",
                        "",
                        "1",
                        "test",
                        "2",
                    ),
                ],
            )
            conn.executemany(
                """
                INSERT INTO exercise_books (
                  book_id, book_name, book_type, math_type, plan_stage,
                  is_registration_selectable, is_active, sort_order
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    ("m2_880", "880题", "主选", "数学二", "强化", "1", "1", "1"),
                    ("m2_1000", "1000题", "主选", "数学二", "强化", "1", "1", "2"),
                    ("m2_660", "660题", "备选", "数学二", "强化", "1", "1", "3"),
                ],
            )
            for level in ("A", "B", "C", "D"):
                conn.execute(
                    """
                    INSERT INTO classes (
                      class_id, class_name, math_type, level, group_name,
                      qq_group_number, capacity, is_accepting_new_students
                    ) VALUES (?, ?, '数学二', ?, '', '', '50', '1')
                    """,
                    (f"{level}201", f"{level}班", level),
                )
                conn.execute(
                    """
                    INSERT INTO class_intake_settings (
                      setting_id, math_type, level, active_class_id, is_active
                    ) VALUES (?, '数学二', ?, ?, '1')
                    """,
                    (f"SET-{level}", level, f"{level}201"),
                )
            for index, chapter_id in enumerate(chapter_ids, start=1):
                conn.execute(
                    """
                    INSERT INTO chapters (
                      chapter_id, math_type, subject, chapter_name, is_active
                    ) VALUES (?, '数学二', '高数', ?, '1')
                    """,
                    (chapter_id, f"起始章节{index}"),
                )
                for variant_id in ("math2_mixed", "math2_gaoshu_first"):
                    conn.execute(
                        """
                        INSERT INTO chapter_orders (
                          plan_variant_id, chapter_id, chapter_sort, is_active
                        ) VALUES (?, ?, ?, '1')
                        """,
                        (variant_id, chapter_id, str(index)),
                    )
            conn.commit()

        return {
            "chapter_ids": chapter_ids,
            "created_chapter_table": created_chapter_table,
            "created_order_table": created_order_table,
        }

    def cleanup_onboarding_start_chapter_fixture(self, fixture: dict | None) -> None:
        if not fixture:
            return
        execute_sql("DELETE FROM questions WHERE question_id LIKE 'TEST-ONBOARDING-START-CHAPTER-%'")
        if fixture.get("created_order_table"):
            execute_sql("DROP TABLE IF EXISTS chapter_orders")
        if fixture.get("created_chapter_table"):
            execute_sql("DROP TABLE IF EXISTS chapters")

    def test_auth_me_preserves_legacy_user_id_fallback(self) -> None:
        response = self.client.get("/auth/me")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["user"]["account"]["userId"], "10001")

    def test_internal_test_accounts_are_seeded_active_and_loginable(self) -> None:
        passwords = {
            "stu10001": "98106374",
            "stu10002": "53087197",
            "stu10003": "65674025",
            "stu10004": "25068969",
            "stu10005": "87487965",
            "stu10006": "30253556",
            "stu10007": "74363128",
            "stu10008": "97738287",
            "stu10009": "26989080",
            "stu10010": "39843025",
        }
        response = self.client.get("/admin/students?date=2026-06-10")
        self.assertEqual(response.status_code, 200, response.text)
        students = {item["loginAccount"]: item for item in response.json()["items"]}

        for login_account, password in passwords.items():
            with self.subTest(login_account=login_account):
                student = students[login_account]
                self.assertEqual(student["accountType"], "internal_test")
                self.assertEqual(student["accountStatus"], "active")
                self.assertIn(student["classId"], {"A201", "B201", "C201", "D201"})
                self.assertGreaterEqual(student["paidUntil"], "2026-12-31")

                login_response = self.client.post(
                    "/auth/login",
                    json={"loginAccount": login_account, "password": password},
                )
                self.assertEqual(login_response.status_code, 200, login_response.text)
                self.assertEqual(login_response.json()["user"]["userId"], login_account.replace("stu", "", 1))

    def test_student_session_auto_login_and_single_device_replacement(self) -> None:
        student_id = "10001"
        original = fetch_one(
            "SELECT current_session_id FROM users WHERE user_id = ?",
            (student_id,),
        )
        session_ids: list[str] = []
        try:
            login_a = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": "stu10001",
                    "password": "98106374",
                    "deviceId": "phone-a",
                    "deviceName": "Phone A",
                },
            )
            self.assertEqual(login_a.status_code, 200, login_a.text)
            payload_a = login_a.json()
            session_ids.append(payload_a["sessionId"])

            me_a = self.client.get("/auth/me", headers=auth_headers(payload_a, "phone-a"))
            self.assertEqual(me_a.status_code, 200, me_a.text)
            self.assertEqual(me_a.json()["user"]["userId"], "10001")
            self.assertFalse(me_a.json()["needsOnboarding"])

            plan_a = self.client.get(
                "/students/10001/daily-plan",
                params={"date": "2026-06-10"},
                headers=auth_headers(payload_a, "phone-a"),
            )
            self.assertEqual(plan_a.status_code, 200, plan_a.text)

            login_b = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": "stu10001",
                    "password": "98106374",
                    "deviceId": "phone-b",
                    "deviceName": "Phone B",
                },
            )
            self.assertEqual(login_b.status_code, 200, login_b.text)
            payload_b = login_b.json()
            session_ids.append(payload_b["sessionId"])

            replaced_a = self.client.get("/auth/me", headers=auth_headers(payload_a, "phone-a"))
            self.assertEqual(replaced_a.status_code, 401)
            self.assertEqual(replaced_a.json()["detail"]["code"], "SESSION_REPLACED")

            blocked_plan_a = self.client.get(
                "/students/10001/daily-plan",
                params={"date": "2026-06-10"},
                headers=auth_headers(payload_a, "phone-a"),
            )
            self.assertEqual(blocked_plan_a.status_code, 401)
            self.assertEqual(blocked_plan_a.json()["detail"]["code"], "SESSION_REPLACED")

            me_b = self.client.get("/auth/me", headers=auth_headers(payload_b, "phone-b"))
            self.assertEqual(me_b.status_code, 200, me_b.text)
            self.assertEqual(me_b.json()["sessionId"], payload_b["sessionId"])
        finally:
            for session_id in session_ids:
                execute_sql("DELETE FROM user_sessions WHERE session_id = ?", (session_id,))
            execute_sql(
                "UPDATE users SET current_session_id = ? WHERE user_id = ?",
                (original["current_session_id"], student_id),
            )

    def test_student_mutating_routes_require_session_before_resource_lookup(self) -> None:
        reset_test_clock("2026-06-10")
        try:
            with TestClient(app, raise_server_exceptions=False) as no_raise_client:
                advance = no_raise_client.post("/students/10001/test-day/advance")
            self.assertEqual(advance.status_code, 401)
            self.assertEqual(advance.json()["detail"]["code"], "SESSION_REQUIRED")

            pause = self.client.post("/focus-sessions/FS-NOT-FOUND/pause")
            self.assertEqual(pause.status_code, 401)
            self.assertEqual(pause.json()["detail"]["code"], "SESSION_REQUIRED")

            resume = self.client.post("/focus-sessions/FS-NOT-FOUND/resume")
            self.assertEqual(resume.status_code, 401)
            self.assertEqual(resume.json()["detail"]["code"], "SESSION_REQUIRED")

            finish = self.client.post("/focus-sessions/FS-NOT-FOUND/finish", json={"focusedSeconds": 60})
            self.assertEqual(finish.status_code, 401)
            self.assertEqual(finish.json()["detail"]["code"], "SESSION_REQUIRED")

            valid_session = self.student_auth_headers("10001")
            missing = self.client.post("/focus-sessions/FS-NOT-FOUND/pause", headers=valid_session)
            self.assertEqual(missing.status_code, 400)
        finally:
            reset_test_clock()

    def test_onboarding_session_replaces_initial_login_session(self) -> None:
        user_id = ""
        try:
            short_response = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(short_response.status_code, 200, short_response.text)
            short_payload = short_response.json()
            user_id = short_payload["activatedUserId"]

            login_response = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": short_payload["activatedAccount"],
                    "password": short_payload["initialPassword"],
                    "deviceId": "trial-phone",
                    "deviceName": "Trial Phone",
                },
            )
            self.assertEqual(login_response.status_code, 200, login_response.text)
            login_payload = login_response.json()
            self.assertTrue(login_payload["needsOnboarding"])

            onboarding_response = self.client.post(
                f"/students/{user_id}/onboarding",
                headers=auth_headers(login_payload, "trial-phone"),
                json={
                    "level": "A",
                    "targetScore": 120,
                    "primaryBooks": [
                        {"book": "880", "scope": "both"},
                        {"book": "1000", "scope": "both"},
                    ],
                    "backupBook": "660",
                    "dailyTargetMinutes": 180,
                },
            )
            self.assertEqual(onboarding_response.status_code, 200, onboarding_response.text)
            onboard_payload = onboarding_response.json()
            self.assertNotEqual(onboard_payload["sessionId"], login_payload["sessionId"])

            old_session = self.client.get("/auth/me", headers=auth_headers(login_payload, "trial-phone"))
            self.assertEqual(old_session.status_code, 401)
            self.assertEqual(old_session.json()["detail"]["code"], "SESSION_REPLACED")

            new_session = self.client.get("/auth/me", headers=auth_headers(onboard_payload, "trial-phone"))
            self.assertEqual(new_session.status_code, 200, new_session.text)
            self.assertFalse(new_session.json()["needsOnboarding"])

            stored_session = fetch_one(
                "SELECT device_id, device_name FROM user_sessions WHERE session_id = ?",
                (onboard_payload["sessionId"],),
            )
            self.assertEqual(stored_session["device_id"], "trial-phone")
            self.assertEqual(stored_session["device_name"], "Trial Phone")

            profile = fetch_one(
                "SELECT initial_start_chapter_id, current_type_id FROM student_profiles WHERE user_id = ?",
                (user_id,),
            )
            self.assertEqual(profile["initial_start_chapter_id"], "")
            self.assertTrue(profile["current_type_id"])

            wrong_device = self.client.get("/auth/me", headers=auth_headers(onboard_payload, "other-phone"))
            self.assertEqual(wrong_device.status_code, 401)
            self.assertEqual(wrong_device.json()["detail"]["code"], "SESSION_REPLACED")
        finally:
            if user_id:
                for table, column in (
                    ("student_study_goals", "user_id"),
                    ("student_progress", "user_id"),
                    ("student_profiles", "user_id"),
                    ("user_sessions", "user_id"),
                    ("test_user_dates", "user_id"),
                    ("users", "user_id"),
                ):
                    execute_sql(f"DELETE FROM {table} WHERE {column} = ?", (user_id,))
            reset_test_clock()

    def test_daily_plan_questions_expose_structured_source_display_fields(self) -> None:
        _, open_type = self.open_plan_type("10001", "2026-06-10")

        first_question = open_type["questions"][0]

        for field in (
            "bookDisplayLabel",
            "sourceChapterLabel",
            "sourceSectionLabel",
            "sourceSubsectionLabel",
            "sourceQuestionLabel",
        ):
            self.assertIn(field, first_question)
        self.assertTrue(first_question["bookDisplayLabel"])
        self.assertTrue(first_question["sourceQuestionLabel"])
        self.assertIn("sourceGroups", open_type)

    def test_question_preview_uses_only_two_primary_books_and_their_scopes(self) -> None:
        user_id = "10001"
        type_id = "G-01-Z-01"
        test_rows = (
            ("TEST-PREVIEW-880-BASIC", "880", "\u57fa\u7840", "1"),
            ("TEST-PREVIEW-880-ADVANCED", "880", "\u5f3a\u5316", "2"),
            ("TEST-PREVIEW-1000-BASIC", "1000", "\u57fa\u7840", "3"),
            ("TEST-PREVIEW-1000-ADVANCED", "1000", "\u5f3a\u5316", "4"),
            ("TEST-PREVIEW-660-BACKUP", "660", "\u5f3a\u5316", "5"),
            ("TEST-PREVIEW-YANXUAN-OTHER", "yanxuan", "\u5f3a\u5316", "6"),
        )
        original_profile = fetch_one(
            """
            SELECT primary_book_1, primary_book_2, backup_book,
                   primary_book_1_scope, primary_book_2_scope, current_stage
            FROM student_profiles
            WHERE user_id = ?
            """,
            (user_id,),
        )
        try:
            for question_id, source, section, order in test_rows:
                execute_sql(
                    """
                    INSERT OR REPLACE INTO questions (
                        question_id, type_id, type_name_check, question_note,
                        applicable_levels, plan_stage, source,
                        original_position_text, question_order, display_permission,
                        book_display_label, source_chapter_label,
                        source_section_label, source_subsection_label,
                        source_question_label
                    ) VALUES (?, ?, '', '', 'A,B,C,D', '', ?, ?, ?, '', ?,
                              'Chapter 1', ?, 'Choice', ?)
                    """,
                    (
                        question_id,
                        type_id,
                        source,
                        f"{source} preview question {order}",
                        str(900 + int(order)),
                        source,
                        section,
                        order,
                    ),
                )
            execute_sql(
                """
                UPDATE student_profiles
                SET primary_book_1 = 'm2_880', primary_book_2 = 'm2_1000',
                    backup_book = 'm2_660', primary_book_1_scope = 'basic',
                    primary_book_2_scope = 'advanced', current_stage = ''
                WHERE user_id = ?
                """,
                (user_id,),
            )

            headers = self.student_auth_headers(user_id)
            before_progress = fetch_one(
                "SELECT * FROM student_progress WHERE user_id = ?",
                (user_id,),
            )
            before_plan_count = fetch_one(
                "SELECT COUNT(*) AS count FROM daily_plans WHERE user_id = ?",
                (user_id,),
            )["count"]

            missing_auth = self.client.get(f"/students/{user_id}/question-preview/options")
            self.assertEqual(missing_auth.status_code, 401)

            options_response = self.client.get(
                f"/students/{user_id}/question-preview/options",
                headers=headers,
            )
            self.assertEqual(options_response.status_code, 200, options_response.text)
            options = options_response.json()
            self.assertTrue(options["chapters"])
            option_type_ids = {
                item["typeId"]
                for chapter in options["chapters"]
                for item in chapter["types"]
            }
            self.assertIn(type_id, option_type_ids)
            self.assertEqual(
                [item["bookId"] for item in options["primaryBooks"]],
                ["m2_880", "m2_1000"],
            )
            self.assertEqual(
                [item["scope"] for item in options["primaryBooks"]],
                ["basic", "advanced"],
            )

            preview_response = self.client.get(
                f"/students/{user_id}/question-preview/types/{type_id}",
                headers=headers,
            )
            self.assertEqual(preview_response.status_code, 200, preview_response.text)
            preview = preview_response.json()
            returned_ids = {item["questionId"] for item in preview["questions"]}
            self.assertIn("TEST-PREVIEW-880-BASIC", returned_ids)
            self.assertIn("TEST-PREVIEW-1000-ADVANCED", returned_ids)
            self.assertNotIn("TEST-PREVIEW-880-ADVANCED", returned_ids)
            self.assertNotIn("TEST-PREVIEW-1000-BASIC", returned_ids)
            self.assertNotIn("TEST-PREVIEW-660-BACKUP", returned_ids)
            self.assertNotIn("TEST-PREVIEW-YANXUAN-OTHER", returned_ids)
            self.assertEqual(
                {group["sourceName"] for group in preview["sourceGroups"]},
                {"880", "1000"},
            )

            self.assertEqual(
                fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,)),
                before_progress,
            )
            self.assertEqual(
                fetch_one(
                    "SELECT COUNT(*) AS count FROM daily_plans WHERE user_id = ?",
                    (user_id,),
                )["count"],
                before_plan_count,
            )
        finally:
            execute_sql(
                "DELETE FROM questions WHERE question_id LIKE 'TEST-PREVIEW-%'"
            )
            execute_sql(
                """
                UPDATE student_profiles
                SET primary_book_1 = ?, primary_book_2 = ?, backup_book = ?,
                    primary_book_1_scope = ?, primary_book_2_scope = ?,
                    current_stage = ?
                WHERE user_id = ?
                """,
                (
                    original_profile["primary_book_1"],
                    original_profile["primary_book_2"],
                    original_profile["backup_book"],
                    original_profile["primary_book_1_scope"],
                    original_profile["primary_book_2_scope"],
                    original_profile["current_stage"],
                    user_id,
                ),
            )

    def test_daily_plan_extra_practice_uses_student_backup_book(self) -> None:
        user_id = "backuppractice"
        plan_date = "2026-06-10"
        backup_question_id = "TEST-BACKUP-660-Q1"
        try:
            execute_sql(
                """
                INSERT OR REPLACE INTO questions (
                    question_id, type_id, type_name_check, question_note, applicable_levels,
                    plan_stage, source, original_position_text, question_order, display_permission,
                    book_display_label, source_chapter_label, source_section_label,
                    source_subsection_label, source_question_label
                ) VALUES (?, 'G-01-Z-01', '', '', 'A,B,C,D', '', '660',
                    '《660》第1章基础第254题', '1', '',
                    '660', '第一章', '基础', '选择', '254')
                """,
                (backup_question_id,),
            )
            self.create_plan_override_student(
                user_id,
                "G-01-Z-01",
                primary_book_1="m2_880",
                primary_book_2="",
            )
            execute_sql("UPDATE student_profiles SET current_stage = '' WHERE user_id = ?", (user_id,))
            execute_sql("UPDATE student_progress SET plan_stage = '' WHERE user_id = ?", (user_id,))

            payload = services.get_daily_plan(user_id, plan_date)

            extras = payload["plan"].get("extraPracticeTypes") or []
            self.assertTrue(extras)
            self.assertEqual(extras[0]["typeId"], "G-01-Z-01")
            source_names = {
                group.get("sourceName")
                for item in extras
                for group in item.get("sourceGroups", [])
            }
            self.assertEqual(source_names, {"660"})
            self.assertFalse(
                any(
                    group.get("sourceName") == "1000"
                    for item in extras
                    for group in item.get("sourceGroups", [])
                )
            )
        finally:
            self.delete_temp_student(user_id)
            execute_sql("DELETE FROM questions WHERE question_id = ?", (backup_question_id,))

    def test_study_room_is_bound_to_student_class(self) -> None:
        profile = fetch_one("SELECT class_id FROM student_profiles WHERE user_id = ?", ("10001",))
        response = self.client.get(
            "/students/10001/study-room",
            params={"date": "2026-06-10"},
            headers=self.student_auth_headers("10001"),
        )

        self.assertEqual(response.status_code, 200, response.text)
        room_id = response.json()["studyRoom"]["roomId"]
        room = fetch_one("SELECT class_id FROM study_rooms WHERE room_id = ?", (room_id,))
        self.assertEqual(room["class_id"], profile["class_id"])

    def test_admin_student_patch_persists_primary_book_scopes(self) -> None:
        student_id = "10001"
        original = fetch_one(
            """
            SELECT primary_book_1, primary_book_2, backup_book,
                   primary_book_1_scope, primary_book_2_scope
            FROM student_profiles
            WHERE user_id = ?
            """,
            (student_id,),
        )
        try:
            response = self.client.patch(
                f"/admin/students/{student_id}",
                json={
                    "actor": "adm001",
                    "primaryBooks": [
                        {"book": "880", "scope": "basic"},
                        {"book": "1000", "scope": "advanced"},
                    ],
                    "backupBook": "660",
                },
            )

            self.assertEqual(response.status_code, 200, response.text)
            student = response.json()["student"]
            self.assertEqual(student["primaryBooks"], ["880", "1000"])
            self.assertEqual(student["primaryBookScopes"], ["basic", "advanced"])
            profile = fetch_one(
                """
                SELECT primary_book_1, primary_book_2,
                       primary_book_1_scope, primary_book_2_scope
                FROM student_profiles
                WHERE user_id = ?
                """,
                (student_id,),
            )
            self.assertEqual(profile["primary_book_1"], "m2_880")
            self.assertEqual(profile["primary_book_1_scope"], "basic")
            self.assertEqual(profile["primary_book_2"], "m2_1000")
            self.assertEqual(profile["primary_book_2_scope"], "advanced")
        finally:
            execute_sql(
                """
                UPDATE student_profiles
                SET primary_book_1 = ?, primary_book_2 = ?, backup_book = ?,
                    primary_book_1_scope = ?, primary_book_2_scope = ?
                WHERE user_id = ?
                """,
                (
                    original["primary_book_1"],
                    original["primary_book_2"],
                    original["backup_book"],
                    original["primary_book_1_scope"],
                    original["primary_book_2_scope"],
                    student_id,
                ),
            )

    def test_admin_chapter_options_returns_public_chapters(self) -> None:
        created_chapter_table = fetch_optional(
            "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'chapters'"
        ) is None
        original_chapter_rows = []
        if not created_chapter_table:
            original_chapter_rows = fetch_all("SELECT * FROM chapters")
        if created_chapter_table:
            execute_sql(
                """
                CREATE TABLE chapters (
                    chapter_id TEXT PRIMARY KEY,
                    math_type TEXT NOT NULL,
                    subject TEXT NOT NULL DEFAULT '',
                    chapter_name TEXT NOT NULL,
                    chapter_alias TEXT NOT NULL DEFAULT '',
                    is_active TEXT NOT NULL DEFAULT '1',
                    source_pdf TEXT NOT NULL DEFAULT '',
                    source_pdf_page TEXT NOT NULL DEFAULT '',
                    note TEXT NOT NULL DEFAULT ''
                )
                """
            )
        try:
            execute_sql("DELETE FROM chapters WHERE math_type = ?", ("数学二",))
            for index in range(1, 16):
                chapter_id = f"T-CH-{index:02d}"
                execute_sql(
                    """
                    INSERT INTO chapters (
                        chapter_id, math_type, subject, chapter_name, is_active
                    ) VALUES (?, '数学二', '测试', ?, '1')
                    """,
                    (chapter_id, f"测试章节{index:02d}"),
                )

            response = self.client.get("/admin/chapter-options", params={"mathType": "数学二"})

            self.assertEqual(response.status_code, 200, response.text)
            items = response.json()["items"]
            self.assertEqual(len(items), 15)
            self.assertEqual(items[0]["chapterId"], "T-CH-01")
            self.assertEqual(items[0]["chapterName"], "测试章节01")
            self.assertEqual(items[-1]["chapterId"], "T-CH-15")
        finally:
            if created_chapter_table:
                execute_sql("DROP TABLE IF EXISTS chapters")
            else:
                execute_sql("DELETE FROM chapters WHERE chapter_id LIKE 'T-CH-%'")
                for row in original_chapter_rows:
                    execute_sql(
                        f"""
                        INSERT OR REPLACE INTO chapters ({", ".join(row.keys())})
                        VALUES ({", ".join("?" for _ in row)})
                        """,
                        tuple(row.values()),
                    )

    def test_public_registration_options_include_start_chapters_by_review_order(self) -> None:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as temp_dir:
            seed_path = Path(temp_dir) / "public-seed.sqlite"
            with sqlite3.connect(seed_path) as conn:
                conn.executescript(
                    """
                    CREATE TABLE plan_variants (
                      plan_variant_id TEXT PRIMARY KEY,
                      math_type TEXT NOT NULL,
                      review_order TEXT NOT NULL,
                      variant_name TEXT NOT NULL,
                      student_label TEXT NOT NULL,
                      description TEXT,
                      is_active TEXT NOT NULL,
                      version TEXT NOT NULL,
                      sort_order TEXT NOT NULL
                    );
                    CREATE TABLE exercise_books (
                      book_id TEXT PRIMARY KEY,
                      book_name TEXT NOT NULL,
                      book_type TEXT NOT NULL,
                      math_type TEXT,
                      plan_stage TEXT,
                      is_registration_selectable TEXT NOT NULL,
                      is_active TEXT NOT NULL,
                      sort_order TEXT NOT NULL
                    );
                    CREATE TABLE classes (
                      class_id TEXT PRIMARY KEY,
                      class_name TEXT NOT NULL,
                      math_type TEXT NOT NULL,
                      level TEXT NOT NULL,
                      group_name TEXT,
                      qq_group_number TEXT,
                      capacity TEXT,
                      is_accepting_new_students TEXT NOT NULL
                    );
                    CREATE TABLE class_intake_settings (
                      setting_id TEXT PRIMARY KEY,
                      math_type TEXT NOT NULL,
                      level TEXT NOT NULL,
                      active_class_id TEXT NOT NULL,
                      is_active TEXT NOT NULL
                    );
                    CREATE TABLE chapters (
                      chapter_id TEXT PRIMARY KEY,
                      math_type TEXT NOT NULL,
                      subject TEXT NOT NULL,
                      chapter_name TEXT NOT NULL,
                      is_active TEXT NOT NULL
                    );
                    CREATE TABLE chapter_orders (
                      plan_variant_id TEXT NOT NULL,
                      chapter_id TEXT NOT NULL,
                      chapter_sort TEXT NOT NULL,
                      is_active TEXT NOT NULL
                    );
                    """
                )
                conn.executemany(
                    """
                    INSERT INTO plan_variants (
                      plan_variant_id, math_type, review_order, variant_name,
                      student_label, description, is_active, version, sort_order
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        ("math2_mixed", "数学二", "mixed", "数学二穿插版", "高数线代穿插", "", "1", "test", "1"),
                        (
                            "math2_gaoshu_first",
                            "数学二",
                            "gaoshu_first",
                            "数学二先高数版",
                            "先高数后线代",
                            "",
                            "1",
                            "test",
                            "2",
                        ),
                    ],
                )
                conn.executemany(
                    """
                    INSERT INTO exercise_books (
                      book_id, book_name, book_type, math_type, plan_stage,
                      is_registration_selectable, is_active, sort_order
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        ("m2_880", "880题", "主选", "数学二", "强化", "1", "1", "1"),
                        ("m2_1000", "1000题", "主选", "数学二", "强化", "1", "1", "2"),
                        ("m2_660", "660题", "备选", "数学二", "强化", "1", "1", "3"),
                    ],
                )
                chapters = [
                    (
                        f"M2-CH-{index:02d}",
                        "数学二",
                        "高数" if index <= 10 else "线代",
                        f"测试章节{index:02d}",
                        "1",
                    )
                    for index in range(1, 16)
                ]
                conn.executemany(
                    """
                    INSERT INTO chapters (
                      chapter_id, math_type, subject, chapter_name, is_active
                    ) VALUES (?, ?, ?, ?, ?)
                    """,
                    chapters,
                )
                orders = [
                    (variant_id, f"M2-CH-{index:02d}", str(index), "1")
                    for variant_id in ("math2_mixed", "math2_gaoshu_first")
                    for index in range(1, 16)
                ]
                conn.executemany(
                    """
                    INSERT INTO chapter_orders (
                      plan_variant_id, chapter_id, chapter_sort, is_active
                    ) VALUES (?, ?, ?, ?)
                    """,
                    orders,
                )
                conn.commit()

            with patch.dict(os.environ, {"PUBLIC_SEED_SQLITE_PATH": str(seed_path)}):
                public_seed.reset_public_seed_reader()
                response = self.client.get("/public/registration-options", params={"mathType": "数学二"})
                public_seed.reset_public_seed_reader()

        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        options_by_order = payload.get("startChapterOptionsByReviewOrder")
        self.assertIsInstance(options_by_order, dict)
        self.assertIn("gaoshu_first", options_by_order)
        self.assertIn("mixed", options_by_order)
        self.assertEqual(len(options_by_order["gaoshu_first"]), 15)
        self.assertEqual(len(options_by_order["mixed"]), 15)
        expected_sort_order = list(range(1, 16))
        self.assertEqual(
            [item["sortOrder"] for item in options_by_order["gaoshu_first"]],
            expected_sort_order,
        )
        self.assertEqual(
            [item["sortOrder"] for item in options_by_order["mixed"]],
            expected_sort_order,
        )
        first = options_by_order["gaoshu_first"][0]
        self.assertTrue(first["chapterId"])
        self.assertTrue(first["chapterName"])
        self.assertIn(first["subject"], {"高数", "线代", ""})
        self.assertEqual(first["sortOrder"], 1)

    def test_admin_book_change_regenerates_today_unsubmitted_plan_and_clears_old_leftovers(self) -> None:
        user_id = "planoverridebook"
        plan_date = "2026-06-10"
        first_type = self.first_assignable_type()
        try:
            self.create_plan_override_student(user_id, first_type["type_id"], primary_book_1="m2_880")
            services.get_daily_plan(user_id, plan_date)
            open_row = fetch_one(
                """
                SELECT type_id
                FROM daily_plans
                WHERE user_id = ? AND plan_date = ? AND feedback_open_status = 'open'
                ORDER BY CAST(type_index AS INTEGER)
                LIMIT 1
                """,
                (user_id, plan_date),
            )
            questions = fetch_all(
                """
                SELECT question_id
                FROM daily_plans
                WHERE user_id = ? AND plan_date = ? AND type_id = ?
                ORDER BY CAST(question_order AS INTEGER)
                """,
                (user_id, plan_date, open_row["type_id"]),
            )
            answers = [
                {
                    "questionId": row["question_id"],
                    "result": "unattempted" if index == 0 else "correct",
                    "unattemptedAction": "defer" if index == 0 else "",
                }
                for index, row in enumerate(questions)
            ]
            services.submit_feedback(user_id, plan_date, open_row["type_id"], answers, unattempted_action="defer")
            before_progress = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(before_progress["current_type_status"], "unfinished")
            self.assertEqual(before_progress["unfinished_type_id"], open_row["type_id"])

            response = self.client.patch(
                f"/admin/students/{user_id}",
                json={
                    "actor": "adm001",
                    "primaryBooks": [{"book": "1000", "scope": "both"}],
                    "backupBook": "660",
                },
            )

            self.assertEqual(response.status_code, 200, response.text)
            remaining_sources = {
                row["source"]
                for row in fetch_all(
                    """
                    SELECT DISTINCT source
                    FROM daily_plans
                    WHERE user_id = ?
                      AND plan_date = ?
                      AND feedback_open_status <> 'submitted'
                    """,
                    (user_id, plan_date),
                )
            }
            self.assertEqual(remaining_sources, {"1000"})
            submitted_types = fetch_all(
                """
                SELECT DISTINCT type_id
                FROM daily_plans
                WHERE user_id = ? AND plan_date = ? AND feedback_open_status = 'submitted'
                """,
                (user_id, plan_date),
            )
            self.assertEqual([row["type_id"] for row in submitted_types], [open_row["type_id"]])
            after_progress = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(after_progress["current_type_status"], "active")
            self.assertEqual(after_progress["unfinished_type_id"], "")
            self.assertEqual(after_progress["carryover_type_id"], "")
        finally:
            self.delete_temp_student(user_id)

    def test_next_day_plan_fills_after_unsubmitted_or_deferred_type(self) -> None:
        plan_date = "2026-06-10"
        next_date = "2026-06-11"
        fixture_days = {
            "G-01-Z-01": "0.3",
            "G-01-Z-02": "0.1",
            "G-01-T-01-01": "0.6",
            "G-01-T-01-02": "1.0",
        }
        original_days = fetch_all(
            f"""
            SELECT type_id, suggested_days
            FROM question_types
            WHERE type_id IN ({", ".join("?" for _ in fixture_days)})
            """,
            tuple(fixture_days),
        )
        question_ids = [
            "TEST-FILL-T1-Q1",
            "TEST-FILL-T2-Q1",
            "TEST-FILL-T3-Q1",
            "TEST-FILL-T3-Q2",
            *[f"TEST-FILL-T4-Q{index:02d}" for index in range(1, 11)],
        ]

        def insert_question(question_id: str, type_id: str, order: int) -> None:
            execute_sql(
                """
                INSERT OR REPLACE INTO questions (
                    question_id, type_id, type_name_check, question_note, applicable_levels,
                    plan_stage, source, original_position_text, question_order, display_permission
                ) VALUES (?, ?, '', '', 'A,B,C,D', '', '880', ?, ?, '')
                """,
                (question_id, type_id, question_id, str(order)),
            )

        def submit_open_type(user_id: str, result: str = "correct", action: str = "") -> str:
            _, open_type = self.open_plan_type(user_id, plan_date)
            rows = fetch_all(
                """
                SELECT question_id
                FROM daily_plans
                WHERE user_id = ? AND plan_date = ? AND type_id = ?
                ORDER BY CAST(question_order AS INTEGER)
                """,
                (user_id, plan_date, open_type["typeId"]),
            )
            answers = []
            for row in rows:
                item = {"questionId": row["question_id"], "result": result}
                if action:
                    item["unattemptedAction"] = action
                answers.append(item)
            services.submit_feedback(
                user_id,
                plan_date,
                open_type["typeId"],
                answers,
                unattempted_action=action or None,
            )
            return open_type["typeId"]

        def prepare_student(user_id: str) -> None:
            self.create_plan_override_student(user_id, "G-01-Z-01", primary_book_1="m2_880", primary_book_2="")
            execute_sql("UPDATE student_profiles SET current_stage = '' WHERE user_id = ?", (user_id,))
            execute_sql("UPDATE student_progress SET plan_stage = '' WHERE user_id = ?", (user_id,))

        try:
            for type_id, suggested_days in fixture_days.items():
                execute_sql(
                    "UPDATE question_types SET suggested_days = ? WHERE type_id = ?",
                    (suggested_days, type_id),
                )
            insert_question("TEST-FILL-T1-Q1", "G-01-Z-01", 1)
            insert_question("TEST-FILL-T2-Q1", "G-01-Z-02", 1)
            insert_question("TEST-FILL-T3-Q1", "G-01-T-01-01", 1)
            insert_question("TEST-FILL-T3-Q2", "G-01-T-01-01", 2)
            for index in range(1, 11):
                insert_question(f"TEST-FILL-T4-Q{index:02d}", "G-01-T-01-02", index)

            for user_id, should_defer_third in (("planfillopen", False), ("planfilldefer", True)):
                prepare_student(user_id)
                first_day = services.get_daily_plan(user_id, plan_date)
                self.assertEqual(
                    [row["typeId"] for row in first_day["plan"]["types"][:3]],
                    ["G-01-Z-01", "G-01-Z-02", "G-01-T-01-01"],
                )
                submit_open_type(user_id)
                submit_open_type(user_id)
                if should_defer_third:
                    submit_open_type(user_id, result="unattempted", action="defer")

                next_plan = services.get_daily_plan(user_id, next_date)
                next_type_ids = [row["typeId"] for row in next_plan["plan"]["types"]]
                self.assertGreaterEqual(len(next_type_ids), 2)
                self.assertEqual(next_type_ids[:2], ["G-01-T-01-01", "G-01-T-01-02"])
                t4_rows = fetch_all(
                    """
                    SELECT question_id
                    FROM daily_plans
                    WHERE user_id = ? AND plan_date = ? AND type_id = 'G-01-T-01-02'
                    """,
                    (user_id, next_date),
                )
                self.assertGreater(len(t4_rows), 0)
                self.assertLess(len(t4_rows), 10)
        finally:
            for row in original_days:
                execute_sql(
                    "UPDATE question_types SET suggested_days = ? WHERE type_id = ?",
                    (row["suggested_days"], row["type_id"]),
                )
            for question_id in question_ids:
                execute_sql("DELETE FROM questions WHERE question_id = ?", (question_id,))
            self.delete_temp_student("planfillopen")
            self.delete_temp_student("planfilldefer")

    def test_split_type_carryover_survives_same_day_previous_type_submission(self) -> None:
        user_id = "carrysplit"
        plan_date = "2026-06-10"
        next_date = "2026-06-11"
        lead_type_id = "TEST-CARRY-SPLIT-001"
        split_type_id = "TEST-CARRY-SPLIT-002"
        lead_question_id = "TEST-CARRY-SPLIT-001-Q01"
        split_question_ids = [f"TEST-CARRY-SPLIT-002-Q{index:02d}" for index in range(1, 11)]
        type_ids = (lead_type_id, split_type_id)

        def insert_question(question_id: str, type_id: str, order: int) -> None:
            execute_sql(
                """
                INSERT INTO questions (
                    question_id, type_id, type_name_check, question_note, applicable_levels,
                    plan_stage, source, original_position_text, question_order, display_permission
                ) VALUES (?, ?, '', '', 'A,B,C,D', '', '880', ?, ?, '')
                """,
                (question_id, type_id, question_id, str(order)),
            )

        try:
            execute_sql(
                f"DELETE FROM questions WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                f"DELETE FROM question_types WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                """
                INSERT INTO question_types (
                    type_id, sort_code, subject, chapter_id, chapter_name, content_block,
                    type_name, parent_type_id, suggested_days, direct_assignable,
                    lecture_image_url, input_note
                ) VALUES
                    (?, 'ZZZ-CARRY-SPLIT-001', 'test', 'TEST-CARRY-SPLIT', 'Carry split chapter', 'test',
                     'Carry split lead type', '', '0.1', '1', '', ''),
                    (?, 'ZZZ-CARRY-SPLIT-002', 'test', 'TEST-CARRY-SPLIT', 'Carry split chapter', 'test',
                     'Carry split long type', '', '1.6', '1', '', '')
                """,
                type_ids,
            )
            insert_question(lead_question_id, lead_type_id, 1)
            for order, question_id in enumerate(split_question_ids, start=1):
                insert_question(question_id, split_type_id, order)

            self.create_plan_override_student(user_id, lead_type_id, primary_book_1="m2_880", primary_book_2="")
            execute_sql(
                "UPDATE student_profiles SET current_stage = '', plan_variant_id = '', review_order = '' WHERE user_id = ?",
                (user_id,),
            )
            execute_sql("UPDATE student_progress SET plan_stage = '' WHERE user_id = ?", (user_id,))

            plan = services.get_daily_plan(user_id, plan_date)
            self.assertEqual([row["typeId"] for row in plan["plan"]["types"][:2]], [lead_type_id, split_type_id])
            self.assertEqual(
                [row["questionId"] for row in plan["plan"]["types"][1]["questions"]],
                split_question_ids[:7],
            )
            progress_after_plan = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(progress_after_plan["carryover_type_id"], split_type_id)
            self.assertEqual(json.loads(progress_after_plan["carryover_question_ids"]), split_question_ids[7:])

            services.submit_feedback(
                user_id,
                plan_date,
                lead_type_id,
                [{"questionId": lead_question_id, "result": "correct"}],
            )
            progress_after_lead = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(progress_after_lead["current_type_id"], split_type_id)
            self.assertEqual(progress_after_lead["current_type_status"], "active")
            self.assertEqual(progress_after_lead["carryover_type_id"], split_type_id)
            self.assertEqual(json.loads(progress_after_lead["carryover_question_ids"]), split_question_ids[7:])

            services.submit_feedback(
                user_id,
                plan_date,
                split_type_id,
                [{"questionId": question_id, "result": "correct"} for question_id in split_question_ids[:7]],
            )
            progress_after_split = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(progress_after_split["current_type_id"], split_type_id)
            self.assertEqual(progress_after_split["current_type_status"], "carryover")
            self.assertEqual(progress_after_split["carryover_type_id"], split_type_id)
            self.assertEqual(json.loads(progress_after_split["carryover_question_ids"]), split_question_ids[7:])

            next_plan = services.get_daily_plan(user_id, next_date)
            self.assertEqual(next_plan["plan"]["types"][0]["typeId"], split_type_id)
            self.assertEqual(
                [row["questionId"] for row in next_plan["plan"]["types"][0]["questions"]],
                split_question_ids[7:],
            )
        finally:
            self.delete_temp_student(user_id)
            execute_sql(
                f"DELETE FROM questions WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                f"DELETE FROM question_types WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )

    def test_deferred_split_type_keeps_unassigned_carryover_questions(self) -> None:
        user_id = "defercarrysplit"
        plan_date = "2026-06-10"
        next_date = "2026-06-11"
        lead_type_id = "TEST-DEFER-CARRY-SPLIT-001"
        split_type_id = "TEST-DEFER-CARRY-SPLIT-002"
        lead_question_id = "TEST-DEFER-CARRY-SPLIT-001-Q01"
        split_question_ids = [f"TEST-DEFER-CARRY-SPLIT-002-Q{index:02d}" for index in range(1, 11)]
        type_ids = (lead_type_id, split_type_id)

        def insert_question(question_id: str, type_id: str, order: int) -> None:
            execute_sql(
                """
                INSERT INTO questions (
                    question_id, type_id, type_name_check, question_note, applicable_levels,
                    plan_stage, source, original_position_text, question_order, display_permission
                ) VALUES (?, ?, '', '', 'A,B,C,D', '', '880', ?, ?, '')
                """,
                (question_id, type_id, question_id, str(order)),
            )

        try:
            execute_sql(
                f"DELETE FROM questions WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                f"DELETE FROM question_types WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                """
                INSERT INTO question_types (
                    type_id, sort_code, subject, chapter_id, chapter_name, content_block,
                    type_name, parent_type_id, suggested_days, direct_assignable,
                    lecture_image_url, input_note
                ) VALUES
                    (?, 'ZZZ-DEFER-CARRY-SPLIT-001', 'test', 'TEST-DEFER-CARRY-SPLIT', 'Defer carry split chapter', 'test',
                     'Defer carry lead type', '', '0.1', '1', '', ''),
                    (?, 'ZZZ-DEFER-CARRY-SPLIT-002', 'test', 'TEST-DEFER-CARRY-SPLIT', 'Defer carry split chapter', 'test',
                     'Defer carry long type', '', '1.6', '1', '', '')
                """,
                type_ids,
            )
            insert_question(lead_question_id, lead_type_id, 1)
            for order, question_id in enumerate(split_question_ids, start=1):
                insert_question(question_id, split_type_id, order)

            self.create_plan_override_student(user_id, lead_type_id, primary_book_1="m2_880", primary_book_2="")
            execute_sql(
                "UPDATE student_profiles SET current_stage = '', plan_variant_id = '', review_order = '' WHERE user_id = ?",
                (user_id,),
            )
            execute_sql("UPDATE student_progress SET plan_stage = '' WHERE user_id = ?", (user_id,))

            plan = services.get_daily_plan(user_id, plan_date)
            self.assertEqual([row["typeId"] for row in plan["plan"]["types"][:2]], [lead_type_id, split_type_id])
            self.assertEqual(
                [row["questionId"] for row in plan["plan"]["types"][1]["questions"]],
                split_question_ids[:7],
            )
            progress_after_plan = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(progress_after_plan["carryover_type_id"], split_type_id)
            self.assertEqual(json.loads(progress_after_plan["carryover_question_ids"]), split_question_ids[7:])

            services.submit_feedback(
                user_id,
                plan_date,
                lead_type_id,
                [{"questionId": lead_question_id, "result": "correct"}],
            )
            split_answers = [
                {
                    "questionId": question_id,
                    "result": "unattempted" if question_id == split_question_ids[2] else "correct",
                    "unattemptedAction": "defer" if question_id == split_question_ids[2] else "",
                }
                for question_id in split_question_ids[:7]
            ]
            services.submit_feedback(
                user_id,
                plan_date,
                split_type_id,
                split_answers,
                unattempted_action="defer",
            )
            progress_after_defer = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(progress_after_defer["current_type_id"], split_type_id)
            self.assertEqual(progress_after_defer["current_type_status"], "unfinished")
            self.assertEqual(progress_after_defer["unfinished_type_id"], split_type_id)
            self.assertEqual(
                json.loads(progress_after_defer["unfinished_question_ids"]),
                [split_question_ids[2], *split_question_ids[7:]],
            )

            next_plan = services.get_daily_plan(user_id, next_date)
            self.assertEqual(next_plan["plan"]["types"][0]["typeId"], split_type_id)
            self.assertEqual(
                [row["questionId"] for row in next_plan["plan"]["types"][0]["questions"]],
                [split_question_ids[2], *split_question_ids[7:]],
            )
        finally:
            self.delete_temp_student(user_id)
            execute_sql(
                f"DELETE FROM questions WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                f"DELETE FROM question_types WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )

    def test_partial_second_type_carryover_keeps_all_unassigned_questions(self) -> None:
        user_id = "partialcarry"
        plan_date = "2026-06-10"
        lead_type_id = "TEST-PARTIAL-CARRY-001"
        split_type_id = "TEST-PARTIAL-CARRY-002"
        lead_question_id = "TEST-PARTIAL-CARRY-001-Q01"
        split_question_ids = [f"TEST-PARTIAL-CARRY-002-Q{index:02d}" for index in range(1, 79)]
        type_ids = (lead_type_id, split_type_id)

        def insert_question(question_id: str, type_id: str, order: int) -> None:
            execute_sql(
                """
                INSERT INTO questions (
                    question_id, type_id, type_name_check, question_note, applicable_levels,
                    plan_stage, source, original_position_text, question_order, display_permission
                ) VALUES (?, ?, '', '', 'A,B,C,D', '', '880', ?, ?, '')
                """,
                (question_id, type_id, question_id, str(order)),
            )

        try:
            execute_sql(
                f"DELETE FROM questions WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                f"DELETE FROM question_types WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                """
                INSERT INTO question_types (
                    type_id, sort_code, subject, chapter_id, chapter_name, content_block,
                    type_name, parent_type_id, suggested_days, direct_assignable,
                    lecture_image_url, input_note
                ) VALUES
                    (?, 'ZZZ-PARTIAL-CARRY-001', 'test', 'TEST-PARTIAL-CARRY', 'Partial carry chapter', 'test',
                     'Partial carry lead type', '', '0.7', '1', '', ''),
                    (?, 'ZZZ-PARTIAL-CARRY-002', 'test', 'TEST-PARTIAL-CARRY', 'Partial carry chapter', 'test',
                     'Partial carry long type', '', '1.6', '1', '', '')
                """,
                type_ids,
            )
            insert_question(lead_question_id, lead_type_id, 1)
            for order, question_id in enumerate(split_question_ids, start=1):
                insert_question(question_id, split_type_id, order)

            self.create_plan_override_student(user_id, lead_type_id, primary_book_1="m2_880", primary_book_2="")
            execute_sql(
                "UPDATE student_profiles SET current_stage = '', plan_variant_id = '', review_order = '' WHERE user_id = ?",
                (user_id,),
            )
            execute_sql("UPDATE student_progress SET plan_stage = '' WHERE user_id = ?", (user_id,))

            plan = services.get_daily_plan(user_id, plan_date)
            self.assertEqual([row["typeId"] for row in plan["plan"]["types"][:2]], [lead_type_id, split_type_id])
            self.assertEqual(
                [row["questionId"] for row in plan["plan"]["types"][1]["questions"]],
                split_question_ids[:15],
            )
            progress = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(progress["carryover_type_id"], split_type_id)
            self.assertEqual(json.loads(progress["carryover_question_ids"]), split_question_ids[15:])
        finally:
            self.delete_temp_student(user_id)
            execute_sql(
                f"DELETE FROM questions WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )
            execute_sql(
                f"DELETE FROM question_types WHERE type_id IN ({', '.join('?' for _ in type_ids)})",
                type_ids,
            )

    def test_admin_next_plan_start_chapter_applies_when_next_plan_is_generated(self) -> None:
        user_id = "planoverridechapter"
        plan_date = "2026-06-11"
        first_type = self.first_assignable_type()
        target_type = self.first_assignable_type_in_later_chapter(first_type["chapter_id"])
        try:
            self.create_plan_override_student(user_id, first_type["type_id"], primary_book_1="m2_880")
            execute_sql("DELETE FROM questions WHERE question_id = ?", ("TEST-OVERRIDE-G02-001",))
            execute_sql(
                """
                INSERT INTO questions (
                    question_id, type_id, type_name_check, question_note, applicable_levels,
                    plan_stage, source, original_position_text, question_order, display_permission
                ) VALUES (?, ?, ?, '', 'A,B,C,D', '强化', '880', '《880》测试章节第1题', '1', '仅显示题号')
                """,
                ("TEST-OVERRIDE-G02-001", target_type["type_id"], target_type["type_name"]),
            )

            response = self.client.patch(
                f"/admin/students/{user_id}",
                json={
                    "actor": "adm001",
                    "nextPlanStartChapterId": target_type["chapter_id"],
                },
            )

            self.assertEqual(response.status_code, 200, response.text)
            student = response.json()["student"]
            self.assertIn("nextPlanStartChapterId", student)
            self.assertEqual(student["nextPlanStartChapterId"], target_type["chapter_id"])
            services.get_daily_plan(user_id, plan_date)
            first_plan_type = fetch_one(
                """
                SELECT type_id, chapter_id, feedback_open_status
                FROM daily_plans
                WHERE user_id = ? AND plan_date = ?
                ORDER BY CAST(type_index AS INTEGER), CAST(question_order AS INTEGER)
                LIMIT 1
                """,
                (user_id, plan_date),
            )
            self.assertEqual(first_plan_type["chapter_id"], target_type["chapter_id"])
            self.assertEqual(first_plan_type["feedback_open_status"], "open")
            progress = fetch_one("SELECT * FROM student_progress WHERE user_id = ?", (user_id,))
            self.assertEqual(progress["next_plan_start_status"], "applied")
        finally:
            execute_sql("DELETE FROM questions WHERE question_id = ?", ("TEST-OVERRIDE-G02-001",))
            self.delete_temp_student(user_id)

    def test_reset_password_returns_copyable_password_that_can_login(self) -> None:
        student_id = "10001"
        original = fetch_one(
            """
            SELECT login_account, password_hash, login_password, password_changed_at
            FROM users
            WHERE user_id = ?
            """,
            (student_id,),
        )
        try:
            admin_payload = self.admin_auth_payload(device_id="admin-reset-student-password")
            headers = auth_headers(admin_payload, "admin-reset-student-password")
            self.verify_admin_risk_authorization(headers, admin_payload["user"]["userId"])
            response = self.client.post(f"/admin/accounts/{student_id}/reset-password", headers=headers)

            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertTrue(payload["temporaryPassword"])
            self.assertEqual(payload["newPassword"], payload["temporaryPassword"])

            login_response = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": original["login_account"],
                    "password": payload["newPassword"],
                },
            )
            self.assertEqual(login_response.status_code, 200, login_response.text)
            self.assertEqual(login_response.json()["user"]["userId"], student_id)
        finally:
            execute_sql(
                """
                UPDATE users
                SET password_hash = ?, login_password = ?, password_changed_at = ?
                WHERE user_id = ?
                """,
                (
                    original["password_hash"],
                    original["login_password"],
                    original["password_changed_at"],
                    student_id,
                ),
            )

    def test_admin_update_student_daily_math_goal_hours_updates_study_goal(self) -> None:
        student_id = "10001"
        runtime = fetch_optional("SELECT state_value FROM test_runtime_state WHERE state_key = 'current_test_date'")
        stat_date = runtime["state_value"] if runtime else services.CURRENT_DATE
        original_goal = fetch_optional(
            "SELECT * FROM student_study_goals WHERE user_id = ? AND study_date = ?",
            (student_id, stat_date),
        )
        try:
            response = self.client.patch(
                f"/admin/students/{student_id}",
                json={
                    "dailyMathGoalHours": 4.5,
                    "actor": "adm001",
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["student"]["dailyMathGoalHours"], 4.5)
            goal = fetch_one(
                "SELECT target_minutes FROM student_study_goals WHERE user_id = ? AND study_date = ?",
                (student_id, stat_date),
            )
            self.assertEqual(goal["target_minutes"], "270")
        finally:
            if original_goal:
                execute_sql(
                    """
                    UPDATE student_study_goals
                    SET room_id = ?, target_minutes = ?, completed_minutes = ?,
                        completion_rate = ?, current_focus_minutes = ?,
                        current_task_type_id = ?, current_task_name = ?,
                        focus_status = ?, current_round_started_at = ?,
                        created_at = ?, updated_at = ?
                    WHERE goal_id = ?
                    """,
                    (
                        original_goal["room_id"],
                        original_goal["target_minutes"],
                        original_goal["completed_minutes"],
                        original_goal["completion_rate"],
                        original_goal["current_focus_minutes"],
                        original_goal["current_task_type_id"],
                        original_goal["current_task_name"],
                        original_goal["focus_status"],
                        original_goal["current_round_started_at"],
                        original_goal["created_at"],
                        original_goal["updated_at"],
                        original_goal["goal_id"],
                    ),
                )
            else:
                execute_sql(
                    "DELETE FROM student_study_goals WHERE user_id = ? AND study_date = ?",
                    (student_id, stat_date),
                )

    def test_test_clock_api_and_student_next_day_share_runtime_date(self) -> None:
        reset_test_clock("2026-06-10")
        try:
            response = self.client.get("/api/test-clock")
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["currentDate"], "2026-06-10")

            next_day = self.client.post("/api/test-clock/next-day")
            self.assertEqual(next_day.status_code, 200, next_day.text)
            self.assertEqual(next_day.json()["currentDate"], "2026-06-11")
            self.assertEqual(next_day.json()["offsetDays"], 1)

            student_next_day = self.client.post("/test-day/advance", params={"userId": "10001"})
            self.assertEqual(student_next_day.status_code, 200, student_next_day.text)
            self.assertEqual(student_next_day.json()["currentTestDate"], "2026-06-12")

            clock = self.client.get("/api/test-clock")
            self.assertEqual(clock.status_code, 200, clock.text)
            self.assertEqual(clock.json()["currentDate"], "2026-06-12")
        finally:
            reset_test_clock()

    def test_renew_student_service_uses_shared_test_clock_for_expired_student(self) -> None:
        student_id = "10001"
        original_profile = fetch_one(
            """
            SELECT paid_until
            FROM student_profiles
            WHERE user_id = ?
            """,
            (student_id,),
        )
        original_user = fetch_one(
            """
            SELECT account_status
            FROM users
            WHERE user_id = ?
            """,
            (student_id,),
        )
        reset_test_clock("2026-06-12")
        try:
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                ("2026-06-10", student_id),
            )

            response = self.client.post(
                f"/admin/students/{student_id}/renew-service",
                json={"days": 30, "actor": "adm001"},
            )

            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["student"]["paidUntil"], "2026-07-12")
        finally:
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                (original_profile["paid_until"], student_id),
            )
            execute_sql(
                "UPDATE users SET account_status = ? WHERE user_id = ?",
                (original_user["account_status"], student_id),
            )
            reset_test_clock()

    def test_admin_students_marks_expired_trial_as_pending_review(self) -> None:
        student_id = "10011"
        original_profile = fetch_one(
            """
            SELECT paid_until
            FROM student_profiles
            WHERE user_id = ?
            """,
            (student_id,),
        )
        original_user = fetch_one(
            """
            SELECT account_status
            FROM users
            WHERE user_id = ?
            """,
            (student_id,),
        )
        reset_test_clock("2026-06-11")
        try:
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                ("2026-06-10", student_id),
            )
            execute_sql(
                "UPDATE users SET account_status = ? WHERE user_id = ?",
                ("active", student_id),
            )

            response = self.client.get("/admin/students?date=2026-06-11")

            self.assertEqual(response.status_code, 200, response.text)
            student = next(item for item in response.json()["items"] if item["userId"] == student_id)
            self.assertEqual(student["accountType"], "trial")
            self.assertEqual(student["accountStatus"], "expired_pending_review")
        finally:
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                (original_profile["paid_until"], student_id),
            )
            execute_sql(
                "UPDATE users SET account_status = ? WHERE user_id = ?",
                (original_user["account_status"], student_id),
            )
            reset_test_clock()

    def test_admin_students_does_not_soft_delete_expired_trial_after_retention_window(self) -> None:
        student_id = "10011"
        original_profile = fetch_one(
            """
            SELECT paid_until
            FROM student_profiles
            WHERE user_id = ?
            """,
            (student_id,),
        )
        original_user = fetch_one(
            """
            SELECT account_status, account_type
            FROM users
            WHERE user_id = ?
            """,
            (student_id,),
        )
        reset_test_clock("2026-06-15")
        try:
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                ("2026-06-10", student_id),
            )
            execute_sql(
                "UPDATE users SET account_status = ?, account_type = ? WHERE user_id = ?",
                ("active", "trial", student_id),
            )

            response = self.client.get("/admin/students?date=2026-06-15")

            self.assertEqual(response.status_code, 200, response.text)
            student = next(item for item in response.json()["items"] if item["userId"] == student_id)
            self.assertEqual(student["accountType"], "trial")
            self.assertEqual(student["accountStatus"], "expired_pending_review")
            user = fetch_one("SELECT account_status FROM users WHERE user_id = ?", (student_id,))
            self.assertEqual(user["account_status"], "active")
        finally:
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                (original_profile["paid_until"], student_id),
            )
            execute_sql(
                "UPDATE users SET account_status = ?, account_type = ? WHERE user_id = ?",
                (original_user["account_status"], original_user["account_type"], student_id),
            )
            reset_test_clock()

    def test_admin_paid_until_future_reactivates_expired_trial_student(self) -> None:
        student_id = "10011"
        original_profile = fetch_one(
            """
            SELECT paid_until
            FROM student_profiles
            WHERE user_id = ?
            """,
            (student_id,),
        )
        original_user = fetch_one(
            """
            SELECT account_status, account_type
            FROM users
            WHERE user_id = ?
            """,
            (student_id,),
        )
        reset_test_clock("2026-06-15")
        try:
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                ("2026-06-10", student_id),
            )
            execute_sql(
                "UPDATE users SET account_status = ?, account_type = ? WHERE user_id = ?",
                ("deleted", "trial", student_id),
            )

            response = self.client.patch(
                f"/admin/students/{student_id}",
                json={"actor": "adm001", "paidUntil": "2026-06-20"},
            )

            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["student"]["accountStatus"], "active")
            user = fetch_one("SELECT account_status FROM users WHERE user_id = ?", (student_id,))
            self.assertEqual(user["account_status"], "active")
            audit = fetch_one(
                """
                SELECT action, after_value
                FROM admin_audit_logs
                WHERE target = ?
                ORDER BY created_at DESC
                LIMIT 1
                """,
                (student_id,),
            )
            self.assertEqual(audit["action"], "student.update")
            self.assertIn("active", audit["after_value"])
        finally:
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                (original_profile["paid_until"], student_id),
            )
            execute_sql(
                "UPDATE users SET account_status = ?, account_type = ? WHERE user_id = ?",
                (original_user["account_status"], original_user["account_type"], student_id),
            )
            reset_test_clock()

    def test_expired_student_pages_freeze_on_paid_until_snapshot(self) -> None:
        student_id = "19981"
        self.create_plan_override_student(student_id, self.first_assignable_type()["type_id"])
        headers = self.student_auth_headers(student_id, password="12345678")
        reset_test_clock("2026-06-15")
        try:
            execute_sql(
                "UPDATE users SET account_type = ?, account_status = ? WHERE user_id = ?",
                ("trial", "active", student_id),
            )
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                ("2026-06-10", student_id),
            )
            insert_daily_activity(student_id, "2026-06-10", 120)
            insert_daily_activity(student_id, "2026-06-15", 360)

            profile_response = self.client.get(f"/students/{student_id}/profile?date=2026-06-15", headers=headers)
            study_room_response = self.client.get(f"/students/{student_id}/study-room?date=2026-06-15", headers=headers)
            curve_response = self.client.get(f"/students/{student_id}/curve?date=2026-06-15", headers=headers)

            self.assertEqual(profile_response.status_code, 200, profile_response.text)
            self.assertEqual(study_room_response.status_code, 200, study_room_response.text)
            self.assertEqual(curve_response.status_code, 200, curve_response.text)
            self.assertEqual(profile_response.json()["profile"]["todayStudyHours"], 2.0)
            self.assertEqual(study_room_response.json()["studyRoom"]["completedMinutes"], 120)
            self.assertEqual(curve_response.json()["curve"]["activity"][-1]["rawDate"], "2026-06-10")
            focus_response = self.client.post(
                "/focus-sessions/start",
                headers=headers,
                json={
                    "userId": student_id,
                    "planDate": "2026-06-15",
                    "taskTypeId": self.first_assignable_type()["type_id"],
                    "targetMinutes": 30,
                },
            )
            self.assertEqual(focus_response.status_code, 400, focus_response.text)
            generated_today_goal = fetch_one(
                """
                SELECT COUNT(*) AS count
                FROM student_study_goals
                WHERE user_id = ? AND study_date = ?
                """,
                (student_id, "2026-06-15"),
            )
            self.assertEqual(generated_today_goal["count"], 0)
        finally:
            self.delete_temp_student(student_id)
            reset_test_clock()

    def test_admin_wrong_question_export_includes_only_real_wrong_records(self) -> None:
        original_feedbacks = fetch_all(
            "SELECT * FROM question_feedbacks WHERE user_id = ?",
            ("10001",),
        )
        try:
            execute_sql("DELETE FROM question_feedbacks WHERE user_id = ?", ("10001",))
            rows = [
                (
                    "T-EXPORT-CORRECT",
                    "880-01-JC1-004",
                    "correct",
                    "0",
                    "0",
                    "",
                    "",
                    "",
                ),
                (
                    "T-EXPORT-DIRTY-CORRECT-WRONG-RECORD",
                    "880-01-JC1-004",
                    "correct",
                    "1",
                    "0",
                    "calculation_error",
                    "计算错误",
                    "",
                ),
                (
                    "T-EXPORT-WRONG",
                    "880-01-JC1-005",
                    "wrong",
                    "1",
                    "0",
                    "calculation_error",
                    "计算错误",
                    "",
                ),
                (
                    "T-EXPORT-MARK-WRONG",
                    "880-01-JC1-010",
                    "unattempted",
                    "1",
                    "0",
                    "knowledge_gap",
                    "知识漏洞",
                    "mark_wrong",
                ),
                (
                    "T-EXPORT-DEFER",
                    "1000-01-QH-017",
                    "unattempted",
                    "0",
                    "1",
                    "",
                    "",
                    "defer",
                ),
                (
                    "T-EXPORT-DIRTY-BLOCKER",
                    "1000-01-QH-017",
                    "wrong",
                    "1",
                    "1",
                    "method_error",
                    "方法错误",
                    "",
                ),
            ]
            for feedback_id, question_id, result, is_wrong, is_blocker, reason_code, reason_name, action in rows:
                execute_sql(
                    """
                    INSERT INTO question_feedbacks (
                      feedback_id, plan_id, user_id, plan_date, type_id, question_id,
                      result, is_wrong_record, is_unfinished_blocker, submitted_at,
                      created_at, updated_at, wrong_reason_code, wrong_reason_name, unattempted_action
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        feedback_id,
                        "P20260610-10001",
                        "10001",
                        "2026-06-10",
                        "G-01-Z-01",
                        question_id,
                        result,
                        is_wrong,
                        is_blocker,
                        "2026-06-10 18:00:00",
                        "2026-06-10 18:00:00",
                        "2026-06-10 18:00:00",
                        reason_code,
                        reason_name,
                        action,
                    ),
                )

            response = self.client.get("/admin/students/10001/wrong-questions.xlsx")

            self.assertEqual(response.status_code, 200, response.text)
            self.assertIn(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                response.headers["content-type"],
            )
            self.assertIn(
                "attachment",
                response.headers.get("content-disposition", "").lower(),
            )
            workbook = load_workbook(BytesIO(response.content))
            self.assertEqual(workbook.sheetnames, ["错题清单"])
            sheet = workbook["错题清单"]
            self.assertEqual(
                [cell.value for cell in sheet[1]],
                [
                    "学生账号",
                    "学生姓名",
                    "班级",
                    "做题日期",
                    "章节",
                    "题型",
                    "资料书",
                    "资料章节",
                    "资料小节",
                    "题号",
                    "原始题目位置",
                    "错因",
                    "提交时间",
                ],
            )
            exported_question_ids = [
                str(row[10].value)
                for row in sheet.iter_rows(min_row=2)
            ]
            self.assertEqual(len(exported_question_ids), 2)
            self.assertEqual(
                set(exported_question_ids),
                {"《880》第1章基础一第5题", "《880》第1章基础一第10题"},
            )
        finally:
            execute_sql("DELETE FROM question_feedbacks WHERE user_id = ?", ("10001",))
            for row in original_feedbacks:
                execute_sql(
                    f"""
                    INSERT INTO question_feedbacks ({", ".join(row.keys())})
                    VALUES ({", ".join("?" for _ in row)})
                    """,
                    tuple(row.values()),
                )

    def test_admin_wrong_questions_export_returns_404_for_missing_student(self) -> None:
        response = self.client.get("/admin/students/999999/wrong-questions.xlsx")

        self.assertEqual(response.status_code, 404, response.text)

    def test_admin_can_activate_short_term_account_that_student_can_login(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        try:
            response = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )

            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            created_user_id = payload["activatedUserId"]
            self.assertTrue(payload["activatedAccount"].startswith("stu"))
            self.assertTrue(payload["initialPassword"])
            self.assertEqual(payload["student"]["accountType"], "trial")
            self.assertEqual(payload["student"]["paidUntil"], "")
            self.assertTrue(payload["startsAfterOnboarding"])
            self.assertEqual(payload["trialDurationDays"], 7)
            profile_count = fetch_one(
                "SELECT COUNT(*) AS count FROM student_profiles WHERE user_id = ?",
                (created_user_id,),
            )
            self.assertEqual(profile_count["count"], 0)

            login_response = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": payload["activatedAccount"],
                    "password": payload["initialPassword"],
                },
            )
            self.assertEqual(login_response.status_code, 200, login_response.text)
            self.assertEqual(login_response.json()["user"]["userId"], created_user_id)
            self.assertTrue(login_response.json()["needsOnboarding"])

            onboarding_response = self.client.post(
                f"/students/{created_user_id}/onboarding",
                headers=auth_headers(login_response.json()),
                json={
                    "level": "A",
                    "targetScore": 120,
                    "primaryBooks": [
                        {"book": "880", "scope": "both"},
                        {"book": "1000", "scope": "both"},
                    ],
                    "backupBook": "660",
                    "dailyTargetMinutes": 180,
                },
            )
            self.assertEqual(onboarding_response.status_code, 200, onboarding_response.text)
            profile = fetch_one(
                "SELECT paid_until, class_id FROM student_profiles WHERE user_id = ?",
                (created_user_id,),
            )
            self.assertEqual(profile["paid_until"], "2026-06-17")
            study_goal = fetch_one(
                """
                SELECT g.study_date, r.class_id
                FROM student_study_goals g
                JOIN study_rooms r ON r.room_id = g.room_id
                WHERE g.user_id = ?
                """,
                (created_user_id,),
            )
            self.assertEqual(study_goal["study_date"], CURRENT_DATE)
            self.assertEqual(study_goal["class_id"], profile["class_id"])
        finally:
            if created_user_id:
                for table, column in (
                    ("student_study_goals", "user_id"),
                    ("student_progress", "user_id"),
                    ("student_profiles", "user_id"),
                    ("user_sessions", "user_id"),
                    ("test_user_dates", "user_id"),
                    ("users", "user_id"),
                ):
                    execute_sql(f"DELETE FROM {table} WHERE {column} = ?", (created_user_id,))
            reset_test_clock()

    def test_short_term_formal_onboarding_payload_keeps_seven_day_service(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        try:
            activate = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(activate.status_code, 200, activate.text)
            activated = activate.json()
            created_user_id = activated["activatedUserId"]

            login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": activated["activatedAccount"],
                    "password": activated["initialPassword"],
                    "deviceId": "trial-formal-payload-phone",
                    "deviceName": "Trial Formal Payload Phone",
                },
            )
            self.assertEqual(login.status_code, 200, login.text)

            onboarding = self.client.post(
                f"/students/{created_user_id}/onboarding",
                headers=auth_headers(login.json(), "trial-formal-payload-phone"),
                json={
                    "nickname": "试用生",
                    "gender": "female",
                    "qq": "123456789",
                    "newPassword": "abc12345",
                    "confirmPassword": "abc12345",
                    "mathType": "数学二",
                    "targetScore": 120,
                    "foundationLevel": "基础一般",
                    "primaryBooks": [
                        {"book": "m2_880", "scope": "both"},
                        {"book": "m2_1000", "scope": "both"},
                    ],
                    "backupBookId": "m2_660",
                    "backupBook": "660",
                    "dailyStudyHours": 3,
                    "intensiveDays": 60,
                    "reviewOrder": "mixed",
                    "diagnosticAnswers": [
                        {"questionId": "limit-equivalent-infinitesimal", "answer": "A"},
                        {"questionId": "derivative-absolute-value", "answer": "A"},
                    ],
                },
            )
            self.assertEqual(onboarding.status_code, 200, onboarding.text)

            profile = fetch_one(
                "SELECT paid_until FROM student_profiles WHERE user_id = ?",
                (created_user_id,),
            )
            self.assertEqual(profile["paid_until"], "2026-06-17")
        finally:
            if created_user_id:
                self.delete_temp_student(created_user_id)
            reset_test_clock()

    def test_formal_onboarding_assigns_female_avatar_from_g_pool(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        try:
            activate = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(activate.status_code, 200, activate.text)
            activated = activate.json()
            created_user_id = activated["activatedUserId"]

            login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": activated["activatedAccount"],
                    "password": activated["initialPassword"],
                    "deviceId": "female-avatar-phone",
                    "deviceName": "Female Avatar Phone",
                },
            )
            self.assertEqual(login.status_code, 200, login.text)

            onboarding = self.client.post(
                f"/students/{created_user_id}/onboarding",
                headers=auth_headers(login.json(), "female-avatar-phone"),
                json={
                    "nickname": "Alice",
                    "gender": "female",
                    "qq": "123456789",
                    "newPassword": "abc12345",
                    "confirmPassword": "abc12345",
                    "mathType": "\u6570\u5b66\u4e8c",
                    "targetScore": 120,
                    "foundationLevel": "\u57fa\u7840\u4e00\u822c",
                    "primaryBooks": [
                        {"book": "m2_880", "scope": "both"},
                        {"book": "m2_1000", "scope": "both"},
                    ],
                    "backupBookId": "m2_660",
                    "backupBook": "660",
                    "dailyStudyHours": 3,
                    "intensiveDays": 60,
                    "reviewOrder": "mixed",
                    "diagnosticAnswers": [
                        {"questionId": "limit-equivalent-infinitesimal", "answer": "A"},
                        {"questionId": "derivative-absolute-value", "answer": "A"},
                    ],
                },
            )
            self.assertEqual(onboarding.status_code, 200, onboarding.text)

            user = fetch_one(
                "SELECT gender, avatar_url FROM users WHERE user_id = ?",
                (created_user_id,),
            )
            self.assertEqual(user["gender"], "female")
            self.assertRegex(user["avatar_url"], r"^/assets/images/avatars/user/G_kaoyan_avatar_\d+\.webp$")
        finally:
            if created_user_id:
                self.delete_temp_student(created_user_id)
            reset_test_clock()

    def test_onboarding_start_chapter_controls_first_daily_plan(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        fixture = None
        try:
            temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
            self.addCleanup(temp_dir.cleanup)
            seed_path = Path(temp_dir.name) / "public-seed.sqlite"
            fixture = self.prepare_onboarding_start_chapter_fixture(seed_path)
            activate = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(activate.status_code, 200, activate.text)
            activated = activate.json()
            created_user_id = activated["activatedUserId"]

            with patch.dict(os.environ, {"PUBLIC_SEED_SQLITE_PATH": str(seed_path)}):
                public_seed.reset_public_seed_reader()
                options = self.client.get("/public/registration-options", params={"mathType": "数学二"})
                public_seed.reset_public_seed_reader()
            self.assertEqual(options.status_code, 200, options.text)
            chapters = options.json()["startChapterOptionsByReviewOrder"]["gaoshu_first"]
            self.assertGreaterEqual(len(chapters), 2)
            target_chapter = chapters[1]["chapterId"]

            login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": activated["activatedAccount"],
                    "password": activated["initialPassword"],
                    "deviceId": "start-chapter-phone",
                    "deviceName": "Start Chapter Phone",
                },
            )
            self.assertEqual(login.status_code, 200, login.text)

            onboarding = self.client.post(
                f"/students/{created_user_id}/onboarding",
                headers=auth_headers(login.json(), "start-chapter-phone"),
                json={
                    "nickname": "起点生",
                    "gender": "male",
                    "qq": "123456789",
                    "newPassword": "abc12345",
                    "confirmPassword": "abc12345",
                    "mathType": "数学二",
                    "targetScore": 120,
                    "foundationLevel": "基础一般",
                    "primaryBooks": [
                        {"book": "m2_880", "scope": "both"},
                        {"book": "m2_1000", "scope": "both"},
                    ],
                    "backupBookId": "m2_660",
                    "backupBook": "660",
                    "dailyStudyHours": 3,
                    "intensiveDays": 60,
                    "reviewOrder": "gaoshu_first",
                    "startChapterId": target_chapter,
                    "diagnosticAnswers": [
                        {"questionId": "limit-equivalent-infinitesimal", "answer": "A"},
                        {"questionId": "derivative-absolute-value", "answer": "A"},
                    ],
                },
            )
            self.assertEqual(onboarding.status_code, 200, onboarding.text)

            profile = fetch_one(
                "SELECT initial_start_chapter_id FROM student_profiles WHERE user_id = ?",
                (created_user_id,),
            )
            self.assertEqual(profile["initial_start_chapter_id"], target_chapter)

            progress = fetch_one(
                "SELECT current_type_id FROM student_progress WHERE user_id = ?",
                (created_user_id,),
            )
            qtype = fetch_one(
                "SELECT chapter_id FROM question_types WHERE type_id = ?",
                (progress["current_type_id"],),
            )
            self.assertEqual(qtype["chapter_id"], target_chapter)

            plan = self.client.get(
                f"/students/{created_user_id}/daily-plan",
                params={"date": "2026-06-10"},
                headers=auth_headers(onboarding.json(), "start-chapter-phone"),
            )
            self.assertEqual(plan.status_code, 200, plan.text)
            first_type_id = plan.json()["plan"]["types"][0]["typeId"]
            first_type = fetch_one("SELECT chapter_id FROM question_types WHERE type_id = ?", (first_type_id,))
            self.assertEqual(first_type["chapter_id"], target_chapter)
        finally:
            if created_user_id:
                self.delete_temp_student(created_user_id)
            self.cleanup_onboarding_start_chapter_fixture(fixture)
            reset_test_clock()

    def test_onboarding_rejects_invalid_start_chapter(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        try:
            activate = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(activate.status_code, 200, activate.text)
            activated = activate.json()
            created_user_id = activated["activatedUserId"]

            login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": activated["activatedAccount"],
                    "password": activated["initialPassword"],
                    "deviceId": "invalid-start-chapter-phone",
                    "deviceName": "Invalid Start Chapter Phone",
                },
            )
            self.assertEqual(login.status_code, 200, login.text)

            onboarding = self.client.post(
                f"/students/{created_user_id}/onboarding",
                headers=auth_headers(login.json(), "invalid-start-chapter-phone"),
                json={
                    "nickname": "非法章",
                    "gender": "male",
                    "qq": "123456789",
                    "newPassword": "abc12345",
                    "confirmPassword": "abc12345",
                    "mathType": "数学二",
                    "targetScore": 120,
                    "foundationLevel": "基础一般",
                    "primaryBooks": [
                        {"book": "m2_880", "scope": "both"},
                        {"book": "m2_1000", "scope": "both"},
                    ],
                    "backupBookId": "m2_660",
                    "backupBook": "660",
                    "dailyStudyHours": 3,
                    "intensiveDays": 60,
                    "reviewOrder": "gaoshu_first",
                    "startChapterId": "NOT-A-CHAPTER",
                    "diagnosticAnswers": [
                        {"questionId": "limit-equivalent-infinitesimal", "answer": "A"},
                        {"questionId": "derivative-absolute-value", "answer": "A"},
                    ],
                },
            )
            self.assertEqual(onboarding.status_code, 400)
            self.assertIn("起始章节", onboarding.json()["detail"])
        finally:
            if created_user_id:
                self.delete_temp_student(created_user_id)
            reset_test_clock()

    def test_onboarding_rejects_start_chapter_missing_from_review_order(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        target_chapter = ""
        deleted_order_rows: list[dict] = []
        try:
            target_type = fetch_one(
                """
                SELECT *
                FROM question_types
                WHERE direct_assignable = '1'
                  AND chapter_id <> ''
                ORDER BY sort_code
                LIMIT 1 OFFSET 1
                """
            )
            target_chapter = target_type["chapter_id"]
            if fetch_optional("SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'chapter_orders'"):
                deleted_order_rows = fetch_all("SELECT * FROM chapter_orders WHERE chapter_id = ?", (target_chapter,))
                execute_sql("DELETE FROM chapter_orders WHERE chapter_id = ?", (target_chapter,))
            execute_sql("DELETE FROM questions WHERE question_id = ?", ("TEST-ONBOARDING-UNORDERED-CHAPTER",))
            execute_sql(
                """
                INSERT INTO questions (
                    question_id, type_id, type_name_check, question_note, applicable_levels,
                    plan_stage, source, original_position_text, question_order, display_permission
                ) VALUES (?, ?, ?, '', 'A,B,C,D', '强化', '880', '未排序章节测试题', '1', '仅显示题号')
                """,
                ("TEST-ONBOARDING-UNORDERED-CHAPTER", target_type["type_id"], target_type["type_name"]),
            )

            activate = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(activate.status_code, 200, activate.text)
            activated = activate.json()
            created_user_id = activated["activatedUserId"]

            login = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": activated["activatedAccount"],
                    "password": activated["initialPassword"],
                    "deviceId": "unordered-start-chapter-phone",
                    "deviceName": "Unordered Start Chapter Phone",
                },
            )
            self.assertEqual(login.status_code, 200, login.text)

            onboarding = self.client.post(
                f"/students/{created_user_id}/onboarding",
                headers=auth_headers(login.json(), "unordered-start-chapter-phone"),
                json={
                    "nickname": "未排序",
                    "gender": "male",
                    "qq": "123456789",
                    "newPassword": "abc12345",
                    "confirmPassword": "abc12345",
                    "mathType": "数学二",
                    "targetScore": 120,
                    "foundationLevel": "基础一般",
                    "primaryBooks": [
                        {"book": "m2_880", "scope": "both"},
                        {"book": "m2_1000", "scope": "both"},
                    ],
                    "backupBookId": "m2_660",
                    "backupBook": "660",
                    "dailyStudyHours": 3,
                    "intensiveDays": 60,
                    "reviewOrder": "gaoshu_first",
                    "startChapterId": target_chapter,
                    "diagnosticAnswers": [
                        {"questionId": "limit-equivalent-infinitesimal", "answer": "A"},
                        {"questionId": "derivative-absolute-value", "answer": "A"},
                    ],
                },
            )
            self.assertEqual(onboarding.status_code, 400)
            self.assertIn("起始章节", onboarding.json()["detail"])
        finally:
            execute_sql("DELETE FROM questions WHERE question_id = ?", ("TEST-ONBOARDING-UNORDERED-CHAPTER",))
            for row in deleted_order_rows:
                execute_sql(
                    f"""
                    INSERT OR REPLACE INTO chapter_orders ({", ".join(row.keys())})
                    VALUES ({", ".join("?" for _ in row)})
                    """,
                    tuple(row.values()),
                )
            if created_user_id:
                self.delete_temp_student(created_user_id)
            reset_test_clock()

    def test_admin_short_term_accounts_list_and_detail_use_real_statuses(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        try:
            activate = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(activate.status_code, 200, activate.text)
            created_user_id = activate.json()["activatedUserId"]

            before_onboarding = self.client.get("/admin/short-term-accounts", params={"date": "2026-06-10"})
            self.assertEqual(before_onboarding.status_code, 200, before_onboarding.text)
            before_payload = before_onboarding.json()
            before_summary = before_payload["summary"]
            pending = next(item for item in before_payload["items"] if item["userId"] == created_user_id)
            self.assertEqual(pending["shortTermStatus"], "pending_onboarding")
            self.assertEqual(pending["paidUntil"], "")
            self.assertEqual(pending["classId"], "")

            onboarding = self.client.post(
                f"/students/{created_user_id}/onboarding",
                headers=self.student_auth_headers(created_user_id),
                json={
                    "level": "A",
                    "targetScore": 120,
                    "primaryBooks": [
                        {"book": "880", "scope": "both"},
                        {"book": "1000", "scope": "both"},
                    ],
                    "backupBook": "660",
                    "dailyTargetMinutes": 180,
                },
            )
            self.assertEqual(onboarding.status_code, 200, onboarding.text)

            active_list = self.client.get("/admin/short-term-accounts", params={"date": "2026-06-10"})
            self.assertEqual(active_list.status_code, 200, active_list.text)
            active = next(item for item in active_list.json()["items"] if item["userId"] == created_user_id)
            self.assertEqual(active["shortTermStatus"], "active_trial")
            self.assertEqual(active["paidUntil"], "2026-06-17")
            self.assertTrue(active["classId"])
            self.assertEqual(active_list.json()["summary"]["activeTrialCount"], before_summary["activeTrialCount"] + 1)
            self.assertEqual(
                active_list.json()["summary"]["pendingOnboardingCount"],
                before_summary["pendingOnboardingCount"] - 1,
            )

            detail = self.client.get(f"/admin/short-term-accounts/{created_user_id}", params={"date": "2026-06-10"})
            self.assertEqual(detail.status_code, 200, detail.text)
            payload = detail.json()
            self.assertEqual(payload["student"]["userId"], created_user_id)
            self.assertEqual(payload["student"]["accountType"], "trial")
            self.assertEqual(payload["student"]["shortTermStatus"], "active_trial")
            self.assertIn("formalDetail", payload)
            self.assertEqual(payload["formalDetail"]["classId"], payload["student"]["classId"])
            self.assertTrue(payload["conversionPreview"]["loginAccount"].startswith("stu4"))
        finally:
            if created_user_id:
                for table, column in (
                    ("question_feedbacks", "user_id"),
                    ("daily_plans", "user_id"),
                    ("student_daily_activity", "user_id"),
                    ("student_stats", "user_id"),
                    ("student_study_goals", "user_id"),
                    ("student_progress", "user_id"),
                    ("student_profiles", "user_id"),
                    ("user_sessions", "user_id"),
                    ("test_user_dates", "user_id"),
                    ("users", "user_id"),
                ):
                    execute_sql(f"DELETE FROM {table} WHERE {column} = ?", (created_user_id,))
            reset_test_clock()

    def test_admin_short_term_accounts_expiring_and_expired_statuses(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_ids: list[str] = []
        try:
            for paid_until in ("2026-06-12", "2026-06-09"):
                activate = self.client.post(
                    "/admin/short-term-accounts/activate-next",
                    json={"adminUserId": "adm001", "durationDays": 7},
                )
                self.assertEqual(activate.status_code, 200, activate.text)
                user_id = activate.json()["activatedUserId"]
                created_user_ids.append(user_id)

                onboarding = self.client.post(
                    f"/students/{user_id}/onboarding",
                    headers=self.student_auth_headers(user_id),
                    json={
                        "level": "A",
                        "targetScore": 120,
                        "primaryBooks": [
                            {"book": "880", "scope": "both"},
                            {"book": "1000", "scope": "both"},
                        ],
                        "backupBook": "660",
                        "dailyTargetMinutes": 180,
                    },
                )
                self.assertEqual(onboarding.status_code, 200, onboarding.text)
                execute_sql(
                    "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                    (paid_until, user_id),
                )

            response = self.client.get("/admin/short-term-accounts", params={"date": "2026-06-10"})
            self.assertEqual(response.status_code, 200, response.text)
            items = {item["userId"]: item for item in response.json()["items"] if item["userId"] in created_user_ids}
            self.assertEqual(items[created_user_ids[0]]["shortTermStatus"], "expiring_soon")
            self.assertEqual(items[created_user_ids[0]]["daysRemaining"], "2")
            self.assertEqual(items[created_user_ids[1]]["shortTermStatus"], "expired_pending_review")
            self.assertEqual(items[created_user_ids[1]]["daysRemaining"], "-1")
            self.assertGreaterEqual(response.json()["summary"]["expiringSoonCount"], 1)
            self.assertGreaterEqual(response.json()["summary"]["expiredPendingReviewCount"], 1)
        finally:
            for user_id in created_user_ids:
                for table, column in (
                    ("student_study_goals", "user_id"),
                    ("student_progress", "user_id"),
                    ("student_profiles", "user_id"),
                    ("user_sessions", "user_id"),
                    ("test_user_dates", "user_id"),
                    ("users", "user_id"),
                ):
                    execute_sql(f"DELETE FROM {table} WHERE {column} = ?", (user_id,))
            reset_test_clock()

    def test_admin_short_term_accounts_malformed_paid_until_is_safe(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        try:
            activate = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(activate.status_code, 200, activate.text)
            created_user_id = activate.json()["activatedUserId"]

            onboarding = self.client.post(
                f"/students/{created_user_id}/onboarding",
                headers=self.student_auth_headers(created_user_id),
                json={
                    "level": "A",
                    "targetScore": 120,
                    "primaryBooks": [
                        {"book": "880", "scope": "both"},
                        {"book": "1000", "scope": "both"},
                    ],
                    "backupBook": "660",
                    "dailyTargetMinutes": 180,
                },
            )
            self.assertEqual(onboarding.status_code, 200, onboarding.text)
            execute_sql(
                "UPDATE student_profiles SET paid_until = ? WHERE user_id = ?",
                ("not-a-date", created_user_id),
            )

            with TestClient(app, raise_server_exceptions=False) as client:
                list_response = client.get("/admin/short-term-accounts", params={"date": "2026-06-10"})
                self.assertEqual(list_response.status_code, 200, list_response.text)
                item = next(item for item in list_response.json()["items"] if item["userId"] == created_user_id)
                self.assertEqual(item["shortTermStatus"], "expired_pending_review")
                self.assertEqual(item["paidUntil"], "not-a-date")
                self.assertEqual(item["daysRemaining"], "")

                detail = client.get(f"/admin/short-term-accounts/{created_user_id}", params={"date": "2026-06-10"})
                self.assertEqual(detail.status_code, 200, detail.text)
                self.assertEqual(detail.json()["student"]["shortTermStatus"], "expired_pending_review")
                self.assertEqual(detail.json()["student"]["daysRemaining"], "")
        finally:
            if created_user_id:
                for table, column in (
                    ("question_feedbacks", "user_id"),
                    ("daily_plans", "user_id"),
                    ("student_daily_activity", "user_id"),
                    ("student_stats", "user_id"),
                    ("student_study_goals", "user_id"),
                    ("student_progress", "user_id"),
                    ("student_profiles", "user_id"),
                    ("user_sessions", "user_id"),
                    ("test_user_dates", "user_id"),
                    ("users", "user_id"),
                ):
                    execute_sql(f"DELETE FROM {table} WHERE {column} = ?", (created_user_id,))
            reset_test_clock()

    def test_admin_short_term_pending_detail_has_stable_fields(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        try:
            activate = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(activate.status_code, 200, activate.text)
            created_user_id = activate.json()["activatedUserId"]

            detail = self.client.get(f"/admin/short-term-accounts/{created_user_id}", params={"date": "2026-06-10"})
            self.assertEqual(detail.status_code, 200, detail.text)
            payload = detail.json()
            expected_fields = (
                "classId",
                "className",
                "targetScore",
                "todayStatus",
                "primaryBooks",
                "backupBooks",
                "mathType",
                "level",
                "reviewOrder",
                "paidUntil",
                "contactQQ",
                "currentTypeId",
                "openTypeName",
                "manualNextTypeId",
            )
            for key in ("student", "formalDetail"):
                row = payload[key]
                for field in expected_fields:
                    self.assertIn(field, row)
                self.assertEqual(row["targetScore"], 0)
                self.assertEqual(row["classId"], "")
                self.assertEqual(row["className"], "")
                self.assertEqual(row["todayStatus"], "pending_onboarding")
                self.assertEqual(row["primaryBooks"], [])
                self.assertEqual(row["backupBooks"], [])
        finally:
            if created_user_id:
                for table, column in (
                    ("user_sessions", "user_id"),
                    ("test_user_dates", "user_id"),
                    ("users", "user_id"),
                ):
                    execute_sql(f"DELETE FROM {table} WHERE {column} = ?", (created_user_id,))
            reset_test_clock()

    def test_admin_can_deactivate_short_term_account_without_deleting_data(self) -> None:
        reset_test_clock("2026-06-10")
        created_user_id = ""
        try:
            activate = self.client.post("/admin/short-term-accounts/activate-next", json={"adminUserId": "adm001"})
            self.assertEqual(activate.status_code, 200, activate.text)
            created_user_id = activate.json()["activatedUserId"]

            deactivate = self.client.post(
                f"/admin/short-term-accounts/{created_user_id}/deactivate",
                json={"actor": "adm001"},
            )
            self.assertEqual(deactivate.status_code, 200, deactivate.text)
            self.assertEqual(deactivate.json()["student"]["shortTermStatus"], "deactivated")

            user = fetch_one("SELECT account_status FROM users WHERE user_id = ?", (created_user_id,))
            self.assertEqual(user["account_status"], "deactivated")
        finally:
            if created_user_id:
                execute_sql("DELETE FROM users WHERE user_id = ?", (created_user_id,))
            reset_test_clock()

    def test_trial_conversion_uses_separate_40001_pool_without_moving_formal_pointer(self) -> None:
        reset_test_clock("2026-06-10")
        source_user_id = ""
        target_user_id = ""
        try:
            short_response = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(short_response.status_code, 200, short_response.text)
            short_payload = short_response.json()
            source_user_id = short_payload["activatedUserId"]
            login_response = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": short_payload["activatedAccount"],
                    "password": short_payload["initialPassword"],
                },
            )
            self.assertEqual(login_response.status_code, 200, login_response.text)

            onboarding_response = self.client.post(
                f"/students/{source_user_id}/onboarding",
                headers=auth_headers(login_response.json()),
                json={
                    "level": "A",
                    "targetScore": 120,
                    "primaryBooks": [
                        {"book": "880", "scope": "both"},
                        {"book": "1000", "scope": "both"},
                    ],
                    "backupBook": "660",
                    "dailyTargetMinutes": 180,
                },
            )
            self.assertEqual(onboarding_response.status_code, 200, onboarding_response.text)

            execute_sql(
                """
                INSERT INTO question_feedbacks (
                  feedback_id, plan_id, user_id, plan_date, type_id, question_id,
                  result, is_wrong_record, is_unfinished_blocker, submitted_at,
                  created_at, updated_at, wrong_reason_code, wrong_reason_name, unattempted_action
                ) VALUES (?, ?, ?, ?, ?, ?, 'wrong', '1', '0', ?, ?, ?, 'calculation', 'calculation error', '')
                """,
                (
                    f"TRIAL-FB-{source_user_id}",
                    f"P20260610-{source_user_id}",
                    source_user_id,
                    "2026-06-10",
                    "M2-QT-0001",
                    "880-01-JC1-004",
                    "2026-06-10 18:00:00",
                    "2026-06-10 18:00:00",
                    "2026-06-10 18:00:00",
                ),
            )

            preview_before = self.client.get(
                "/admin/formal-accounts/next-preview",
                params={"adminUserId": "adm001"},
            )
            self.assertEqual(preview_before.status_code, 200, preview_before.text)
            formal_preview_login = preview_before.json()["credential"]["loginAccount"]
            self.assertTrue(formal_preview_login.startswith("stu6"))

            migrate_response = self.client.post(
                f"/admin/short-term-accounts/{source_user_id}/migrate-to-formal",
                json={"actor": "adm001"},
            )
            self.assertEqual(migrate_response.status_code, 200, migrate_response.text)
            target = migrate_response.json()["targetStudent"]
            target_user_id = target["userId"]
            self.assertEqual(target["loginAccount"], "stu40001")
            self.assertEqual(target["accountType"], "formal")

            copied_feedbacks = fetch_all(
                "SELECT user_id, wrong_reason_code FROM question_feedbacks WHERE user_id = ?",
                (target_user_id,),
            )
            self.assertEqual(len(copied_feedbacks), 1)
            self.assertEqual(copied_feedbacks[0]["wrong_reason_code"], "calculation")
            old_login = self.client.post(
                "/auth/login",
                json={"loginAccount": short_payload["activatedAccount"], "password": short_payload["initialPassword"]},
            )
            self.assertEqual(old_login.status_code, 400)

            preview_after = self.client.get(
                "/admin/formal-accounts/next-preview",
                params={"adminUserId": "adm001"},
            )
            self.assertEqual(preview_after.status_code, 200, preview_after.text)
            self.assertEqual(preview_after.json()["credential"]["loginAccount"], formal_preview_login)
        finally:
            for user_id in {source_user_id, target_user_id, "stu40001"} - {""}:
                for table, column in (
                    ("question_feedbacks", "user_id"),
                    ("daily_plans", "user_id"),
                    ("student_daily_activity", "user_id"),
                    ("student_stats", "user_id"),
                    ("student_study_goals", "user_id"),
                    ("student_progress", "user_id"),
                    ("student_profiles", "user_id"),
                    ("user_sessions", "user_id"),
                    ("test_user_dates", "user_id"),
                    ("users", "user_id"),
                ):
                    execute_sql(f"DELETE FROM {table} WHERE {column} = ?", (user_id,))
            execute_sql("DELETE FROM formal_account_previews WHERE admin_user_id = ?", ("adm001",))
            reset_test_clock()

    def test_trial_conversion_pool_is_monotonic_and_does_not_backfill_gaps(self) -> None:
        reset_test_clock("2026-06-10")
        source_user_id = ""
        target_user_id = ""
        reserved_user_id = "stu40002"
        try:
            execute_sql(
                """
                INSERT INTO users (
                  user_id, login_account, password_hash, login_password, role, nickname,
                  avatar_url, QQ, account_status, current_session_id, last_login_at,
                  created_at, updated_at, account_type, activated_at
                ) VALUES (?, ?, ?, ?, 'student', ?, '', '', 'active', '', '', ?, ?, 'formal', ?)
                """,
                (
                    reserved_user_id,
                    reserved_user_id,
                    "reserved-hash",
                    "",
                    reserved_user_id,
                    "2026-06-10 00:00:00",
                    "2026-06-10 00:00:00",
                    "2026-06-10 00:00:00",
                ),
            )
            short_response = self.client.post(
                "/admin/short-term-accounts/activate-next",
                json={"adminUserId": "adm001", "durationDays": 7},
            )
            self.assertEqual(short_response.status_code, 200, short_response.text)
            short_payload = short_response.json()
            source_user_id = short_payload["activatedUserId"]
            login_response = self.client.post(
                "/auth/login",
                json={
                    "loginAccount": short_payload["activatedAccount"],
                    "password": short_payload["initialPassword"],
                },
            )
            self.assertEqual(login_response.status_code, 200, login_response.text)

            onboarding_response = self.client.post(
                f"/students/{source_user_id}/onboarding",
                headers=auth_headers(login_response.json()),
                json={
                    "level": "A",
                    "targetScore": 120,
                    "primaryBooks": [
                        {"book": "880", "scope": "both"},
                        {"book": "1000", "scope": "both"},
                    ],
                    "backupBook": "660",
                    "dailyTargetMinutes": 180,
                },
            )
            self.assertEqual(onboarding_response.status_code, 200, onboarding_response.text)

            migrate_response = self.client.post(
                f"/admin/experience-accounts/{source_user_id}/migrate-to-formal",
                json={"actor": "adm001"},
            )
            self.assertEqual(migrate_response.status_code, 200, migrate_response.text)
            target = migrate_response.json()["targetStudent"]
            target_user_id = target["userId"]
            self.assertEqual(target["loginAccount"], "stu40003")
        finally:
            for user_id in {source_user_id, target_user_id, reserved_user_id, "stu40003"} - {""}:
                for table, column in (
                    ("student_study_goals", "user_id"),
                    ("student_progress", "user_id"),
                    ("student_profiles", "user_id"),
                    ("user_sessions", "user_id"),
                    ("test_user_dates", "user_id"),
                    ("users", "user_id"),
                ):
                    execute_sql(f"DELETE FROM {table} WHERE {column} = ?", (user_id,))
            reset_test_clock()

    def test_mark_wrong_unattempted_counts_as_wrong_without_blocking_progress(self) -> None:
        plan, open_type = self.open_plan_type("10010", "2026-06-11")
        answers = [
            {
                "questionId": item["questionId"],
                "result": "unattempted",
                "wrongReason": "calculation_error",
            }
            for item in open_type["questions"]
        ]

        response = self.client.post(
            "/feedback/submit",
            headers=self.student_auth_headers("10010"),
            json={
                "userId": "10010",
                "planDate": plan["plan"]["planDate"],
                "typeId": open_type["typeId"],
                "answers": answers,
                "unattemptedAction": "mark_wrong",
            },
        )

        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["unattemptedAction"], "mark_wrong")
        self.assertIn(payload["planStatus"], {"in_progress", "completed"})
        feedback_stats = fetch_one(
            """
            SELECT
              COUNT(*) AS row_count,
              SUM(CASE WHEN is_wrong_record = '1' THEN 1 ELSE 0 END) AS wrong_rows,
              SUM(CASE WHEN is_unfinished_blocker = '1' THEN 1 ELSE 0 END) AS blocker_rows
            FROM question_feedbacks
            WHERE user_id = ? AND plan_date = ? AND type_id = ?
            """,
            ("10010", plan["plan"]["planDate"], open_type["typeId"]),
        )
        self.assertEqual(feedback_stats["wrong_rows"], feedback_stats["row_count"])
        self.assertEqual(feedback_stats["blocker_rows"], 0)

    def test_curve_returns_weekly_metrics_without_priority_recommendations(self) -> None:
        response = self.client.get(
            "/students/10001/curve",
            params={"date": "2026-06-10"},
            headers=self.student_auth_headers("10001"),
        )

        self.assertEqual(response.status_code, 200, response.text)
        curve = response.json()["curve"]
        self.assertNotIn("priorities", curve)
        weekly = curve["weekly"]
        self.assertIn("completionRate", weekly)
        self.assertIn("accuracyRate", weekly)
        self.assertIn("questionCount", weekly)
        self.assertIn("completedTypeCount", weekly)

    def test_curve_creates_fixed_weekly_target_for_remaining_week(self) -> None:
        execute_sql("DELETE FROM student_weekly_targets WHERE user_id = ?", ("10001",))

        response = self.client.get(
            "/students/10001/curve",
            params={"date": "2026-06-10"},
            headers=self.student_auth_headers("10001"),
        )

        self.assertEqual(response.status_code, 200, response.text)
        weekly = response.json()["curve"]["weekly"]
        target = fetch_one(
            """
            SELECT *
            FROM student_weekly_targets
            WHERE user_id = ? AND week_start_date = ?
            """,
            ("10001", "2026-06-08"),
        )
        self.assertEqual(target["week_end_date"], "2026-06-14")
        self.assertEqual(target["basis_date"], "2026-06-10")
        self.assertEqual(target["target_days"], "5")
        self.assertGreater(int(target["target_question_count"]), 5)
        self.assertEqual(weekly["plannedQuestionCount"], int(target["target_question_count"]))

    def test_curve_weekly_accuracy_excludes_deferred_unattempted_questions(self) -> None:
        execute_sql("DELETE FROM question_feedbacks WHERE user_id = ?", ("10001",))
        execute_sql("DELETE FROM student_weekly_targets WHERE user_id = ?", ("10001",))
        rows = [
            ("T-WEEKLY-CORRECT", "880-01-JC1-004", "correct", "0", "0", ""),
            ("T-WEEKLY-WRONG", "880-01-JC1-005", "wrong", "1", "0", ""),
            ("T-WEEKLY-MARK-WRONG", "880-01-JC1-010", "unattempted", "1", "0", "mark_wrong"),
            ("T-WEEKLY-DEFER", "1000-01-QH-017", "unattempted", "0", "1", "defer"),
        ]
        for feedback_id, question_id, result, is_wrong, is_blocker, action in rows:
            execute_sql(
                """
                INSERT INTO question_feedbacks (
                  feedback_id, plan_id, user_id, plan_date, type_id, question_id,
                  result, is_wrong_record, is_unfinished_blocker, submitted_at,
                  created_at, updated_at, wrong_reason_code, wrong_reason_name, unattempted_action
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    feedback_id,
                    "P20260610-10001",
                    "10001",
                    "2026-06-10",
                    "G-01-Z-01",
                    question_id,
                    result,
                    is_wrong,
                    is_blocker,
                    "2026-06-10 18:00:00",
                    "2026-06-10 18:00:00",
                    "2026-06-10 18:00:00",
                    "calculation_error" if is_wrong == "1" else "",
                    "计算错误" if is_wrong == "1" else "",
                    action,
                ),
            )

        response = self.client.get(
            "/students/10001/curve",
            params={"date": "2026-06-10"},
            headers=self.student_auth_headers("10001"),
        )

        self.assertEqual(response.status_code, 200, response.text)
        weekly = response.json()["curve"]["weekly"]
        self.assertEqual(weekly["questionCount"], 3)
        self.assertEqual(weekly["accuracyRate"], 33)

    def test_curve_mastery_uses_display_chapter_total_question_denominator(self) -> None:
        execute_sql("DELETE FROM question_feedbacks WHERE user_id = ?", ("10001",))
        execute_sql("DELETE FROM student_weekly_targets WHERE user_id = ?", ("10001",))
        execute_sql(
            """
            CREATE TABLE IF NOT EXISTS display_chapters (
              display_chapter_id TEXT PRIMARY KEY,
              math_type TEXT NOT NULL,
              subject TEXT NOT NULL DEFAULT '',
              display_chapter_name TEXT NOT NULL,
              art_label TEXT NOT NULL DEFAULT '',
              svg_filename TEXT NOT NULL DEFAULT '',
              sort_order INTEGER NOT NULL DEFAULT 0,
              is_active TEXT NOT NULL DEFAULT '1'
            )
            """
        )
        execute_sql(
            """
            CREATE TABLE IF NOT EXISTS chapter_display_mappings (
              math_type TEXT NOT NULL,
              chapter_id TEXT NOT NULL,
              display_chapter_id TEXT NOT NULL,
              display_sort_in_group INTEGER NOT NULL DEFAULT 0,
              is_active TEXT NOT NULL DEFAULT '1',
              PRIMARY KEY (math_type, chapter_id)
            )
            """
        )
        execute_sql(
            """
            INSERT OR REPLACE INTO display_chapters (
              display_chapter_id, math_type, subject, display_chapter_name, sort_order, is_active
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            ("TEST-DISPLAY-MATH2", "数学二", "高等数学", "测试合并大章节", 1, "1"),
        )
        for order, chapter_id in enumerate(("G-01", "G-02"), start=1):
            execute_sql(
                """
                INSERT OR REPLACE INTO chapter_display_mappings (
                  math_type, chapter_id, display_chapter_id, display_sort_in_group, is_active
                ) VALUES (?, ?, ?, ?, ?)
                """,
                ("数学二", chapter_id, "TEST-DISPLAY-MATH2", order, "1"),
            )
        execute_sql(
            """
            INSERT OR REPLACE INTO questions (
              question_id, type_id, type_name_check, question_note, applicable_levels,
              plan_stage, source, original_position_text, question_order, display_permission
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "TEST-G02-Q01",
                "G-02-Z-01",
                "导数应用",
                "",
                "A,B,C,D",
                "强化",
                "880",
                "《880》测试第1题",
                "1",
                "仅显示题号",
            ),
        )
        for feedback_id, type_id, question_id in (
            ("T-MASTERY-CORRECT-G01", "G-01-Z-01", "880-01-JC1-004"),
            ("T-MASTERY-CORRECT-G02", "G-02-Z-01", "TEST-G02-Q01"),
        ):
            execute_sql(
                """
                INSERT INTO question_feedbacks (
                  feedback_id, plan_id, user_id, plan_date, type_id, question_id,
                  result, is_wrong_record, is_unfinished_blocker, submitted_at,
                  created_at, updated_at, wrong_reason_code, wrong_reason_name, unattempted_action
                ) VALUES (?, ?, ?, ?, ?, ?, 'correct', '0', '0', ?, ?, ?, '', '', '')
                """,
                (
                    feedback_id,
                    "P20260610-10001",
                    "10001",
                    "2026-06-10",
                    type_id,
                    question_id,
                    "2026-06-10 18:30:00",
                    "2026-06-10 18:30:00",
                    "2026-06-10 18:30:00",
                ),
            )

        response = self.client.get(
            "/students/10001/curve",
            params={"date": "2026-06-10"},
            headers=self.student_auth_headers("10001"),
        )

        self.assertEqual(response.status_code, 200, response.text)
        display_mastery = next(
            row
            for row in response.json()["curve"]["mastery"]
            if row["chapterId"] == "TEST-DISPLAY-MATH2"
        )
        self.assertEqual(display_mastery["submittedQuestionCount"], 2)
        self.assertEqual(display_mastery["correctQuestionCount"], 2)
        self.assertEqual(display_mastery["totalQuestionCount"], 9)
        self.assertEqual(display_mastery["rate"], 22)
        execute_sql("DELETE FROM chapter_display_mappings WHERE display_chapter_id = ?", ("TEST-DISPLAY-MATH2",))
        execute_sql("DELETE FROM display_chapters WHERE display_chapter_id = ?", ("TEST-DISPLAY-MATH2",))
        execute_sql("DELETE FROM questions WHERE question_id = ?", ("TEST-G02-Q01",))
        execute_sql(
            "DELETE FROM question_feedbacks WHERE feedback_id IN (?, ?)",
            ("T-MASTERY-CORRECT-G01", "T-MASTERY-CORRECT-G02"),
        )

    def test_curve_mastery_shows_minimum_progress_after_first_correct_question(self) -> None:
        user_id = "masterymin"
        question_ids = [f"TEST-MIN-MASTERY-Q{index:03d}" for index in range(1, 251)]
        try:
            self.create_plan_override_student(user_id, "G-01-Z-01", primary_book_1="m2_880", primary_book_2="")
            execute_sql("UPDATE student_profiles SET current_stage = '' WHERE user_id = ?", (user_id,))
            execute_sql("UPDATE student_progress SET plan_stage = '' WHERE user_id = ?", (user_id,))
            for index, question_id in enumerate(question_ids, start=1):
                execute_sql(
                    """
                    INSERT OR REPLACE INTO questions (
                        question_id, type_id, type_name_check, question_note, applicable_levels,
                        plan_stage, source, original_position_text, question_order, display_permission
                    ) VALUES (?, 'G-01-Z-01', '', '', 'A,B,C,D', '', '880', ?, ?, '')
                    """,
                    (question_id, question_id, str(1000 + index)),
                )
            execute_sql(
                """
                INSERT INTO question_feedbacks (
                  feedback_id, plan_id, user_id, plan_date, type_id, question_id,
                  result, is_wrong_record, is_unfinished_blocker, submitted_at,
                  created_at, updated_at, wrong_reason_code, wrong_reason_name, unattempted_action
                ) VALUES (?, ?, ?, ?, ?, ?, 'correct', '0', '0', ?, ?, ?, '', '', '')
                """,
                (
                    "T-MIN-MASTERY-CORRECT",
                    "P20260610-masterymin",
                    user_id,
                    "2026-06-10",
                    "G-01-Z-01",
                    question_ids[0],
                    "2026-06-10 18:30:00",
                    "2026-06-10 18:30:00",
                    "2026-06-10 18:30:00",
                ),
            )

            response = self.client.get(
                f"/students/{user_id}/curve",
                params={"date": "2026-06-10"},
                headers=self.student_auth_headers(user_id, password="12345678"),
            )

            self.assertEqual(response.status_code, 200, response.text)
            limit_mastery = next(
                row
                for row in response.json()["curve"]["mastery"]
                if row["chapterId"] == "G-01"
            )
            self.assertGreater(limit_mastery["totalQuestionCount"], 200)
            self.assertEqual(limit_mastery["submittedQuestionCount"], 1)
            self.assertEqual(limit_mastery["correctQuestionCount"], 1)
            self.assertEqual(limit_mastery["rate"], 1)
        finally:
            self.delete_temp_student(user_id)
            for question_id in question_ids:
                execute_sql("DELETE FROM questions WHERE question_id = ?", (question_id,))

    def test_test_day_advance_skips_removed_priority_recommendations_table(self) -> None:
        execute_sql("DROP TABLE IF EXISTS priority_recommendations")
        reset_test_clock("2026-06-10")

        response = self.client.post("/test-day/advance", params={"userId": "10001"})

        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        next_date = payload["currentTestDate"]
        self.assertEqual(next_date, "2026-06-11")

        daily_plan = self.client.get(
            "/students/10001/daily-plan",
            params={"date": next_date},
            headers=self.student_auth_headers("10001"),
        )
        self.assertEqual(daily_plan.status_code, 200, daily_plan.text)
        self.assertEqual(daily_plan.json()["plan"]["planDate"], next_date)
        reset_test_clock()

    def test_rankings_carry_completed_checkin_streak_into_next_test_day(self) -> None:
        user_id = "10001"
        dates = ("2026-06-10", "2026-06-11")
        original_rows = fetch_all(
            """
            SELECT *
            FROM student_stats
            WHERE user_id = ? AND stat_date IN (?, ?)
            """,
            (user_id, *dates),
        )
        columns = [
            "user_id",
            "stat_date",
            "continuous_days",
            "today_completion_rate",
            "weekly_completion_rate",
            "overall_accuracy",
            "repeat_wrong_rate",
            "current_badge_name",
            "next_badge_target",
            "badge_progress_value",
            "badge_progress_rate",
            "rank_in_class",
            "updated_at",
            "today_checkin_order",
            "today_checkin_at",
            "weekly_study_days",
        ]
        try:
            with sqlite3.connect(DB_PATH) as conn:
                conn.row_factory = sqlite3.Row
                services.ensure_student_stats_ranking_columns(conn)
                conn.commit()
            reset_test_clock("2026-06-10")
            execute_sql(
                "DELETE FROM student_stats WHERE user_id = ? AND stat_date IN (?, ?)",
                (user_id, *dates),
            )
            execute_sql(
                f"""
                INSERT INTO student_stats ({", ".join(columns)})
                VALUES ({", ".join("?" for _ in columns)})
                """,
                (
                    user_id,
                    "2026-06-10",
                    "0",
                    "1.0000",
                    "0.0000",
                    "1.0000",
                    "0.0000",
                    "暂无勋章",
                    "7天连续学习勋章",
                    "0",
                    "0.0000",
                    "0",
                    "2026-06-10 20:00:00",
                    "1",
                    "2026-06-10 20:00:00",
                    "0",
                ),
            )

            advance = self.client.post("/test-day/advance", params={"userId": user_id})
            self.assertEqual(advance.status_code, 200, advance.text)
            self.assertEqual(advance.json()["currentTestDate"], "2026-06-11")

            rankings = self.client.get(
                f"/students/{user_id}/rankings",
                params={"date": "2026-06-11"},
                headers=self.student_auth_headers(user_id),
            )
            self.assertEqual(rankings.status_code, 200, rankings.text)
            streak_board = rankings.json()["boards"]["streak"]
            self.assertEqual(streak_board["my"]["continuousDays"], 1)
        finally:
            execute_sql(
                "DELETE FROM student_stats WHERE user_id = ? AND stat_date IN (?, ?)",
                (user_id, *dates),
            )
            for row in original_rows:
                execute_sql(
                    f"""
                    INSERT INTO student_stats ({", ".join(row.keys())})
                    VALUES ({", ".join("?" for _ in row)})
                    """,
                    tuple(row.values()),
                )
            reset_test_clock()

    def test_rankings_use_completed_daily_plan_when_stats_missing_and_small_class_offsets_rank(self) -> None:
        user_id = "rank-small-001"
        class_id = "RANK-SMALL"
        plan_date = "2026-06-12"
        try:
            execute_sql("DELETE FROM daily_plans WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM student_stats WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM student_daily_activity WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM student_profiles WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM users WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM classes WHERE class_id = ?", (class_id,))
            execute_sql(
                """
                INSERT INTO classes (
                    class_id, class_name, math_type, target_exam_year, level,
                    teacher_user_id, group_name, qq_group_number, capacity,
                    is_accepting_new_students, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    class_id,
                    "排行榜小班",
                    "数学二",
                    "2026",
                    "A",
                    "90001",
                    "排行榜小班群",
                    "100000001",
                    "150",
                    "1",
                    "2026-06-12 08:00:00",
                    "2026-06-12 08:00:00",
                ),
            )
            execute_sql(
                """
                INSERT INTO users (
                    user_id, login_account, password_hash, login_password, role,
                    nickname, avatar_url, QQ, account_status, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    user_id,
                    user_id,
                    "",
                    "testpass",
                    "student",
                    "排行榜小班学生",
                    "assets/images/avatars/user/rank-small.webp",
                    "100000001",
                    "active",
                    "2026-06-12 08:00:00",
                    "2026-06-12 08:00:00",
                ),
            )
            execute_sql(
                """
                INSERT INTO student_profiles (
                    user_id, math_type, target_exam_year, target_exam_date, target_score,
                    class_id, level, plan_enabled, paid_until, current_stage,
                    created_at, updated_at, review_order, plan_variant_id
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    user_id,
                    "数学二",
                    "2026",
                    "2026-12-19",
                    "120",
                    class_id,
                    "A",
                    "1",
                    "2026-12-31",
                    "强化",
                    "2026-06-12 08:00:00",
                    "2026-06-12 08:00:00",
                    "mixed",
                    "math2_mixed",
                ),
            )
            for question_index in range(1, 3):
                execute_sql(
                    """
                    INSERT INTO daily_plans (
                        plan_id, user_id, plan_date, plan_status, slogan, exam_date,
                        days_to_exam, type_index, type_id, type_name, question_id,
                        source, original_position_text, question_order,
                        feedback_open_status, created_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        f"RANK-SMALL-{question_index}",
                        user_id,
                        plan_date,
                        "completed",
                        "",
                        "2026-12-19",
                        "190",
                        "1",
                        "M2-QT-0001",
                        "测试题型",
                        f"RANK-SMALL-Q{question_index}",
                        "880",
                        f"第{question_index}题",
                        str(question_index),
                        "submitted",
                        "2026-06-12 08:00:00",
                        "2026-06-12 20:00:00",
                    ),
                )
            insert_daily_activity(user_id, plan_date, 30)

            response = self.client.get(
                f"/students/{user_id}/rankings",
                params={"date": plan_date},
                headers=self.student_auth_headers(user_id),
            )

            self.assertEqual(response.status_code, 200, response.text)
            rankings = response.json()
            self.assertEqual(rankings["summary"]["continuousDays"], 1)
            self.assertEqual(rankings["summary"]["todayRank"], 4)
            self.assertEqual(rankings["summary"]["continuousRank"], 4)
            self.assertEqual(rankings["boards"]["today"]["myRank"], 4)
            self.assertEqual(rankings["boards"]["today"]["my"]["todayStudyMinutes"], 30)
            self.assertTrue(rankings["boards"]["today"]["my"]["metric"].endswith("1h"))
            self.assertEqual(rankings["boards"]["today"]["my"]["todayCompletionRate"], 1)
            self.assertEqual(rankings["boards"]["streak"]["myRank"], 4)
            self.assertEqual(rankings["boards"]["streak"]["my"]["continuousDays"], 1)
        finally:
            execute_sql("DELETE FROM daily_plans WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM student_stats WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM student_daily_activity WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM student_profiles WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM users WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM classes WHERE class_id = ?", (class_id,))

    def test_rankings_do_not_offset_rank_when_class_has_at_least_three_students(self) -> None:
        class_id = "RANK-FULL"
        user_ids = [f"rank-full-00{index}" for index in range(1, 5)]
        plan_date = "2026-06-12"
        try:
            for user_id in user_ids:
                execute_sql("DELETE FROM daily_plans WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM student_stats WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM student_daily_activity WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM student_profiles WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM users WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM classes WHERE class_id = ?", (class_id,))
            execute_sql(
                """
                INSERT INTO classes (
                    class_id, class_name, math_type, target_exam_year, level,
                    teacher_user_id, group_name, qq_group_number, capacity,
                    is_accepting_new_students, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    class_id,
                    "排行榜正常班",
                    "数学二",
                    "2026",
                    "A",
                    "90001",
                    "排行榜正常班群",
                    "100000002",
                    "150",
                    "1",
                    "2026-06-12 08:00:00",
                    "2026-06-12 08:00:00",
                ),
            )
            for index, user_id in enumerate(user_ids, start=1):
                execute_sql(
                    """
                    INSERT INTO users (
                        user_id, login_account, password_hash, login_password, role,
                        nickname, avatar_url, QQ, account_status, created_at, updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        user_id,
                        user_id,
                        "",
                        "testpass",
                        "student",
                        f"排行榜正常学生{index}",
                        f"assets/images/avatars/user/rank-full-{index}.webp",
                        f"10000000{index}",
                        "active",
                        "2026-06-12 08:00:00",
                        "2026-06-12 08:00:00",
                    ),
                )
                execute_sql(
                    """
                    INSERT INTO student_profiles (
                        user_id, math_type, target_exam_year, target_exam_date, target_score,
                        class_id, level, plan_enabled, paid_until, current_stage,
                        created_at, updated_at, review_order, plan_variant_id
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        user_id,
                        "数学二",
                        "2026",
                        "2026-12-19",
                        "120",
                        class_id,
                        "A",
                        "1",
                        "2026-12-31",
                        "强化",
                        "2026-06-12 08:00:00",
                        "2026-06-12 08:00:00",
                        "mixed",
                        "math2_mixed",
                    ),
                )
                execute_sql(
                    """
                    INSERT INTO student_stats (
                        user_id, stat_date, continuous_days, today_completion_rate,
                        weekly_completion_rate, overall_accuracy, repeat_wrong_rate,
                        current_badge_name, next_badge_target, badge_progress_value,
                        badge_progress_rate, rank_in_class, updated_at,
                        today_checkin_order, today_checkin_at, weekly_study_days
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        user_id,
                        plan_date,
                        "1",
                        "1.0000",
                        "0.0000",
                        "1.0000",
                        "0.0000",
                        "暂无勋章",
                        "7天连续学习勋章",
                        "1",
                        "0.1429",
                        str(index),
                        f"2026-06-12 20:0{index}:00",
                        str(index),
                        f"2026-06-12 20:0{index}:00",
                        "0",
                    ),
                )
                insert_daily_activity(user_id, plan_date, [61, 119, 60, 1][index - 1])

            response = self.client.get(
                f"/students/{user_ids[-1]}/rankings",
                params={"date": plan_date},
                headers=self.student_auth_headers(user_ids[-1]),
            )

            self.assertEqual(response.status_code, 200, response.text)
            rankings = response.json()
            today = rankings["boards"]["today"]
            streak = rankings["boards"]["streak"]
            self.assertEqual(today["myRank"], 4)
            self.assertEqual(streak["myRank"], 4)
            self.assertEqual([row["userId"] for row in today["top"]], [user_ids[1], user_ids[0], user_ids[2]])
            self.assertTrue(today["top"][0]["metric"].endswith("2h"))
            self.assertTrue(today["top"][1]["metric"].endswith("2h"))
            self.assertTrue(today["top"][2]["metric"].endswith("1h"))
            self.assertEqual([row["userId"] for row in streak["top"]], user_ids[:3])
        finally:
            for user_id in user_ids:
                execute_sql("DELETE FROM daily_plans WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM student_stats WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM student_daily_activity WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM student_profiles WHERE user_id = ?", (user_id,))
                execute_sql("DELETE FROM users WHERE user_id = ?", (user_id,))
            execute_sql("DELETE FROM classes WHERE class_id = ?", (class_id,))


if __name__ == "__main__":
    unittest.main()
