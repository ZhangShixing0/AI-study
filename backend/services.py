from __future__ import annotations

import base64
import hashlib
import hmac
import json
import math
import os
import secrets
import struct
import time
import uuid
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook

from . import public_seed
from .db import (
    assert_runtime_database_ready,
    get_connection,
    row_to_dict,
    rows_to_dicts,
    table_exists_query,
)
from .settings import APP_ENV, CURRENT_DATE

DAILY_PLAN_TARGET_LOAD = 1.0
DAILY_PLAN_SPLIT_THRESHOLD = 1.2
FORMAL_ACCOUNT_START = 60001
FORMAL_ACCOUNT_END = 99999
SHORT_TERM_ACCOUNT_START = 10011
SHORT_TERM_ACCOUNT_END = 39999
TRIAL_CONVERSION_FORMAL_ACCOUNT_START = 40001
TRIAL_CONVERSION_FORMAL_ACCOUNT_END = 59999
FORMAL_ACCOUNT_PREFIX = "stu"
FORMAL_PREVIEW_TTL_MINUTES = 30
FORMAL_OPENINGS_PER_TOTP = 30
TOTP_ISSUER = "AI Tutor"
TOTP_PERIOD_SECONDS = 30
TOTP_SETUP_WINDOW_MINUTES = 10
RISK_AUTHORIZATION_MINUTES = 10
DEFAULT_ADMIN_USER_ID = "adm001"
DEFAULT_ADMIN_LOGIN_ACCOUNT = "adm001"
DEFAULT_LOCAL_ADMIN_CREDENTIAL = "admin123"
DEFAULT_ADMIN_NICKNAME = "全局运营管理员"
FORMAL_EXAM_DATE = "2026-12-19"
CLASS_INTENSIVE_DAYS = {"A": 45, "B": 60, "C": 60, "D": 75}
SCOPED_PRIMARY_BOOKS = {"880", "1000"}
PRIMARY_BOOK_SCOPES = {"basic", "advanced", "both"}
BOOK_PREFIX_BY_MATH_TYPE = {"数学一": "m1", "数学二": "m2", "数学三": "m3"}
BOOK_CODE_ALIASES = {
    "880题": "880",
    "1000题": "1000",
    "660题": "660",
    "严选题": "yanxuan",
    "严选题基础篇": "yanxuan_basic",
}
BOOK_DISPLAY_CODES = {
    "yanxuan": "严选题",
    "yanxuan_basic": "严选题基础篇",
}
INITIAL_CREDENTIAL_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
DIAGNOSTIC_LIMIT_QUESTION_ID = "limit-equivalent-infinitesimal"
DIAGNOSTIC_DERIVATIVE_QUESTION_ID = "derivative-absolute-value"
DIAGNOSTIC_ANSWERS = {
    DIAGNOSTIC_LIMIT_QUESTION_ID: "B",
    DIAGNOSTIC_DERIVATIVE_QUESTION_ID: "B",
}
MALE_ONBOARDING_AVATAR_POOL = [
    "/assets/images/avatars/user/B_kaoyan_avatar_01.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_03.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_05.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_07.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_13.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_16.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_18.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_24.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_26.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_28.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_30.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_32.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_35.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_38.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_46.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_54.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_62.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_71.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_72.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_73.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_74.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_75.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_76.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_77.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_78.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_79.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_80.webp",
    "/assets/images/avatars/user/B_kaoyan_avatar_81.webp",
]

FEMALE_ONBOARDING_AVATAR_POOL = [
    "/assets/images/avatars/user/G_kaoyan_avatar_02.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_04.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_06.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_08.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_10.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_11.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_12.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_15.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_17.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_19.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_25.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_27.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_29.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_34.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_39.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_41.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_43.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_45.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_47.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_49.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_51.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_53.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_57.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_59.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_61.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_63.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_67.webp",
    "/assets/images/avatars/user/G_kaoyan_avatar_83.webp",
]


def now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def is_production_runtime() -> bool:
    return os.environ.get("APP_ENV", "").strip().lower() == "production"


def current_date_text() -> str:
    if is_production_runtime():
        return date.today().isoformat()
    return CURRENT_DATE


def resolve_requested_date(value: Any = "") -> str:
    return str(value or current_date_text())


def sha256_password(password: str) -> str:
    return "$sha256$" + hashlib.sha256(str(password).encode("utf-8")).hexdigest()


class SessionAuthError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def auth_token_hash(token: str) -> str:
    return "$sha256$" + hashlib.sha256(str(token or "").strip().encode("utf-8")).hexdigest()


def normalize_device_id(value: Any) -> str:
    text = str(value or "").strip()
    return text[:128]


def password_matches(user: dict[str, Any], password: str) -> bool:
    password_hash = str(user.get("password_hash") or "")
    if password_hash.startswith("$sha256$"):
        return password_hash == sha256_password(password)
    return str(password) == str(user.get("login_password") or "")


def password_is_new(user: dict[str, Any], password: str) -> bool:
    return not password_matches(user, password)


def set_user_password_values(values: dict[str, Any], password: str) -> None:
    values["password_hash"] = sha256_password(password)
    values["login_password"] = password


def pct(value: Any) -> int:
    try:
        return round(float(value) * 100)
    except (TypeError, ValueError):
        return 0


def as_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def is_service_expired(profile: dict[str, Any], plan_date: str) -> bool:
    paid_until = str(profile.get("paid_until") or "")
    return bool(paid_until) and paid_until < plan_date


def as_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def days_between(start_date: str, end_date: str) -> int:
    try:
        start = date.fromisoformat(start_date)
        end = date.fromisoformat(end_date)
    except ValueError:
        return 0
    return max(0, (end - start).days)


def iso_date_or_none(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value or ""))
    except ValueError:
        return None


def effective_service_display_date(profile: dict[str, Any], requested_date: str) -> str:
    requested = str(requested_date or CURRENT_DATE)
    requested_day = iso_date_or_none(requested)
    paid_until = str(profile.get("paid_until") or "")
    paid_until_day = iso_date_or_none(paid_until)
    if paid_until_day and requested_day and paid_until_day < requested_day:
        return paid_until_day.isoformat()
    return requested


def assert_service_active_for_write(conn, profile: dict[str, Any], action_date: str) -> None:
    runtime_date = get_runtime_test_date(conn)
    if is_service_expired(profile, action_date) or is_service_expired(profile, runtime_date):
        raise ValueError("服务已到期，请联系老师续费后继续学习。")


def build_student_access_context(
    user: dict[str, Any],
    profile: dict[str, Any],
    requested_date: str,
    not_found_message: str = "未找到该学生",
) -> dict[str, Any]:
    if not user or not profile:
        raise ValueError(not_found_message)
    requested = str(requested_date or CURRENT_DATE)
    effective = effective_service_display_date(profile, requested)
    expired = effective != requested
    return {
        "user": user,
        "profile": profile,
        "requestedDate": requested,
        "effectiveDate": effective,
        "isExpired": expired,
        "canWrite": not expired,
    }


def resolve_student_access_context(
    conn,
    user_id: str,
    requested_date: str,
    not_found_message: str = "未找到该学生",
) -> dict[str, Any]:
    user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
    profile = row_to_dict(
        conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
    )
    return build_student_access_context(user, profile, requested_date, not_found_message)


def level_matches(applicable_levels: str, level: str) -> bool:
    levels = {item.strip() for item in str(applicable_levels or "").split(",") if item.strip()}
    return not levels or level in levels


def book_prefix_for_math_type(math_type: Any) -> str:
    return BOOK_PREFIX_BY_MATH_TYPE.get(str(math_type or "").strip(), "")


def legacy_book_code(book_key: Any) -> str:
    text = str(book_key or "").strip()
    if len(text) > 3 and text[0] == "m" and text[1].isdigit() and text[2] == "_":
        return text.split("_", 1)[1]
    return text


def canonical_profile_book_id(book_key: Any, math_type: Any) -> str:
    text = str(book_key or "").strip()
    if not text:
        return ""
    if len(text) > 3 and text[0] == "m" and text[1].isdigit() and text[2] == "_":
        return text
    if text.startswith("《") and text.endswith("》"):
        text = text[1:-1].strip()
    text = BOOK_CODE_ALIASES.get(text, text)
    prefix = book_prefix_for_math_type(math_type)
    return f"{prefix}_{text}" if prefix else text


def profile_book_display_code(book_key: Any) -> str:
    code = legacy_book_code(book_key)
    return BOOK_DISPLAY_CODES.get(code, code)


def profile_book_input_value(item: Any) -> Any:
    if isinstance(item, dict):
        return item.get("bookId") or item.get("book")
    return item


def expanded_book_keys(book_key: Any, math_type: Any) -> set[str]:
    text = str(book_key or "").strip()
    if not text:
        return set()
    legacy = legacy_book_code(text)
    keys = {text, legacy}
    prefix = book_prefix_for_math_type(math_type)
    if prefix and legacy:
        keys.add(f"{prefix}_{legacy}")
    return keys


def selected_book_keys(profile: dict[str, Any]) -> set[str]:
    keys: set[str] = set()
    math_type = profile.get("math_type") or "数学二"
    for field in ("primary_book_1", "primary_book_2"):
        keys.update(expanded_book_keys(profile.get(field), math_type))
    keys.discard("")
    return keys


def selected_primary_book_scopes(profile: dict[str, Any]) -> dict[str, str]:
    scopes: dict[str, str] = {}
    math_type = profile.get("math_type") or "数学二"
    for index in (1, 2):
        book = profile.get(f"primary_book_{index}")
        if not book:
            continue
        scope = normalize_primary_book_scope(
            book,
            profile.get(f"primary_book_{index}_scope"),
            require_scoped_selection=False,
        )
        for key in expanded_book_keys(book, math_type):
            if key:
                scopes[key] = scope
        canonical = canonical_profile_book_id(book, math_type)
        for key in expanded_book_keys(canonical, math_type):
            if key:
                scopes[key] = scope
    return scopes


def question_is_basic_source(row: dict[str, Any]) -> bool:
    source_section = str(row.get("source_section_label") or "").strip()
    if source_section:
        return source_section.startswith("基础")
    return "基础" in str(row.get("original_position_text") or "")


def question_matches_primary_book_scope(
    row: dict[str, Any],
    profile: dict[str, Any],
    scope_map: dict[str, str] | None = None,
) -> bool:
    book_key = question_book_key(row)
    if legacy_book_code(book_key) not in SCOPED_PRIMARY_BOOKS:
        return True
    scope = (scope_map or selected_primary_book_scopes(profile)).get(book_key)
    if not scope:
        return True
    if scope in {"all", "both"}:
        return True
    is_basic = question_is_basic_source(row)
    if scope == "basic":
        return is_basic
    if scope == "advanced":
        return not is_basic
    return True


def canonicalize_primary_books(primary_books: list[dict[str, str]], math_type: Any) -> list[dict[str, str]]:
    normalized: list[dict[str, str]] = []
    seen: set[str] = set()
    for item in primary_books:
        book_id = canonical_profile_book_id(item.get("book"), math_type)
        if not book_id:
            raise ValueError("主选习题册不能为空")
        if book_id in seen:
            raise ValueError("两本主选习题册不能重复")
        seen.add(book_id)
        normalized.append({"book": book_id, "scope": item.get("scope") or "all"})
    return normalized


def canonicalize_student_profile_book_ids(conn, updated_at: str | None = None) -> dict[str, int]:
    columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(student_profiles)").fetchall()
    }
    book_fields = ("primary_book_1", "primary_book_2", "backup_book")
    if not {"user_id", *book_fields}.issubset(columns):
        return {"updated_rows": 0, "updated_values": 0}

    math_type_select = "math_type" if "math_type" in columns else "'' AS math_type"
    rows = rows_to_dicts(
        conn.execute(
            f"""
            SELECT user_id, {math_type_select}, primary_book_1, primary_book_2, backup_book
            FROM student_profiles
            """
        ).fetchall()
    )
    update_time = updated_at or now_text()
    updated_rows = 0
    updated_values = 0
    for row in rows:
        math_type = row.get("math_type") or "数学二"
        assignments: list[str] = []
        values: list[Any] = []
        changed_values = 0
        for field in book_fields:
            current_value = str(row.get(field) or "").strip()
            canonical_value = canonical_profile_book_id(current_value, math_type)
            if canonical_value != current_value:
                assignments.append(f"{field} = ?")
                values.append(canonical_value)
                changed_values += 1
        if not assignments:
            continue
        if "updated_at" in columns:
            assignments.append("updated_at = ?")
            values.append(update_time)
        values.append(row["user_id"])
        conn.execute(
            f"UPDATE student_profiles SET {', '.join(assignments)} WHERE user_id = ?",
            values,
        )
        updated_rows += 1
        updated_values += changed_values
    return {"updated_rows": updated_rows, "updated_values": updated_values}


def question_book_key(row: dict[str, Any]) -> str:
    return str(row.get("book_id") or row.get("source") or "").strip()


def question_source_label(row: dict[str, Any]) -> str:
    return str(row.get("source") or legacy_book_code(row.get("book_id")) or "").strip()


QUESTION_SOURCE_DISPLAY_FIELDS = (
    "book_display_label",
    "source_chapter_label",
    "source_section_label",
    "source_subsection_label",
    "source_question_label",
)


def ensure_question_source_display_columns(conn) -> None:
    columns = table_columns(conn, "questions")
    for field in QUESTION_SOURCE_DISPLAY_FIELDS:
        if field not in columns:
            conn.execute(f"ALTER TABLE questions ADD COLUMN {field} TEXT NOT NULL DEFAULT ''")


DAILY_PLAN_SEQUENCE_FIELDS = {
    "chapter_id": "TEXT NOT NULL DEFAULT ''",
    "student_type_sort_in_chapter": "TEXT NOT NULL DEFAULT '0'",
    "display_type_label": "TEXT NOT NULL DEFAULT ''",
    "is_chapter_boundary_after": "TEXT NOT NULL DEFAULT '0'",
}


def ensure_daily_plan_sequence_columns(conn) -> None:
    if not table_exists(conn, "daily_plans"):
        return
    columns = table_columns(conn, "daily_plans")
    for field, definition in DAILY_PLAN_SEQUENCE_FIELDS.items():
        if field not in columns:
            conn.execute(f"ALTER TABLE daily_plans ADD COLUMN {field} {definition}")
            columns = table_columns(conn, "daily_plans")


def display_type_label(sort_in_chapter: int) -> str:
    return f"\u9898\u578b{max(sort_in_chapter, 1)}"


def book_display_label(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if text.startswith("《") and text.endswith("》"):
        return text
    return f"《{profile_book_display_code(text)}》"


def source_display_value(value: Any) -> str:
    text = str(value or "").strip()
    return "" if text == "空" else text


def question_source_display(row: dict[str, Any]) -> dict[str, str]:
    source = row.get("source") or row.get("book_id")
    return {
        "bookDisplayLabel": book_display_label(row.get("book_display_label") or source),
        "sourceChapterLabel": source_display_value(row.get("source_chapter_label")),
        "sourceSectionLabel": source_display_value(row.get("source_section_label")),
        "sourceSubsectionLabel": source_display_value(row.get("source_subsection_label")),
        "sourceQuestionLabel": (
            source_display_value(row.get("source_question_label"))
            or str(row.get("original_position_text") or "").strip()
            or str(row.get("question_id") or "").strip()
        ),
    }


def table_columns(conn, table_name: str) -> set[str]:
    return {
        row["name"]
        for row in conn.execute(f"PRAGMA table_info({table_name})").fetchall()
    }


DEFAULT_STUDY_ROOMS = (
    (
        "SR-M2-2026-A201",
        "26\u8003\u7814\u6570\u5b66\u4e8cA01\u73ed\u81ea\u4e60\u5ba4",
        "A201",
    ),
    (
        "SR-M2-2026-B201",
        "26\u8003\u7814\u6570\u5b66\u4e8cB01\u73ed\u81ea\u4e60\u5ba4",
        "B201",
    ),
    (
        "SR-M2-2026-C201",
        "26\u8003\u7814\u6570\u5b66\u4e8cC01\u73ed\u81ea\u4e60\u5ba4",
        "C201",
    ),
    (
        "SR-M2-2026-D201",
        "26\u8003\u7814\u6570\u5b66\u4e8cD01\u73ed\u81ea\u4e60\u5ba4",
        "D201",
    ),
)
DEFAULT_STUDY_ROOM_MATH_TYPE = "\u6570\u5b66\u4e8c"
DEFAULT_STUDY_ROOM_RULE = "\u4e0d\u95f2\u804a\u3001\u4e0d\u5237\u5c4f\u3001\u6309\u8ba1\u5212\u4e13\u6ce8\u81ea\u4e60"


def ensure_study_room_tables(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS study_rooms (
          room_id TEXT PRIMARY KEY,
          room_name TEXT NOT NULL,
          class_id TEXT NOT NULL DEFAULT '',
          math_type TEXT NOT NULL DEFAULT '',
          target_exam_year TEXT NOT NULL DEFAULT '',
          online_count TEXT NOT NULL DEFAULT '0',
          quiet_mode_enabled TEXT NOT NULL DEFAULT '0',
          room_rule TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL DEFAULT '',
          updated_at TEXT NOT NULL DEFAULT ''
        )
        """
    )
    columns = table_columns(conn, "study_rooms")
    if "class_id" not in columns:
        conn.execute("ALTER TABLE study_rooms ADD COLUMN class_id TEXT NOT NULL DEFAULT ''")
    conn.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_study_rooms_class_id
        ON study_rooms(class_id)
        WHERE class_id <> ''
        """
    )
    timestamp = now_text()
    for room_id, room_name, class_id in DEFAULT_STUDY_ROOMS:
        conn.execute(
            """
            INSERT OR IGNORE INTO study_rooms (
              room_id, room_name, class_id, math_type, target_exam_year,
              online_count, quiet_mode_enabled, room_rule, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                room_id,
                room_name,
                class_id,
                DEFAULT_STUDY_ROOM_MATH_TYPE,
                "2026",
                "0",
                "1",
                DEFAULT_STUDY_ROOM_RULE,
                timestamp,
                timestamp,
            ),
        )


QUESTION_TYPE_INTRO_FIELDS = (
    "card_type_label",
    "card_title",
    "importance_tag",
    "question_kind",
    "intro_title",
    "intro_body",
)


def ensure_question_type_card_label_column(conn) -> None:
    columns = table_columns(conn, "question_types")
    if "card_type_label" not in columns:
        conn.execute(
            "ALTER TABLE question_types "
            "ADD COLUMN card_type_label TEXT"
        )


def question_type_intro_select_clause(conn) -> str:
    columns = table_columns(conn, "question_types")
    return ",\n              ".join(
        f"qt.{field} AS {field}" if field in columns else f"'' AS {field}"
        for field in QUESTION_TYPE_INTRO_FIELDS
    )


def question_type_intro_payload(row: dict[str, Any]) -> dict[str, Any]:
    card_type_label = str(row.get("card_type_label") or "").strip()
    card_title = str(row.get("card_title") or row.get("type_name") or "").strip()
    importance_tag = str(row.get("importance_tag") or "").strip()
    question_kind = str(row.get("question_kind") or "").strip()
    intro_title = str(row.get("intro_title") or "").strip()
    intro_body = str(row.get("intro_body") or "").strip()
    summary_prefix = "｜".join(item for item in (importance_tag, question_kind) if item)
    summary_title = intro_title or card_title
    if summary_prefix and summary_title:
        summary_title = f"{summary_prefix}：{summary_title}"
    elif summary_prefix:
        summary_title = summary_prefix
    return {
        "cardTypeLabel": card_type_label,
        "cardTitle": card_title,
        "importanceTag": importance_tag,
        "questionKind": question_kind,
        "introTitle": intro_title,
        "introBody": intro_body,
        "summary": {
            "title": summary_title,
            "body": intro_body,
        },
    }


def ensure_student_chapter_completions_table(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS student_chapter_completions (
          completion_id TEXT PRIMARY KEY,
          user_id TEXT NOT NULL,
          raw_chapter_id TEXT NOT NULL,
          chapter_name TEXT NOT NULL DEFAULT '',
          math_type TEXT NOT NULL DEFAULT '',
          plan_stage TEXT NOT NULL DEFAULT '',
          started_at TEXT NOT NULL DEFAULT '',
          completed_at TEXT NOT NULL DEFAULT '',
          spent_days TEXT NOT NULL DEFAULT '1',
          completed_type_count TEXT NOT NULL DEFAULT '0',
          total_type_count TEXT NOT NULL DEFAULT '0',
          submitted_question_count TEXT NOT NULL DEFAULT '0',
          wrong_question_count TEXT NOT NULL DEFAULT '0',
          accuracy_rate TEXT NOT NULL DEFAULT '0',
          main_loss_reasons TEXT NOT NULL DEFAULT '',
          review_type_ids TEXT NOT NULL DEFAULT '',
          review_type_names TEXT NOT NULL DEFAULT '',
          summary_template_id TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )


def question_type_chapter_info(conn, type_id: str) -> dict[str, str]:
    row = row_to_dict(
        conn.execute("SELECT * FROM question_types WHERE type_id = ?", (type_id,)).fetchone()
    ) or {}
    return {
        "chapter_id": str(row.get("chapter_id") or row.get("raw_chapter_id") or "").strip(),
        "chapter_name": str(row.get("chapter_name") or row.get("display_chapter_name") or "").strip(),
        "type_name": str(row.get("type_name") or row.get("card_title") or "").strip(),
    }


def find_chapter_summary_template(conn, math_type: str, chapter_id: str, plan_stage: str) -> dict[str, Any]:
    columns = table_columns(conn, "chapter_summary_templates")
    if not columns:
        return {}
    chapter_column = "chapter_id" if "chapter_id" in columns else "raw_chapter_id"
    active_clause = "AND is_active = '1'" if "is_active" in columns else ""
    return row_to_dict(
        conn.execute(
            f"""
            SELECT *
            FROM chapter_summary_templates
            WHERE math_type = ?
              AND {chapter_column} = ?
              AND plan_stage = ?
              {active_clause}
            LIMIT 1
            """,
            (math_type, chapter_id, plan_stage),
        ).fetchone()
    ) or {}


def chapter_feedback_stats(conn, user_id: str, plan_date: str, chapter_id: str) -> dict[str, Any]:
    feedbacks = rows_to_dicts(
        conn.execute(
            """
            SELECT f.*, dp.type_name
            FROM question_feedbacks f
            JOIN daily_plans dp
              ON dp.user_id = f.user_id
             AND dp.plan_date = f.plan_date
             AND dp.type_id = f.type_id
             AND dp.question_id = f.question_id
            JOIN question_types qt ON qt.type_id = f.type_id
            WHERE f.user_id = ?
              AND f.plan_date = ?
              AND qt.chapter_id = ?
            ORDER BY f.submitted_at, f.type_id, f.question_id
            """,
            (user_id, plan_date, chapter_id),
        ).fetchall()
    )
    submitted_count = len(feedbacks)
    wrong_rows = [row for row in feedbacks if str(row.get("result") or "") in {"wrong", "unattempted"}]
    wrong_count = len(wrong_rows)
    correct_count = max(0, submitted_count - wrong_count)
    reason_names = []
    for row in wrong_rows:
        reason_name = str(row.get("wrong_reason_name") or "").strip()
        if reason_name and reason_name not in reason_names:
            reason_names.append(reason_name)
    return {
        "submitted_count": submitted_count,
        "wrong_count": wrong_count,
        "accuracy_rate": (correct_count / submitted_count) if submitted_count else 0,
        "main_loss_reasons": ", ".join(reason_names),
    }


def record_chapter_completion_if_needed(
    conn,
    user_id: str,
    plan_date: str,
    type_id: str,
    next_type_id: str,
    profile: dict[str, Any],
    submitted_at: str,
) -> None:
    current_chapter = question_type_chapter_info(conn, type_id)
    chapter_id = current_chapter["chapter_id"]
    if not chapter_id:
        return
    if next_type_id:
        next_chapter = question_type_chapter_info(conn, next_type_id)
        if next_chapter["chapter_id"] == chapter_id:
            return

    ensure_student_chapter_completions_table(conn)
    math_type = str(profile.get("math_type") or "").strip()
    plan_stage = str(profile.get("current_stage") or "强化").strip()
    template = find_chapter_summary_template(conn, math_type, chapter_id, plan_stage)
    stats = chapter_feedback_stats(conn, user_id, plan_date, chapter_id)
    total_type_count = row_to_dict(
        conn.execute(
            """
            SELECT COUNT(DISTINCT dp.type_id) AS count_value
            FROM daily_plans dp
            JOIN question_types qt ON qt.type_id = dp.type_id
            WHERE dp.user_id = ?
              AND dp.plan_date = ?
              AND qt.chapter_id = ?
            """,
            (user_id, plan_date, chapter_id),
        ).fetchone()
    ) or {}
    completed_type_count = row_to_dict(
        conn.execute(
            """
            SELECT COUNT(DISTINCT dp.type_id) AS count_value
            FROM daily_plans dp
            JOIN question_types qt ON qt.type_id = dp.type_id
            WHERE dp.user_id = ?
              AND dp.plan_date = ?
              AND qt.chapter_id = ?
              AND dp.feedback_open_status = 'submitted'
            """,
            (user_id, plan_date, chapter_id),
        ).fetchone()
    ) or {}
    completion_id = f"SCC-{user_id}-{chapter_id}-{plan_date}"
    conn.execute(
        """
        DELETE FROM student_chapter_completions
        WHERE user_id = ? AND raw_chapter_id = ? AND plan_stage = ?
        """,
        (user_id, chapter_id, plan_stage),
    )
    conn.execute(
        """
        INSERT INTO student_chapter_completions (
          completion_id, user_id, raw_chapter_id, chapter_name, math_type, plan_stage,
          started_at, completed_at, spent_days, completed_type_count, total_type_count,
          submitted_question_count, wrong_question_count, accuracy_rate,
          main_loss_reasons, review_type_ids, review_type_names, summary_template_id,
          created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            completion_id,
            user_id,
            chapter_id,
            current_chapter["chapter_name"],
            math_type,
            plan_stage,
            plan_date,
            submitted_at,
            "1",
            str(as_int(completed_type_count.get("count_value"))),
            str(as_int(total_type_count.get("count_value"))),
            str(stats["submitted_count"]),
            str(stats["wrong_count"]),
            f"{stats['accuracy_rate']:.4f}",
            stats["main_loss_reasons"],
            "",
            "",
            template.get("summary_template_id") or "",
            submitted_at,
            submitted_at,
        ),
    )


def latest_chapter_summary_for_plan(conn, user_id: str, plan_date: str) -> dict[str, Any] | None:
    columns = table_columns(conn, "student_chapter_completions")
    if not columns:
        return None
    completion = row_to_dict(
        conn.execute(
            """
            SELECT *
            FROM student_chapter_completions
            WHERE user_id = ? AND started_at = ?
            ORDER BY updated_at DESC, completion_id DESC
            LIMIT 1
            """,
            (user_id, plan_date),
        ).fetchone()
    ) or {}
    if not completion:
        return None
    template = {}
    if completion.get("summary_template_id"):
        template = row_to_dict(
            conn.execute(
                "SELECT * FROM chapter_summary_templates WHERE summary_template_id = ?",
                (completion["summary_template_id"],),
            ).fetchone()
        ) or {}
    if not template:
        template = find_chapter_summary_template(
            conn,
            str(completion.get("math_type") or ""),
            str(completion.get("raw_chapter_id") or ""),
            str(completion.get("plan_stage") or ""),
        )
    review_advice = str(template.get("review_suggestions") or completion.get("main_loss_reasons") or "").strip()
    return {
        "chapterId": completion.get("raw_chapter_id") or template.get("chapter_id") or "",
        "chapterName": completion.get("chapter_name") or "",
        "summaryTitle": template.get("summary_title") or "",
        "summaryText": template.get("summary_text") or "",
        "examFocus": template.get("exam_focus") or "",
        "keyPoints": template.get("key_points") or "",
        "commonMistakes": template.get("common_mistakes") or "",
        "reviewAdvice": review_advice,
        "accuracyRate": pct(completion.get("accuracy_rate")),
        "submittedQuestionCount": as_int(completion.get("submitted_question_count")),
        "wrongQuestionCount": as_int(completion.get("wrong_question_count")),
        "mainLossReasons": completion.get("main_loss_reasons") or "",
    }


def sync_feedback_growth_stats(conn, user_id: str, plan_date: str, type_id: str, updated_at: str) -> None:
    ensure_growth_snapshot_tables(conn)
    feedbacks = rows_to_dicts(
        conn.execute(
            """
            SELECT f.*, dp.type_name, dp.type_index
            FROM question_feedbacks f
            JOIN daily_plans dp
              ON dp.user_id = f.user_id
             AND dp.plan_date = f.plan_date
             AND dp.type_id = f.type_id
             AND dp.question_id = f.question_id
            WHERE f.user_id = ?
              AND f.plan_date = ?
              AND f.type_id = ?
            ORDER BY CAST(dp.question_order AS INTEGER), f.question_id
            """,
            (user_id, plan_date, type_id),
        ).fetchall()
    )
    if feedbacks:
        submitted_count = len(feedbacks)
        wrong_count = sum(1 for row in feedbacks if str(row.get("result") or "") in {"wrong", "unattempted"})
        mastery_rate = (submitted_count - wrong_count) / submitted_count if submitted_count else 0
        status_tag = "良好" if mastery_rate >= 0.75 else "待提升" if mastery_rate >= 0.6 else "薄弱项"
        conn.execute(
            "DELETE FROM type_mastery_stats WHERE user_id = ? AND stat_date = ? AND type_id = ?",
            (user_id, plan_date, type_id),
        )
        conn.execute(
            """
            INSERT INTO type_mastery_stats (
              mastery_id, user_id, stat_date, type_id, type_name, mastery_rate,
              status_tag, submitted_question_count, wrong_question_count, sort_order, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f"M-{user_id}-{plan_date}-{type_id}",
                user_id,
                plan_date,
                type_id,
                feedbacks[0].get("type_name") or "",
                f"{mastery_rate:.4f}",
                status_tag,
                str(submitted_count),
                str(wrong_count),
                str(as_int(feedbacks[0].get("type_index"))),
                updated_at,
            ),
        )

    reason_rows = rows_to_dicts(
        conn.execute(
            """
            SELECT wrong_reason_code, wrong_reason_name, COUNT(*) AS wrong_count
            FROM question_feedbacks
            WHERE user_id = ?
              AND plan_date = ?
              AND result = 'wrong'
              AND COALESCE(wrong_reason_code, '') <> ''
            GROUP BY wrong_reason_code, wrong_reason_name
            ORDER BY COUNT(*) DESC, wrong_reason_code
            """,
            (user_id, plan_date),
        ).fetchall()
    )
    total_reason_count = sum(as_int(row.get("wrong_count")) for row in reason_rows)
    conn.execute(
        "DELETE FROM loss_reason_stats WHERE user_id = ? AND stat_date = ?",
        (user_id, plan_date),
    )
    for index, row in enumerate(reason_rows, start=1):
        wrong_count = as_int(row.get("wrong_count"))
        reason_rate = wrong_count / total_reason_count if total_reason_count else 0
        reason_code = str(row.get("wrong_reason_code") or "").strip()
        conn.execute(
            """
            INSERT INTO loss_reason_stats (
              loss_reason_id, user_id, stat_date, reason_code, reason_name,
              reason_rate, wrong_question_count, sort_order, updated_at, trend_delta
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f"LR-{user_id}-{plan_date}-{reason_code}",
                user_id,
                plan_date,
                reason_code,
                row.get("wrong_reason_name") or WRONG_REASON_LABELS.get(reason_code, reason_code),
                f"{reason_rate:.4f}",
                str(wrong_count),
                str(index),
                updated_at,
                f"{reason_rate:.4f}",
            ),
        )


DEFAULT_TEST_ACCESS_CODES = [
    ("demo-a01", "10001", "claimed"),
    ("demo-b01", "10002", "claimed"),
    ("new-a01", "", "unused"),
    ("new-b01", "", "unused"),
]
TEST_RUNTIME_CURRENT_DATE_KEY = "current_test_date"

ANNOUNCEMENT_AUDIENCE_ACTIVATED = "activated_students"
ANNOUNCEMENT_STATUS_PUBLISHED = "published"
ANNOUNCEMENT_STATUS_HIDDEN = "hidden"
ANNOUNCEMENT_MAX_IMAGE_BYTES = 1 * 1024 * 1024
ANNOUNCEMENT_ALLOWED_IMAGE_TYPES = {
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/webp": ".webp",
}

WRONG_REASON_LABELS = {
    "method_error": "方法选择错误",
    "calculation_error": "计算失误",
    "knowledge_gap": "知识点不熟/概念混淆",
    "condition_missed": "条件或边界漏考虑",
}


def normalize_test_code(code: str) -> str:
    return str(code or "").strip().lower()


def ensure_question_feedbacks_reason_columns(conn) -> None:
    columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(question_feedbacks)").fetchall()
    }
    if "wrong_reason_code" not in columns:
        conn.execute("ALTER TABLE question_feedbacks ADD COLUMN wrong_reason_code TEXT NOT NULL DEFAULT ''")
    if "wrong_reason_name" not in columns:
        conn.execute("ALTER TABLE question_feedbacks ADD COLUMN wrong_reason_name TEXT NOT NULL DEFAULT ''")
    if "unattempted_action" not in columns:
        conn.execute("ALTER TABLE question_feedbacks ADD COLUMN unattempted_action TEXT NOT NULL DEFAULT ''")


def ensure_student_progress_tracking_columns(conn) -> None:
    columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(student_progress)").fetchall()
    }
    field_definitions = {
        "current_type_status": "TEXT NOT NULL DEFAULT 'active'",
        "carryover_type_id": "TEXT NOT NULL DEFAULT ''",
        "carryover_question_ids": "TEXT NOT NULL DEFAULT '[]'",
        "next_plan_start_chapter_id": "TEXT NOT NULL DEFAULT ''",
        "next_plan_start_effective_date": "TEXT NOT NULL DEFAULT ''",
        "next_plan_start_status": "TEXT NOT NULL DEFAULT ''",
        "next_plan_start_updated_by": "TEXT NOT NULL DEFAULT ''",
        "next_plan_start_updated_at": "TEXT NOT NULL DEFAULT ''",
        "next_plan_start_applied_at": "TEXT NOT NULL DEFAULT ''",
        "mastery_scope_start_type_id": "TEXT NOT NULL DEFAULT ''",
        "mastery_scope_started_at": "TEXT NOT NULL DEFAULT ''",
        "mastery_scope_reason": "TEXT NOT NULL DEFAULT ''",
    }
    for field, definition in field_definitions.items():
        if field not in columns:
            conn.execute(f"ALTER TABLE student_progress ADD COLUMN {field} {definition}")
            columns.add(field)


def ensure_student_profile_target_school_columns(conn) -> None:
    columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(student_profiles)").fetchall()
    }
    field_definitions = {
        "target_school_id": "TEXT NOT NULL DEFAULT ''",
        "target_school_name": "TEXT NOT NULL DEFAULT ''",
        "initial_start_chapter_id": "TEXT NOT NULL DEFAULT ''",
    }
    for field, definition in field_definitions.items():
        if field not in columns:
            conn.execute(f"ALTER TABLE student_profiles ADD COLUMN {field} {definition}")
            columns.add(field)


def ensure_student_stats_ranking_columns(conn) -> None:
    columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(student_stats)").fetchall()
    }
    if "today_checkin_order" not in columns:
        conn.execute("ALTER TABLE student_stats ADD COLUMN today_checkin_order TEXT NOT NULL DEFAULT '0'")
    if "today_checkin_at" not in columns:
        conn.execute("ALTER TABLE student_stats ADD COLUMN today_checkin_at TEXT NOT NULL DEFAULT ''")
    if "weekly_study_days" not in columns:
        conn.execute("ALTER TABLE student_stats ADD COLUMN weekly_study_days TEXT NOT NULL DEFAULT '0'")


def ensure_growth_snapshot_tables(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS student_weekly_targets (
          weekly_target_id TEXT PRIMARY KEY,
          user_id TEXT NOT NULL,
          week_start_date TEXT NOT NULL,
          week_end_date TEXT NOT NULL,
          basis_date TEXT NOT NULL,
          target_days TEXT NOT NULL DEFAULT '0',
          target_question_count TEXT NOT NULL DEFAULT '0',
          target_type_count TEXT NOT NULL DEFAULT '0',
          basis_type_ids TEXT NOT NULL DEFAULT '[]',
          basis_note TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS type_mastery_stats (
          mastery_id TEXT PRIMARY KEY,
          user_id TEXT NOT NULL,
          stat_date TEXT NOT NULL,
          type_id TEXT NOT NULL,
          type_name TEXT NOT NULL DEFAULT '',
          mastery_rate TEXT NOT NULL DEFAULT '0',
          status_tag TEXT NOT NULL DEFAULT '',
          submitted_question_count TEXT NOT NULL DEFAULT '0',
          wrong_question_count TEXT NOT NULL DEFAULT '0',
          sort_order TEXT NOT NULL DEFAULT '0',
          updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS loss_reason_stats (
          loss_reason_id TEXT PRIMARY KEY,
          user_id TEXT NOT NULL,
          stat_date TEXT NOT NULL,
          reason_code TEXT NOT NULL,
          reason_name TEXT NOT NULL DEFAULT '',
          reason_rate TEXT NOT NULL DEFAULT '0',
          wrong_question_count TEXT NOT NULL DEFAULT '0',
          sort_order TEXT NOT NULL DEFAULT '0',
          updated_at TEXT NOT NULL,
          trend_delta TEXT NOT NULL DEFAULT '0'
        )
        """
    )
    loss_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(loss_reason_stats)").fetchall()
    }
    if "trend_delta" not in loss_columns:
        conn.execute("ALTER TABLE loss_reason_stats ADD COLUMN trend_delta TEXT NOT NULL DEFAULT '0'")
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS student_chapter_completions (
          completion_id TEXT PRIMARY KEY,
          user_id TEXT NOT NULL,
          raw_chapter_id TEXT NOT NULL,
          chapter_name TEXT NOT NULL DEFAULT '',
          math_type TEXT NOT NULL DEFAULT '',
          plan_stage TEXT NOT NULL DEFAULT '',
          started_at TEXT NOT NULL DEFAULT '',
          completed_at TEXT NOT NULL,
          spent_days TEXT NOT NULL DEFAULT '0',
          completed_type_count TEXT NOT NULL DEFAULT '0',
          total_type_count TEXT NOT NULL DEFAULT '0',
          submitted_question_count TEXT NOT NULL DEFAULT '0',
          wrong_question_count TEXT NOT NULL DEFAULT '0',
          accuracy_rate TEXT NOT NULL DEFAULT '0',
          main_loss_reasons TEXT NOT NULL DEFAULT '',
          review_type_ids TEXT NOT NULL DEFAULT '',
          review_type_names TEXT NOT NULL DEFAULT '',
          summary_template_id TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS chapter_summary_templates (
          summary_template_id TEXT PRIMARY KEY,
          raw_chapter_id TEXT NOT NULL,
          chapter_name TEXT NOT NULL DEFAULT '',
          math_type TEXT NOT NULL DEFAULT '',
          plan_stage TEXT NOT NULL DEFAULT '',
          summary_title TEXT NOT NULL DEFAULT '',
          summary_text TEXT NOT NULL DEFAULT '',
          exam_focus TEXT NOT NULL DEFAULT '',
          key_points TEXT NOT NULL DEFAULT '',
          common_mistakes TEXT NOT NULL DEFAULT '',
          review_suggestions TEXT NOT NULL DEFAULT '',
          next_step_text TEXT NOT NULL DEFAULT '',
          is_active TEXT NOT NULL DEFAULT '1',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_weekly_targets_user_week ON student_weekly_targets(user_id, week_start_date)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_mastery_user_date ON type_mastery_stats(user_id, stat_date)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_losses_user_date ON loss_reason_stats(user_id, stat_date)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_feedbacks_user_date_type ON question_feedbacks(user_id, plan_date, type_id)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_feedbacks_user_reason ON question_feedbacks(user_id, wrong_reason_code)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_daily_plans_user_type ON daily_plans(user_id, type_id)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_chapter_completions_user_chapter ON student_chapter_completions(user_id, raw_chapter_id, plan_stage)"
    )
    summary_columns = table_columns(conn, "chapter_summary_templates")
    if "exam_focus" not in summary_columns:
        conn.execute("ALTER TABLE chapter_summary_templates ADD COLUMN exam_focus TEXT NOT NULL DEFAULT ''")
        summary_columns = table_columns(conn, "chapter_summary_templates")
    summary_chapter_column = "chapter_id" if "chapter_id" in summary_columns else "raw_chapter_id"
    conn.execute(
        f"CREATE INDEX IF NOT EXISTS idx_chapter_templates_lookup ON chapter_summary_templates(math_type, {summary_chapter_column}, plan_stage)"
    )


def ensure_announcement_tables(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS announcements (
          announcement_id TEXT PRIMARY KEY,
          title TEXT NOT NULL,
          content TEXT NOT NULL DEFAULT '',
          audience TEXT NOT NULL DEFAULT 'activated_students',
          status TEXT NOT NULL DEFAULT 'published',
          image_url TEXT NOT NULL DEFAULT '',
          image_filename TEXT NOT NULL DEFAULT '',
          image_original_name TEXT NOT NULL DEFAULT '',
          image_content_type TEXT NOT NULL DEFAULT '',
          image_size_bytes TEXT NOT NULL DEFAULT '0',
          created_by TEXT NOT NULL DEFAULT 'adm001',
          published_at TEXT NOT NULL,
          hidden_at TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS announcement_reads (
          announcement_id TEXT NOT NULL,
          user_id TEXT NOT NULL,
          read_at TEXT NOT NULL,
          PRIMARY KEY (announcement_id, user_id)
        )
        """
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_announcements_status_published ON announcements(status, published_at)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_announcement_reads_user ON announcement_reads(user_id, announcement_id)"
    )


def ensure_formal_activation_tables(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS formal_account_previews (
          preview_id TEXT PRIMARY KEY,
          login_account TEXT NOT NULL,
          initial_password_hash TEXT NOT NULL,
          initial_password_plain TEXT NOT NULL,
          status TEXT NOT NULL,
          admin_user_id TEXT NOT NULL,
          expires_at TEXT NOT NULL,
          activated_user_id TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS admin_totp_state (
          admin_user_id TEXT PRIMARY KEY,
          formal_activations_since_totp TEXT NOT NULL DEFAULT '0',
          totp_secret TEXT NOT NULL DEFAULT '',
          totp_bound_at TEXT NOT NULL DEFAULT '',
          setup_token_hash TEXT NOT NULL DEFAULT '',
          setup_token_expires_at TEXT NOT NULL DEFAULT '',
          last_verified_at TEXT NOT NULL DEFAULT '',
          updated_at TEXT NOT NULL
        )
        """
    )
    totp_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(admin_totp_state)").fetchall()
    }
    if "totp_bound_at" not in totp_columns:
        conn.execute("ALTER TABLE admin_totp_state ADD COLUMN totp_bound_at TEXT NOT NULL DEFAULT ''")
    if "setup_token_hash" not in totp_columns:
        conn.execute("ALTER TABLE admin_totp_state ADD COLUMN setup_token_hash TEXT NOT NULL DEFAULT ''")
    if "setup_token_expires_at" not in totp_columns:
        conn.execute("ALTER TABLE admin_totp_state ADD COLUMN setup_token_expires_at TEXT NOT NULL DEFAULT ''")
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_formal_preview_account ON formal_account_previews(login_account)"
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_formal_preview_status ON formal_account_previews(status, expires_at)"
    )
    user_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(users)").fetchall()
    }
    if "gender" not in user_columns:
        conn.execute("ALTER TABLE users ADD COLUMN gender TEXT NOT NULL DEFAULT ''")
    if "password_changed_at" not in user_columns:
        conn.execute("ALTER TABLE users ADD COLUMN password_changed_at TEXT NOT NULL DEFAULT ''")
    if "account_type" not in user_columns:
        conn.execute("ALTER TABLE users ADD COLUMN account_type TEXT NOT NULL DEFAULT ''")
    if "activated_at" not in user_columns:
        conn.execute("ALTER TABLE users ADD COLUMN activated_at TEXT NOT NULL DEFAULT ''")
    profile_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(student_profiles)").fetchall()
    }
    if "intensive_days" not in profile_columns:
        conn.execute("ALTER TABLE student_profiles ADD COLUMN intensive_days TEXT NOT NULL DEFAULT '60'")
    if "review_order" not in profile_columns:
        conn.execute("ALTER TABLE student_profiles ADD COLUMN review_order TEXT NOT NULL DEFAULT 'mixed'")
    if "plan_variant_id" not in profile_columns:
        conn.execute("ALTER TABLE student_profiles ADD COLUMN plan_variant_id TEXT NOT NULL DEFAULT ''")
    if "diagnostic_json" not in profile_columns:
        conn.execute("ALTER TABLE student_profiles ADD COLUMN diagnostic_json TEXT NOT NULL DEFAULT '{}'")
    if "primary_book_1_scope" not in profile_columns:
        conn.execute("ALTER TABLE student_profiles ADD COLUMN primary_book_1_scope TEXT NOT NULL DEFAULT 'all'")
    if "primary_book_2_scope" not in profile_columns:
        conn.execute("ALTER TABLE student_profiles ADD COLUMN primary_book_2_scope TEXT NOT NULL DEFAULT 'all'")


def configured_default_admin_password() -> tuple[str, bool]:
    password = (
        os.environ.get("ADMIN_INITIAL_PASSWORD")
        or os.environ.get("INITIAL_ADMIN_PASSWORD")
        or ""
    ).strip()
    if password:
        return password, True
    if APP_ENV == "production":
        raise RuntimeError("ADMIN_INITIAL_PASSWORD is required when APP_ENV=production")
    return DEFAULT_LOCAL_ADMIN_CREDENTIAL, False


def ensure_default_backoffice_admin(conn) -> None:
    ensure_formal_activation_tables(conn)
    password, password_from_env = configured_default_admin_password()
    password_values: dict[str, Any] = {}
    set_user_password_values(password_values, password)
    created_at = now_text()
    existing = row_to_dict(
        conn.execute(
            """
            SELECT * FROM users
            WHERE user_id = ? OR login_account = ?
            ORDER BY CASE WHEN user_id = ? THEN 0 ELSE 1 END
            LIMIT 1
            """,
            (DEFAULT_ADMIN_USER_ID, DEFAULT_ADMIN_LOGIN_ACCOUNT, DEFAULT_ADMIN_USER_ID),
        ).fetchone()
    )
    if not existing:
        conn.execute(
            """
            INSERT INTO users (
              user_id, login_account, password_hash, login_password, role, nickname,
              avatar_url, QQ, account_status, current_session_id, last_login_at,
              created_at, updated_at, gender, password_changed_at, account_type, activated_at
            ) VALUES (?, ?, ?, ?, 'admin', ?, '', '', 'active', '', '', ?, ?, '', '', 'admin', ?)
            """,
            (
                DEFAULT_ADMIN_USER_ID,
                DEFAULT_ADMIN_LOGIN_ACCOUNT,
                password_values["password_hash"],
                password_values["login_password"],
                DEFAULT_ADMIN_NICKNAME,
                created_at,
                created_at,
                created_at,
            ),
        )
    elif (
        existing.get("user_id") == DEFAULT_ADMIN_USER_ID
        and existing.get("role") == "admin"
        and password_from_env
        and password_matches(existing, DEFAULT_LOCAL_ADMIN_CREDENTIAL)
    ):
        conn.execute(
            """
            UPDATE users
            SET password_hash = ?, login_password = ?, updated_at = ?
            WHERE user_id = ?
            """,
            (
                password_values["password_hash"],
                password_values["login_password"],
                created_at,
                DEFAULT_ADMIN_USER_ID,
            ),
        )
    conn.execute(
        """
        INSERT OR IGNORE INTO admin_totp_state (
          admin_user_id, formal_activations_since_totp, totp_secret,
          totp_bound_at, setup_token_hash, setup_token_expires_at,
          last_verified_at, updated_at
        ) VALUES (?, '0', '', '', '', '', '', ?)
        """,
        (DEFAULT_ADMIN_USER_ID, created_at),
    )


def ensure_admin_console_tables(conn) -> None:
    class_columns = {
        row["name"]
        for row in conn.execute("PRAGMA table_info(classes)").fetchall()
    }
    if "capacity" not in class_columns:
        conn.execute("ALTER TABLE classes ADD COLUMN capacity TEXT NOT NULL DEFAULT '60'")
    if "is_accepting_new_students" not in class_columns:
        conn.execute("ALTER TABLE classes ADD COLUMN is_accepting_new_students TEXT NOT NULL DEFAULT '1'")
    if "qq_group_number" not in class_columns:
        conn.execute("ALTER TABLE classes ADD COLUMN qq_group_number TEXT NOT NULL DEFAULT ''")
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS admin_audit_logs (
          log_id TEXT PRIMARY KEY,
          actor TEXT NOT NULL DEFAULT 'adm001',
          action TEXT NOT NULL,
          target TEXT NOT NULL DEFAULT '',
          before_value TEXT NOT NULL DEFAULT '',
          after_value TEXT NOT NULL DEFAULT '',
          confirmed TEXT NOT NULL DEFAULT '1',
          created_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS class_intake_settings (
          math_type TEXT NOT NULL,
          level TEXT NOT NULL,
          active_class_id TEXT NOT NULL DEFAULT '',
          status TEXT NOT NULL DEFAULT '待建班',
          updated_by TEXT NOT NULL DEFAULT 'system',
          updated_at TEXT NOT NULL,
          PRIMARY KEY (math_type, level)
        )
        """
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_admin_audit_logs_created ON admin_audit_logs(created_at)"
    )
    seed_class_intake_settings(conn)


def ensure_test_access_codes_table(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS test_access_codes (
          code TEXT PRIMARY KEY,
          user_id TEXT NOT NULL DEFAULT '',
          status TEXT NOT NULL DEFAULT 'unused',
          created_at TEXT NOT NULL,
          claimed_at TEXT NOT NULL DEFAULT '',
          updated_at TEXT NOT NULL
        )
        """
    )
    created_at = now_text()
    for code, user_id, status in DEFAULT_TEST_ACCESS_CODES:
        conn.execute(
            """
            INSERT OR IGNORE INTO test_access_codes (
              code, user_id, status, created_at, claimed_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (code, user_id, status, created_at, created_at if user_id else "", created_at),
        )


def ensure_test_user_dates_table(conn) -> None:
    ensure_test_runtime_state_table(conn)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS test_user_dates (
          user_id TEXT PRIMARY KEY,
          current_test_date TEXT NOT NULL,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )


def ensure_test_runtime_state_table(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS test_runtime_state (
          state_key TEXT PRIMARY KEY,
          state_value TEXT NOT NULL,
          updated_at TEXT NOT NULL
        )
        """
    )
    conn.execute(
        """
        INSERT OR IGNORE INTO test_runtime_state (state_key, state_value, updated_at)
        VALUES ('current_test_date', ?, ?)
        """,
        (CURRENT_DATE, now_text()),
    )


def run_startup_migrations() -> None:
    with get_connection() as conn:
        ensure_question_type_card_label_column(conn)
        ensure_question_feedbacks_reason_columns(conn)
        ensure_daily_plan_sequence_columns(conn)
        ensure_student_progress_tracking_columns(conn)
        ensure_student_profile_target_school_columns(conn)
        ensure_student_stats_ranking_columns(conn)
        ensure_growth_snapshot_tables(conn)
        ensure_announcement_tables(conn)
        ensure_formal_activation_tables(conn)
        ensure_default_backoffice_admin(conn)
        ensure_admin_console_tables(conn)
        ensure_study_room_tables(conn)
        canonicalize_student_profile_book_ids(conn)
        ensure_test_access_codes_table(conn)
        ensure_test_user_dates_table(conn)
        conn.commit()


def get_current_test_date(conn, user_id: str) -> str:
    if is_production_runtime():
        return current_date_text()
    ensure_test_user_dates_table(conn)
    runtime_date = get_runtime_test_date(conn)
    created_at = now_text()
    conn.execute(
        """
        INSERT INTO test_user_dates (user_id, current_test_date, created_at, updated_at)
        VALUES (?, ?, ?, ?)
        ON CONFLICT(user_id) DO UPDATE SET
          current_test_date = excluded.current_test_date,
          updated_at = excluded.updated_at
        """,
        (user_id, runtime_date, created_at, created_at),
    )
    return runtime_date


def get_runtime_test_date(conn) -> str:
    if is_production_runtime():
        return current_date_text()
    ensure_test_runtime_state_table(conn)
    row = row_to_dict(
        conn.execute(
            "SELECT state_value FROM test_runtime_state WHERE state_key = 'current_test_date'"
        ).fetchone()
    )
    return row.get("state_value") if row else CURRENT_DATE


def test_clock_payload(current_date: str) -> dict[str, Any]:
    if is_production_runtime():
        current_date = current_date_text()
        return {
            "currentDate": current_date,
            "offsetDays": 0,
        }
    try:
        offset_days = (date.fromisoformat(current_date) - date.fromisoformat(CURRENT_DATE)).days
    except ValueError:
        current_date = CURRENT_DATE
        offset_days = 0
    return {
        "currentDate": current_date,
        "offsetDays": offset_days,
    }


def get_test_clock() -> dict[str, Any]:
    with get_connection() as conn:
        return test_clock_payload(get_runtime_test_date(conn))


def advance_shared_test_clock(conn) -> str:
    if is_production_runtime():
        raise ValueError("测试日期推进接口在正式环境已禁用")
    ensure_test_user_dates_table(conn)
    current_test_date = get_runtime_test_date(conn)
    try:
        next_test_date = (date.fromisoformat(current_test_date) + timedelta(days=1)).isoformat()
    except ValueError as exc:
        raise ValueError("测试日期格式应为 YYYY-MM-DD") from exc

    updated_at = now_text()
    conn.execute(
        """
        UPDATE test_runtime_state
        SET state_value = ?, updated_at = ?
        WHERE state_key = 'current_test_date'
        """,
        (next_test_date, updated_at),
    )
    conn.execute(
        """
        INSERT OR IGNORE INTO test_user_dates (user_id, current_test_date, created_at, updated_at)
        SELECT user_id, ?, ?, ?
        FROM users
        WHERE role = 'student'
        """,
        (next_test_date, updated_at, updated_at),
    )
    conn.execute(
        """
        UPDATE test_user_dates
        SET current_test_date = ?, updated_at = ?
        WHERE user_id IN (SELECT user_id FROM users WHERE role = 'student')
        """,
        (next_test_date, updated_at),
    )
    return next_test_date


def advance_test_clock() -> dict[str, Any]:
    with get_connection() as conn:
        next_test_date = advance_shared_test_clock(conn)
        conn.commit()
    return {
        "success": True,
        **test_clock_payload(next_test_date),
    }


def max_numeric_user_id(conn) -> int:
    rows = rows_to_dicts(conn.execute("SELECT user_id FROM users").fetchall())
    max_id = 0
    for row in rows:
        text = str(row.get("user_id") or "").strip()
        if text.isdigit():
            max_id = max(max_id, int(text))
    return max_id


def next_numeric_user_id(conn) -> str:
    return str(max_numeric_user_id(conn) + 1)


def create_test_user(conn, code: str) -> dict[str, Any]:
    user_id = next_numeric_user_id(conn)
    created_at = now_text()
    nickname = f"Test {user_id}"
    conn.execute(
        """
        INSERT INTO users (
          user_id, login_account, password_hash, login_password, role, nickname,
          avatar_url, QQ, account_status, current_session_id, last_login_at,
          created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            user_id,
            code,
            "",
            "123456",
            "student",
            nickname,
            "",
            "",
            "active",
            "",
            "",
            created_at,
            created_at,
        ),
    )
    return row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())


def formal_account_from_number(number: int) -> str:
    if number < FORMAL_ACCOUNT_START or number > FORMAL_ACCOUNT_END:
        raise ValueError("formal account range exhausted")
    return f"{FORMAL_ACCOUNT_PREFIX}{number:05d}"


def formal_account_number(login_account: str) -> int:
    text = str(login_account or "")
    if not text.startswith(FORMAL_ACCOUNT_PREFIX):
        return 0
    try:
        return int(text[len(FORMAL_ACCOUNT_PREFIX):])
    except ValueError:
        return 0


def next_formal_account(conn) -> str:
    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT login_account FROM users
            WHERE role = 'student'
              AND login_account LIKE 'stu%'
            """
        ).fetchall()
    )
    used_numbers = {
        formal_account_number(row["login_account"])
        for row in rows
        if FORMAL_ACCOUNT_START <= formal_account_number(row["login_account"]) <= FORMAL_ACCOUNT_END
    }
    current = FORMAL_ACCOUNT_START
    while current in used_numbers:
        current += 1
    return formal_account_from_number(current)


def trial_conversion_formal_account_from_number(number: int) -> str:
    if number < TRIAL_CONVERSION_FORMAL_ACCOUNT_START or number > TRIAL_CONVERSION_FORMAL_ACCOUNT_END:
        raise ValueError("trial conversion formal account range exhausted")
    return f"{FORMAL_ACCOUNT_PREFIX}{number:05d}"


def next_trial_conversion_formal_account(conn) -> str:
    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT login_account FROM users
            WHERE role = 'student'
              AND login_account LIKE 'stu%'
            """
        ).fetchall()
    )
    max_used = TRIAL_CONVERSION_FORMAL_ACCOUNT_START - 1
    for row in rows:
        number = formal_account_number(row["login_account"])
        if TRIAL_CONVERSION_FORMAL_ACCOUNT_START <= number <= TRIAL_CONVERSION_FORMAL_ACCOUNT_END:
            max_used = max(max_used, number)
    current = max_used + 1
    return trial_conversion_formal_account_from_number(current)


def short_term_account_from_number(number: int) -> str:
    if number < SHORT_TERM_ACCOUNT_START or number > SHORT_TERM_ACCOUNT_END:
        raise ValueError("short-term account range exhausted")
    return f"{FORMAL_ACCOUNT_PREFIX}{number:05d}"


def next_short_term_account(conn) -> str:
    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT login_account FROM users
            WHERE role = 'student'
              AND login_account LIKE 'stu%'
            """
        ).fetchall()
    )
    used_numbers = {
        formal_account_number(row["login_account"])
        for row in rows
        if SHORT_TERM_ACCOUNT_START <= formal_account_number(row["login_account"]) <= SHORT_TERM_ACCOUNT_END
    }
    current = SHORT_TERM_ACCOUNT_START
    while current in used_numbers:
        current += 1
    return short_term_account_from_number(current)


def generate_initial_password(login_account: str) -> str:
    number_part = f"{formal_account_number(login_account):05d}"
    suffix = "".join(secrets.choice(INITIAL_CREDENTIAL_ALPHABET) for _ in range(2))
    return f"S{number_part}{suffix}"


def preview_expiry_text(created_at: str | None = None) -> str:
    base = datetime.strptime(created_at or now_text(), "%Y-%m-%d %H:%M:%S")
    return (base + timedelta(minutes=FORMAL_PREVIEW_TTL_MINUTES)).strftime("%Y-%m-%d %H:%M:%S")


def serialize_formal_preview(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "previewId": row["preview_id"],
        "loginAccount": row["login_account"],
        "initialPassword": row["initial_password_plain"],
        "expiresAt": row["expires_at"],
        "status": row["status"],
    }


def serialize_pending_formal_onboarding(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "previewId": row.get("preview_id") or "",
        "loginAccount": row.get("login_account") or "",
        "userId": row.get("activated_user_id") or "",
        "nickname": row.get("user_nickname") or row.get("login_account") or "",
        "initialPassword": row.get("initial_password_plain") or "",
        "activatedAt": row.get("activated_at") or row.get("updated_at") or "",
        "accountStatus": row.get("account_status") or "",
        "status": "pending_onboarding",
        "needsOnboarding": True,
    }


def get_or_create_formal_account_preview(admin_user_id: str = "adm001") -> dict[str, Any]:
    admin_user_id = str(admin_user_id or "adm001").strip() or "adm001"
    with get_connection() as conn:
        ensure_formal_activation_tables(conn)
        now = now_text()
        existing = row_to_dict(
            conn.execute(
                """
                SELECT * FROM formal_account_previews
                WHERE admin_user_id = ?
                  AND status = 'previewed'
                  AND expires_at > ?
                ORDER BY created_at DESC
                LIMIT 1
                """,
                (admin_user_id, now),
            ).fetchone()
        )
        if existing:
            return {"credential": serialize_formal_preview(existing)}

        login_account = next_formal_account(conn)
        password = generate_initial_password(login_account)
        preview_id = f"FAP-{uuid.uuid4().hex[:12].upper()}"
        expires_at = preview_expiry_text(now)
        conn.execute(
            """
            INSERT INTO formal_account_previews (
              preview_id, login_account, initial_password_hash,
              initial_password_plain, status, admin_user_id,
              expires_at, activated_user_id, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                preview_id,
                login_account,
                sha256_password(password),
                password,
                "previewed",
                admin_user_id,
                expires_at,
                "",
                now,
                now,
            ),
        )
        conn.commit()
        row = row_to_dict(
            conn.execute("SELECT * FROM formal_account_previews WHERE preview_id = ?", (preview_id,)).fetchone()
        )
        return {"credential": serialize_formal_preview(row)}


def refresh_formal_account_preview(admin_user_id: str, preview_id: str) -> dict[str, Any]:
    admin_user_id = str(admin_user_id or "adm001").strip() or "adm001"
    preview_id = str(preview_id or "").strip()
    with get_connection() as conn:
        ensure_formal_activation_tables(conn)
        now = now_text()
        row = row_to_dict(
            conn.execute(
                """
                SELECT * FROM formal_account_previews
                WHERE preview_id = ? AND admin_user_id = ? AND status = 'previewed'
                """,
                (preview_id, admin_user_id),
            ).fetchone()
        )
        if not row:
            raise ValueError("preview credential not found")
        password = generate_initial_password(row["login_account"])
        while password == row["initial_password_plain"]:
            password = generate_initial_password(row["login_account"])
        expires_at = preview_expiry_text(now)
        conn.execute(
            """
            UPDATE formal_account_previews
            SET initial_password_hash = ?, initial_password_plain = ?,
                expires_at = ?, updated_at = ?
            WHERE preview_id = ?
            """,
            (sha256_password(password), password, expires_at, now, preview_id),
        )
        conn.commit()
        updated = row_to_dict(
            conn.execute("SELECT * FROM formal_account_previews WHERE preview_id = ?", (preview_id,)).fetchone()
        )
        return {"credential": serialize_formal_preview(updated)}


def list_pending_formal_onboarding(admin_user_id: str = "adm001") -> dict[str, Any]:
    admin_user_id = str(admin_user_id or "adm001").strip() or "adm001"
    with get_connection() as conn:
        ensure_formal_activation_tables(conn)
        rows = rows_to_dicts(
            conn.execute(
                """
                SELECT
                  f.preview_id, f.login_account, f.initial_password_plain,
                  f.admin_user_id, f.activated_user_id, f.updated_at,
                  u.nickname AS user_nickname, u.account_status, u.activated_at
                FROM formal_account_previews f
                JOIN users u ON u.user_id = f.activated_user_id
                LEFT JOIN student_profiles p ON p.user_id = u.user_id
                WHERE f.admin_user_id = ?
                  AND f.status = 'activated'
                  AND p.user_id IS NULL
                ORDER BY f.updated_at DESC
                """,
                (admin_user_id,),
            ).fetchall()
        )
    return {"items": [serialize_pending_formal_onboarding(row) for row in rows]}


def generate_totp_secret() -> str:
    return base64.b32encode(secrets.token_bytes(20)).decode("ascii").rstrip("=")


def normalize_totp_secret(secret: str) -> bytes:
    clean = str(secret or "").strip().replace(" ", "").upper()
    padding = "=" * ((8 - len(clean) % 8) % 8)
    return base64.b32decode(clean + padding)


def hotp_code(secret: str, counter: int, digits: int = 6) -> str:
    key = normalize_totp_secret(secret)
    message = struct.pack(">Q", counter)
    digest = hmac.new(key, message, hashlib.sha1).digest()
    offset = digest[-1] & 0x0F
    value = struct.unpack(">I", digest[offset : offset + 4])[0] & 0x7FFFFFFF
    return str(value % (10**digits)).zfill(digits)


def totp_code(secret: str, at_time: int | None = None) -> str:
    timestamp = int(time.time() if at_time is None else at_time)
    return hotp_code(secret, timestamp // TOTP_PERIOD_SECONDS)


def verify_totp_code(secret: str, code: str, at_time: int | None = None, window: int = 1) -> bool:
    clean_code = str(code or "").strip()
    if len(clean_code) != 6 or not clean_code.isdigit() or not secret:
        return False
    timestamp = int(time.time() if at_time is None else at_time)
    counter = timestamp // TOTP_PERIOD_SECONDS
    return any(
        hmac.compare_digest(hotp_code(secret, counter + offset), clean_code)
        for offset in range(-window, window + 1)
    )


def verify_admin_totp_code(secret: str, code: str) -> bool:
    return verify_totp_code(secret, code)


def admin_totp_otpauth_url(login_account: str, secret: str) -> str:
    label = quote(f"{TOTP_ISSUER}:{login_account}")
    issuer = quote(TOTP_ISSUER)
    return f"otpauth://totp/{label}?secret={secret}&issuer={issuer}&period={TOTP_PERIOD_SECONDS}&digits=6"


def generate_setup_token() -> str:
    return secrets.token_urlsafe(32)


def add_minutes_from_now_text(minutes: int) -> str:
    return (datetime.now() + timedelta(minutes=minutes)).strftime("%Y-%m-%d %H:%M:%S")


def datetime_text_or_none(value: Any) -> datetime | None:
    try:
        return datetime.strptime(str(value or ""), "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return None


def admin_totp_state(conn, admin_user_id: str) -> dict[str, Any]:
    state = row_to_dict(
        conn.execute(
            "SELECT * FROM admin_totp_state WHERE admin_user_id = ?",
            (admin_user_id,),
        ).fetchone()
    )
    if state:
        return state
    updated_at = now_text()
    conn.execute(
        """
        INSERT INTO admin_totp_state (
          admin_user_id, formal_activations_since_totp, totp_secret,
          totp_bound_at, setup_token_hash, setup_token_expires_at,
          last_verified_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (admin_user_id, "0", "", "", "", "", "", updated_at),
    )
    return {
        "admin_user_id": admin_user_id,
        "formal_activations_since_totp": "0",
        "totp_secret": "",
        "totp_bound_at": "",
        "setup_token_hash": "",
        "setup_token_expires_at": "",
        "last_verified_at": "",
        "updated_at": updated_at,
    }


def risk_authorized_until_text(verified_at: Any) -> str:
    verified_time = datetime_text_or_none(verified_at)
    if not verified_time:
        return ""
    return (verified_time + timedelta(minutes=RISK_AUTHORIZATION_MINUTES)).strftime("%Y-%m-%d %H:%M:%S")


def admin_risk_authorization_active(state: dict[str, Any]) -> bool:
    authorized_until = risk_authorized_until_text(state.get("last_verified_at"))
    return bool(authorized_until and authorized_until > now_text())


def verify_admin_risk_authorization(admin_user_id: str, totp_code: str) -> dict[str, Any]:
    admin_user_id = str(admin_user_id or DEFAULT_ADMIN_USER_ID).strip() or DEFAULT_ADMIN_USER_ID
    with get_connection() as conn:
        ensure_formal_activation_tables(conn)
        admin = row_to_dict(
            conn.execute(
                "SELECT * FROM users WHERE user_id = ? AND role = 'admin'",
                (admin_user_id,),
            ).fetchone()
        )
        if not admin or admin.get("account_status") != "active":
            raise ValueError("admin account not found")
        state = admin_totp_state(conn, admin_user_id)
        if not state.get("totp_bound_at") or not state.get("totp_secret"):
            raise ValueError("admin TOTP not bound")
        if not verify_admin_totp_code(state.get("totp_secret", ""), str(totp_code or "").strip()):
            raise ValueError("TOTP code invalid")
        verified_at = now_text()
        conn.execute(
            """
            UPDATE admin_totp_state
            SET last_verified_at = ?, updated_at = ?
            WHERE admin_user_id = ?
            """,
            (verified_at, verified_at, admin_user_id),
        )
        conn.commit()
    return {
        "ok": True,
        "authorizedUntil": risk_authorized_until_text(verified_at),
        "riskAuthorizationMinutes": RISK_AUTHORIZATION_MINUTES,
    }


def require_recent_admin_risk_authorization(conn, admin_user_id: str) -> None:
    ensure_formal_activation_tables(conn)
    state = admin_totp_state(conn, admin_user_id)
    if not admin_risk_authorization_active(state):
        raise ValueError("TOTP_REQUIRED")


def activate_next_formal_account(admin_user_id: str, preview_id: str, totp_code: str = "") -> dict[str, Any]:
    admin_user_id = str(admin_user_id or "adm001").strip() or "adm001"
    preview_id = str(preview_id or "").strip()
    with get_connection() as conn:
        ensure_formal_activation_tables(conn)
        now = now_text()
        state = admin_totp_state(conn, admin_user_id)
        requires_totp = as_int(state.get("formal_activations_since_totp")) >= FORMAL_OPENINGS_PER_TOTP
        if requires_totp and not verify_admin_totp_code(state.get("totp_secret", ""), totp_code):
            conn.commit()
            return {
                "requiresTotp": True,
                "message": "本轮已开通 30 个正式学员，请完成管理员 TOTP 验证后继续。",
            }

        preview = row_to_dict(
            conn.execute(
                """
                SELECT * FROM formal_account_previews
                WHERE preview_id = ? AND admin_user_id = ? AND status = 'previewed'
                """,
                (preview_id, admin_user_id),
            ).fetchone()
        )
        if not preview:
            raise ValueError("preview credential not found")
        if preview["expires_at"] <= now:
            raise ValueError("preview credential expired")
        login_account = preview["login_account"]
        if login_account != next_formal_account(conn):
            raise ValueError("formal account pointer changed")
        if row_to_dict(conn.execute("SELECT * FROM users WHERE login_account = ?", (login_account,)).fetchone()):
            raise ValueError("formal account already exists")

        user_id = str(formal_account_number(login_account))
        user_values = {
            "user_id": user_id,
            "login_account": login_account,
            "role": "student",
            "nickname": login_account,
            "avatar_url": "",
            "QQ": "",
            "account_status": "active",
            "current_session_id": "",
            "last_login_at": "",
            "created_at": now,
            "updated_at": now,
            "gender": "",
            "password_changed_at": "",
            "account_type": "formal",
            "activated_at": now,
        }
        set_user_password_values(user_values, preview["initial_password_plain"])
        conn.execute(
            """
            INSERT INTO users (
              user_id, login_account, password_hash, login_password, role, nickname,
              avatar_url, QQ, account_status, current_session_id, last_login_at,
              created_at, updated_at, gender, password_changed_at, account_type, activated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                user_values["user_id"],
                user_values["login_account"],
                user_values["password_hash"],
                user_values["login_password"],
                user_values["role"],
                user_values["nickname"],
                user_values["avatar_url"],
                user_values["QQ"],
                user_values["account_status"],
                user_values["current_session_id"],
                user_values["last_login_at"],
                user_values["created_at"],
                user_values["updated_at"],
                user_values["gender"],
                user_values["password_changed_at"],
                user_values["account_type"],
                user_values["activated_at"],
            ),
        )
        conn.execute(
            """
            UPDATE formal_account_previews
            SET status = 'activated', activated_user_id = ?, updated_at = ?
            WHERE preview_id = ?
            """,
            (user_id, now, preview_id),
        )
        if requires_totp:
            next_count = "0"
            last_verified_at = now
        else:
            next_count = str(as_int(state.get("formal_activations_since_totp")) + 1)
            last_verified_at = state.get("last_verified_at") or ""
        conn.execute(
            """
            UPDATE admin_totp_state
            SET formal_activations_since_totp = ?, last_verified_at = ?, updated_at = ?
            WHERE admin_user_id = ?
            """,
            (next_count, last_verified_at, now, admin_user_id),
        )
        conn.commit()
        return {
            "requiresTotp": False,
            "activatedAccount": login_account,
            "activatedUserId": user_id,
            "credential": serialize_formal_preview(preview),
            "nextFormalAccount": next_formal_account(conn),
        }


def activate_next_short_term_account(
    admin_user_id: str = "adm001",
    duration_days: int = 7,
    math_type: str = "数学二",
    level: str = "A",
    target_score: int = 120,
) -> dict[str, Any]:
    admin_user_id = str(admin_user_id or "adm001").strip() or "adm001"
    duration_days = clamp(as_int(duration_days, 7), 1, 30)
    math_type = str(math_type or "数学二").strip() or "数学二"
    if math_type not in {"数学一", "数学二", "数学三"}:
        raise ValueError("mathType invalid")
    level = str(level or "A").strip().upper() or "A"
    if level not in {"A", "B", "C", "D"}:
        raise ValueError("level invalid")
    target_score = clamp(as_int(target_score, 120), 1, 150)

    with get_connection() as conn:
        ensure_formal_activation_tables(conn)
        ensure_admin_console_tables(conn)
        now = now_text()
        login_account = next_short_term_account(conn)
        if row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? OR login_account = ?", (login_account, login_account)).fetchone()):
            raise ValueError("short-term account already exists")

        user_id = str(formal_account_number(login_account))
        password = generate_initial_password(login_account)
        password_values: dict[str, Any] = {}
        set_user_password_values(password_values, password)
        conn.execute(
            """
            INSERT INTO users (
              user_id, login_account, password_hash, login_password, role, nickname,
              avatar_url, QQ, account_status, current_session_id, last_login_at,
              created_at, updated_at, gender, password_changed_at, account_type, activated_at
            ) VALUES (?, ?, ?, ?, 'student', ?, '', '', 'active', '', '', ?, ?, '', '', 'trial', ?)
            """,
            (
                user_id,
                login_account,
                password_values["password_hash"],
                password_values["login_password"],
                login_account,
                now,
                now,
                now,
            ),
        )
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        student = {
            **serialize_user_account(user),
            "trialExpiresAt": "",
            "everActivated": True,
            "trialDurationDays": duration_days,
            "startsAfterOnboarding": True,
        }
        log_admin_action(
            conn,
            "short_term.activate_next",
            user_id,
            "",
            {
                "loginAccount": login_account,
                "paidUntil": "",
                "durationDays": duration_days,
                "startsAfterOnboarding": True,
                "mathType": math_type,
                "level": level,
                "targetScore": target_score,
            },
            admin_user_id,
        )
        next_short_term = next_short_term_account(conn)
        conn.commit()
    return {
        "activatedAccount": login_account,
        "activatedUserId": user_id,
        "initialPassword": password,
        "paidUntil": "",
        "trialDurationDays": duration_days,
        "startsAfterOnboarding": True,
        "student": student,
        "nextShortTermAccount": next_short_term,
    }


def user_needs_onboarding(conn, user_id: str) -> bool:
    return (
        conn.execute("SELECT 1 FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        is None
    )


def start_auth_session(
    conn,
    user: dict[str, Any],
    credential: str,
    device_id: str = "",
    device_name: str = "",
    ip_address: str = "",
    user_agent: str = "",
) -> dict[str, str]:
    session_id = f"S-{user['user_id']}-{uuid.uuid4().hex[:12].upper()}"
    token = hashlib.sha256(f"{session_id}:{credential}:{now_text()}".encode("utf-8")).hexdigest()
    if user.get("current_session_id"):
        conn.execute(
            "UPDATE user_sessions SET revoked_at = ?, updated_at = ? WHERE session_id = ?",
            (now_text(), now_text(), user["current_session_id"]),
        )
    conn.execute(
        """
        INSERT INTO user_sessions (
          session_id, user_id, refresh_token_hash, device_id, device_name,
          ip_address, user_agent, expires_at, revoked_at, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            session_id,
            user["user_id"],
            auth_token_hash(token),
            normalize_device_id(device_id) or "unknown-device",
            str(device_name or "Student Browser")[:120],
            str(ip_address or "")[:80],
            str(user_agent or "")[:240],
            "2026-12-31 23:59:59",
            "",
            now_text(),
            now_text(),
        ),
    )
    conn.execute(
        "UPDATE users SET current_session_id = ?, last_login_at = ?, updated_at = ? WHERE user_id = ?",
        (session_id, now_text(), now_text(), user["user_id"]),
    )
    return {"token": token, "sessionId": session_id}


def auth_payload(
    conn,
    user: dict[str, Any],
    token: str,
    session_id: str,
    include_current_test_date: bool = False,
) -> dict[str, Any]:
    payload = {
        "token": token,
        "sessionId": session_id,
        "needsOnboarding": user.get("role") == "student" and user_needs_onboarding(conn, user["user_id"]),
        "user": {
            "userId": user["user_id"],
            "nickname": user["nickname"],
            "role": user["role"],
        },
    }
    if include_current_test_date or user.get("role") == "student":
        payload["user"]["currentTestDate"] = get_current_test_date(conn, user["user_id"])
    return payload


def validate_student_session(
    conn,
    token: str,
    session_id: str,
    device_id: str = "",
    required_user_id: str = "",
) -> dict[str, Any]:
    token = str(token or "").strip()
    session_id = str(session_id or "").strip()
    if not token or not session_id:
        raise SessionAuthError("SESSION_REQUIRED", "请先登录")

    row = row_to_dict(
        conn.execute(
            """
            SELECT s.*, u.current_session_id, u.role, u.account_status, u.nickname, u.login_account
            FROM user_sessions s
            JOIN users u ON u.user_id = s.user_id
            WHERE s.session_id = ?
            """,
            (session_id,),
        ).fetchone()
    )
    if not row or row.get("refresh_token_hash") != auth_token_hash(token):
        raise SessionAuthError("SESSION_EXPIRED", "登录状态已失效，请重新登录")
    if row.get("revoked_at"):
        raise SessionAuthError("SESSION_REPLACED", "账号已在其他设备登录，请重新登录")
    if row.get("current_session_id") != session_id:
        raise SessionAuthError("SESSION_REPLACED", "账号已在其他设备登录，请重新登录")
    if row.get("expires_at") and str(row.get("expires_at")) <= now_text():
        raise SessionAuthError("SESSION_EXPIRED", "登录状态已失效，请重新登录")
    stored_device_id = normalize_device_id(row.get("device_id"))
    request_device_id = normalize_device_id(device_id)
    if stored_device_id and stored_device_id != "unknown-device" and stored_device_id != request_device_id:
        raise SessionAuthError("SESSION_REPLACED", "账号已在其他设备登录，请重新登录")
    if row.get("role") != "student":
        raise SessionAuthError("SESSION_FORBIDDEN", "当前账号不能访问学生端")
    if row.get("account_status") != "active":
        raise SessionAuthError("SESSION_EXPIRED", "账号已停用，请联系管理员")
    if required_user_id and str(row.get("user_id")) != str(required_user_id):
        raise SessionAuthError("SESSION_FORBIDDEN", "不能访问其他学生数据")
    return row


def validate_admin_session(
    conn,
    token: str,
    session_id: str,
    device_id: str = "",
) -> dict[str, Any]:
    token = str(token or "").strip()
    session_id = str(session_id or "").strip()
    if not token or not session_id:
        raise SessionAuthError("SESSION_REQUIRED", "璇峰厛鐧诲綍")

    row = row_to_dict(
        conn.execute(
            """
            SELECT s.*, u.current_session_id, u.role, u.account_status, u.nickname, u.login_account
            FROM user_sessions s
            JOIN users u ON u.user_id = s.user_id
            WHERE s.session_id = ?
            """,
            (session_id,),
        ).fetchone()
    )
    if not row or row.get("refresh_token_hash") != auth_token_hash(token):
        raise SessionAuthError("SESSION_EXPIRED", "鐧诲綍鐘舵€佸凡澶辨晥锛岃閲嶆柊鐧诲綍")
    if row.get("revoked_at"):
        raise SessionAuthError("SESSION_REPLACED", "璐﹀彿宸插湪鍏朵粬璁惧鐧诲綍锛岃閲嶆柊鐧诲綍")
    if row.get("current_session_id") != session_id:
        raise SessionAuthError("SESSION_REPLACED", "璐﹀彿宸插湪鍏朵粬璁惧鐧诲綍锛岃閲嶆柊鐧诲綍")
    if row.get("expires_at") and str(row.get("expires_at")) <= now_text():
        raise SessionAuthError("SESSION_EXPIRED", "鐧诲綍鐘舵€佸凡澶辨晥锛岃閲嶆柊鐧诲綍")
    stored_device_id = normalize_device_id(row.get("device_id"))
    request_device_id = normalize_device_id(device_id)
    if stored_device_id and stored_device_id != "unknown-device" and stored_device_id != request_device_id:
        raise SessionAuthError("SESSION_REPLACED", "璐﹀彿宸插湪鍏朵粬璁惧鐧诲綍锛岃閲嶆柊鐧诲綍")
    if row.get("role") != "admin":
        raise SessionAuthError("SESSION_FORBIDDEN", "admin session required")
    if row.get("account_status") != "active":
        raise SessionAuthError("SESSION_EXPIRED", "璐﹀彿宸插仠鐢紝璇疯仈绯荤鐞嗗憳")
    return row


def validate_teacher_session(
    conn,
    token: str,
    session_id: str,
    device_id: str = "",
) -> dict[str, Any]:
    token = str(token or "").strip()
    session_id = str(session_id or "").strip()
    if not token or not session_id:
        raise SessionAuthError("SESSION_REQUIRED", "login required")

    row = row_to_dict(
        conn.execute(
            """
            SELECT s.*, u.current_session_id, u.role, u.account_status, u.nickname, u.login_account
            FROM user_sessions s
            JOIN users u ON u.user_id = s.user_id
            WHERE s.session_id = ?
            """,
            (session_id,),
        ).fetchone()
    )
    if not row or row.get("refresh_token_hash") != auth_token_hash(token):
        raise SessionAuthError("SESSION_EXPIRED", "login expired")
    if row.get("revoked_at"):
        raise SessionAuthError("SESSION_REPLACED", "account signed in on another device")
    if row.get("current_session_id") != session_id:
        raise SessionAuthError("SESSION_REPLACED", "account signed in on another device")
    if row.get("expires_at") and str(row.get("expires_at")) <= now_text():
        raise SessionAuthError("SESSION_EXPIRED", "login expired")
    stored_device_id = normalize_device_id(row.get("device_id"))
    request_device_id = normalize_device_id(device_id)
    if stored_device_id and stored_device_id != "unknown-device" and stored_device_id != request_device_id:
        raise SessionAuthError("SESSION_REPLACED", "account signed in on another device")
    if row.get("role") != "teacher":
        raise SessionAuthError("SESSION_FORBIDDEN", "teacher session required")
    if row.get("account_status") != "active":
        raise SessionAuthError("SESSION_EXPIRED", "account disabled")
    return row


def auth_me(token: str, session_id: str, device_id: str = "") -> dict[str, Any]:
    with get_connection() as conn:
        session = validate_student_session(conn, token, session_id, device_id)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (session["user_id"],)).fetchone())
        payload = auth_payload(conn, user, token, session_id, include_current_test_date=True)
        payload["sessionId"] = session_id
        return payload


def login_with_test_code(code: str) -> dict[str, Any]:
    normalized_code = normalize_test_code(code)
    if not normalized_code:
        raise ValueError("test code is required")

    with get_connection() as conn:
        ensure_test_access_codes_table(conn)
        access_code = row_to_dict(
            conn.execute(
                "SELECT * FROM test_access_codes WHERE code = ?",
                (normalized_code,),
            ).fetchone()
        )
        if not access_code or access_code["status"] == "disabled":
            raise ValueError("invalid test code")

        user_id = access_code.get("user_id") or ""
        user = None
        if user_id:
            user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
            if not user or user["account_status"] != "active":
                raise ValueError("test code is not linked to an active user")
        else:
            user = create_test_user(conn, normalized_code)
            claimed_at = now_text()
            conn.execute(
                """
                UPDATE test_access_codes
                SET user_id = ?, status = 'claimed', claimed_at = ?, updated_at = ?
                WHERE code = ?
                """,
                (user["user_id"], claimed_at, claimed_at, normalized_code),
            )

        session = start_auth_session(conn, user, normalized_code)
        conn.commit()
        return auth_payload(conn, user, session["token"], session["sessionId"], include_current_test_date=True)


def normalize_primary_books(raw_books: Any, require_scoped_selection: bool = True) -> list[dict[str, str]]:
    if not isinstance(raw_books, list) or len(raw_books) != 2:
        raise ValueError("主选习题册必须选择两本且不能重复")
    normalized: list[dict[str, str]] = []
    for item in raw_books:
        if isinstance(item, dict):
            book = str(profile_book_input_value(item) or "").strip()
            scope = str(item.get("scope") or "").strip()
        else:
            book = str(item or "").strip()
            scope = "all"
        if not book:
            raise ValueError("主选习题册不能为空")
        if legacy_book_code(book) in SCOPED_PRIMARY_BOOKS:
            if not require_scoped_selection and scope == "all":
                scope = "both"
            elif scope not in PRIMARY_BOOK_SCOPES:
                raise ValueError("880 和 1000 必须选择基础题、强化题或基础+强化都布置")
        else:
            scope = "all"
        normalized.append({"book": book, "scope": scope})
    if normalized[0]["book"] == normalized[1]["book"]:
        raise ValueError("两本主选习题册不能重复")
    return normalized


def normalize_primary_book_scope(book_key: Any, scope: Any, require_scoped_selection: bool = False) -> str:
    code = legacy_book_code(book_key)
    normalized_scope = str(scope or "all").strip()
    if code not in SCOPED_PRIMARY_BOOKS:
        return "all"
    if not require_scoped_selection and normalized_scope == "all":
        return "both"
    if normalized_scope not in PRIMARY_BOOK_SCOPES:
        raise ValueError("880 和 1000 必须选择基础题、强化题或基础+强化都布置")
    return normalized_scope


def diagnostic_correct_count(answers: list[dict[str, Any]]) -> int:
    answer_map = {item.get("questionId"): item.get("answer") for item in answers}
    if set(answer_map) != set(DIAGNOSTIC_ANSWERS):
        raise ValueError("基础诊断题必须全部作答")
    return sum(
        1
        for question_id, correct_answer in DIAGNOSTIC_ANSWERS.items()
        if answer_map.get(question_id) == correct_answer
    )


def assign_student_level(target_score: int, intensive_days: int, foundation_level: str, correct_count: int) -> str:
    if target_score >= 125:
        if intensive_days <= 55 and correct_count >= 1 and foundation_level != "基础薄弱":
            return "A"
        return "B" if correct_count >= 1 else "C"
    if target_score >= 100:
        if target_score <= 109 and intensive_days >= 70 and correct_count == 0:
            return "D"
        if intensive_days <= 65 and correct_count >= 1 and foundation_level in {"基础较好", "强化冲刺"}:
            return "B"
        return "C"
    return "D"


def validate_new_password(user: dict[str, Any], new_password: str, confirm_password: str) -> None:
    if new_password != confirm_password:
        raise ValueError("两次输入的新密码不一致")
    if len(new_password) != 8:
        raise ValueError("新密码必须为 8 位")
    if not any(char.isalpha() for char in new_password) or not any(char.isdigit() for char in new_password):
        raise ValueError("新密码必须同时包含字母和数字")
    if not password_is_new(user, new_password):
        raise ValueError("新密码不能和初始密码相同")


def avatar_for_onboarding(user_id: str, gender: str) -> str:
    pool = FEMALE_ONBOARDING_AVATAR_POOL if gender == "female" else MALE_ONBOARDING_AVATAR_POOL
    index = as_int(user_id) % len(pool)
    return pool[index]


def complete_onboarding(
    user_id: str,
    payload: dict[str, Any],
    session_context: dict[str, Any] | None = None,
) -> dict[str, Any]:
    session_context = session_context or {}
    is_formal_payload = bool(payload.get("newPassword") or payload.get("diagnosticAnswers"))
    if is_formal_payload:
        nickname = str(payload.get("nickname") or "").strip()
        gender = str(payload.get("gender") or "").strip()
        qq = str(payload.get("qq") or "").strip()
        math_type = str(payload.get("mathType") or "").strip()
        target_score = as_int(payload.get("targetScore"))
        foundation_level = str(payload.get("foundationLevel") or "").strip()
        primary_books = normalize_primary_books(payload.get("primaryBooks") or [])
        backup_book = str(payload.get("backupBookId") or payload.get("backupBook") or "").strip()
        daily_study_hours = as_float(payload.get("dailyStudyHours"))
        daily_target_minutes = str(round(daily_study_hours * 60))
        expected_intensive_days = as_int(payload.get("intensiveDays"))
        review_order = str(payload.get("reviewOrder") or "").strip()
        start_chapter_id = str(payload.get("startChapterId") or "").strip()
        diagnostic_answers = payload.get("diagnosticAnswers") or []

        if not nickname:
            raise ValueError("姓名不能为空")
        if len(nickname) > 5:
            raise ValueError("姓名最多 5 个字")
        if gender not in {"male", "female"}:
            raise ValueError("性别不正确")
        if not qq.isdigit():
            raise ValueError("QQ号只能填写数字")
        if math_type not in {"数学一", "数学二", "数学三"}:
            raise ValueError("数学类型不正确")
        if target_score <= 0 or target_score > 150:
            raise ValueError("目标分数必须在 1 到 150 之间")
        if foundation_level not in {"基础薄弱", "基础一般", "基础较好", "强化冲刺"}:
            raise ValueError("基础水平不正确")
        if not backup_book:
            raise ValueError("备选习题册不能为空")
        if daily_study_hours <= 0 or daily_study_hours > 12:
            raise ValueError("每日数学学习时长目标必须大于 0 且不超过 12h")
        if expected_intensive_days < 40 or expected_intensive_days > 100:
            raise ValueError("预计强化阶段天数必须在 40 到 100 天之间")
        if review_order not in {"gaoshu_first", "mixed"}:
            raise ValueError("复习顺序不正确")
        correct_count = diagnostic_correct_count(diagnostic_answers)
        level = assign_student_level(target_score, expected_intensive_days, foundation_level, correct_count)
        assigned_intensive_days = CLASS_INTENSIVE_DAYS[level]
    else:
        nickname = ""
        gender = ""
        qq = ""
        math_type = "数学二"
        target_score = as_int(payload.get("targetScore"))
        level = str(payload.get("level") or "").strip().upper()
        if level not in {"A", "B", "C", "D"}:
            raise ValueError("level must be one of A, B, C, D")
        if target_score <= 0 or target_score > 150:
            raise ValueError("targetScore must be between 1 and 150")
        primary_books = normalize_primary_books(payload.get("primaryBooks") or [], require_scoped_selection=False)
        backup_book = str(payload.get("backupBookId") or payload.get("backupBook") or "").strip() or "660"
        daily_target_minutes = str(clamp(as_int(payload.get("dailyTargetMinutes"), 180), 120, 480))
        review_order = "mixed"
        start_chapter_id = str(payload.get("startChapterId") or "").strip()
        diagnostic_answers = []
        correct_count = 0
        expected_intensive_days = CLASS_INTENSIVE_DAYS.get(level, 60)
        assigned_intensive_days = expected_intensive_days

    updated_at = now_text()
    plan_variant_id = public_seed.resolve_plan_variant_id(math_type, review_order)
    primary_books = canonicalize_primary_books(primary_books, math_type)
    backup_book = canonical_profile_book_id(backup_book, math_type)

    with get_connection() as conn:
        ensure_student_progress_tracking_columns(conn)
        ensure_student_profile_target_school_columns(conn)
        ensure_formal_activation_tables(conn)
        ensure_admin_console_tables(conn)
        class_id = resolve_onboarding_class_id(conn, math_type, level)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        if not user or user["account_status"] != "active":
            raise ValueError("student not found")
        app_today = get_runtime_test_date(conn)
        account_type = account_type_for_user(user)
        if account_type == "trial":
            initial_paid_until = add_days_text(app_today, 7)
        elif is_formal_payload:
            initial_paid_until = add_days_text(app_today, 30)
        else:
            initial_paid_until = "2026-12-31"
        profile_for_sequence = {
            "math_type": math_type,
            "primary_book_1": primary_books[0]["book"],
            "primary_book_2": primary_books[1]["book"],
            "primary_book_1_scope": primary_books[0]["scope"],
            "primary_book_2_scope": primary_books[1]["scope"],
            "level": level,
        }
        if start_chapter_id:
            start_chapter_id = validate_start_chapter_for_onboarding(conn, math_type, plan_variant_id, start_chapter_id)
            first_type = first_actual_assignable_type_in_chapter(
                conn,
                profile_for_sequence,
                "强化",
                plan_variant_id,
                start_chapter_id,
            )
        else:
            first_type = next_actual_assignable_type(conn, "", profile_for_sequence, "强化", plan_variant_id)
        if not first_type:
            raise ValueError("当前资料组合下该章节暂无可布置题目，请更换起始章节或习题册。")
        if is_formal_payload:
            validate_new_password(user, str(payload.get("newPassword") or ""), str(payload.get("confirmPassword") or ""))
            password_values: dict[str, Any] = {}
            set_user_password_values(password_values, str(payload.get("newPassword") or ""))
            conn.execute(
                """
                UPDATE users
                SET nickname = ?, gender = ?, QQ = ?, avatar_url = ?,
                    password_hash = ?, login_password = ?, password_changed_at = ?,
                    updated_at = ?
                WHERE user_id = ?
                """,
                (
                    nickname,
                    gender,
                    qq,
                    avatar_for_onboarding(user_id, gender),
                    password_values["password_hash"],
                    password_values["login_password"],
                    updated_at,
                    updated_at,
                    user_id,
                ),
            )
        conn.execute(
            """
            INSERT OR REPLACE INTO student_profiles (
              user_id, math_type, target_exam_year, target_exam_date, target_score,
              class_id, level, plan_enabled, paid_until, primary_book_1, primary_book_2,
              backup_book, current_type_id, current_stage, last_plan_date, created_at,
              updated_at, target_school_id, target_school_name, intensive_days,
              review_order, plan_variant_id, initial_start_chapter_id, diagnostic_json,
              primary_book_1_scope, primary_book_2_scope
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                user_id,
                math_type,
                "2026",
                FORMAL_EXAM_DATE,
                str(target_score),
                class_id,
                level,
                "1",
                initial_paid_until,
                primary_books[0]["book"],
                primary_books[1]["book"],
                backup_book,
                first_type["type_id"],
                "强化",
                "",
                updated_at,
                updated_at,
                "",
                "",
                str(assigned_intensive_days),
                review_order,
                plan_variant_id,
                start_chapter_id,
                json.dumps(
                    {
                        "answers": diagnostic_answers,
                        "correctCount": correct_count,
                        "expectedIntensiveDays": expected_intensive_days,
                    },
                    ensure_ascii=False,
                ),
                primary_books[0]["scope"],
                primary_books[1]["scope"],
            ),
        )
        conn.execute(
            """
            INSERT OR REPLACE INTO student_progress (
              progress_id, user_id, math_type, plan_stage, current_type_id,
              current_type_sort_code, progress_status, current_type_status, last_generated_plan_date,
              last_completed_type_id, unfinished_type_id, unfinished_question_ids,
              carryover_type_id, carryover_question_ids, mastery_scope_start_type_id,
              mastery_scope_started_at, mastery_scope_reason, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f"PR-{user_id}",
                user_id,
                math_type,
                "强化",
                first_type["type_id"],
                first_type["sort_code"],
                "active",
                "active",
                "",
                "",
                "",
                "[]",
                "",
                "[]",
                first_type["type_id"],
                updated_at,
                "onboarding",
                updated_at,
                updated_at,
            ),
        )
        room = study_room_for_profile(conn, {"class_id": class_id, "math_type": math_type, "target_exam_year": "2026"})
        if room:
            conn.execute(
                """
                INSERT OR REPLACE INTO student_study_goals (
                  goal_id, user_id, study_date, room_id, target_minutes,
                  completed_minutes, completion_rate, current_focus_minutes,
                  current_task_type_id, current_task_name, focus_status,
                  current_round_started_at, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    f"SG-{user_id}-{app_today.replace('-', '')}",
                    user_id,
                    app_today,
                    room["room_id"],
                    daily_target_minutes,
                    "0",
                    "0.0000",
                    "45",
                    first_type["type_id"],
                    first_type["type_name"],
                    "idle",
                    "",
                    updated_at,
                    updated_at,
                ),
            )
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        session = start_auth_session(
            conn,
            user,
            "onboarding",
            device_id=session_context.get("device_id") or "",
            device_name=session_context.get("device_name") or "",
            ip_address=session_context.get("ip_address") or "",
            user_agent=session_context.get("user_agent") or "",
        )
        conn.commit()
        result = auth_payload(conn, user, session["token"], session["sessionId"], include_current_test_date=True)
        result["assignedLevel"] = level
        return result


def table_exists(conn, table_name: str) -> bool:
    query, params = table_exists_query(table_name)
    row = row_to_dict(conn.execute(query, params).fetchone())
    if not row:
        return False
    return bool(row.get("table_ref") or row.get("name"))


def _truthy_sql_filter(column_name: str) -> str:
    return f"CAST({column_name} AS TEXT) IN ('1', 'true', 'TRUE', 'yes', 'YES', 'active', 'ACTIVE')"


def _variant_order_available(conn, plan_variant_id: str) -> bool:
    if not plan_variant_id or not table_exists(conn, "question_type_orders"):
        return False
    row = conn.execute(
        f"""
        SELECT 1
        FROM question_type_orders
        WHERE plan_variant_id = ?
          AND {_truthy_sql_filter("is_active")}
        LIMIT 1
        """,
        (plan_variant_id,),
    ).fetchone()
    return row is not None


def variant_duration_field_for_level(level: Any) -> str:
    level_text = str(level or "").strip().upper()
    if level_text == "A":
        return "variant_duration_130_days"
    if level_text == "D":
        return "variant_duration_80_days"
    return "variant_duration_100_days"


def variant_duration_select_clause(conn) -> str:
    columns = set(table_columns(conn, "question_type_orders"))
    if {"duration_130_days", "duration_100_days", "duration_80_days"}.issubset(columns):
        return """
                       o.duration_130_days AS variant_duration_130_days,
                       o.duration_100_days AS variant_duration_100_days,
                       o.duration_80_days AS variant_duration_80_days
        """.strip()
    if "suggested_days" in columns:
        return """
                       o.suggested_days AS variant_duration_130_days,
                       o.suggested_days AS variant_duration_100_days,
                       o.suggested_days AS variant_duration_80_days
        """.strip()
    return """
                       NULL AS variant_duration_130_days,
                       NULL AS variant_duration_100_days,
                       NULL AS variant_duration_80_days
    """.strip()


def _apply_variant_order(row: dict[str, Any] | None, level: Any = "") -> dict[str, Any] | None:
    if not row:
        return None
    ordered = dict(row)
    if "variant_sort_code" in ordered:
        ordered["sort_code"] = ordered.get("variant_sort_code") or ordered.get("sort_code") or ""
    duration_field = variant_duration_field_for_level(level)
    if duration_field in ordered or "variant_duration_100_days" in ordered or "variant_suggested_days" in ordered:
        ordered["suggested_days"] = (
            ordered.get(duration_field)
            or ordered.get("variant_duration_100_days")
            or ordered.get("variant_suggested_days")
            or ordered.get("suggested_days")
            or ""
        )
    if "variant_direct_assignable" in ordered:
        ordered["direct_assignable"] = ordered.get("variant_direct_assignable") or ordered.get("direct_assignable") or ""
    return ordered


def question_type_for_plan_variant(
    conn,
    type_id: str,
    plan_variant_id: str = "",
    level: Any = "",
) -> dict[str, Any] | None:
    type_id = str(type_id or "").strip()
    plan_variant_id = str(plan_variant_id or "").strip()
    if not type_id:
        return None
    if _variant_order_available(conn, plan_variant_id):
        row = row_to_dict(
            conn.execute(
                f"""
                SELECT q.*,
                       o.sort_code AS variant_sort_code,
                       {variant_duration_select_clause(conn)},
                       o.direct_assignable AS variant_direct_assignable
                FROM question_type_orders o
                JOIN question_types q ON q.type_id = o.type_id
                WHERE o.plan_variant_id = ?
                  AND o.type_id = ?
                  AND {_truthy_sql_filter("o.is_active")}
                LIMIT 1
                """,
                (plan_variant_id, type_id),
            ).fetchone()
        )
        if row:
            return _apply_variant_order(row, level)
    return row_to_dict(
        conn.execute("SELECT * FROM question_types WHERE type_id = ?", (type_id,)).fetchone()
    )


def next_assignable_type(
    conn,
    current_type_id: str,
    plan_variant_id: str = "",
    level: Any = "",
) -> dict[str, Any] | None:
    plan_variant_id = str(plan_variant_id or "").strip()
    if _variant_order_available(conn, plan_variant_id):
        current_order = row_to_dict(
            conn.execute(
                f"""
                SELECT sort_code
                FROM question_type_orders
                WHERE plan_variant_id = ?
                  AND type_id = ?
                  AND {_truthy_sql_filter("is_active")}
                LIMIT 1
                """,
                (plan_variant_id, current_type_id),
            ).fetchone()
        )
        current_sort_code = current_order.get("sort_code") if current_order else ""
        row = row_to_dict(
            conn.execute(
                f"""
                SELECT q.*,
                       o.sort_code AS variant_sort_code,
                       {variant_duration_select_clause(conn)},
                       o.direct_assignable AS variant_direct_assignable
                FROM question_type_orders o
                JOIN question_types q ON q.type_id = o.type_id
                WHERE o.plan_variant_id = ?
                  AND {_truthy_sql_filter("o.is_active")}
                  AND {_truthy_sql_filter("o.direct_assignable")}
                  AND (? = '' OR o.sort_code > ?)
                ORDER BY o.sort_code
                LIMIT 1
                """,
                (plan_variant_id, current_sort_code, current_sort_code),
            ).fetchone()
        )
        return _apply_variant_order(row, level)

    current = row_to_dict(
        conn.execute("SELECT * FROM question_types WHERE type_id = ?", (current_type_id,)).fetchone()
    )
    if not current:
        return row_to_dict(
            conn.execute(
                """
                SELECT * FROM question_types
                WHERE direct_assignable = '1'
                ORDER BY sort_code
                LIMIT 1
                """
            ).fetchone()
        )
    return row_to_dict(
        conn.execute(
            """
            SELECT * FROM question_types
            WHERE direct_assignable = '1' AND sort_code > ?
            ORDER BY sort_code
            LIMIT 1
            """,
            (current["sort_code"],),
        ).fetchone()
    )


def get_type_questions(
    conn,
    type_id: str,
    profile: dict[str, Any],
    plan_stage: str,
    question_ids: set[str] | None = None,
) -> list[dict[str, Any]]:
    return get_type_questions_for_book_keys(
        conn,
        type_id,
        profile,
        plan_stage,
        selected_book_keys(profile),
        question_ids,
    )


def get_type_questions_for_book_keys(
    conn,
    type_id: str,
    profile: dict[str, Any],
    plan_stage: str,
    book_keys: set[str],
    question_ids: set[str] | None = None,
) -> list[dict[str, Any]]:
    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT * FROM questions
            WHERE type_id = ?
            ORDER BY CAST(question_order AS INTEGER), question_id
            """,
            (type_id,),
        ).fetchall()
    )
    scope_map = selected_primary_book_scopes(profile)
    filtered = []
    for row in rows:
        if "is_active" in row and not truthy_text(row.get("is_active")):
            continue
        if book_keys and question_book_key(row) not in book_keys:
            continue
        if not question_matches_primary_book_scope(row, profile, scope_map):
            continue
        if plan_stage and row.get("plan_stage") and row["plan_stage"] != plan_stage:
            continue
        if not level_matches(row.get("applicable_levels") or "", profile.get("level") or ""):
            continue
        if question_ids is not None and row["question_id"] not in question_ids:
            continue
        filtered.append(row)
    return filtered


def question_source_group_entry(row: dict[str, Any]) -> dict[str, Any]:
    display = question_source_display(row)
    return {
        "chapterLabel": display["sourceChapterLabel"],
        "phaseLabel": display["sourceSectionLabel"],
        "questionTypeLabel": display["sourceSubsectionLabel"],
        "questionTypeShort": display["sourceSubsectionLabel"].replace("题", ""),
        "title": " · ".join(
            item
            for item in (
                display["sourceChapterLabel"],
                display["sourceSectionLabel"],
                display["sourceSubsectionLabel"],
            )
            if item
        ),
        "continuationTitle": " · ".join(
            item
            for item in (
                display["sourceChapterLabel"],
                display["sourceSectionLabel"],
                display["sourceSubsectionLabel"],
            )
            if item
        ),
        "numbers": [display["sourceQuestionLabel"]],
    }


def question_rows_source_groups(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    source_group_map: dict[str, dict[str, Any]] = {}
    for row in rows:
        display = question_source_display(row)
        source_name = profile_book_display_code(question_book_key(row) or question_source_label(row))
        source_label = display["bookDisplayLabel"] or book_display_label(source_name)
        group = source_group_map.setdefault(
            source_label,
            {
                "sourceName": source_name,
                "sourceLabel": source_label,
                "entries": [],
            },
        )
        group["entries"].append(question_source_group_entry(row))
    return list(source_group_map.values())


def question_preview_primary_books(profile: dict[str, Any]) -> list[dict[str, str]]:
    math_type = profile.get("math_type") or "数学二"
    books: list[dict[str, str]] = []
    seen: set[str] = set()
    for index in (1, 2):
        raw_book_id = profile.get(f"primary_book_{index}")
        book_id = canonical_profile_book_id(raw_book_id, math_type)
        if not book_id or book_id in seen:
            continue
        seen.add(book_id)
        scope = normalize_primary_book_scope(
            raw_book_id,
            profile.get(f"primary_book_{index}_scope"),
            require_scoped_selection=False,
        )
        books.append(
            {
                "bookId": book_id,
                "bookCode": profile_book_display_code(book_id),
                "bookLabel": book_display_label(book_id),
                "scope": scope,
            }
        )
    return books


def get_question_preview_options(user_id: str) -> dict[str, Any]:
    with get_connection() as conn:
        profile = row_to_dict(
            conn.execute(
                "SELECT * FROM student_profiles WHERE user_id = ?",
                (user_id,),
            ).fetchone()
        )
        if not profile:
            raise ValueError("未找到该学生")

        math_type = str(profile.get("math_type") or "").strip()
        plan_variant_id = str(profile.get("plan_variant_id") or "").strip()
        if not plan_variant_id:
            plan_variant_id = public_seed.resolve_plan_variant_id(
                math_type,
                str(profile.get("review_order") or "").strip(),
            )

        question_type_columns = set(table_columns(conn, "question_types"))
        type_active_clause = (
            f"AND {_truthy_sql_filter('qt.is_active')}"
            if "is_active" in question_type_columns
            else ""
        )
        has_variant_order = (
            table_exists(conn, "chapter_orders")
            and _variant_order_available(conn, plan_variant_id)
        )
        if has_variant_order:
            chapter_rows = rows_to_dicts(
                conn.execute(
                    f"""
                    SELECT c.chapter_id, c.chapter_name, c.subject,
                           co.chapter_sort
                    FROM chapter_orders co
                    JOIN chapters c ON c.chapter_id = co.chapter_id
                    WHERE co.plan_variant_id = ?
                      AND c.math_type = ?
                      AND {_truthy_sql_filter("co.is_active")}
                      AND {_truthy_sql_filter("c.is_active")}
                    ORDER BY CAST(co.chapter_sort AS INTEGER), c.chapter_id
                    """,
                    (plan_variant_id, math_type),
                ).fetchall()
            )
            type_rows = rows_to_dicts(
                conn.execute(
                    f"""
                    SELECT qt.type_id, qt.chapter_id, qt.type_name,
                           qto.type_sort_in_chapter, qto.final_type_label,
                           qto.sort_code,
                           {question_type_intro_select_clause(conn)}
                    FROM question_type_orders qto
                    JOIN question_types qt ON qt.type_id = qto.type_id
                    JOIN chapters c ON c.chapter_id = qt.chapter_id
                    WHERE qto.plan_variant_id = ?
                      AND c.math_type = ?
                      AND {_truthy_sql_filter("qto.is_active")}
                      AND {_truthy_sql_filter("c.is_active")}
                      {type_active_clause}
                    ORDER BY qto.sort_code, qt.type_id
                    """,
                    (plan_variant_id, math_type),
                ).fetchall()
            )
        else:
            chapter_rows = rows_to_dicts(
                conn.execute(
                    f"""
                    SELECT c.chapter_id, c.chapter_name, c.subject,
                           c.chapter_id AS chapter_sort
                    FROM chapters c
                    WHERE c.math_type = ?
                      AND {_truthy_sql_filter("c.is_active")}
                    ORDER BY c.chapter_id
                    """,
                    (math_type,),
                ).fetchall()
            )
            type_rows = rows_to_dicts(
                conn.execute(
                    f"""
                    SELECT qt.type_id, qt.chapter_id, qt.type_name,
                           0 AS type_sort_in_chapter,
                           '' AS final_type_label,
                           qt.sort_code AS sort_code,
                           {question_type_intro_select_clause(conn)}
                    FROM question_types qt
                    JOIN chapters c ON c.chapter_id = qt.chapter_id
                    WHERE c.math_type = ?
                      AND {_truthy_sql_filter("c.is_active")}
                      {type_active_clause}
                    ORDER BY qt.sort_code, qt.type_id
                    """,
                    (math_type,),
                ).fetchall()
            )

        chapter_map: dict[str, dict[str, Any]] = {}
        chapters: list[dict[str, Any]] = []
        for index, row in enumerate(chapter_rows, start=1):
            chapter_id = str(row.get("chapter_id") or "").strip()
            if not chapter_id or chapter_id in chapter_map:
                continue
            chapter = {
                "chapterId": chapter_id,
                "chapterName": str(row.get("chapter_name") or chapter_id).strip(),
                "subject": str(row.get("subject") or "").strip(),
                "sortOrder": as_int(row.get("chapter_sort"), index),
                "types": [],
            }
            chapter_map[chapter_id] = chapter
            chapters.append(chapter)

        type_chapter_map: dict[str, str] = {}
        for row in type_rows:
            chapter_id = str(row.get("chapter_id") or "").strip()
            chapter = chapter_map.get(chapter_id)
            type_id = str(row.get("type_id") or "").strip()
            if not chapter or not type_id:
                continue
            intro = question_type_intro_payload(row)
            type_name = str(row.get("type_name") or intro.get("cardTitle") or type_id).strip()
            display_type_label = str(
                row.get("final_type_label") or intro.get("cardTypeLabel") or ""
            ).strip()
            chapter["types"].append(
                {
                    "typeId": type_id,
                    "typeName": type_name,
                    "displayTypeLabel": display_type_label,
                    "cardTypeLabel": intro.get("cardTypeLabel") or display_type_label,
                    "sortOrder": as_int(row.get("type_sort_in_chapter"), 0),
                    "sortCode": str(row.get("sort_code") or "").strip(),
                }
            )
            type_chapter_map[type_id] = chapter_id

        progress = row_to_dict(
            conn.execute(
                "SELECT current_type_id FROM student_progress WHERE user_id = ?",
                (user_id,),
            ).fetchone()
        )
        current_type_id = str(progress.get("current_type_id") or "").strip()
        if current_type_id not in type_chapter_map:
            current_type_id = ""
        current_chapter_id = type_chapter_map.get(current_type_id, "")
        if not current_type_id:
            first_chapter = next((item for item in chapters if item["types"]), None)
            if first_chapter:
                current_chapter_id = first_chapter["chapterId"]
                current_type_id = first_chapter["types"][0]["typeId"]

    return {
        "student": {
            "userId": user_id,
            "mathType": math_type,
            "planVariantId": plan_variant_id,
        },
        "primaryBooks": question_preview_primary_books(profile),
        "currentChapterId": current_chapter_id,
        "currentTypeId": current_type_id,
        "chapters": chapters,
    }


def get_question_preview_type(user_id: str, type_id: str) -> dict[str, Any]:
    type_id = str(type_id or "").strip()
    if not type_id:
        raise ValueError("请选择题型")

    with get_connection() as conn:
        profile = row_to_dict(
            conn.execute(
                "SELECT * FROM student_profiles WHERE user_id = ?",
                (user_id,),
            ).fetchone()
        )
        if not profile:
            raise ValueError("未找到该学生")

        math_type = str(profile.get("math_type") or "").strip()
        plan_variant_id = str(profile.get("plan_variant_id") or "").strip()
        if not plan_variant_id:
            plan_variant_id = public_seed.resolve_plan_variant_id(
                math_type,
                str(profile.get("review_order") or "").strip(),
            )

        question_type_columns = set(table_columns(conn, "question_types"))
        type_active_clause = (
            f"AND {_truthy_sql_filter('qt.is_active')}"
            if "is_active" in question_type_columns
            else ""
        )
        if _variant_order_available(conn, plan_variant_id):
            type_row = row_to_dict(
                conn.execute(
                    f"""
                    SELECT qt.type_id, qt.chapter_id, qt.type_name,
                           c.chapter_name, c.subject,
                           qto.type_sort_in_chapter, qto.final_type_label,
                           qto.sort_code,
                           {question_type_intro_select_clause(conn)}
                    FROM question_type_orders qto
                    JOIN question_types qt ON qt.type_id = qto.type_id
                    JOIN chapters c ON c.chapter_id = qt.chapter_id
                    WHERE qto.plan_variant_id = ?
                      AND qto.type_id = ?
                      AND c.math_type = ?
                      AND {_truthy_sql_filter("qto.is_active")}
                      AND {_truthy_sql_filter("c.is_active")}
                      {type_active_clause}
                    LIMIT 1
                    """,
                    (plan_variant_id, type_id, math_type),
                ).fetchone()
            )
        else:
            type_row = row_to_dict(
                conn.execute(
                    f"""
                    SELECT qt.type_id, qt.chapter_id, qt.type_name,
                           c.chapter_name, c.subject,
                           0 AS type_sort_in_chapter,
                           '' AS final_type_label,
                           qt.sort_code AS sort_code,
                           {question_type_intro_select_clause(conn)}
                    FROM question_types qt
                    JOIN chapters c ON c.chapter_id = qt.chapter_id
                    WHERE qt.type_id = ?
                      AND c.math_type = ?
                      AND {_truthy_sql_filter("c.is_active")}
                      {type_active_clause}
                    LIMIT 1
                    """,
                    (type_id, math_type),
                ).fetchone()
            )
        if not type_row:
            raise ValueError("题型不属于当前学生的数学类型或复习方案")

        primary_books = question_preview_primary_books(profile)
        primary_book_keys: set[str] = set()
        for item in primary_books:
            primary_book_keys.update(expanded_book_keys(item["bookId"], math_type))
        rows = (
            get_type_questions_for_book_keys(
                conn,
                type_id,
                profile,
                str(profile.get("current_stage") or "").strip(),
                primary_book_keys,
            )
            if primary_book_keys
            else []
        )
        intro = question_type_intro_payload(type_row)

    return {
        "student": {"userId": user_id, "mathType": math_type},
        "primaryBooks": primary_books,
        "type": {
            "typeId": type_id,
            "typeName": str(type_row.get("type_name") or type_id).strip(),
            "chapterId": str(type_row.get("chapter_id") or "").strip(),
            "chapterName": str(type_row.get("chapter_name") or "").strip(),
            "subject": str(type_row.get("subject") or "").strip(),
            "displayTypeLabel": str(
                type_row.get("final_type_label") or intro.get("cardTypeLabel") or ""
            ).strip(),
            **intro,
        },
        "sourceGroups": question_rows_source_groups(rows),
        "questions": [
            {
                "questionId": str(row.get("question_id") or "").strip(),
                "bookId": canonical_profile_book_id(question_book_key(row), math_type),
                "label": str(row.get("original_position_text") or "").strip(),
                "order": as_int(row.get("question_order"), 0),
                "questionNote": str(row.get("question_note") or "").strip(),
                **question_source_display(row),
            }
            for row in rows
        ],
    }


def extra_practice_rows_by_type(
    conn,
    profile: dict[str, Any],
    plan_rows: list[dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    math_type = profile.get("math_type") or ""
    backup_keys = expanded_book_keys(profile.get("backup_book"), math_type)
    backup_keys.discard("")
    if not backup_keys:
        return {}

    plan_stage = profile.get("current_stage") or ""
    planned_question_ids = {str(row.get("question_id") or "") for row in plan_rows}
    result: dict[str, list[dict[str, Any]]] = {}
    for type_id in sorted({str(row.get("type_id") or "") for row in plan_rows}):
        if not type_id:
            continue
        rows = [
            row
            for row in get_type_questions_for_book_keys(conn, type_id, profile, plan_stage, backup_keys)
            if str(row.get("question_id") or "") not in planned_question_ids
        ]
        if rows:
            result[type_id] = rows
    return result


def build_extra_practice_types(
    plan_types: list[dict[str, Any]],
    rows_by_type: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    extras: list[dict[str, Any]] = []
    for plan_type in plan_types:
        type_id = str(plan_type.get("typeId") or "")
        rows = rows_by_type.get(type_id) or []
        if not rows:
            continue
        sources: dict[str, list[str]] = {}
        for row in rows:
            sources.setdefault(question_source_label(row), []).append(row["original_position_text"])
        extras.append(
            {
                "typeIndex": plan_type.get("typeIndex"),
                "typeId": type_id,
                "typeName": plan_type.get("typeName") or "",
                "chapterId": plan_type.get("chapterId") or "",
                "studentTypeSortInChapter": plan_type.get("studentTypeSortInChapter") or 0,
                "displayTypeLabel": plan_type.get("displayTypeLabel") or "",
                "cardTypeLabel": plan_type.get("cardTypeLabel") or plan_type.get("displayTypeLabel") or "",
                "hideSummary": True,
                "sources": [
                    {"source": source, "items": items}
                    for source, items in sources.items()
                ],
                "sourceGroups": question_rows_source_groups(rows),
                "questions": [
                    {
                        "questionId": row["question_id"],
                        "label": row["original_position_text"],
                        "order": as_int(row["question_order"]),
                        "questionNote": row.get("question_note") or "",
                        **question_source_display(row),
                    }
                    for row in rows
                ],
            }
        )
    return extras


def actual_assignable_type_sequence(
    conn,
    profile: dict[str, Any],
    plan_stage: str,
    plan_variant_id: str = "",
) -> list[dict[str, Any]]:
    sequence: list[dict[str, Any]] = []
    chapter_counts: dict[str, int] = {}
    seen: set[str] = set()
    level = profile.get("level") if profile else ""
    cursor = next_assignable_type(conn, "", plan_variant_id, level)
    while cursor:
        type_id = str(cursor.get("type_id") or "").strip()
        if not type_id or type_id in seen:
            break
        seen.add(type_id)
        if get_type_questions(conn, type_id, profile, plan_stage):
            chapter_id = str(cursor.get("chapter_id") or "").strip()
            chapter_counts[chapter_id] = chapter_counts.get(chapter_id, 0) + 1
            enriched = dict(cursor)
            enriched["chapter_id"] = chapter_id
            enriched["student_type_sort_in_chapter"] = str(chapter_counts[chapter_id])
            enriched["display_type_label"] = display_type_label(chapter_counts[chapter_id])
            enriched["actual_sequence_index"] = str(len(sequence) + 1)
            sequence.append(enriched)
        cursor = next_assignable_type(conn, type_id, plan_variant_id, level)

    for index, row in enumerate(sequence):
        next_row = sequence[index + 1] if index + 1 < len(sequence) else None
        row["next_actual_type_id"] = str(next_row.get("type_id") or "") if next_row else ""
        row["is_chapter_boundary_after"] = (
            "1"
            if not next_row or str(next_row.get("chapter_id") or "") != str(row.get("chapter_id") or "")
            else "0"
        )
    return sequence


def sequence_start_index_for_type(
    sequence: list[dict[str, Any]],
    current_type: dict[str, Any] | None,
) -> int:
    if not sequence:
        return 0
    if not current_type:
        return 0
    current_type_id = str(current_type.get("type_id") or "").strip()
    for index, row in enumerate(sequence):
        if str(row.get("type_id") or "") == current_type_id:
            return index
    current_sort_code = str(current_type.get("sort_code") or "").strip()
    if current_sort_code:
        for index, row in enumerate(sequence):
            if str(row.get("sort_code") or "").strip() > current_sort_code:
                return index
    return len(sequence)


def next_actual_assignable_type(
    conn,
    current_type_id: str,
    profile: dict[str, Any],
    plan_stage: str,
    plan_variant_id: str = "",
) -> dict[str, Any] | None:
    sequence = actual_assignable_type_sequence(conn, profile, plan_stage, plan_variant_id)
    if not sequence:
        return None
    current_type = question_type_for_plan_variant(conn, current_type_id, plan_variant_id, profile.get("level"))
    start_index = sequence_start_index_for_type(sequence, current_type)
    if start_index < len(sequence) and str(sequence[start_index].get("type_id") or "") == str(current_type_id or ""):
        next_index = start_index + 1
    else:
        next_index = start_index
    return sequence[next_index] if next_index < len(sequence) else None


def first_actual_assignable_type_in_chapter(
    conn,
    profile: dict[str, Any],
    plan_stage: str,
    plan_variant_id: str,
    chapter_id: str,
) -> dict[str, Any] | None:
    chapter_id = str(chapter_id or "").strip()
    if not chapter_id:
        return None
    for row in actual_assignable_type_sequence(conn, profile, plan_stage, plan_variant_id):
        if str(row.get("chapter_id") or "").strip() == chapter_id:
            return row
    return None


def first_chapter_for_plan_variant(conn, math_type: str, plan_variant_id: str) -> str:
    if not table_exists(conn, "chapters") or not table_exists(conn, "chapter_orders"):
        raise ValueError("起始章节配置缺失，无法校验当前数学类型或复习顺序")
    row = row_to_dict(
        conn.execute(
            f"""
            SELECT o.chapter_id
            FROM chapter_orders o
            JOIN chapters c ON c.chapter_id = o.chapter_id
            WHERE c.math_type = ?
              AND {_truthy_sql_filter("c.is_active")}
              AND o.plan_variant_id = ?
              AND {_truthy_sql_filter("o.is_active")}
            ORDER BY CAST(o.chapter_sort AS INTEGER), o.chapter_id
            LIMIT 1
            """,
            (math_type, plan_variant_id),
        ).fetchone()
    )
    return str((row or {}).get("chapter_id") or "").strip()


def validate_start_chapter_for_onboarding(
    conn,
    math_type: str,
    plan_variant_id: str,
    start_chapter_id: str,
) -> str:
    start_chapter_id = str(start_chapter_id or "").strip()
    if not start_chapter_id:
        start_chapter_id = first_chapter_for_plan_variant(conn, math_type, plan_variant_id)
    if not table_exists(conn, "chapters") or not table_exists(conn, "chapter_orders"):
        raise ValueError("起始章节配置缺失，无法校验当前数学类型或复习顺序")
    row = row_to_dict(
        conn.execute(
            f"""
            SELECT c.chapter_id
            FROM chapter_orders o
            JOIN chapters c ON c.chapter_id = o.chapter_id
            WHERE c.math_type = ?
              AND c.chapter_id = ?
              AND {_truthy_sql_filter("c.is_active")}
              AND o.plan_variant_id = ?
              AND {_truthy_sql_filter("o.is_active")}
            LIMIT 1
            """,
            (math_type, start_chapter_id, plan_variant_id),
        ).fetchone()
    )
    if not row:
        raise ValueError("起始章节不属于当前数学类型或复习顺序")
    return start_chapter_id


def get_system_message(key: str, default: str) -> str:
    with get_connection() as conn:
        row = row_to_dict(
            conn.execute(
                "SELECT * FROM system_messages WHERE message_key = ?",
                (key,),
            ).fetchone()
        )
    return row.get("message_text") if row else default


def require_activated_student(conn, user_id: str) -> dict[str, Any]:
    user = row_to_dict(
        conn.execute(
            """
            SELECT *
            FROM users
            WHERE user_id = ? AND role = 'student' AND account_status = 'active'
            """,
            (user_id,),
        ).fetchone()
    )
    if not user:
        raise ValueError("student not found")
    return user


def serialize_announcement(row: dict[str, Any], read_at: str = "") -> dict[str, Any]:
    return {
        "announcementId": row["announcement_id"],
        "title": row["title"],
        "content": row["content"],
        "audience": row["audience"],
        "status": row["status"],
        "imageUrl": row.get("image_url") or "",
        "imageFilename": row.get("image_filename") or "",
        "imageOriginalName": row.get("image_original_name") or "",
        "imageContentType": row.get("image_content_type") or "",
        "imageSizeBytes": as_int(row.get("image_size_bytes"), 0),
        "publishedAt": row["published_at"],
        "createdAt": row["created_at"],
        "updatedAt": row["updated_at"],
        "isRead": bool(read_at),
    }


def validate_announcement_text(title: str, content: str) -> tuple[str, str]:
    clean_title = str(title or "").strip()
    clean_content = str(content or "").strip()
    if not clean_title:
        raise ValueError("请输入公告标题")
    if len(clean_title) > 80:
        raise ValueError("公告标题不能超过 80 个字")
    if not clean_content:
        raise ValueError("请输入公告内容")
    if len(clean_content) > 2000:
        raise ValueError("公告内容不能超过 2000 个字")
    return clean_title, clean_content


def validate_announcement_image_meta(image: dict[str, Any] | None) -> dict[str, Any]:
    if not image:
        return {}
    content_type = str(image.get("content_type") or "").lower()
    size = as_int(image.get("size"), 0)
    if content_type not in ANNOUNCEMENT_ALLOWED_IMAGE_TYPES:
        raise ValueError("公告图片仅支持 JPG / PNG / WebP")
    if size <= 0:
        raise ValueError("公告图片为空")
    if size > ANNOUNCEMENT_MAX_IMAGE_BYTES:
        raise ValueError("公告图片不能超过 1MB")
    return image


def create_announcement(
    title: str,
    content: str,
    image: dict[str, Any] | None = None,
    created_by: str = "adm001",
) -> dict[str, Any]:
    clean_title, clean_content = validate_announcement_text(title, content)
    image_meta = validate_announcement_image_meta(image)
    created_at = now_text()
    announcement_id = f"ann-{uuid.uuid4().hex[:12]}"
    with get_connection() as conn:
        ensure_announcement_tables(conn)
        conn.execute(
            """
            INSERT INTO announcements (
              announcement_id, title, content, audience, status,
              image_url, image_filename, image_original_name, image_content_type, image_size_bytes,
              created_by, published_at, hidden_at, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '', ?, ?)
            """,
            (
                announcement_id,
                clean_title,
                clean_content,
                ANNOUNCEMENT_AUDIENCE_ACTIVATED,
                ANNOUNCEMENT_STATUS_PUBLISHED,
                image_meta.get("url", ""),
                image_meta.get("filename", ""),
                image_meta.get("original_name", ""),
                image_meta.get("content_type", ""),
                str(image_meta.get("size", 0) or 0),
                str(created_by or "adm001").strip() or "adm001",
                created_at,
                created_at,
                created_at,
            ),
        )
        conn.commit()
        row = row_to_dict(
            conn.execute(
                "SELECT * FROM announcements WHERE announcement_id = ?",
                (announcement_id,),
            ).fetchone()
        )
    return {"announcement": serialize_announcement(row)}


def list_admin_announcements() -> dict[str, Any]:
    with get_connection() as conn:
        ensure_announcement_tables(conn)
        rows = rows_to_dicts(
            conn.execute(
                """
                SELECT *
                FROM announcements
                ORDER BY published_at DESC, created_at DESC
                """
            ).fetchall()
        )
    return {"items": [serialize_announcement(row) for row in rows]}


def hide_announcement(announcement_id: str) -> dict[str, Any]:
    updated_at = now_text()
    with get_connection() as conn:
        ensure_announcement_tables(conn)
        row = row_to_dict(
            conn.execute(
                "SELECT * FROM announcements WHERE announcement_id = ?",
                (announcement_id,),
            ).fetchone()
        )
        if not row:
            raise ValueError("announcement not found")
        conn.execute(
            """
            UPDATE announcements
            SET status = ?, hidden_at = ?, updated_at = ?
            WHERE announcement_id = ?
            """,
            (ANNOUNCEMENT_STATUS_HIDDEN, updated_at, updated_at, announcement_id),
        )
        conn.commit()
        updated = row_to_dict(
            conn.execute(
                "SELECT * FROM announcements WHERE announcement_id = ?",
                (announcement_id,),
            ).fetchone()
        )
    return {"announcement": serialize_announcement(updated)}


def student_notification_rows(conn, user_id: str) -> list[dict[str, Any]]:
    require_activated_student(conn, user_id)
    return rows_to_dicts(
        conn.execute(
            """
            SELECT a.*, COALESCE(r.read_at, '') AS read_at
            FROM announcements a
            LEFT JOIN announcement_reads r
              ON r.announcement_id = a.announcement_id AND r.user_id = ?
            WHERE a.status = ? AND a.audience = ?
            ORDER BY a.published_at DESC, a.created_at DESC
            """,
            (user_id, ANNOUNCEMENT_STATUS_PUBLISHED, ANNOUNCEMENT_AUDIENCE_ACTIVATED),
        ).fetchall()
    )


def get_student_notification_status(user_id: str) -> dict[str, bool]:
    with get_connection() as conn:
        ensure_announcement_tables(conn)
        rows = student_notification_rows(conn, user_id)
    return {"hasUnread": any(not row.get("read_at") for row in rows)}


def get_student_notifications(user_id: str) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_announcement_tables(conn)
        rows = student_notification_rows(conn, user_id)
    return {
        "hasUnread": any(not row.get("read_at") for row in rows),
        "items": [serialize_announcement(row, row.get("read_at") or "") for row in rows],
    }


def mark_student_notifications_read(user_id: str) -> dict[str, bool]:
    read_at = now_text()
    with get_connection() as conn:
        ensure_announcement_tables(conn)
        rows = student_notification_rows(conn, user_id)
        for row in rows:
            if row.get("read_at"):
                continue
            conn.execute(
                """
                INSERT OR IGNORE INTO announcement_reads (announcement_id, user_id, read_at)
                VALUES (?, ?, ?)
                """,
                (row["announcement_id"], user_id, read_at),
            )
        conn.commit()
    return {"hasUnread": False}


def truthy_text(value: Any) -> bool:
    return str(value or "").strip().lower() not in {"", "0", "false", "no", "disabled"}


def add_days_text(date_text: str, days: int) -> str:
    try:
        base = date.fromisoformat(date_text)
    except ValueError:
        base = date.fromisoformat(CURRENT_DATE)
    return (base + timedelta(days=days)).isoformat()


def account_type_for_user(user: dict[str, Any]) -> str:
    explicit = str(user.get("account_type") or "").strip()
    if explicit:
        return explicit
    if user.get("role") != "student":
        return user.get("role") or "user"
    number = formal_account_number(user.get("login_account") or user.get("user_id") or "")
    if 10001 <= number <= 10010:
        return "internal_test"
    if SHORT_TERM_ACCOUNT_START <= number <= SHORT_TERM_ACCOUNT_END:
        return "trial"
    if TRIAL_CONVERSION_FORMAL_ACCOUNT_START <= number <= TRIAL_CONVERSION_FORMAL_ACCOUNT_END:
        return "formal"
    if FORMAL_ACCOUNT_START <= number <= FORMAL_ACCOUNT_END:
        return "formal"
    return "formal"


def class_has_open_capacity(conn, class_row: dict[str, Any] | None) -> bool:
    if not class_row or not truthy_text(class_row.get("is_accepting_new_students")):
        return False
    capacity = as_int(class_row.get("capacity"), 0)
    if capacity <= 0:
        return True
    return class_member_count(conn, class_row.get("class_id") or "") < capacity


def resolve_onboarding_class_id(conn, math_type: str, level: str) -> str:
    math_type = str(math_type or "").strip()
    level = str(level or "").strip().upper()
    active_class = row_to_dict(
        conn.execute(
            """
            SELECT c.*
            FROM class_intake_settings s
            JOIN classes c ON c.class_id = s.active_class_id
            WHERE s.math_type = ? AND s.level = ?
            ORDER BY s.updated_at DESC, c.class_id
            LIMIT 1
            """,
            (math_type, level),
        ).fetchone()
    )
    if class_has_open_capacity(conn, active_class):
        return active_class.get("class_id") or ""

    class_rows = rows_to_dicts(
        conn.execute(
            """
            SELECT *
            FROM classes
            WHERE math_type = ? AND level = ?
            ORDER BY class_id
            """,
            (math_type, level),
        ).fetchall()
    )
    for class_row in class_rows:
        if class_has_open_capacity(conn, class_row):
            return class_row.get("class_id") or ""

    return public_seed.resolve_active_class_id(math_type, level) or f"{level}201"


def class_accepting_label(class_row: dict[str, Any] | None) -> str:
    if not class_row:
        return "待建班"
    return "正常接收" if truthy_text(class_row.get("is_accepting_new_students")) else "已指定，需先开启接收"


def seed_class_intake_settings(conn) -> None:
    existing = row_to_dict(conn.execute("SELECT COUNT(*) AS total FROM class_intake_settings").fetchone())
    if as_int(existing.get("total") if existing else 0) > 0:
        return
    updated_at = now_text()
    for math_type in ("数学一", "数学二", "数学三"):
        for level in ("A", "B", "C", "D"):
            class_row = row_to_dict(
                conn.execute(
                    """
                    SELECT *
                    FROM classes
                    WHERE math_type = ? AND level = ?
                    ORDER BY is_accepting_new_students DESC, class_id
                    LIMIT 1
                    """,
                    (math_type, level),
                ).fetchone()
            )
            active_class_id = class_row.get("class_id") if class_row else ""
            conn.execute(
                """
                INSERT OR IGNORE INTO class_intake_settings (
                  math_type, level, active_class_id, status, updated_by, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    math_type,
                    level,
                    active_class_id,
                    class_accepting_label(class_row),
                    "system",
                    updated_at,
                ),
            )


def log_admin_action(
    conn,
    action: str,
    target: str = "",
    before: Any = "",
    after: Any = "",
    actor: str = "adm001",
    confirmed: bool = True,
) -> None:
    ensure_admin_console_tables(conn)
    def encode(value: Any) -> str:
        if isinstance(value, str):
            return value
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    conn.execute(
        """
        INSERT INTO admin_audit_logs (
          log_id, actor, action, target, before_value, after_value, confirmed, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            f"AUD-{uuid.uuid4().hex[:12].upper()}",
            str(actor or "adm001"),
            action,
            str(target or ""),
            encode(before),
            encode(after),
            "1" if confirmed else "0",
            now_text(),
        ),
    )


def serialize_user_account(user: dict[str, Any], profile: dict[str, Any] | None = None) -> dict[str, Any]:
    profile = profile or {}
    return {
        "userId": user.get("user_id") or "",
        "loginAccount": user.get("login_account") or "",
        "role": user.get("role") or "",
        "nickname": user.get("nickname") or "",
        "avatarUrl": user.get("avatar_url") or "",
        "contactQQ": user.get("QQ") or "",
        "accountType": account_type_for_user(user),
        "accountStatus": user.get("account_status") or "",
        "classId": profile.get("class_id") or "",
        "mathType": profile.get("math_type") or "",
        "level": profile.get("level") or "",
        "targetScore": as_int(profile.get("target_score")),
        "paidUntil": profile.get("paid_until") or "",
        "activatedAt": user.get("activated_at") or user.get("created_at") or "",
        "lastLoginAt": user.get("last_login_at") or "",
    }


def student_daily_target_minutes(conn, user_id: str, stat_date: str = CURRENT_DATE) -> int:
    goal = row_to_dict(
        conn.execute(
            "SELECT target_minutes FROM student_study_goals WHERE user_id = ? AND study_date = ?",
            (user_id, stat_date),
        ).fetchone()
    )
    if not goal:
        goal = row_to_dict(
            conn.execute(
                """
                SELECT target_minutes
                FROM student_study_goals
                WHERE user_id = ?
                ORDER BY study_date DESC, updated_at DESC
                LIMIT 1
                """,
                (user_id,),
            ).fetchone()
        )
    return clamp(as_int(goal.get("target_minutes") if goal else 180), 120, 480)


def serialize_teacher(conn, user: dict[str, Any]) -> dict[str, Any]:
    classes = [
        row["class_id"]
        for row in conn.execute(
            "SELECT class_id FROM classes WHERE teacher_user_id = ? ORDER BY class_id",
            (user.get("user_id") or "",),
        ).fetchall()
    ]
    return {
        "userId": user.get("user_id") or "",
        "loginAccount": user.get("login_account") or "",
        "nickname": user.get("nickname") or "",
        "status": user.get("account_status") or "",
        "contactQQ": user.get("QQ") or "",
        "classes": classes,
        "lastLoginAt": user.get("last_login_at") or "-",
    }


def serialize_class(conn, row: dict[str, Any]) -> dict[str, Any]:
    return {
        "classId": row.get("class_id") or "",
        "className": row.get("class_name") or "",
        "mathType": row.get("math_type") or "",
        "level": row.get("level") or "",
        "teacherId": row.get("teacher_user_id") or "",
        "teacherUserId": row.get("teacher_user_id") or "",
        "qqGroupNumber": row.get("qq_group_number") or row.get("group_name") or "",
        "capacity": as_int(row.get("capacity"), 60),
        "activeStudents": class_member_count(conn, row.get("class_id") or ""),
        "isTestClass": str(row.get("class_id") or "").upper().startswith("TEST") or row.get("level") == "T",
        "isAcceptingNewStudents": truthy_text(row.get("is_accepting_new_students")),
    }


def today_status_for_student(conn, user_id: str, stat_date: str = CURRENT_DATE) -> dict[str, Any]:
    plan_rows = rows_to_dicts(
        conn.execute(
            """
            SELECT plan_status, feedback_open_status, type_name
            FROM daily_plans
            WHERE user_id = ? AND plan_date = ?
            ORDER BY type_index
            """,
            (user_id, stat_date),
        ).fetchall()
    )
    if not plan_rows:
        return {
            "todayStatus": "unsubmitted",
            "planStatus": "not_started",
            "openTypeName": "",
            "submittedTypes": 0,
            "totalTypes": 0,
            "unfinishedCount": 0,
        }
    statuses = {row.get("plan_status") for row in plan_rows}
    open_row = next((row for row in plan_rows if row.get("feedback_open_status") == "open"), plan_rows[0])
    if "blocked_by_unfinished" in statuses:
        today_status = "unfinished"
    elif statuses == {"completed"}:
        today_status = "completed"
    elif "expired" in statuses:
        today_status = "abnormal"
    else:
        today_status = "unsubmitted"
    return {
        "todayStatus": today_status,
        "planStatus": open_row.get("plan_status") or "",
        "openTypeName": open_row.get("type_name") or "",
        "submittedTypes": sum(1 for row in plan_rows if row.get("plan_status") == "completed"),
        "totalTypes": len(plan_rows),
        "unfinishedCount": sum(1 for row in plan_rows if row.get("plan_status") == "blocked_by_unfinished"),
    }


def serialize_student(conn, user: dict[str, Any], profile: dict[str, Any] | None = None, stat_date: str = CURRENT_DATE) -> dict[str, Any]:
    if profile is None:
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user.get("user_id"),)).fetchone()
        )
    profile = profile or {}
    progress = row_to_dict(
        conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user.get("user_id"),)).fetchone()
    ) or {}
    status = today_status_for_student(conn, user.get("user_id") or "", stat_date)
    account = serialize_user_account(user, profile)
    if (
        account["accountType"] == "trial"
        and account["accountStatus"] == "active"
        and is_service_expired(profile, stat_date)
    ):
        account["accountStatus"] = "expired_pending_review"
    return {
        **account,
        "primaryBooks": [
            profile_book_display_code(item)
            for item in [profile.get("primary_book_1"), profile.get("primary_book_2")]
            if item
        ],
        "primaryBookScopes": [profile.get("primary_book_1_scope") or "all", profile.get("primary_book_2_scope") or "all"],
        "backupBook": profile_book_display_code(profile.get("backup_book")),
        "reviewOrder": profile.get("review_order") or "mixed",
        "planVariantId": profile.get("plan_variant_id") or "",
        "manualNextTypeId": progress.get("current_type_id") or "",
        "currentTypeId": progress.get("current_type_id") or "",
        "nextPlanStartChapterId": progress.get("next_plan_start_chapter_id") or "",
        "nextPlanStartEffectiveDate": progress.get("next_plan_start_effective_date") or "",
        "nextPlanStartStatus": progress.get("next_plan_start_status") or "",
        "dailyMathGoalHours": hours_from_minutes(student_daily_target_minutes(conn, user.get("user_id") or "", stat_date)),
        "trialExpiresAt": profile.get("paid_until") if account["accountType"] == "trial" else "",
        "everActivated": bool(user.get("activated_at") or profile),
        "flags": [],
        **status,
    }


def serialize_intake_setting(conn, row: dict[str, Any]) -> dict[str, Any]:
    class_row = row_to_dict(
        conn.execute("SELECT * FROM classes WHERE class_id = ?", (row.get("active_class_id") or "",)).fetchone()
    )
    status = row.get("status") or class_accepting_label(class_row)
    return {
        "mathType": row.get("math_type") or "",
        "level": row.get("level") or "",
        "activeClassId": row.get("active_class_id") or "",
        "status": status,
        "updatedBy": row.get("updated_by") or "",
        "updatedAt": row.get("updated_at") or "",
    }


def soft_delete_expired_short_term_accounts(conn, stat_date: str) -> None:
    # Historical compatibility hook: expired trial accounts are now a computed
    # display state, never an automatic database mutation.
    return None


def serialize_audit_log(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "logId": row.get("log_id") or "",
        "time": row.get("created_at") or "",
        "actor": row.get("actor") or "",
        "action": row.get("action") or "",
        "target": row.get("target") or "",
        "before": row.get("before_value") or "",
        "after": row.get("after_value") or "",
        "confirmed": truthy_text(row.get("confirmed")),
    }


def admin_users_with_profiles(conn, where_clause: str = "", params: tuple[Any, ...] = ()) -> list[dict[str, Any]]:
    query = f"""
        SELECT
          u.*,
          p.math_type, p.target_score, p.class_id, p.level, p.paid_until,
          p.primary_book_1, p.primary_book_2, p.backup_book
        FROM users u
        LEFT JOIN student_profiles p ON p.user_id = u.user_id
        {where_clause}
        ORDER BY u.role, u.user_id
    """
    return rows_to_dicts(conn.execute(query, params).fetchall())


def list_admin_accounts() -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        rows = admin_users_with_profiles(conn)
    return {"items": [serialize_user_account(row, row) for row in rows]}


def get_admin_account(user_id: str) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        row = row_to_dict(
            conn.execute(
                """
                SELECT u.*, p.*
                FROM users u
                LEFT JOIN student_profiles p ON p.user_id = u.user_id
                WHERE u.user_id = ? OR u.login_account = ?
                """,
                (user_id, user_id),
            ).fetchone()
        )
        if not row:
            raise ValueError("account not found")
    return {"account": serialize_user_account(row, row)}


def revoke_user_sessions(conn, user_id: str, updated_at: str, keep_session_id: str = "") -> None:
    if keep_session_id:
        conn.execute(
            """
            UPDATE user_sessions
            SET revoked_at = ?, updated_at = ?
            WHERE user_id = ? AND session_id <> ? AND revoked_at = ''
            """,
            (updated_at, updated_at, user_id, keep_session_id),
        )
        return
    conn.execute(
        """
        UPDATE user_sessions
        SET revoked_at = ?, updated_at = ?
        WHERE user_id = ? AND revoked_at = ''
        """,
        (updated_at, updated_at, user_id),
    )


def set_account_status(user_id: str, status: str, actor: str = "adm001") -> dict[str, Any]:
    if status not in {"active", "disabled", "archived"}:
        raise ValueError("unsupported account status")
    updated_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? OR login_account = ?", (user_id, user_id)).fetchone())
        if not user:
            raise ValueError("account not found")
        before = {"accountStatus": user.get("account_status")}
        current_session_id = user.get("current_session_id") or ""
        next_session_id = current_session_id if status == "active" else ""
        if status != "active":
            revoke_user_sessions(conn, user["user_id"], updated_at)
        conn.execute(
            "UPDATE users SET account_status = ?, current_session_id = ?, updated_at = ? WHERE user_id = ?",
            (status, next_session_id, updated_at, user["user_id"]),
        )
        log_admin_action(conn, f"account.{status}", user["user_id"], before, {"accountStatus": status}, actor)
        conn.commit()
        updated = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user["user_id"],)).fetchone())
    return {"account": serialize_user_account(updated)}


def generate_reset_password() -> str:
    return "R" + "".join(secrets.choice(INITIAL_CREDENTIAL_ALPHABET) for _ in range(7))


def generate_teacher_initial_password() -> str:
    return "T" + "".join(secrets.choice(INITIAL_CREDENTIAL_ALPHABET) for _ in range(7))


def reset_account_password(user_id: str, actor: str = "adm001") -> dict[str, Any]:
    password = generate_reset_password()
    updated_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? OR login_account = ?", (user_id, user_id)).fetchone())
        if not user:
            raise ValueError("account not found")
        values: dict[str, Any] = {}
        set_user_password_values(values, password)
        conn.execute(
            """
            UPDATE users
            SET password_hash = ?, login_password = ?, password_changed_at = '',
                current_session_id = '', updated_at = ?
            WHERE user_id = ?
            """,
            (values["password_hash"], values["login_password"], updated_at, user["user_id"]),
        )
        revoke_user_sessions(conn, user["user_id"], updated_at)
        log_admin_action(conn, "account.reset_password", user["user_id"], "", {"temporaryPasswordIssued": True}, actor)
        conn.commit()
        updated = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user["user_id"],)).fetchone())
    return {"account": serialize_user_account(updated), "temporaryPassword": password, "newPassword": password}


def list_admin_teachers() -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        rows = rows_to_dicts(conn.execute("SELECT * FROM users WHERE role = 'teacher' ORDER BY user_id").fetchall())
        items = [serialize_teacher(conn, row) for row in rows]
    return {"items": items}


def next_teacher_user_id(conn) -> str:
    return str(max(90000, max_numeric_user_id(conn)) + 1)


def create_admin_teacher(payload: dict[str, Any], actor: str = "adm001") -> dict[str, Any]:
    login_account = str(payload.get("loginAccount") or "").strip()
    nickname = str(payload.get("nickname") or "").strip() or "新老师"
    contact_qq = str(payload.get("contactQQ") or "").strip()
    status = str(payload.get("status") or "active").strip()
    if not login_account:
        raise ValueError("teacher loginAccount required")
    if status not in {"active", "disabled"}:
        raise ValueError("teacher status invalid")
    created_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        if row_to_dict(conn.execute("SELECT * FROM users WHERE login_account = ?", (login_account,)).fetchone()):
            raise ValueError("teacher loginAccount already exists")
        user_id = next_teacher_user_id(conn)
        initial_password = generate_teacher_initial_password()
        values: dict[str, Any] = {}
        set_user_password_values(values, initial_password)
        conn.execute(
            """
            INSERT INTO users (
              user_id, login_account, password_hash, login_password, role,
              nickname, avatar_url, QQ, account_status, current_session_id,
              last_login_at, created_at, updated_at, account_type
            ) VALUES (?, ?, ?, ?, 'teacher', ?, '', ?, ?, '', '', ?, ?, 'teacher')
            """,
            (
                user_id,
                login_account,
                values["password_hash"],
                values["login_password"],
                nickname,
                contact_qq,
                status,
                created_at,
                created_at,
            ),
        )
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        serialized = serialize_teacher(conn, user)
        log_admin_action(conn, "teacher.create", user_id, "", serialized, actor)
        conn.commit()
    return {"teacher": serialized, "initialPassword": initial_password}


def update_teacher_classes(user_id: str, class_ids: list[str], actor: str = "adm001") -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        teacher = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? AND role = 'teacher'", (user_id,)).fetchone())
        if not teacher:
            raise ValueError("teacher not found")
        valid_ids = {
            row["class_id"]
            for row in conn.execute("SELECT class_id FROM classes WHERE class_id IN (%s)" % ",".join("?" for _ in class_ids), class_ids).fetchall()
        } if class_ids else set()
        if set(class_ids) != valid_ids:
            raise ValueError("classIds contain unknown class")
        before = serialize_teacher(conn, teacher)
        conn.execute("UPDATE classes SET teacher_user_id = '', updated_at = ? WHERE teacher_user_id = ?", (now_text(), user_id))
        for class_id in class_ids:
            conn.execute("UPDATE classes SET teacher_user_id = ?, updated_at = ? WHERE class_id = ?", (user_id, now_text(), class_id))
        updated = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        after = serialize_teacher(conn, updated)
        log_admin_action(conn, "teacher.assign_classes", user_id, before, after, actor)
        conn.commit()
    return {"teacher": after}


def update_teacher_status(user_id: str, status: str, actor: str = "adm001", nickname: str | None = None) -> dict[str, Any]:
    if status not in {"active", "disabled", "archived"}:
        raise ValueError("unsupported account status")
    updated_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        teacher = row_to_dict(
            conn.execute(
                "SELECT * FROM users WHERE user_id = ? AND role = 'teacher'",
                (user_id,),
            ).fetchone()
        )
        if not teacher:
            raise ValueError("teacher not found")
        if status == "archived":
            assigned = rows_to_dicts(
                conn.execute(
                    "SELECT class_id FROM classes WHERE teacher_user_id = ? ORDER BY class_id",
                    (user_id,),
                ).fetchall()
            )
            if assigned:
                raise ValueError("teacher still owns classes; transfer classes before archive")
        before = serialize_teacher(conn, teacher)
        current_session_id = teacher.get("current_session_id") or ""
        next_session_id = current_session_id if status == "active" else ""
        if status != "active":
            revoke_user_sessions(conn, teacher["user_id"], updated_at)
        next_nickname = str(nickname).strip() if nickname is not None else ""
        if not next_nickname:
            next_nickname = teacher.get("nickname") or ""
        conn.execute(
            """
            UPDATE users
            SET account_status = ?, nickname = ?, current_session_id = ?, updated_at = ?
            WHERE user_id = ?
            """,
            (status, next_nickname, next_session_id, updated_at, teacher["user_id"]),
        )
        updated = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (teacher["user_id"],)).fetchone())
        after = serialize_teacher(conn, updated)
        log_admin_action(conn, "teacher.update_status", user_id, before, after, actor)
        conn.commit()
    return {"account": serialize_user_account(updated), "teacher": after}


def get_teacher_by_id(user_id: str) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? AND role = 'teacher'", (user_id,)).fetchone())
        if not user:
            raise ValueError("teacher not found")
        return serialize_teacher(conn, user)


def change_teacher_password(
    teacher_user_id: str,
    current_password: str,
    new_password: str,
    confirm_password: str,
    session_id: str,
) -> dict[str, Any]:
    teacher_user_id = str(teacher_user_id or "").strip()
    session_id = str(session_id or "").strip()
    updated_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        user = row_to_dict(
            conn.execute(
                "SELECT * FROM users WHERE user_id = ? AND role = 'teacher'",
                (teacher_user_id,),
            ).fetchone()
        )
        if not user:
            raise ValueError("teacher not found")
        if not password_matches(user, current_password):
            raise ValueError("current password invalid")
        validate_new_password(user, new_password, confirm_password)
        values: dict[str, Any] = {}
        set_user_password_values(values, new_password)
        revoke_user_sessions(conn, teacher_user_id, updated_at, keep_session_id=session_id)
        conn.execute(
            """
            UPDATE users
            SET password_hash = ?, login_password = ?, password_changed_at = ?,
                current_session_id = ?, updated_at = ?
            WHERE user_id = ?
            """,
            (
                values["password_hash"],
                values["login_password"],
                updated_at,
                session_id,
                updated_at,
                teacher_user_id,
            ),
        )
        conn.execute(
            "UPDATE user_sessions SET updated_at = ? WHERE session_id = ? AND user_id = ?",
            (updated_at, session_id, teacher_user_id),
        )
        conn.commit()
        updated = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (teacher_user_id,)).fetchone())
    return {
        "ok": True,
        "sessionId": session_id,
        "passwordChangedAt": updated.get("password_changed_at") or updated_at,
        "teacher": {
            "userId": updated.get("user_id") or "",
            "loginAccount": updated.get("login_account") or "",
            "nickname": updated.get("nickname") or "",
        },
    }


def update_teacher_profile(teacher_user_id: str, contact_qq: str) -> dict[str, Any]:
    teacher_user_id = str(teacher_user_id or "").strip()
    contact_qq = str(contact_qq or "").strip()
    if not contact_qq:
        raise ValueError("QQ号不能为空")
    if not contact_qq.isdigit():
        raise ValueError("QQ号只能填写数字")
    updated_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        user = row_to_dict(
            conn.execute(
                "SELECT * FROM users WHERE user_id = ? AND role = 'teacher'",
                (teacher_user_id,),
            ).fetchone()
        )
        if not user:
            raise ValueError("teacher not found")
        conn.execute(
            "UPDATE users SET QQ = ?, updated_at = ? WHERE user_id = ?",
            (contact_qq, updated_at, teacher_user_id),
        )
        conn.commit()
        updated = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (teacher_user_id,)).fetchone())
        teacher = serialize_teacher(conn, updated)
    return {"teacher": teacher}


def list_admin_classes() -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        rows = rows_to_dicts(conn.execute("SELECT * FROM classes ORDER BY class_id").fetchall())
        items = [serialize_class(conn, row) for row in rows]
    return {"items": items}


def create_admin_class(payload: dict[str, Any], actor: str = "adm001") -> dict[str, Any]:
    class_id = str(payload.get("classId") or "").strip()
    class_name = str(payload.get("className") or class_id).strip()
    math_type = str(payload.get("mathType") or "数学二").strip()
    level = str(payload.get("level") or "A").strip().upper()
    teacher_user_id = str(payload.get("teacherUserId") or payload.get("teacherId") or "").strip()
    qq_group_number = str(payload.get("qqGroupNumber") or "").strip()
    capacity = str(as_int(payload.get("capacity"), 60))
    accepting = "1" if bool(payload.get("isAcceptingNewStudents", True)) else "0"
    if not class_id:
        raise ValueError("classId required")
    if math_type not in {"数学一", "数学二", "数学三"} or level not in {"A", "B", "C", "D"}:
        raise ValueError("class mathType or level invalid")
    created_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        if row_to_dict(conn.execute("SELECT * FROM classes WHERE class_id = ?", (class_id,)).fetchone()):
            raise ValueError("class already exists")
        conn.execute(
            """
            INSERT INTO classes (
              class_id, class_name, math_type, target_exam_year, level,
              teacher_user_id, group_name, created_at, updated_at,
              capacity, is_accepting_new_students, qq_group_number
            ) VALUES (?, ?, ?, '2026', ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                class_id,
                class_name,
                math_type,
                level,
                teacher_user_id,
                qq_group_number,
                created_at,
                created_at,
                capacity,
                accepting,
                qq_group_number,
            ),
        )
        row = row_to_dict(conn.execute("SELECT * FROM classes WHERE class_id = ?", (class_id,)).fetchone())
        serialized = serialize_class(conn, row)
        log_admin_action(conn, "class.create", class_id, "", serialized, actor)
        conn.commit()
    return {"class": serialized}


def patch_admin_class(class_id: str, payload: dict[str, Any], actor: str = "adm001") -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        row = row_to_dict(conn.execute("SELECT * FROM classes WHERE class_id = ?", (class_id,)).fetchone())
        if not row:
            raise ValueError("class not found")
        before = serialize_class(conn, row)
        next_values = {
            "class_name": str(payload.get("className") or row.get("class_name") or class_id).strip(),
            "teacher_user_id": str(payload.get("teacherUserId") or payload.get("teacherId") or row.get("teacher_user_id") or "").strip(),
            "qq_group_number": str(payload.get("qqGroupNumber") or row.get("qq_group_number") or "").strip(),
            "group_name": str(payload.get("qqGroupNumber") or row.get("group_name") or "").strip(),
            "capacity": str(as_int(payload.get("capacity"), as_int(row.get("capacity"), 60))),
            "is_accepting_new_students": "1" if bool(payload.get("isAcceptingNewStudents", truthy_text(row.get("is_accepting_new_students")))) else "0",
            "updated_at": now_text(),
        }
        conn.execute(
            """
            UPDATE classes
            SET class_name = ?, teacher_user_id = ?, qq_group_number = ?, group_name = ?,
                capacity = ?, is_accepting_new_students = ?, updated_at = ?
            WHERE class_id = ?
            """,
            (
                next_values["class_name"],
                next_values["teacher_user_id"],
                next_values["qq_group_number"],
                next_values["group_name"],
                next_values["capacity"],
                next_values["is_accepting_new_students"],
                next_values["updated_at"],
                class_id,
            ),
        )
        updated = row_to_dict(conn.execute("SELECT * FROM classes WHERE class_id = ?", (class_id,)).fetchone())
        serialized = serialize_class(conn, updated)
        log_admin_action(conn, "class.update", class_id, before, serialized, actor)
        conn.commit()
    return {"class": serialized}


def list_admin_students(stat_date: str | None = None) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        stat_date = stat_date or get_runtime_test_date(conn)
        rows = rows_to_dicts(
            conn.execute(
                """
                SELECT u.*
                FROM users u
                JOIN student_profiles p ON p.user_id = u.user_id
                WHERE u.role = 'student'
                ORDER BY u.user_id
                """
            ).fetchall()
        )
        items = [serialize_student(conn, row, stat_date=stat_date) for row in rows]
        conn.commit()
    return {"items": items}


def list_admin_chapter_options(math_type: str = "数学二", plan_variant_id: str = "") -> dict[str, Any]:
    math_type = str(math_type or "数学二").strip() or "数学二"
    plan_variant_id = str(plan_variant_id or "").strip()
    with get_connection() as conn:
        if table_exists(conn, "chapters"):
            if plan_variant_id and table_exists(conn, "chapter_orders"):
                rows = rows_to_dicts(
                    conn.execute(
                        f"""
                        SELECT
                          c.chapter_id,
                          c.chapter_name,
                          c.subject,
                          COALESCE(co.chapter_sort, 9999) AS sort_order
                        FROM chapters c
                        LEFT JOIN chapter_orders co
                          ON co.chapter_id = c.chapter_id
                         AND co.plan_variant_id = ?
                         AND {_truthy_sql_filter("co.is_active")}
                        WHERE c.math_type = ?
                          AND {_truthy_sql_filter("c.is_active")}
                        ORDER BY CAST(COALESCE(co.chapter_sort, 9999) AS INTEGER), c.chapter_id
                        """,
                        (plan_variant_id, math_type),
                    ).fetchall()
                )
            else:
                rows = rows_to_dicts(
                    conn.execute(
                        f"""
                        SELECT
                          chapter_id,
                          chapter_name,
                          subject,
                          chapter_id AS sort_order
                        FROM chapters
                        WHERE math_type = ?
                          AND {_truthy_sql_filter("is_active")}
                        ORDER BY chapter_id
                        """,
                        (math_type,),
                    ).fetchall()
                )
        else:
            rows = raw_chapter_type_rows(conn, math_type)

        seen: set[str] = set()
        items: list[dict[str, Any]] = []
        for index, row in enumerate(rows, start=1):
            chapter_id = str(row.get("chapter_id") or "").strip()
            if not chapter_id or chapter_id in seen:
                continue
            seen.add(chapter_id)
            items.append(
                {
                    "chapterId": chapter_id,
                    "chapterName": str(row.get("chapter_name") or chapter_id).strip(),
                    "subject": str(row.get("subject") or "").strip(),
                    "sortOrder": as_int(row.get("sort_order"), index),
                }
            )
    return {"items": items}


def set_student_progress_plan_start(
    conn,
    user_id: str,
    type_row: dict[str, Any] | None,
    updated_at: str,
    progress_status: str = "active",
) -> None:
    type_id = str((type_row or {}).get("type_id") or "").strip()
    sort_code = str((type_row or {}).get("sort_code") or "").strip()
    conn.execute(
        """
        UPDATE student_progress
        SET current_type_id = ?, current_type_sort_code = ?,
            progress_status = ?, current_type_status = 'active',
            unfinished_type_id = '', unfinished_question_ids = '[]',
            carryover_type_id = '', carryover_question_ids = '[]',
            updated_at = ?
        WHERE user_id = ?
        """,
        (type_id, sort_code, progress_status, updated_at, user_id),
    )


def set_mastery_scope_start(
    conn,
    user_id: str,
    type_row: dict[str, Any] | None,
    updated_at: str,
    reason: str,
) -> None:
    if not type_row:
        return
    type_id = str(type_row.get("type_id") or "").strip()
    if not type_id:
        return
    conn.execute(
        """
        UPDATE student_progress
        SET mastery_scope_start_type_id = ?,
            mastery_scope_started_at = ?,
            mastery_scope_reason = ?,
            updated_at = ?
        WHERE user_id = ?
        """,
        (type_id, updated_at, reason, updated_at, user_id),
    )


def primary_book_selection_changed(
    profile: dict[str, Any],
    first_book: str,
    first_scope: str,
    second_book: str,
    second_scope: str,
) -> bool:
    return (
        str(profile.get("primary_book_1") or "") != first_book
        or str(profile.get("primary_book_2") or "") != second_book
        or str(profile.get("primary_book_1_scope") or "all") != first_scope
        or str(profile.get("primary_book_2_scope") or "all") != second_scope
    )


def regenerate_today_unsubmitted_plan_after_book_change(
    conn,
    user: dict[str, Any],
    profile: dict[str, Any],
    progress: dict[str, Any],
    plan_date: str,
    updated_at: str,
) -> None:
    ensure_daily_plan_sequence_columns(conn)
    plan_rows = rows_to_dicts(
        conn.execute(
            """
            SELECT *
            FROM daily_plans
            WHERE user_id = ? AND plan_date = ?
            ORDER BY CAST(type_index AS INTEGER), CAST(question_order AS INTEGER)
            """,
            (user["user_id"], plan_date),
        ).fetchall()
    )
    plan_stage = progress.get("plan_stage") or profile.get("current_stage") or "强化"
    plan_variant_id = str(profile.get("plan_variant_id") or "").strip()
    submitted_type_ids: list[str] = []
    for row in plan_rows:
        type_id = str(row.get("type_id") or "").strip()
        if row.get("feedback_open_status") == "submitted" and type_id and type_id not in submitted_type_ids:
            submitted_type_ids.append(type_id)

    if plan_rows and len(submitted_type_ids) == len({str(row.get("type_id") or "") for row in plan_rows if row.get("type_id")}):
        conn.execute(
            """
            UPDATE daily_plans
            SET plan_status = 'completed', updated_at = ?
            WHERE user_id = ? AND plan_date = ?
            """,
            (updated_at, user["user_id"], plan_date),
        )
        current_type = question_type_for_plan_variant(conn, progress.get("current_type_id"), plan_variant_id, profile.get("level"))
        set_student_progress_plan_start(conn, user["user_id"], current_type, updated_at)
        set_mastery_scope_start(conn, user["user_id"], current_type, updated_at, "admin_book_change")
        return

    if submitted_type_ids:
        start_type = next_actual_assignable_type(
            conn,
            submitted_type_ids[-1],
            profile,
            plan_stage,
            plan_variant_id,
        )
    else:
        active_type_id = str(progress.get("current_type_id") or "").strip()
        start_type = question_type_for_plan_variant(conn, active_type_id, plan_variant_id, profile.get("level")) or next_assignable_type(
            conn,
            "",
            plan_variant_id,
            profile.get("level"),
        )
        if start_type:
            sequence = actual_assignable_type_sequence(conn, profile, plan_stage, plan_variant_id)
            start_index = sequence_start_index_for_type(sequence, start_type)
            start_type = sequence[start_index] if start_index < len(sequence) else None

    conn.execute(
        """
        DELETE FROM daily_plans
        WHERE user_id = ? AND plan_date = ? AND feedback_open_status <> 'submitted'
        """,
        (user["user_id"], plan_date),
    )

    if not start_type:
        conn.execute(
            """
            UPDATE daily_plans
            SET plan_status = 'completed', updated_at = ?
            WHERE user_id = ? AND plan_date = ?
            """,
            (updated_at, user["user_id"], plan_date),
        )
        set_student_progress_plan_start(conn, user["user_id"], None, updated_at, progress_status="completed")
        return

    set_student_progress_plan_start(conn, user["user_id"], start_type, updated_at)
    set_mastery_scope_start(conn, user["user_id"], start_type, updated_at, "admin_book_change")
    refreshed_progress = row_to_dict(
        conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user["user_id"],)).fetchone()
    ) or {}
    generate_daily_plan(
        conn,
        user,
        profile,
        refreshed_progress,
        plan_date,
        type_index_offset=len(submitted_type_ids),
        commit=False,
    )
    conn.execute(
        """
        UPDATE daily_plans
        SET plan_status = 'in_progress', updated_at = ?
        WHERE user_id = ? AND plan_date = ?
        """,
        (updated_at, user["user_id"], plan_date),
    )


def upsert_student_daily_target_minutes(
    conn,
    user_id: str,
    stat_date: str,
    profile: dict[str, Any],
    target_minutes: int,
    updated_at: str,
) -> None:
    target = clamp(as_int(target_minutes, 180), 120, 480)
    existing = row_to_dict(
        conn.execute(
            "SELECT goal_id FROM student_study_goals WHERE user_id = ? AND study_date = ?",
            (user_id, stat_date),
        ).fetchone()
    )
    if existing:
        conn.execute(
            """
            UPDATE student_study_goals
            SET target_minutes = ?, updated_at = ?
            WHERE user_id = ? AND study_date = ?
            """,
            (str(target), updated_at, user_id, stat_date),
        )
        return

    room = study_room_for_profile(conn, profile)
    plan_row = row_to_dict(
        conn.execute(
            """
            SELECT type_id, type_name
            FROM daily_plans
            WHERE user_id = ? AND plan_date = ?
            ORDER BY type_index
            LIMIT 1
            """,
            (user_id, stat_date),
        ).fetchone()
    ) or {}
    conn.execute(
        """
        INSERT INTO student_study_goals (
          goal_id, user_id, study_date, room_id, target_minutes, completed_minutes,
          completion_rate, current_focus_minutes, current_task_type_id, current_task_name,
          focus_status, current_round_started_at, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, '0', '0', '45', ?, ?, 'idle', '', ?, ?)
        """,
        (
            f"SG-{user_id}-{stat_date.replace('-', '')}",
            user_id,
            stat_date,
            room["room_id"] if room else "",
            str(target),
            plan_row.get("type_id") or "",
            plan_row.get("type_name") or "",
            updated_at,
            updated_at,
        ),
    )


def patch_admin_student(user_id: str, payload: dict[str, Any], actor: str = "adm001") -> dict[str, Any]:
    primary_books = payload.get("primaryBooks")
    if primary_books is not None and (not isinstance(primary_books, list) or len(primary_books) < 1 or len(primary_books) > 2):
        raise ValueError("primaryBooks invalid")
    updated_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        ensure_formal_activation_tables(conn)
        ensure_student_progress_tracking_columns(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? AND role = 'student'", (user_id,)).fetchone())
        profile = row_to_dict(conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone())
        if not user or not profile:
            raise ValueError("student not found")
        before = serialize_student(conn, user, profile)
        contact_qq = str(payload.get("contactQQ") or user.get("QQ") or "").strip()
        review_order = str(payload.get("reviewOrder") or profile.get("review_order") or "mixed").strip()
        if review_order not in {"gaoshu_first", "mixed"}:
            raise ValueError("复习顺序不正确")
        level = str(payload.get("level") or profile.get("level") or "").strip()
        math_type = str(payload.get("mathType") or profile.get("math_type") or "").strip()
        plan_variant_id = public_seed.resolve_plan_variant_id(math_type, review_order)
        conn.execute("UPDATE users SET QQ = ?, updated_at = ? WHERE user_id = ?", (contact_qq, updated_at, user_id))
        next_primary = primary_books or [
            {"book": profile.get("primary_book_1"), "scope": profile.get("primary_book_1_scope")},
            {"book": profile.get("primary_book_2"), "scope": profile.get("primary_book_2_scope")},
        ]

        def primary_book_and_scope(index: int) -> tuple[str, str]:
            if index >= len(next_primary):
                return "", "all"
            item = next_primary[index]
            book = canonical_profile_book_id(profile_book_input_value(item), math_type)
            if not book:
                return "", "all"
            raw_scope = item.get("scope") if isinstance(item, dict) else profile.get(f"primary_book_{index + 1}_scope")
            scope = normalize_primary_book_scope(book, raw_scope, require_scoped_selection=False)
            return book, scope

        first_book, first_scope = primary_book_and_scope(0)
        second_book, second_scope = primary_book_and_scope(1)
        if first_book and second_book and first_book == second_book:
            raise ValueError("两本主选习题册不能重复")
        books_changed = primary_books is not None and primary_book_selection_changed(
            profile,
            first_book,
            first_scope,
            second_book,
            second_scope,
        )
        backup_book = canonical_profile_book_id(
            payload.get("backupBookId") or payload.get("backupBook") or profile.get("backup_book") or "",
            math_type,
        )
        conn.execute(
            """
            UPDATE student_profiles
            SET level = ?, class_id = ?, math_type = ?, target_score = ?, paid_until = ?,
                primary_book_1 = ?, primary_book_2 = ?, backup_book = ?,
                primary_book_1_scope = ?, primary_book_2_scope = ?,
                review_order = ?, plan_variant_id = ?, updated_at = ?
            WHERE user_id = ?
            """,
            (
                level,
                str(payload.get("classId") or profile.get("class_id") or "").strip(),
                math_type,
                str(payload.get("targetScore") or profile.get("target_score") or "").strip(),
                str(payload.get("paidUntil") or profile.get("paid_until") or "").strip(),
                first_book,
                second_book,
                backup_book,
                first_scope,
                second_scope,
                review_order,
                plan_variant_id,
                updated_at,
                user_id,
            ),
        )
        updated_profile = row_to_dict(conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone())
        progress = row_to_dict(conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user_id,)).fetchone()) or {}
        runtime_date = get_runtime_test_date(conn)
        if "paidUntil" in payload:
            paid_until_day = iso_date_or_none(updated_profile.get("paid_until"))
            runtime_day = iso_date_or_none(runtime_date)
            stored_status = str(user.get("account_status") or "")
            if (
                paid_until_day
                and runtime_day
                and paid_until_day >= runtime_day
                and stored_status in {"deleted", "deactivated", "expired_pending_review"}
            ):
                conn.execute(
                    "UPDATE users SET account_status = 'active', updated_at = ? WHERE user_id = ?",
                    (updated_at, user_id),
                )

        if any(key in payload for key in ("dailyMathGoalHours", "dailyStudyHours", "dailyTargetMinutes")):
            if "dailyTargetMinutes" in payload:
                target_minutes = as_int(payload.get("dailyTargetMinutes"), 0)
                if target_minutes < 120 or target_minutes > 480:
                    raise ValueError("dailyTargetMinutes must be between 120 and 480")
            else:
                raw_hours = as_float(payload.get("dailyMathGoalHours", payload.get("dailyStudyHours")), -1)
                if raw_hours < 2 or raw_hours > 8:
                    raise ValueError("dailyMathGoalHours must be between 2 and 8")
                target_minutes = round(raw_hours * 60)
            upsert_student_daily_target_minutes(
                conn,
                user_id,
                runtime_date,
                updated_profile,
                target_minutes,
                updated_at,
            )

        if "nextPlanStartChapterId" in payload:
            next_chapter_id = str(payload.get("nextPlanStartChapterId") or "").strip()
            if next_chapter_id:
                plan_stage = progress.get("plan_stage") or updated_profile.get("current_stage") or "强化"
                first_type_in_chapter = first_actual_assignable_type_in_chapter(
                    conn,
                    updated_profile,
                    plan_stage,
                    plan_variant_id,
                    next_chapter_id,
                )
                if not first_type_in_chapter:
                    raise ValueError("该章节在当前学生选书下没有可布置题目")
                effective_date = add_days_text(runtime_date, 1)
                conn.execute(
                    """
                    UPDATE student_progress
                    SET next_plan_start_chapter_id = ?,
                        next_plan_start_effective_date = ?,
                        next_plan_start_status = 'pending',
                        next_plan_start_updated_by = ?,
                        next_plan_start_updated_at = ?,
                        next_plan_start_applied_at = '',
                        updated_at = ?
                    WHERE user_id = ?
                    """,
                    (
                        next_chapter_id,
                        effective_date,
                        actor,
                        updated_at,
                        updated_at,
                        user_id,
                    ),
                )
            else:
                conn.execute(
                    """
                    UPDATE student_progress
                    SET next_plan_start_chapter_id = '',
                        next_plan_start_effective_date = '',
                        next_plan_start_status = '',
                        next_plan_start_updated_by = ?,
                        next_plan_start_updated_at = ?,
                        next_plan_start_applied_at = '',
                        updated_at = ?
                    WHERE user_id = ?
                    """,
                    (actor, updated_at, updated_at, user_id),
                )

        progress = row_to_dict(conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user_id,)).fetchone()) or {}
        if books_changed:
            regenerate_today_unsubmitted_plan_after_book_change(
                conn,
                user,
                updated_profile,
                progress,
                runtime_date,
                updated_at,
            )

        manual_next = str(payload.get("manualNextTypeId") or "").strip()
        if manual_next and "nextPlanStartChapterId" not in payload:
            conn.execute(
                """
                UPDATE student_progress
                    SET current_type_id = ?, current_type_sort_code = COALESCE((SELECT sort_code FROM question_types WHERE type_id = ?), current_type_sort_code),
                    updated_at = ?
                WHERE user_id = ?
                """,
                (manual_next, manual_next, updated_at, user_id),
            )
        updated_user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        updated_profile = row_to_dict(conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone())
        serialized = serialize_student(conn, updated_user, updated_profile, stat_date=runtime_date)
        log_admin_action(conn, "student.update", user_id, before, serialized, actor)
        conn.commit()
    return {"student": serialized}


def admin_wrong_question_export_rows(conn, user_id: str) -> list[dict[str, Any]]:
    student = row_to_dict(
        conn.execute(
            "SELECT user_id FROM users WHERE user_id = ? AND role = 'student'",
            (user_id,),
        ).fetchone()
    )
    if not student:
        raise ValueError("student not found")

    question_type_columns = table_columns(conn, "question_types")
    legacy_chapter_name = "qt.chapter_name" if "chapter_name" in question_type_columns else "''"
    if table_exists(conn, "chapters"):
        chapter_name_select = (
            f"COALESCE(NULLIF(c.chapter_name, ''), NULLIF({legacy_chapter_name}, ''), "
            "NULLIF(qt.chapter_id, ''), '') AS chapter_name"
        )
        chapter_join = "LEFT JOIN chapters c ON c.chapter_id = qt.chapter_id"
    else:
        chapter_name_select = (
            f"COALESCE(NULLIF({legacy_chapter_name}, ''), NULLIF(qt.chapter_id, ''), '') "
            "AS chapter_name"
        )
        chapter_join = ""

    return rows_to_dicts(
        conn.execute(
            f"""
            SELECT
              u.login_account,
              u.nickname,
              p.class_id,
              f.plan_date,
              {chapter_name_select},
              COALESCE(NULLIF(qt.type_name, ''), f.type_id, '') AS type_name,
              q.book_display_label,
              q.source_chapter_label,
              q.source_section_label,
              q.source_subsection_label,
              q.source_question_label,
              q.original_position_text,
              COALESCE(NULLIF(f.wrong_reason_name, ''), f.wrong_reason_code, '') AS wrong_reason,
              f.submitted_at,
              f.feedback_id
            FROM question_feedbacks f
            JOIN users u ON u.user_id = f.user_id
            JOIN student_profiles p ON p.user_id = f.user_id
            LEFT JOIN questions q ON q.question_id = f.question_id
            LEFT JOIN question_types qt ON qt.type_id = f.type_id
            {chapter_join}
            WHERE f.user_id = ?
              AND u.role = 'student'
              AND COALESCE(f.result, '') != 'correct'
              AND CAST(COALESCE(f.is_unfinished_blocker, '0') AS TEXT) != '1'
              AND (
                f.result = 'wrong'
                OR (f.result = 'unattempted' AND COALESCE(f.unattempted_action, '') = 'mark_wrong')
                OR CAST(COALESCE(f.is_wrong_record, '0') AS TEXT) = '1'
              )
              AND NOT (
                f.result = 'unattempted'
                AND COALESCE(f.unattempted_action, '') = 'defer'
              )
            ORDER BY f.plan_date, f.submitted_at, f.feedback_id
            """,
            (user_id,),
        ).fetchall()
    )


def build_admin_wrong_question_export(user_id: str) -> tuple[bytes, str]:
    headers = [
        "学生账号",
        "学生姓名",
        "班级",
        "做题日期",
        "章节",
        "题型",
        "资料书",
        "资料章节",
        "资料小节",
        "题号",
        "原始题目位置",
        "错因",
        "提交时间",
    ]
    with get_connection() as conn:
        rows = admin_wrong_question_export_rows(conn, user_id)

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("错题清单")
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [
                row.get("login_account") or "",
                row.get("nickname") or "",
                row.get("class_id") or "",
                row.get("plan_date") or "",
                row.get("chapter_name") or "",
                row.get("type_name") or "",
                row.get("book_display_label") or "",
                row.get("source_chapter_label") or "",
                row.get("source_section_label") or row.get("source_subsection_label") or "",
                row.get("source_question_label") or "",
                row.get("original_position_text") or "",
                row.get("wrong_reason") or "",
                row.get("submitted_at") or "",
            ]
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), f"wrong_questions_{user_id}.xlsx"


def renew_student_service(user_id: str, days: int = 30, actor: str = "adm001") -> dict[str, Any]:
    days = clamp(as_int(days, 30), 1, 365)
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? AND role = 'student'", (user_id,)).fetchone())
        profile = row_to_dict(conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone())
        if not user or not profile:
            raise ValueError("student not found")
        app_today = get_runtime_test_date(conn)
        before = {"paidUntil": profile.get("paid_until")}
        base = profile.get("paid_until") if profile.get("paid_until") and profile.get("paid_until") >= app_today else app_today
        paid_until = add_days_text(base, days)
        conn.execute(
            "UPDATE student_profiles SET paid_until = ?, updated_at = ? WHERE user_id = ?",
            (paid_until, now_text(), user_id),
        )
        conn.execute("UPDATE users SET account_status = 'active', updated_at = ? WHERE user_id = ?", (now_text(), user_id))
        updated_user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        updated_profile = row_to_dict(conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone())
        serialized = serialize_student(conn, updated_user, updated_profile, stat_date=app_today)
        log_admin_action(conn, "student.renew_service", user_id, before, {"paidUntil": paid_until, "days": days}, actor)
        conn.commit()
    return {"student": serialized}


def list_expired_experience_accounts(stat_date: str | None = None) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        stat_date = stat_date or get_runtime_test_date(conn)
        rows = rows_to_dicts(
            conn.execute(
                """
                SELECT u.*
                FROM users u
                JOIN student_profiles p ON p.user_id = u.user_id
                WHERE u.role = 'student' AND p.paid_until < ?
                ORDER BY p.paid_until, u.user_id
                """,
                (stat_date,),
            ).fetchall()
        )
        items = [
            serialize_student(conn, row, stat_date=stat_date)
            for row in rows
            if account_type_for_user(row) == "trial"
        ]
    return {"items": items}


def short_term_status(user: dict[str, Any], profile: dict[str, Any] | None, stat_date: str) -> str:
    if account_type_for_user(user) != "trial":
        return ""
    account_status = str(user.get("account_status") or "")
    if account_status == "converted_to_formal":
        return "migrated_to_formal"
    if account_status in {"disabled", "deactivated", "deleted"}:
        return "deactivated"
    if not profile:
        return "pending_onboarding"
    paid_until = str(profile.get("paid_until") or "")
    if not paid_until:
        return "pending_onboarding"
    paid_until_date = iso_date_or_none(paid_until)
    if not paid_until_date:
        return "expired_pending_review"
    stat_day = iso_date_or_none(stat_date) or date.fromisoformat(CURRENT_DATE)
    days_left = (paid_until_date - stat_day).days
    if days_left < 0:
        return "expired_pending_review"
    if days_left <= 2:
        return "expiring_soon"
    return "active_trial"


def serialize_short_term_account(
    conn,
    user: dict[str, Any],
    profile: dict[str, Any] | None,
    stat_date: str,
) -> dict[str, Any]:
    base = serialize_student(conn, user, profile, stat_date=stat_date) if profile else serialize_user_account(user)
    status = short_term_status(user, profile, stat_date)
    paid_until = str(profile.get("paid_until") or "") if profile else ""
    days_remaining = ""
    paid_until_date = iso_date_or_none(paid_until)
    if paid_until_date:
        stat_day = iso_date_or_none(stat_date) or date.fromisoformat(CURRENT_DATE)
        days_remaining = str((paid_until_date - stat_day).days)
    base.update(
        {
            "accountType": "trial",
            "shortTermStatus": status,
            "trialStartedAt": str(profile.get("created_at") or user.get("activated_at") or "") if profile else "",
            "trialExpiresAt": paid_until,
            "paidUntil": paid_until,
            "daysRemaining": days_remaining,
            "startsAfterOnboarding": True,
        }
    )
    if not profile:
        base.update(
            {
                "classId": "",
                "className": "",
                "targetScore": 0,
                "todayStatus": "pending_onboarding",
                "primaryBooks": [],
                "backupBook": "",
                "backupBooks": [],
                "mathType": "",
                "level": "",
                "reviewOrder": "mixed",
                "paidUntil": "",
                "contactQQ": user.get("QQ") or "",
                "planVariantId": "",
                "manualNextTypeId": "",
                "currentTypeId": "",
                "openTypeName": "",
            }
        )
    return base


def short_term_account_sort_key(user: dict[str, Any]) -> tuple[int, str]:
    account = str(user.get("login_account") or user.get("user_id") or "")
    number = formal_account_number(account)
    return (number if number else 10**12, account)


def list_short_term_accounts(stat_date: str | None = None) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        stat_date = stat_date or get_runtime_test_date(conn)
        users = sorted(
            (
                row
                for row in rows_to_dicts(
                    conn.execute(
                        """
                        SELECT *
                        FROM users
                        WHERE role = 'student'
                        """
                    ).fetchall()
                )
                if account_type_for_user(row) == "trial"
            ),
            key=short_term_account_sort_key,
        )
        items: list[dict[str, Any]] = []
        for user in users:
            profile = row_to_dict(
                conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user["user_id"],)).fetchone()
            )
            items.append(serialize_short_term_account(conn, user, profile, stat_date))
        summary = {
            "pendingOnboardingCount": sum(1 for item in items if item["shortTermStatus"] == "pending_onboarding"),
            "activeTrialCount": sum(1 for item in items if item["shortTermStatus"] == "active_trial"),
            "expiringSoonCount": sum(1 for item in items if item["shortTermStatus"] == "expiring_soon"),
            "expiredPendingReviewCount": sum(
                1 for item in items if item["shortTermStatus"] == "expired_pending_review"
            ),
            "migratedOrDeactivatedCount": sum(
                1 for item in items if item["shortTermStatus"] in {"migrated_to_formal", "deactivated"}
            ),
            "nextShortTermAccount": next_short_term_account(conn),
            "nextConversionAccount": next_trial_conversion_formal_account(conn),
        }
    return {"summary": summary, "items": items}


def get_short_term_account_detail(user_id: str, stat_date: str | None = None) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        stat_date = stat_date or get_runtime_test_date(conn)
        user = row_to_dict(
            conn.execute("SELECT * FROM users WHERE user_id = ? OR login_account = ?", (user_id, user_id)).fetchone()
        )
        if not user or account_type_for_user(user) != "trial":
            raise ValueError("short-term student not found")
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user["user_id"],)).fetchone()
        )
        student = serialize_short_term_account(conn, user, profile, stat_date)
        formal_detail = serialize_student(conn, user, profile, stat_date=stat_date) if profile else student
        return {
            "student": student,
            "formalDetail": formal_detail,
            "conversionPreview": {
                "loginAccount": next_trial_conversion_formal_account(conn),
                "pool": "stu40001-stu59999",
            },
        }


def deactivate_short_term_account(user_id: str, actor: str = "adm001") -> dict[str, Any]:
    updated_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        stat_date = get_runtime_test_date(conn)
        user = row_to_dict(
            conn.execute("SELECT * FROM users WHERE user_id = ? OR login_account = ?", (user_id, user_id)).fetchone()
        )
        if not user or account_type_for_user(user) != "trial":
            raise ValueError("short-term student not found")
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user["user_id"],)).fetchone()
        )
        before = serialize_short_term_account(conn, user, profile, stat_date)
        conn.execute(
            "UPDATE users SET account_status = 'deactivated', updated_at = ? WHERE user_id = ?",
            (updated_at, user["user_id"]),
        )
        updated = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user["user_id"],)).fetchone())
        after = serialize_short_term_account(conn, updated, profile, stat_date)
        log_admin_action(conn, "short_term.deactivate", user["user_id"], before, after, actor)
        conn.commit()
    return {"student": after}


def cloned_identifier_value(
    value: Any,
    source_user_id: str,
    source_login_account: str,
    target_user_id: str,
) -> Any:
    if value is None:
        return value
    text = str(value)
    for source_value in sorted({source_user_id, source_login_account} - {""}, key=len, reverse=True):
        if source_value in text:
            return text.replace(source_value, target_user_id)
    return f"{target_user_id}-{text}" if text else text


def clone_user_rows(
    conn,
    table_name: str,
    source_user_id: str,
    target_user_id: str,
    id_columns: tuple[str, ...] = (),
    source_login_account: str = "",
) -> None:
    if not table_exists(conn, table_name):
        return
    columns = table_columns(conn, table_name)
    if "user_id" not in columns:
        return
    rows = rows_to_dicts(conn.execute(f"SELECT * FROM {table_name} WHERE user_id = ?", (source_user_id,)).fetchall())
    for row in rows:
        values = {column: row.get(column, "") for column in columns}
        values["user_id"] = target_user_id
        for id_column in id_columns:
            if id_column in values and values[id_column]:
                values[id_column] = cloned_identifier_value(
                    values[id_column],
                    source_user_id,
                    source_login_account,
                    target_user_id,
                )
        insert_table_row(conn, table_name, values)


def clone_student_behavior_rows(conn, source_user: dict[str, Any], target_user_id: str) -> None:
    source_user_id = str(source_user.get("user_id") or "")
    source_login_account = str(source_user.get("login_account") or "")
    for table_name, id_columns in (
        ("daily_plans", ("plan_id",)),
        ("question_feedbacks", ("feedback_id", "plan_id")),
        ("student_study_goals", ("goal_id",)),
        ("student_daily_activity", ("activity_id",)),
        ("student_stats", ("stat_id",)),
        ("loss_reason_stats", ("loss_reason_id", "stat_id")),
        ("type_mastery_stats", ("mastery_id", "stat_id")),
        ("chapter_completion_records", ("record_id", "completion_id")),
        ("student_chapter_completions", ("completion_id",)),
    ):
        clone_user_rows(
            conn,
            table_name,
            source_user_id,
            target_user_id,
            id_columns,
            source_login_account=source_login_account,
        )


def migrate_experience_to_formal(user_id: str, actor: str = "adm001") -> dict[str, Any]:
    updated_at = now_text()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        ensure_formal_activation_tables(conn)
        source_user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? OR login_account = ?", (user_id, user_id)).fetchone())
        source_profile = row_to_dict(conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (source_user.get("user_id") if source_user else "",)).fetchone())
        if not source_user or not source_profile or account_type_for_user(source_user) != "trial":
            raise ValueError("source trial student not found")
        app_today = get_runtime_test_date(conn)
        login_account = next_trial_conversion_formal_account(conn)
        if row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? OR login_account = ?", (login_account, login_account)).fetchone()):
            raise ValueError("formal account already exists")
        paid_until_base = source_profile.get("paid_until") if source_profile.get("paid_until") and source_profile.get("paid_until") >= app_today else app_today
        paid_until = add_days_text(paid_until_base, 30)
        conn.execute(
            """
            INSERT INTO users (
              user_id, login_account, password_hash, login_password, role, nickname,
              avatar_url, QQ, account_status, current_session_id, last_login_at,
              created_at, updated_at, account_type, activated_at
            ) VALUES (?, ?, ?, ?, 'student', ?, ?, ?, 'active', '', '', ?, ?, 'formal', ?)
            """,
            (
                login_account,
                login_account,
                source_user.get("password_hash") or sha256_password(source_user.get("login_password") or "123456"),
                source_user.get("login_password") or "",
                source_user.get("nickname") or login_account,
                source_user.get("avatar_url") or "",
                source_user.get("QQ") or "",
                updated_at,
                updated_at,
                updated_at,
            ),
        )
        profile_values = dict(source_profile)
        profile_values["user_id"] = login_account
        profile_values["paid_until"] = paid_until
        profile_values["updated_at"] = updated_at
        columns = [
            row["name"]
            for row in conn.execute("PRAGMA table_info(student_profiles)").fetchall()
            if row["name"] in profile_values
        ]
        conn.execute(
            f"INSERT INTO student_profiles ({', '.join(columns)}) VALUES ({', '.join('?' for _ in columns)})",
            [profile_values[column] for column in columns],
        )
        progress = row_to_dict(conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (source_user["user_id"],)).fetchone())
        if progress:
            progress_values = dict(progress)
            progress_values["user_id"] = login_account
            progress_values["progress_id"] = f"PR-{login_account}"
            progress_values["updated_at"] = updated_at
            progress_columns = [
                row["name"]
                for row in conn.execute("PRAGMA table_info(student_progress)").fetchall()
                if row["name"] in progress_values
            ]
            conn.execute(
                f"INSERT INTO student_progress ({', '.join(progress_columns)}) VALUES ({', '.join('?' for _ in progress_columns)})",
                [progress_values[column] for column in progress_columns],
            )
        clone_student_behavior_rows(conn, source_user, login_account)
        conn.execute(
            "UPDATE users SET account_status = 'converted_to_formal', updated_at = ? WHERE user_id = ?",
            (updated_at, source_user["user_id"]),
        )
        updated_source = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (source_user["user_id"],)).fetchone())
        target_user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (login_account,)).fetchone())
        source_serialized = serialize_student(conn, updated_source, source_profile)
        target_serialized = serialize_student(conn, target_user, profile_values)
        log_admin_action(conn, "trial.migrate_to_formal", source_user["user_id"], source_serialized, target_serialized, actor)
        conn.commit()
    return {"sourceStudent": source_serialized, "targetStudent": target_serialized}


def list_class_intake_settings() -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        rows = rows_to_dicts(conn.execute("SELECT * FROM class_intake_settings ORDER BY math_type, level").fetchall())
        items = [serialize_intake_setting(conn, row) for row in rows]
    return {"items": items}


def update_class_intake_setting(math_type: str, level: str, payload: dict[str, Any], actor: str = "adm001") -> dict[str, Any]:
    active_class_id = str(payload.get("activeClassId") or "").strip()
    updated_by = str(payload.get("updatedBy") or actor or "adm001").strip()
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        class_row = row_to_dict(conn.execute("SELECT * FROM classes WHERE class_id = ?", (active_class_id,)).fetchone())
        if not class_row:
            raise ValueError("active class not found")
        if class_row.get("math_type") != math_type or class_row.get("level") != level:
            raise ValueError("active class does not match math type and level")
        before = row_to_dict(
            conn.execute(
                "SELECT * FROM class_intake_settings WHERE math_type = ? AND level = ?",
                (math_type, level),
            ).fetchone()
        )
        updated_at = now_text()
        conn.execute(
            """
            INSERT OR REPLACE INTO class_intake_settings (
              math_type, level, active_class_id, status, updated_by, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (math_type, level, active_class_id, class_accepting_label(class_row), updated_by, updated_at),
        )
        row = row_to_dict(
            conn.execute(
                "SELECT * FROM class_intake_settings WHERE math_type = ? AND level = ?",
                (math_type, level),
            ).fetchone()
        )
        serialized = serialize_intake_setting(conn, row)
        log_admin_action(conn, "intake.update", f"{math_type}-{level}", before or "", serialized, actor)
        conn.commit()
    return {"setting": serialized}


def list_admin_audit_logs(limit: int = 100) -> dict[str, Any]:
    safe_limit = clamp(as_int(limit, 100), 1, 200)
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        rows = rows_to_dicts(
            conn.execute(
                "SELECT * FROM admin_audit_logs ORDER BY created_at DESC, log_id DESC LIMIT ?",
                (safe_limit,),
            ).fetchall()
        )
    return {"items": [serialize_audit_log(row) for row in rows]}


def teacher_user_or_default(conn, teacher_user_id: str = "") -> dict[str, Any]:
    teacher = row_to_dict(
        conn.execute(
            "SELECT * FROM users WHERE role = 'teacher' AND (user_id = ? OR login_account = ?)",
            (teacher_user_id, teacher_user_id),
        ).fetchone()
    )
    if teacher:
        return teacher
    return row_to_dict(conn.execute("SELECT * FROM users WHERE role = 'teacher' ORDER BY user_id LIMIT 1").fetchone())


def list_teacher_classes(teacher_user_id: str = "") -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        teacher = teacher_user_or_default(conn, teacher_user_id)
        if not teacher:
            raise ValueError("teacher not found")
        rows = rows_to_dicts(
            conn.execute("SELECT * FROM classes WHERE teacher_user_id = ? ORDER BY class_id", (teacher["user_id"],)).fetchall()
        )
        items = [serialize_class(conn, row) for row in rows]
        serialized_teacher = serialize_teacher(conn, teacher)
    return {"teacher": serialized_teacher, "items": items}


def teacher_class_feedback_status(class_id: str, teacher_user_id: str = "", stat_date: str = CURRENT_DATE) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        teacher = teacher_user_or_default(conn, teacher_user_id)
        class_row = row_to_dict(conn.execute("SELECT * FROM classes WHERE class_id = ?", (class_id,)).fetchone())
        if not teacher or not class_row:
            raise ValueError("teacher class not found")
        if class_row.get("teacher_user_id") != teacher["user_id"]:
            raise ValueError("teacher cannot access this class")
        users = rows_to_dicts(
            conn.execute(
                """
                SELECT u.*
                FROM users u
                JOIN student_profiles p ON p.user_id = u.user_id
                WHERE p.class_id = ? AND u.role = 'student'
                ORDER BY u.user_id
                """,
                (class_id,),
            ).fetchall()
        )
        students = [serialize_student(conn, user, stat_date=stat_date) for user in users]
        serialized_class = serialize_class(conn, class_row)
    groups = {
        "unsubmitted": [row for row in students if row["todayStatus"] == "unsubmitted"],
        "unfinished": [row for row in students if row["todayStatus"] == "unfinished"],
        "completed": [row for row in students if row["todayStatus"] == "completed"],
        "abnormal": [row for row in students if row["todayStatus"] == "abnormal"],
    }
    return {"class": serialized_class, "date": stat_date, "groups": groups}


def teacher_student_daily_status(student_id: str, teacher_user_id: str = "", stat_date: str = CURRENT_DATE) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_admin_console_tables(conn)
        teacher = teacher_user_or_default(conn, teacher_user_id)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ? OR login_account = ?", (student_id, student_id)).fetchone())
        profile = row_to_dict(conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user.get("user_id") if user else "",)).fetchone())
        if not teacher or not user or not profile:
            raise ValueError("student not found")
        class_row = row_to_dict(conn.execute("SELECT * FROM classes WHERE class_id = ?", (profile.get("class_id") or "",)).fetchone())
        if not class_row or class_row.get("teacher_user_id") != teacher["user_id"]:
            raise ValueError("teacher cannot access this student")
        student = serialize_student(conn, user, profile, stat_date)
    return {"student": student}


def completion_message() -> str:
    return get_system_message(
        "daily_plan_completed",
        "恭喜你完成今日全部题型！今天的坚持很漂亮，继续保持。",
    )


def parse_json_id_list(value: Any) -> list[str]:
    try:
        parsed = json.loads(value or "[]")
    except (TypeError, json.JSONDecodeError):
        return []
    if not isinstance(parsed, list):
        return []
    return [str(item) for item in parsed if str(item)]


def split_plan_questions(
    all_questions: list[dict[str, Any]],
    available_questions: list[dict[str, Any]],
    suggested_days: float,
) -> tuple[list[dict[str, Any]], list[str], float]:
    if not available_questions:
        return [], [], 0.0

    total_count = max(len(all_questions), 1)
    unit_load = suggested_days * len(available_questions) / total_count
    if unit_load <= DAILY_PLAN_SPLIT_THRESHOLD:
        return available_questions, [], unit_load

    today_count = max(1, math.ceil(len(available_questions) / unit_load))
    today_count = min(today_count, len(available_questions))
    today_questions = available_questions[:today_count]
    carryover_question_ids = [
        question["question_id"]
        for question in available_questions[today_count:]
    ]
    return today_questions, carryover_question_ids, DAILY_PLAN_TARGET_LOAD


def split_plan_questions_to_load(
    all_questions: list[dict[str, Any]],
    available_questions: list[dict[str, Any]],
    suggested_days: float,
    target_load: float,
) -> tuple[list[dict[str, Any]], list[str], float]:
    if not available_questions or target_load <= 0:
        return [], [], 0.0

    total_count = max(len(all_questions), 1)
    available_load = suggested_days * len(available_questions) / total_count
    if available_load <= target_load:
        return available_questions, [], available_load

    exact_count = len(available_questions) * target_load / available_load
    candidate_counts = {
        max(1, min(len(available_questions) - 1, math.floor(exact_count))),
        max(1, min(len(available_questions) - 1, math.ceil(exact_count))),
    }
    today_count = min(
        candidate_counts,
        key=lambda count: abs((suggested_days * count / total_count) - target_load),
    )
    today_questions = available_questions[:today_count]
    carryover_question_ids = [
        question["question_id"]
        for question in available_questions[today_count:]
    ]
    unit_load = suggested_days * len(today_questions) / total_count
    return today_questions, carryover_question_ids, unit_load


def generate_daily_plan(
    conn,
    user: dict[str, Any],
    profile: dict[str, Any],
    progress: dict[str, Any],
    plan_date: str,
    type_index_offset: int = 0,
    commit: bool = True,
) -> list[dict[str, Any]]:
    ensure_daily_plan_sequence_columns(conn)
    if str(profile.get("plan_enabled") or "") != "1":
        raise ValueError("该学生今日计划已暂停")
    if str(progress.get("progress_status") or "active") != "active":
        raise ValueError("该学生学习进度已暂停")

    target_load = DAILY_PLAN_TARGET_LOAD
    plan_stage = progress.get("plan_stage") or profile.get("current_stage") or "强化"
    plan_variant_id = str(profile.get("plan_variant_id") or "").strip()
    current_type_status = progress.get("current_type_status") or (
        "unfinished" if progress.get("unfinished_type_id") else "active"
    )
    current_type_id = (
        progress.get("unfinished_type_id")
        if current_type_status == "unfinished" and progress.get("unfinished_type_id")
        else progress.get("carryover_type_id")
        if current_type_status == "carryover" and progress.get("carryover_type_id")
        else progress.get("current_type_id")
    )
    actual_sequence = actual_assignable_type_sequence(conn, profile, plan_stage, plan_variant_id)
    pending_chapter_id = str(progress.get("next_plan_start_chapter_id") or "").strip()
    pending_status = str(progress.get("next_plan_start_status") or "").strip()
    pending_effective_date = str(progress.get("next_plan_start_effective_date") or "").strip()
    pending_type = None
    if (
        pending_chapter_id
        and pending_status == "pending"
        and (not pending_effective_date or plan_date >= pending_effective_date)
    ):
        pending_type = next(
            (
                row
                for row in actual_sequence
                if str(row.get("chapter_id") or "").strip() == pending_chapter_id
            ),
            None,
        )
        if pending_type:
            applied_at = now_text()
            current_type_id = pending_type["type_id"]
            current_type_status = "active"
            progress = {
                **progress,
                "current_type_id": current_type_id,
                "current_type_status": "active",
                "unfinished_type_id": "",
                "unfinished_question_ids": "[]",
                "carryover_type_id": "",
                "carryover_question_ids": "[]",
                "next_plan_start_status": "applied",
                "next_plan_start_applied_at": applied_at,
            }
            conn.execute(
                """
                UPDATE student_progress
                SET current_type_id = ?, current_type_sort_code = ?,
                    current_type_status = 'active',
                    unfinished_type_id = '', unfinished_question_ids = '[]',
                    carryover_type_id = '', carryover_question_ids = '[]',
                    next_plan_start_status = 'applied',
                    next_plan_start_applied_at = ?,
                    updated_at = ?
                WHERE user_id = ?
                """,
                (
                    pending_type.get("type_id") or "",
                    pending_type.get("sort_code") or "",
                    applied_at,
                    applied_at,
                    user["user_id"],
                ),
            )
            set_mastery_scope_start(conn, user["user_id"], pending_type, applied_at, "admin_chapter_start")

    current_type = pending_type or question_type_for_plan_variant(conn, current_type_id, plan_variant_id, profile.get("level")) or next_assignable_type(
        conn,
        "",
        plan_variant_id,
        profile.get("level"),
    )
    if not current_type and actual_sequence:
        current_type = actual_sequence[0]
    start_index = sequence_start_index_for_type(actual_sequence, current_type)
    if not current_type:
        raise ValueError("题型表中没有可布置题型")

    if start_index >= len(actual_sequence):
        raise ValueError("未找到适合该学生的题目")

    unfinished_ids: set[str] = set()
    if current_type_status == "unfinished" and progress.get("unfinished_type_id") == current_type["type_id"]:
        unfinished_ids = set(parse_json_id_list(progress.get("unfinished_question_ids")))

    carryover_ids: set[str] = set()
    if current_type_status == "carryover" and progress.get("carryover_type_id") == current_type["type_id"]:
        carryover_ids = set(parse_json_id_list(progress.get("carryover_question_ids")))

    selected_units: list[dict[str, Any]] = []
    load = 0.0
    cursor_index = start_index
    first_unit = True
    next_carryover_type_id = ""
    next_carryover_question_ids: list[str] = []

    while cursor_index < len(actual_sequence):
        cursor = actual_sequence[cursor_index]
        all_questions = get_type_questions(conn, cursor["type_id"], profile, plan_stage)
        if not all_questions:
            cursor_index += 1
            first_unit = False
            continue

        suggested_days = as_float(cursor.get("suggested_days"), 0.3)
        questions = all_questions
        if first_unit and unfinished_ids:
            questions = get_type_questions(conn, cursor["type_id"], profile, plan_stage, unfinished_ids)
        elif first_unit and carryover_ids:
            questions = get_type_questions(conn, cursor["type_id"], profile, plan_stage, carryover_ids)

        available_questions_for_unit = questions
        questions, carryover_question_ids, unit_load = split_plan_questions(
            all_questions,
            available_questions_for_unit,
            suggested_days,
        )
        if not questions:
            cursor_index += 1
            first_unit = False
            continue

        if not selected_units:
            selected_units.append({"type": cursor, "questions": questions, "has_carryover": bool(carryover_question_ids)})
            load = unit_load
            next_carryover_type_id = cursor["type_id"] if carryover_question_ids else ""
            next_carryover_question_ids = carryover_question_ids
        else:
            new_load = load + unit_load
            if abs(new_load - target_load) < abs(load - target_load):
                selected_units.append({"type": cursor, "questions": questions, "has_carryover": bool(carryover_question_ids)})
                load = new_load
                if carryover_question_ids:
                    next_carryover_type_id = cursor["type_id"]
                    next_carryover_question_ids = carryover_question_ids
                    break
            else:
                remaining_load = target_load - load
                partial_questions, partial_carryover_ids, partial_unit_load = split_plan_questions_to_load(
                    all_questions,
                    available_questions_for_unit,
                    suggested_days,
                    remaining_load,
                )
                partial_new_load = load + partial_unit_load
                if (
                    partial_questions
                    and partial_carryover_ids
                    and abs(partial_new_load - target_load) < abs(load - target_load)
                ):
                    selected_units.append(
                        {
                            "type": cursor,
                            "questions": partial_questions,
                            "has_carryover": True,
                        }
                    )
                    load = partial_new_load
                    next_carryover_type_id = cursor["type_id"]
                    next_carryover_question_ids = partial_carryover_ids
                break

        if carryover_question_ids:
            break

        cursor_index += 1
        first_unit = False

    if not selected_units:
        raise ValueError("未找到适合该学生的题目")

    created_at = now_text()
    plan_id = f"P{plan_date.replace('-', '')}-{user['user_id']}"
    days_to_exam = days_between(plan_date, profile.get("target_exam_date") or plan_date)
    rows_to_insert = []
    for local_type_index, unit in enumerate(selected_units, start=1):
        type_index = type_index_offset + local_type_index
        type_row = unit["type"]
        feedback_status = "open" if local_type_index == 1 else "locked"
        boundary_after = "1" if str(type_row.get("is_chapter_boundary_after") or "") == "1" and not unit.get("has_carryover") else "0"
        for question_order, question in enumerate(unit["questions"], start=1):
            rows_to_insert.append(
                (
                    plan_id,
                    user["user_id"],
                    plan_date,
                    "in_progress",
                    "稳住节奏，今天也只需要完成眼前这一组。",
                    profile.get("target_exam_date") or "",
                    str(days_to_exam),
                    str(type_index),
                    type_row["type_id"],
                    type_row["type_name"],
                    question["question_id"],
                    question_source_label(question),
                    question["original_position_text"],
                    str(question_order),
                    feedback_status,
                    str(type_row.get("chapter_id") or ""),
                    str(type_row.get("student_type_sort_in_chapter") or "0"),
                    str(type_row.get("display_type_label") or ""),
                    boundary_after,
                    created_at,
                    created_at,
                )
            )

    conn.executemany(
        """
        INSERT INTO daily_plans (
          plan_id, user_id, plan_date, plan_status, slogan, exam_date,
          days_to_exam, type_index, type_id, type_name, question_id,
          source, original_position_text, question_order, feedback_open_status,
          chapter_id, student_type_sort_in_chapter, display_type_label,
          is_chapter_boundary_after,
          created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows_to_insert,
    )
    conn.execute(
        """
        UPDATE student_progress
        SET last_generated_plan_date = ?, current_type_status = ?,
            carryover_type_id = ?, carryover_question_ids = ?, updated_at = ?
        WHERE user_id = ?
        """,
        (
            plan_date,
            "carryover" if next_carryover_type_id else current_type_status,
            next_carryover_type_id,
            json.dumps(next_carryover_question_ids, ensure_ascii=False),
            created_at,
            user["user_id"],
        ),
    )
    if commit:
        conn.commit()

    return rows_to_dicts(
        conn.execute(
            """
            SELECT * FROM daily_plans
            WHERE user_id = ? AND plan_date = ?
            ORDER BY CAST(type_index AS INTEGER), CAST(question_order AS INTEGER)
            """,
            (user["user_id"], plan_date),
        ).fetchall()
    )


def table_columns(conn, table_name: str) -> list[str]:
    return [row["name"] for row in conn.execute(f"PRAGMA table_info({table_name})").fetchall()]


def insert_table_row(conn, table_name: str, values: dict[str, Any]) -> None:
    columns = [column for column in table_columns(conn, table_name) if column in values]
    placeholders = ", ".join("?" for _ in columns)
    column_clause = ", ".join(columns)
    conn.execute(
        f"INSERT INTO {table_name} ({column_clause}) VALUES ({placeholders})",
        [values[column] for column in columns],
    )


def copy_latest_user_date_rows(
    conn,
    table_name: str,
    user_id: str,
    date_column: str,
    target_date: str,
    id_column: str,
    id_prefix: str,
    overrides: dict[str, Any] | None = None,
) -> bool:
    if not table_exists(conn, table_name):
        return False

    existing_count = row_to_dict(
        conn.execute(
            f"SELECT COUNT(*) AS row_count FROM {table_name} WHERE user_id = ? AND {date_column} = ?",
            (user_id, target_date),
        ).fetchone()
    )
    if as_int(existing_count.get("row_count") if existing_count else 0) > 0:
        return True

    source = row_to_dict(
        conn.execute(
            f"""
            SELECT MAX({date_column}) AS source_date
            FROM {table_name}
            WHERE user_id = ? AND {date_column} < ?
            """,
            (user_id, target_date),
        ).fetchone()
    )
    source_date = source.get("source_date") if source else ""
    if not source_date:
        return False

    rows = rows_to_dicts(
        conn.execute(
            f"SELECT * FROM {table_name} WHERE user_id = ? AND {date_column} = ?",
            (user_id, source_date),
        ).fetchall()
    )
    if not rows:
        return False

    columns = table_columns(conn, table_name)
    updated_at = now_text()
    for index, row in enumerate(rows, start=1):
        values = {column: row.get(column, "") for column in columns}
        values[date_column] = target_date
        if id_column in values:
            values[id_column] = f"{id_prefix}-{user_id}-{target_date.replace('-', '')}-{index}"
        if "updated_at" in values:
            values["updated_at"] = updated_at
        if overrides:
            values.update(overrides)
        insert_table_row(conn, table_name, values)
    return True


def stats_row_has_completed_checkin(row: dict[str, Any] | None) -> bool:
    if not row:
        return False
    return bool(str(row.get("today_checkin_at") or "").strip()) or as_float(row.get("today_completion_rate")) >= 1


def plan_completion_has_completed(row: dict[str, Any] | None) -> bool:
    if not row:
        return False
    total_rows = as_int(row.get("plan_row_count"))
    submitted_rows = as_int(row.get("submitted_plan_row_count"))
    return total_rows > 0 and (submitted_rows >= total_rows or as_int(row.get("has_completed_plan_status")) > 0)


def plan_completion_rate(row: dict[str, Any] | None) -> float:
    if not row:
        return 0
    total_rows = as_int(row.get("plan_row_count"))
    if total_rows <= 0:
        return 0
    return min(1, as_int(row.get("submitted_plan_row_count")) / total_rows)


def daily_plan_completion_rows_by_date(conn, user_id: str, stat_date: str) -> dict[str, dict[str, Any]]:
    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT
              plan_date,
              COUNT(*) AS plan_row_count,
              SUM(CASE WHEN feedback_open_status = 'submitted' OR plan_status = 'completed' THEN 1 ELSE 0 END) AS submitted_plan_row_count,
              MAX(CASE WHEN plan_status = 'completed' THEN 1 ELSE 0 END) AS has_completed_plan_status,
              MAX(CASE WHEN feedback_open_status = 'submitted' OR plan_status = 'completed' THEN updated_at ELSE '' END) AS plan_completed_at
            FROM daily_plans
            WHERE user_id = ? AND plan_date <= ?
            GROUP BY plan_date
            """,
            (user_id, stat_date),
        ).fetchall()
    )
    return {str(row.get("plan_date") or ""): row for row in rows}


def calculate_continuous_checkin_days(conn, user_id: str, stat_date: str) -> int:
    try:
        target_date = date.fromisoformat(stat_date)
    except ValueError:
        return 0

    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT stat_date, today_completion_rate, today_checkin_at
            FROM student_stats
            WHERE user_id = ? AND stat_date <= ?
            ORDER BY stat_date DESC
            """,
            (user_id, stat_date),
        ).fetchall()
    )
    rows_by_date = {str(row.get("stat_date") or ""): row for row in rows}
    plan_rows_by_date = daily_plan_completion_rows_by_date(conn, user_id, stat_date)

    def has_completed_checkin(date_text: str) -> bool:
        return stats_row_has_completed_checkin(rows_by_date.get(date_text)) or plan_completion_has_completed(
            plan_rows_by_date.get(date_text)
        )

    cursor_date = target_date
    if not has_completed_checkin(cursor_date.isoformat()):
        cursor_date -= timedelta(days=1)

    continuous_days = 0
    while True:
        if not has_completed_checkin(cursor_date.isoformat()):
            break
        continuous_days += 1
        cursor_date -= timedelta(days=1)
    return continuous_days


def sync_student_continuous_days(conn, user_id: str, stat_date: str) -> int:
    if "continuous_days" not in table_columns(conn, "student_stats"):
        return 0
    continuous_days = calculate_continuous_checkin_days(conn, user_id, stat_date)
    conn.execute(
        """
        UPDATE student_stats
        SET continuous_days = ?,
            badge_progress_value = ?,
            badge_progress_rate = ?,
            updated_at = ?
        WHERE user_id = ? AND stat_date = ?
        """,
        (
            str(continuous_days),
            str(continuous_days),
            f"{min(1, continuous_days / 7):.4f}",
            now_text(),
            user_id,
            stat_date,
        ),
    )
    return continuous_days


def ensure_student_stats_for_date(conn, user_id: str, plan_date: str) -> None:
    ensure_student_stats_ranking_columns(conn)
    copied = copy_latest_user_date_rows(
        conn,
        "student_stats",
        user_id,
        "stat_date",
        plan_date,
        "",
        "SS",
        {
            "today_completion_rate": "0.0000",
            "today_checkin_order": "0",
            "today_checkin_at": "",
            "updated_at": now_text(),
        },
    )
    if copied:
        sync_student_continuous_days(conn, user_id, plan_date)
        return

    values = {column: "" for column in table_columns(conn, "student_stats")}
    values.update(
        {
            "user_id": user_id,
            "stat_date": plan_date,
            "continuous_days": "0",
            "today_completion_rate": "0.0000",
            "today_checkin_order": "0",
            "today_checkin_at": "",
            "weekly_completion_rate": "0.0000",
            "weekly_study_days": "0",
            "overall_accuracy": "0.0000",
            "repeat_wrong_rate": "0.0000",
            "current_badge_name": "暂无勋章",
            "next_badge_target": "7天连续学习",
            "badge_progress_value": "0",
            "badge_progress_rate": "0.0000",
            "rank_in_class": "0",
            "updated_at": now_text(),
        }
    )
    insert_table_row(conn, "student_stats", values)
    sync_student_continuous_days(conn, user_id, plan_date)


def ensure_daily_activity_for_date(conn, user_id: str, plan_date: str) -> None:
    existing = row_to_dict(
        conn.execute(
            "SELECT activity_id FROM student_daily_activity WHERE user_id = ? AND activity_date = ?",
            (user_id, plan_date),
        ).fetchone()
    )
    if existing:
        return

    created_at = now_text()
    values = {column: "" for column in table_columns(conn, "student_daily_activity")}
    values.update(
        {
            "activity_id": f"ACT-{user_id}-{plan_date.replace('-', '')}",
            "user_id": user_id,
            "activity_date": plan_date,
            "is_studied": "0",
            "intensity_level": "0",
            "target_minutes": "180",
            "completed_minutes": "0",
            "completed_questions": "0",
            "submitted_type_count": "0",
            "checkin_status": "pending",
            "created_at": created_at,
            "updated_at": created_at,
        }
    )
    insert_table_row(conn, "student_daily_activity", values)


def study_minutes_intensity_level(completed_minutes: int) -> str:
    if completed_minutes >= 150:
        return "high"
    if completed_minutes >= 100:
        return "medium"
    if completed_minutes > 0:
        return "low"
    return "none"


def daily_activity_for_date(conn, user_id: str, plan_date: str) -> dict[str, Any]:
    ensure_daily_activity_for_date(conn, user_id, plan_date)
    return row_to_dict(
        conn.execute(
            """
            SELECT *
            FROM student_daily_activity
            WHERE user_id = ? AND activity_date = ?
            """,
            (user_id, plan_date),
        ).fetchone()
    )


def plan_day_timestamp(plan_date: str, timestamp: str) -> str:
    time_part = "23:59:59"
    text = str(timestamp or "")
    if " " in text:
        time_part = text.split(" ", 1)[1][:8] or time_part
    return f"{plan_date} {time_part}"


def mastery_status_tag(rate: float) -> str:
    if rate >= 0.75:
        return "良好"
    if rate >= 0.60:
        return "待提升"
    return "薄弱项"


def split_summary_text(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    for separator in ("\n", "、", ",", "，", ";", "；"):
        if separator in text:
            return [item.strip() for item in text.split(separator) if item.strip()]
    return [text]


def previous_loss_reason_rates(conn, user_id: str, stat_date: str) -> dict[str, float]:
    ensure_growth_snapshot_tables(conn)
    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT reason_code, reason_rate
            FROM loss_reason_stats
            WHERE user_id = ? AND stat_date = ?
            """,
            (user_id, stat_date),
        ).fetchall()
    )
    return {str(row.get("reason_code") or ""): as_float(row.get("reason_rate")) for row in rows}


def recalculate_daily_activity_from_feedback(
    conn,
    user_id: str,
    plan_date: str,
    updated_at: str,
) -> None:
    ensure_daily_activity_for_date(conn, user_id, plan_date)
    activity = daily_activity_for_date(conn, user_id, plan_date)
    completed_minutes = as_int(activity.get("completed_minutes"))
    counts = row_to_dict(
        conn.execute(
            """
            SELECT
              SUM(
                CASE
                  WHEN NOT (result = 'unattempted' AND COALESCE(unattempted_action, '') = 'defer')
                  THEN 1 ELSE 0
                END
              ) AS completed_questions,
              COUNT(
                DISTINCT CASE
                  WHEN NOT (result = 'unattempted' AND COALESCE(unattempted_action, '') = 'defer')
                  THEN type_id ELSE NULL
                END
              ) AS submitted_type_count
            FROM question_feedbacks
            WHERE user_id = ? AND plan_date = ?
            """,
            (user_id, plan_date),
        ).fetchone()
    )
    completed_questions = as_int(counts.get("completed_questions") if counts else 0)
    submitted_type_count = as_int(counts.get("submitted_type_count") if counts else 0)
    intensity_level = (
        study_minutes_intensity_level(completed_minutes)
        if completed_minutes > 0
        else "low"
        if completed_questions > 0
        else "none"
    )
    conn.execute(
        """
        UPDATE student_daily_activity
        SET completed_questions = ?,
            submitted_type_count = ?,
            is_studied = ?,
            intensity_level = ?,
            updated_at = ?
        WHERE user_id = ? AND activity_date = ?
        """,
        (
            str(completed_questions),
            str(submitted_type_count),
            "1" if completed_minutes > 0 or completed_questions > 0 else "0",
            intensity_level,
            updated_at,
            user_id,
            plan_date,
        ),
    )


def recalculate_type_mastery_stats(
    conn,
    user_id: str,
    stat_date: str,
    updated_at: str,
) -> None:
    ensure_growth_snapshot_tables(conn)
    profile = row_to_dict(
        conn.execute("SELECT plan_variant_id FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
    ) or {}
    plan_variant_id = str(profile.get("plan_variant_id") or "").strip()
    question_type_columns = table_columns(conn, "question_types")
    daily_plan_sort = "CAST(MIN(CAST(NULLIF(dp.type_index, '') AS INTEGER)) AS TEXT)"
    sort_join = ""
    sort_params: list[Any] = []
    if plan_variant_id and table_exists(conn, "question_type_orders"):
        sort_join = f"""
            LEFT JOIN question_type_orders qto
              ON qto.type_id = f.type_id
             AND qto.plan_variant_id = ?
             AND {_truthy_sql_filter("qto.is_active")}
        """
        sort_params.append(plan_variant_id)
        sort_code_expr = f"COALESCE(NULLIF(MIN(qto.sort_code), ''), {daily_plan_sort}, f.type_id)"
    elif "sort_code" in question_type_columns:
        sort_code_expr = f"COALESCE(NULLIF(MIN(qt.sort_code), ''), {daily_plan_sort}, f.type_id)"
    else:
        sort_code_expr = f"COALESCE({daily_plan_sort}, f.type_id)"
    rows = rows_to_dicts(
        conn.execute(
            f"""
            SELECT
              f.type_id,
              MAX(COALESCE(NULLIF(qt.type_name, ''), NULLIF(dp.type_name, ''), f.type_id)) AS resolved_type_name,
              {sort_code_expr} AS sort_code,
              SUM(
                CASE
                  WHEN NOT (f.result = 'unattempted' AND COALESCE(f.unattempted_action, '') = 'defer')
                  THEN 1 ELSE 0
                END
              ) AS submitted_question_count,
              SUM(
                CASE
                  WHEN CAST(COALESCE(f.is_wrong_record, '0') AS TEXT) = '1'
                    OR f.result = 'wrong'
                  THEN 1 ELSE 0
                END
              ) AS wrong_question_count
            FROM question_feedbacks f
            LEFT JOIN question_types qt ON qt.type_id = f.type_id
            LEFT JOIN daily_plans dp
              ON dp.user_id = f.user_id
             AND dp.plan_date = f.plan_date
             AND dp.type_id = f.type_id
             AND dp.question_id = f.question_id
            {sort_join}
            WHERE f.user_id = ? AND f.plan_date <= ?
            GROUP BY f.type_id
            ORDER BY sort_code, f.type_id
            """,
            [*sort_params, user_id, stat_date],
        ).fetchall()
    )
    conn.execute(
        "DELETE FROM type_mastery_stats WHERE user_id = ? AND stat_date = ?",
        (user_id, stat_date),
    )
    for index, row in enumerate(rows, start=1):
        submitted_count = as_int(row.get("submitted_question_count"))
        wrong_count = as_int(row.get("wrong_question_count"))
        mastery_rate = (submitted_count - wrong_count) / submitted_count if submitted_count else 0
        insert_table_row(
            conn,
            "type_mastery_stats",
            {
                "mastery_id": f"M-{user_id}-{stat_date.replace('-', '')}-{index}",
                "user_id": user_id,
                "stat_date": stat_date,
                "type_id": row.get("type_id") or "",
                "type_name": row.get("resolved_type_name") or "",
                "mastery_rate": f"{mastery_rate:.4f}",
                "status_tag": mastery_status_tag(mastery_rate),
                "submitted_question_count": str(submitted_count),
                "wrong_question_count": str(wrong_count),
                "sort_order": str(index),
                "updated_at": updated_at,
            },
        )


def recalculate_loss_reason_stats(
    conn,
    user_id: str,
    stat_date: str,
    previous_rates: dict[str, float],
    updated_at: str,
) -> None:
    ensure_growth_snapshot_tables(conn)
    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT
              wrong_reason_code AS reason_code,
              COALESCE(NULLIF(wrong_reason_name, ''), wrong_reason_code) AS reason_name,
              COUNT(*) AS wrong_question_count
            FROM question_feedbacks
            WHERE user_id = ?
              AND plan_date <= ?
              AND CAST(COALESCE(is_wrong_record, '0') AS TEXT) = '1'
              AND COALESCE(wrong_reason_code, '') != ''
            GROUP BY wrong_reason_code, reason_name
            """,
            (user_id, stat_date),
        ).fetchall()
    )
    total_wrong_with_reason = sum(as_int(row.get("wrong_question_count")) for row in rows)
    sort_rank = {code: index for index, code in enumerate(WRONG_REASON_LABELS, start=1)}
    rows.sort(key=lambda row: (sort_rank.get(row.get("reason_code"), 99), str(row.get("reason_code") or "")))
    conn.execute(
        "DELETE FROM loss_reason_stats WHERE user_id = ? AND stat_date = ?",
        (user_id, stat_date),
    )
    for index, row in enumerate(rows, start=1):
        code = str(row.get("reason_code") or "")
        count = as_int(row.get("wrong_question_count"))
        rate = count / total_wrong_with_reason if total_wrong_with_reason else 0
        previous_rate = previous_rates.get(code, 0)
        insert_table_row(
            conn,
            "loss_reason_stats",
            {
                "loss_reason_id": f"LR-{user_id}-{stat_date.replace('-', '')}-{index}",
                "user_id": user_id,
                "stat_date": stat_date,
                "reason_code": code,
                "reason_name": WRONG_REASON_LABELS.get(code) or row.get("reason_name") or code,
                "reason_rate": f"{rate:.4f}",
                "wrong_question_count": str(count),
                "sort_order": str(index),
                "updated_at": updated_at,
                "trend_delta": f"{rate - previous_rate:.4f}",
            },
        )


def chapter_assignable_type_rows(
    conn,
    chapter_id: str,
    plan_variant_id: str,
    submitted_type_id: str,
    level: Any = "",
) -> list[dict[str, Any]]:
    variant_rows: list[dict[str, Any]] = []
    if _variant_order_available(conn, plan_variant_id):
        variant_rows = rows_to_dicts(
            conn.execute(
                f"""
                SELECT q.*,
                       o.sort_code AS variant_sort_code,
                       {variant_duration_select_clause(conn)},
                       o.direct_assignable AS variant_direct_assignable
                FROM question_type_orders o
                JOIN question_types q ON q.type_id = o.type_id
                WHERE o.plan_variant_id = ?
                  AND q.chapter_id = ?
                  AND {_truthy_sql_filter("o.is_active")}
                  AND {_truthy_sql_filter("o.direct_assignable")}
                ORDER BY o.sort_code, q.type_id
                """,
                (plan_variant_id, chapter_id),
            ).fetchall()
        )
        if any(row.get("type_id") == submitted_type_id for row in variant_rows):
            return [_apply_variant_order(row, level) or row for row in variant_rows]
    return rows_to_dicts(
        conn.execute(
            f"""
            SELECT *
            FROM question_types
            WHERE chapter_id = ?
              AND {_truthy_sql_filter("direct_assignable")}
            ORDER BY sort_code, type_id
            """,
            (chapter_id,),
        ).fetchall()
    )


def summary_template_for_chapter(
    conn,
    raw_chapter_id: str,
    math_type: str,
    plan_stage: str,
) -> dict[str, Any]:
    ensure_growth_snapshot_tables(conn)
    columns = table_columns(conn, "chapter_summary_templates")
    chapter_column = "chapter_id" if "chapter_id" in columns else "raw_chapter_id"
    active_clause = f"AND {_truthy_sql_filter('is_active')}" if "is_active" in columns else ""
    return row_to_dict(
        conn.execute(
            f"""
            SELECT *
            FROM chapter_summary_templates
            WHERE {chapter_column} = ?
              {active_clause}
              AND (COALESCE(math_type, '') = '' OR math_type = ?)
              AND (COALESCE(plan_stage, '') = '' OR plan_stage = ?)
            ORDER BY
              CASE WHEN math_type = ? THEN 0 ELSE 1 END,
              CASE WHEN plan_stage = ? THEN 0 ELSE 1 END,
              updated_at DESC
            LIMIT 1
            """,
            (raw_chapter_id, math_type, plan_stage, math_type, plan_stage),
        ).fetchone()
    ) or {}


def actual_chapter_type_ids_for_boundary(
    conn,
    profile: dict[str, Any],
    plan_stage: str,
    plan_variant_id: str,
    type_id: str,
    raw_chapter_id: str,
) -> list[str]:
    sequence = actual_assignable_type_sequence(conn, profile, plan_stage, plan_variant_id)
    current_index = next(
        (index for index, row in enumerate(sequence) if str(row.get("type_id") or "") == type_id),
        -1,
    )
    if current_index < 0:
        return []
    current = sequence[current_index]
    if str(current.get("chapter_id") or "") != raw_chapter_id:
        return []
    next_row = sequence[current_index + 1] if current_index + 1 < len(sequence) else None
    if next_row and str(next_row.get("chapter_id") or "") == raw_chapter_id:
        return []
    return [
        str(row.get("type_id") or "")
        for row in sequence[: current_index + 1]
        if str(row.get("chapter_id") or "") == raw_chapter_id and str(row.get("type_id") or "")
    ]


def planned_chapter_type_ids_for_boundary(
    conn,
    user_id: str,
    plan_date: str,
    type_id: str,
    raw_chapter_id: str,
    profile: dict[str, Any],
    plan_stage: str,
    plan_variant_id: str,
) -> tuple[list[str], bool]:
    columns = table_columns(conn, "daily_plans")
    required = {"chapter_id", "student_type_sort_in_chapter", "is_chapter_boundary_after"}
    if not required.issubset(set(columns)):
        return [], False

    boundary_row = row_to_dict(
        conn.execute(
            """
            SELECT chapter_id, student_type_sort_in_chapter, is_chapter_boundary_after
            FROM daily_plans
            WHERE user_id = ? AND plan_date = ? AND type_id = ?
            ORDER BY CAST(question_order AS INTEGER)
            LIMIT 1
            """,
            (user_id, plan_date, type_id),
        ).fetchone()
    )
    if not boundary_row:
        return [], False

    has_sequence_data = bool(
        str(boundary_row.get("chapter_id") or "").strip()
        or as_int(boundary_row.get("student_type_sort_in_chapter")) > 0
    )
    if not has_sequence_data:
        return [], False
    if not truthy_text(boundary_row.get("is_chapter_boundary_after")):
        actual_type_ids = actual_chapter_type_ids_for_boundary(
            conn,
            profile,
            plan_stage,
            plan_variant_id,
            type_id,
            raw_chapter_id,
        )
        return actual_type_ids, True

    boundary_sort = as_int(boundary_row.get("student_type_sort_in_chapter"))
    if boundary_sort <= 0:
        return [], True

    rows = rows_to_dicts(
        conn.execute(
            """
            SELECT
              type_id,
              MIN(CAST(student_type_sort_in_chapter AS INTEGER)) AS type_sort
            FROM daily_plans
            WHERE user_id = ?
              AND plan_date <= ?
              AND chapter_id = ?
              AND CAST(student_type_sort_in_chapter AS INTEGER) > 0
              AND CAST(student_type_sort_in_chapter AS INTEGER) <= ?
            GROUP BY type_id
            ORDER BY type_sort, type_id
            """,
            (user_id, plan_date, raw_chapter_id, boundary_sort),
        ).fetchall()
    )
    type_ids = [str(row.get("type_id") or "") for row in rows if str(row.get("type_id") or "")]
    return type_ids, True


def maybe_upsert_chapter_completion_after_feedback(
    conn,
    user_id: str,
    plan_date: str,
    type_id: str,
    submitted_at: str,
) -> None:
    ensure_growth_snapshot_tables(conn)
    profile = row_to_dict(
        conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
    ) or {}
    plan_variant_id = str(profile.get("plan_variant_id") or "").strip()
    type_row = question_type_for_plan_variant(conn, type_id, plan_variant_id, profile.get("level"))
    if not type_row:
        type_row = row_to_dict(conn.execute("SELECT * FROM question_types WHERE type_id = ?", (type_id,)).fetchone())
    raw_chapter_id = str(type_row.get("chapter_id") or "").strip() if type_row else ""
    if not raw_chapter_id:
        return

    progress = row_to_dict(
        conn.execute(
            """
            SELECT carryover_type_id, carryover_question_ids
            FROM student_progress
            WHERE user_id = ?
            """,
            (user_id,),
        ).fetchone()
    ) or {}
    if (
        str(progress.get("carryover_type_id") or "").strip() == type_id
        and parse_json_id_list(progress.get("carryover_question_ids"))
    ):
        return

    type_ids, has_boundary_plan = planned_chapter_type_ids_for_boundary(
        conn,
        user_id,
        plan_date,
        type_id,
        raw_chapter_id,
        profile,
        str(profile.get("current_stage") or type_row.get("plan_stage") or ""),
        plan_variant_id,
    )
    if not type_ids:
        if has_boundary_plan:
            return
        chapter_types = chapter_assignable_type_rows(conn, raw_chapter_id, plan_variant_id, type_id, profile.get("level"))
        type_ids = [str(row.get("type_id") or "") for row in chapter_types if str(row.get("type_id") or "")]
        if not type_ids:
            return

    placeholders = ", ".join("?" for _ in type_ids)
    submitted_type_rows = rows_to_dicts(
        conn.execute(
            f"""
            SELECT DISTINCT type_id
            FROM question_feedbacks
            WHERE user_id = ?
              AND plan_date <= ?
              AND type_id IN ({placeholders})
            """,
            [user_id, plan_date, *type_ids],
        ).fetchall()
    )
    submitted_type_ids = {row.get("type_id") for row in submitted_type_rows}
    if not set(type_ids).issubset(submitted_type_ids):
        return

    question_type_columns = table_columns(conn, "question_types")
    legacy_chapter_name = "qt.chapter_name" if "chapter_name" in question_type_columns else "''"
    if table_exists(conn, "chapters"):
        chapter_name_select = f"COALESCE(c.chapter_name, {legacy_chapter_name}, '') AS chapter_name"
        chapter_join = "LEFT JOIN chapters c ON c.chapter_id = qt.chapter_id"
    else:
        chapter_name_select = f"{legacy_chapter_name} AS chapter_name"
        chapter_join = ""
    feedback_rows = rows_to_dicts(
        conn.execute(
            f"""
            SELECT f.*, qt.type_name, {chapter_name_select}
            FROM question_feedbacks f
            LEFT JOIN question_types qt ON qt.type_id = f.type_id
            {chapter_join}
            WHERE f.user_id = ?
              AND f.plan_date <= ?
              AND f.type_id IN ({placeholders})
            ORDER BY f.plan_date, f.submitted_at, f.type_id, f.question_id
            """,
            [user_id, plan_date, *type_ids],
        ).fetchall()
    )
    if not feedback_rows:
        return

    wrong_rows = [
        row for row in feedback_rows
        if row.get("result") in {"wrong", "unattempted"} or str(row.get("is_wrong_record") or "0") == "1"
    ]
    wrong_count = len(wrong_rows)
    submitted_count = len(feedback_rows)
    accuracy_rate = (submitted_count - wrong_count) / submitted_count if submitted_count else 0
    reason_counts: dict[str, dict[str, Any]] = {}
    for row in wrong_rows:
        if row.get("result") != "wrong":
            continue
        code = str(row.get("wrong_reason_code") or "").strip()
        if not code:
            continue
        target = reason_counts.setdefault(
            code,
            {"name": WRONG_REASON_LABELS.get(code) or row.get("wrong_reason_name") or code, "count": 0},
        )
        target["count"] += 1
    main_loss_reasons = [
        value["name"]
        for _, value in sorted(reason_counts.items(), key=lambda item: (-item[1]["count"], item[0]))[:2]
    ]
    started_date = min(str(row.get("plan_date") or plan_date) for row in feedback_rows)
    completed_at = plan_day_timestamp(plan_date, submitted_at)
    try:
        spent_days = (date.fromisoformat(plan_date) - date.fromisoformat(started_date)).days + 1
    except ValueError:
        spent_days = 1
    math_type = str(profile.get("math_type") or "")
    plan_stage = str(profile.get("current_stage") or type_row.get("plan_stage") or "")
    template = summary_template_for_chapter(conn, raw_chapter_id, math_type, plan_stage)
    chapter_name = str(
        (feedback_rows[0].get("chapter_name") if feedback_rows else "")
        or type_row.get("chapter_name")
        or template.get("chapter_name")
        or raw_chapter_id
    )
    existing = row_to_dict(
        conn.execute(
            """
            SELECT created_at
            FROM student_chapter_completions
            WHERE user_id = ? AND raw_chapter_id = ? AND plan_stage = ?
            """,
            (user_id, raw_chapter_id, plan_stage),
        ).fetchone()
    ) or {}
    conn.execute(
        "DELETE FROM student_chapter_completions WHERE user_id = ? AND raw_chapter_id = ? AND plan_stage = ?",
        (user_id, raw_chapter_id, plan_stage),
    )
    insert_table_row(
        conn,
        "student_chapter_completions",
        {
            "completion_id": f"SCC-{user_id}-{raw_chapter_id}-{plan_stage or 'stage'}",
            "user_id": user_id,
            "raw_chapter_id": raw_chapter_id,
            "chapter_name": chapter_name,
            "math_type": math_type,
            "plan_stage": plan_stage,
            "started_at": started_date,
            "completed_at": completed_at,
            "spent_days": str(max(spent_days, 1)),
            "completed_type_count": str(len(submitted_type_ids)),
            "total_type_count": str(len(type_ids)),
            "submitted_question_count": str(submitted_count),
            "wrong_question_count": str(wrong_count),
            "accuracy_rate": f"{accuracy_rate:.4f}",
            "main_loss_reasons": ", ".join(main_loss_reasons),
            "summary_template_id": template.get("summary_template_id") or "",
            "created_at": existing.get("created_at") or completed_at,
            "updated_at": completed_at,
        },
    )


def latest_chapter_summary_for_plan(conn, user_id: str, plan_date: str) -> dict[str, Any] | None:
    ensure_growth_snapshot_tables(conn)
    template_columns = table_columns(conn, "chapter_summary_templates")

    def template_select(column: str) -> str:
        return f"t.{column}" if column in template_columns else f"'' AS {column}"

    row = row_to_dict(
        conn.execute(
            f"""
            SELECT
              c.*,
              {template_select('summary_title')},
              {template_select('summary_text')},
              {template_select('exam_focus')},
              {template_select('key_points')},
              {template_select('common_mistakes')},
              {template_select('review_suggestions')},
              {template_select('next_step_text')}
            FROM student_chapter_completions c
            LEFT JOIN chapter_summary_templates t
              ON t.summary_template_id = c.summary_template_id
            WHERE c.user_id = ?
              AND substr(c.completed_at, 1, 10) = ?
            ORDER BY c.completed_at DESC
            LIMIT 1
            """,
            (user_id, plan_date),
        ).fetchone()
    ) or {}
    if not row:
        return None
    return {
        "chapterId": row.get("raw_chapter_id") or "",
        "chapterName": row.get("chapter_name") or "",
        "title": row.get("summary_title") or row.get("chapter_name") or "",
        "summaryTitle": row.get("summary_title") or row.get("chapter_name") or "",
        "summaryText": row.get("summary_text") or "",
        "examFocus": row.get("exam_focus") or "",
        "keyPoints": split_summary_text(row.get("key_points")),
        "commonMistakes": split_summary_text(row.get("common_mistakes")),
        "reviewAdvice": row.get("review_suggestions") or "",
        "reviewSuggestions": row.get("review_suggestions") or "",
        "nextStepText": row.get("next_step_text") or "",
        "accuracyRate": pct(row.get("accuracy_rate")),
        "spentDays": as_int(row.get("spent_days"), 1),
        "submittedQuestionCount": as_int(row.get("submitted_question_count")),
        "wrongQuestionCount": as_int(row.get("wrong_question_count")),
        "mainLossReasons": split_summary_text(row.get("main_loss_reasons")),
    }


def sync_study_time_projections(
    conn,
    user_id: str,
    plan_date: str,
    completed_minutes: int,
    target_minutes: int,
    task_name: str,
    updated_at: str,
) -> None:
    target = clamp(as_int(target_minutes, 180), 120, 480)
    completion_rate = completed_minutes / target if target else 0
    conn.execute(
        """
        UPDATE student_study_goals
        SET completed_minutes = ?, completion_rate = ?, focus_status = 'idle',
            current_round_started_at = '', updated_at = ?
        WHERE user_id = ? AND study_date = ?
        """,
        (str(completed_minutes), f"{completion_rate:.4f}", updated_at, user_id, plan_date),
    )
    conn.execute(
        """
        UPDATE study_room_members
        SET focused_minutes_today = ?, current_task_name = ?, updated_at = ?
        WHERE user_id = ?
        """,
        (str(completed_minutes), task_name, updated_at, user_id),
    )


def ensure_study_goal_for_date(
    conn,
    user_id: str,
    plan_date: str,
    profile: dict[str, Any],
    daily_plan: dict[str, Any],
) -> None:
    existing = row_to_dict(
        conn.execute(
            "SELECT goal_id FROM student_study_goals WHERE user_id = ? AND study_date = ?",
            (user_id, plan_date),
        ).fetchone()
    ) or {}
    if existing:
        return

    room = study_room_for_profile(conn, profile)
    task_options = [
        {"typeId": item["typeId"], "typeName": item["typeName"]}
        for item in daily_plan["plan"]["types"]
    ]
    if room:
        ensure_study_goal(conn, user_id, plan_date, room["room_id"], task_options)


def ensure_test_day_supporting_rows(conn, user_id: str, plan_date: str, daily_plan: dict[str, Any]) -> None:
    user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
    profile = row_to_dict(
        conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
    )
    if not user or not profile:
        raise ValueError("未找到该学生")

    ensure_student_stats_for_date(conn, user_id, plan_date)
    ensure_daily_activity_for_date(conn, user_id, plan_date)
    ensure_study_goal_for_date(conn, user_id, plan_date, profile, daily_plan)
    copy_latest_user_date_rows(
        conn,
        "type_mastery_stats",
        user_id,
        "stat_date",
        plan_date,
        "mastery_id",
        "MS",
        {"updated_at": now_text()},
    )
    copy_latest_user_date_rows(
        conn,
        "loss_reason_stats",
        user_id,
        "stat_date",
        plan_date,
        "loss_reason_id",
        "LR",
        {"updated_at": now_text()},
    )
    copy_latest_user_date_rows(
        conn,
        "priority_recommendations",
        user_id,
        "stat_date",
        plan_date,
        "recommendation_id",
        "PRIO",
        {"updated_at": now_text()},
    )


def time_label(value: Any) -> str:
    text = str(value or "")
    if " " in text:
        return text.split(" ", 1)[1][:5]
    return text[:5]


def rank_rows_for_display(rows: list[dict[str, Any]], current_user_id: str, top_n: int = 5) -> list[dict[str, Any]]:
    if len(rows) <= top_n + 3:
        return rows

    current_index = next((index for index, row in enumerate(rows) if row["userId"] == current_user_id), -1)
    if current_index < 0 or current_index < top_n:
        return rows[:top_n]

    keep: dict[int, dict[str, Any]] = {index: row for index, row in enumerate(rows[:top_n])}
    for index in range(max(0, current_index - 1), min(len(rows), current_index + 2)):
        keep[index] = rows[index]

    display_rows: list[dict[str, Any]] = []
    previous_index = -1
    for index in sorted(keep):
        if previous_index >= 0 and index - previous_index > 1:
            display_rows.append({"isDivider": True})
        display_rows.append(keep[index])
        previous_index = index
    return display_rows


def build_ranking_payload(rows: list[dict[str, Any]], current_user_id: str) -> dict[str, Any]:
    current = next((row for row in rows if row["userId"] == current_user_id), None)
    return {
        "top": rows[:3],
        "rows": rank_rows_for_display(rows, current_user_id),
        "myRank": current["rank"] if current else None,
        "my": current,
    }


def get_rankings(user_id: str, stat_date: str = CURRENT_DATE) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_student_stats_ranking_columns(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not user or not profile:
            raise ValueError("未找到该学生")

        access = build_student_access_context(user, profile, stat_date)
        user = access["user"]
        profile = access["profile"]
        stat_date = access["effectiveDate"]

        class_row = row_to_dict(
            conn.execute("SELECT * FROM classes WHERE class_id = ?", (profile["class_id"],)).fetchone()
        )
        class_students = rows_to_dicts(
            conn.execute(
                """
                SELECT
                  u.user_id, u.nickname, u.avatar_url, u.account_status,
                  p.class_id, p.level, p.math_type,
                  s.stat_date, s.continuous_days, s.today_completion_rate,
                  s.today_checkin_order, s.today_checkin_at,
                  s.weekly_completion_rate, s.weekly_study_days,
                  s.current_badge_name, s.next_badge_target,
                  s.badge_progress_value, s.badge_progress_rate,
                  COALESCE(a.completed_minutes, '0') AS today_focus_minutes,
                  COALESCE(dp.plan_row_count, '0') AS plan_row_count,
                  COALESCE(dp.submitted_plan_row_count, '0') AS submitted_plan_row_count,
                  COALESCE(dp.has_completed_plan_status, '0') AS has_completed_plan_status,
                  COALESCE(dp.plan_completed_at, '') AS plan_completed_at
                FROM student_profiles p
                JOIN users u ON u.user_id = p.user_id
                LEFT JOIN student_stats s ON s.user_id = p.user_id AND s.stat_date = ?
                LEFT JOIN student_daily_activity a ON a.user_id = p.user_id AND a.activity_date = ?
                LEFT JOIN (
                  SELECT
                    user_id,
                    COUNT(*) AS plan_row_count,
                    SUM(CASE WHEN feedback_open_status = 'submitted' OR plan_status = 'completed' THEN 1 ELSE 0 END) AS submitted_plan_row_count,
                    MAX(CASE WHEN plan_status = 'completed' THEN 1 ELSE 0 END) AS has_completed_plan_status,
                    MAX(CASE WHEN feedback_open_status = 'submitted' OR plan_status = 'completed' THEN updated_at ELSE '' END) AS plan_completed_at
                  FROM daily_plans
                  WHERE plan_date = ?
                  GROUP BY user_id
                ) dp ON dp.user_id = p.user_id
                WHERE p.class_id = ? AND u.role = 'student' AND u.account_status = 'active'
                """,
                (stat_date, stat_date, stat_date, profile["class_id"]),
            ).fetchall()
        )
        continuous_days_by_user = {
            str(row.get("user_id") or ""): calculate_continuous_checkin_days(conn, str(row.get("user_id") or ""), stat_date)
            for row in class_students
        }
        small_class_rank_offset = 3 if len(class_students) < 3 else 0

    def base_row(row: dict[str, Any]) -> dict[str, Any]:
        today_study_minutes = as_int(row.get("today_focus_minutes"), 0)
        today_rank_hours = (today_study_minutes + 59) // 60 if today_study_minutes > 0 else 0
        today_completion_rate = max(pct(row.get("today_completion_rate")), plan_completion_rate(row))
        today_checkin_raw = str(row.get("today_checkin_at") or "").strip()
        today_completed = bool(today_checkin_raw) or today_completion_rate >= 1 or plan_completion_has_completed(row)
        return {
            "userId": row["user_id"],
            "name": row["nickname"],
            "isMe": row["user_id"] == user_id,
            "avatarUrl": row["avatar_url"],
            "continuousDays": continuous_days_by_user.get(str(row["user_id"]), as_int(row.get("continuous_days"))),
            "weeklyStudyDays": as_int(row.get("weekly_study_days")),
            "todayStudyMinutes": today_study_minutes,
            "todayStudyHours": hours_from_minutes(today_study_minutes),
            "todayRankHours": today_rank_hours,
            "todayCheckinOrder": as_int(row.get("today_checkin_order"), 9999),
            "todayCheckinRaw": today_checkin_raw,
            "todayCheckinAt": time_label(row.get("today_checkin_at")),
            "todayCompletionRate": today_completion_rate,
            "todayCompleted": today_completed,
            "planCompletedAt": str(row.get("plan_completed_at") or "").strip(),
            "currentBadgeName": row.get("current_badge_name") or "暂无勋章",
            "nextBadgeTarget": row.get("next_badge_target") or "7天连续学习勋章",
            "badgeProgressValue": as_int(row.get("badge_progress_value")),
            "badgeProgressRate": pct(row.get("badge_progress_rate")),
        }

    def today_rank_key(row: dict[str, Any]) -> tuple:
        return (
            -row["todayRankHours"],
            -row["todayStudyMinutes"],
            row["name"],
        )

    today_rows = [
        row
        for row in (base_row(row) for row in class_students)
        if row["todayStudyMinutes"] > 0
    ]
    today_rows.sort(key=today_rank_key)
    for index, row in enumerate(today_rows, start=1):
        row["rank"] = index + small_class_rank_offset
        row["metric"] = f"专注{row['todayRankHours']}h"
        row["status"] = row["metric"]

    weekly_rows = [base_row(row) for row in class_students]
    weekly_rows.sort(key=lambda row: (-row["weeklyStudyDays"], row["name"]))
    for index, row in enumerate(weekly_rows, start=1):
        row["rank"] = index
        row["metric"] = f"本周 {row['weeklyStudyDays']}/7 天"
        row["status"] = "已坚持" if row["weeklyStudyDays"] else "本周未打卡"

    streak_rows = [base_row(row) for row in class_students]
    streak_rows.sort(key=lambda row: (-row["continuousDays"], row["name"]))
    for index, row in enumerate(streak_rows, start=1):
        row["rank"] = index + small_class_rank_offset
        row["metric"] = f"连续 {row['continuousDays']} 天"
        row["status"] = "已打卡" if row["todayCompletionRate"] > 0 else "今日待完成"

    current_stat = next((base_row(row) for row in class_students if row["user_id"] == user_id), {})
    today_me = next((row for row in today_rows if row["userId"] == user_id), None)
    weekly_me = next((row for row in weekly_rows if row["userId"] == user_id), None)
    streak_me = next((row for row in streak_rows if row["userId"] == user_id), None)

    return {
        "student": {
            "userId": user["user_id"],
            "name": user["nickname"],
            "classId": profile["class_id"],
            "className": class_row.get("class_name", profile["class_id"]) if class_row else profile["class_id"],
            "mathType": profile["math_type"],
            "targetScore": as_int(profile["target_score"]),
        },
        "summary": {
            "todayRank": today_me["rank"] if today_me else None,
            "weeklyRank": weekly_me["rank"] if weekly_me else None,
            "continuousRank": streak_me["rank"] if streak_me else None,
            "continuousDays": current_stat.get("continuousDays", 0),
            "weeklyStudyDays": current_stat.get("weeklyStudyDays", 0),
            "badgeProgressValue": current_stat.get("badgeProgressValue", 0),
            "badgeProgressRate": current_stat.get("badgeProgressRate", 0),
            "currentBadgeName": current_stat.get("currentBadgeName", "暂无勋章"),
            "nextBadgeTarget": current_stat.get("nextBadgeTarget", "7天连续学习勋章"),
        },
        "boards": {
            "today": build_ranking_payload(today_rows, user_id),
            "weekly": build_ranking_payload(weekly_rows, user_id),
            "streak": build_ranking_payload(streak_rows, user_id),
        },
    }


def hours_from_minutes(value: Any) -> float:
    return round(as_float(value) / 60, 1)


def activity_level(minutes: float, fallback: Any = "") -> str:
    if minutes <= 0:
        return "none"
    fallback_text = str(fallback or "")
    if fallback_text in {"high", "medium", "low"}:
        return fallback_text
    if minutes >= 150:
        return "high"
    if minutes >= 100:
        return "medium"
    return "low"


def week_bounds(target_date: date) -> tuple[date, date]:
    week_start = target_date - timedelta(days=target_date.weekday())
    return week_start, week_start + timedelta(days=6)


def progress_active_type_id(progress: dict[str, Any]) -> str:
    current_type_status = progress.get("current_type_status") or (
        "unfinished" if progress.get("unfinished_type_id") else "active"
    )
    if current_type_status == "unfinished" and progress.get("unfinished_type_id"):
        return str(progress.get("unfinished_type_id") or "")
    if current_type_status == "carryover" and progress.get("carryover_type_id"):
        return str(progress.get("carryover_type_id") or "")
    return str(progress.get("current_type_id") or "")


def weekly_target_estimate(
    conn,
    profile: dict[str, Any],
    progress: dict[str, Any],
    target_days: int,
) -> dict[str, Any]:
    plan_stage = progress.get("plan_stage") or profile.get("current_stage") or "强化"
    plan_variant_id = str(profile.get("plan_variant_id") or "").strip()
    current_type_status = progress.get("current_type_status") or (
        "unfinished" if progress.get("unfinished_type_id") else "active"
    )
    current_type_id = progress_active_type_id(progress)
    level = profile.get("level") if profile else ""
    cursor = question_type_for_plan_variant(conn, current_type_id, plan_variant_id, level) or next_assignable_type(
        conn,
        "",
        plan_variant_id,
        level,
    )
    remaining_days = max(target_days, 0)
    target_question_count = 0
    basis_type_ids: list[str] = []
    visited_type_ids: set[str] = set()
    first_type = True

    while cursor and remaining_days > 0:
        type_id = str(cursor.get("type_id") or "")
        if not type_id or type_id in visited_type_ids:
            break
        visited_type_ids.add(type_id)

        all_questions = get_type_questions(conn, type_id, profile, plan_stage)
        questions = all_questions
        if first_type and current_type_status == "unfinished" and progress.get("unfinished_type_id") == type_id:
            questions = get_type_questions(
                conn,
                type_id,
                profile,
                plan_stage,
                set(parse_json_id_list(progress.get("unfinished_question_ids"))),
            )
        elif first_type and current_type_status == "carryover" and progress.get("carryover_type_id") == type_id:
            questions = get_type_questions(
                conn,
                type_id,
                profile,
                plan_stage,
                set(parse_json_id_list(progress.get("carryover_question_ids"))),
            )

        if not questions:
            cursor = next_assignable_type(conn, type_id, plan_variant_id, level)
            first_type = False
            continue

        suggested_days = max(as_float(cursor.get("suggested_days"), 1.0), 0.1)
        if remaining_days >= suggested_days:
            take_count = len(questions)
            remaining_days -= suggested_days
        else:
            take_count = max(1, min(len(questions), math.ceil(len(questions) * remaining_days / suggested_days)))
            remaining_days = 0

        target_question_count += take_count
        basis_type_ids.append(type_id)
        cursor = next_assignable_type(conn, type_id, plan_variant_id, level)
        first_type = False

    return {
        "target_question_count": target_question_count,
        "target_type_count": len(basis_type_ids),
        "basis_type_ids": basis_type_ids,
    }


def ensure_student_weekly_target(
    conn,
    user_id: str,
    profile: dict[str, Any],
    progress: dict[str, Any] | None,
    target_date_text: str,
) -> dict[str, Any]:
    ensure_growth_snapshot_tables(conn)
    target_date = date.fromisoformat(target_date_text)
    week_start, week_end = week_bounds(target_date)
    existing = row_to_dict(
        conn.execute(
            """
            SELECT *
            FROM student_weekly_targets
            WHERE user_id = ? AND week_start_date = ?
            """,
            (user_id, week_start.isoformat()),
        ).fetchone()
    )
    if existing:
        return existing

    if not progress:
        progress = row_to_dict(
            conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user_id,)).fetchone()
        ) or {}
    target_days = max(0, (week_end - target_date).days + 1)
    estimate = weekly_target_estimate(conn, profile, progress, target_days) if progress else {
        "target_question_count": 0,
        "target_type_count": 0,
        "basis_type_ids": [],
    }
    created_at = now_text()
    weekly_target_id = f"WT-{user_id}-{week_start.strftime('%Y%m%d')}"
    conn.execute(
        """
        INSERT OR IGNORE INTO student_weekly_targets (
          weekly_target_id, user_id, week_start_date, week_end_date, basis_date,
          target_days, target_question_count, target_type_count, basis_type_ids,
          basis_note, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            weekly_target_id,
            user_id,
            week_start.isoformat(),
            week_end.isoformat(),
            target_date.isoformat(),
            str(target_days),
            str(estimate["target_question_count"]),
            str(estimate["target_type_count"]),
            json.dumps(estimate["basis_type_ids"], ensure_ascii=False),
            "按本周首次请求日至周日，沿当前题型顺序和 duration_100_days 估算",
            created_at,
            created_at,
        ),
    )
    return row_to_dict(
        conn.execute(
            """
            SELECT *
            FROM student_weekly_targets
            WHERE user_id = ? AND week_start_date = ?
            """,
            (user_id, week_start.isoformat()),
        ).fetchone()
    ) or {}


def weekly_feedback_metrics(conn, user_id: str, week_start_date: date, end_date: date) -> dict[str, int]:
    row = row_to_dict(
        conn.execute(
            """
            SELECT
              SUM(
                CASE
                  WHEN NOT (result = 'unattempted' AND COALESCE(unattempted_action, '') = 'defer')
                  THEN 1 ELSE 0
                END
              ) AS completed_questions,
              COUNT(
                DISTINCT CASE
                  WHEN NOT (result = 'unattempted' AND COALESCE(unattempted_action, '') = 'defer')
                  THEN type_id ELSE NULL
                END
              ) AS completed_type_count,
              SUM(CASE WHEN result = 'correct' THEN 1 ELSE 0 END) AS correct_questions,
              SUM(
                CASE
                  WHEN NOT (result = 'unattempted' AND COALESCE(unattempted_action, '') = 'defer')
                   AND (
                     CAST(COALESCE(is_wrong_record, '0') AS TEXT) = '1'
                     OR result = 'wrong'
                     OR (result = 'unattempted' AND COALESCE(unattempted_action, '') = 'mark_wrong')
                   )
                  THEN 1 ELSE 0
                END
              ) AS wrong_questions
            FROM question_feedbacks
            WHERE user_id = ? AND plan_date BETWEEN ? AND ?
            """,
            (user_id, week_start_date.isoformat(), end_date.isoformat()),
        ).fetchone()
    ) or {}
    return {
        "completed_questions": as_int(row.get("completed_questions")),
        "completed_type_count": as_int(row.get("completed_type_count")),
        "correct_questions": as_int(row.get("correct_questions")),
        "wrong_questions": as_int(row.get("wrong_questions")),
    }


def display_chapter_rows(conn, math_type: str) -> tuple[list[dict[str, Any]], dict[str, str]]:
    raw_to_display: dict[str, str] = {}
    rows: list[dict[str, Any]] = []
    if table_exists(conn, "display_chapters") and table_exists(conn, "chapter_display_mappings"):
        rows = rows_to_dicts(
            conn.execute(
                f"""
                SELECT
                  dc.display_chapter_id AS chapter_id,
                  dc.display_chapter_name AS chapter_name,
                  cdm.chapter_id AS raw_chapter_id,
                  dc.sort_order AS sort_order
                FROM display_chapters dc
                JOIN chapter_display_mappings cdm
                  ON cdm.display_chapter_id = dc.display_chapter_id
                 AND cdm.math_type = dc.math_type
                WHERE dc.math_type = ?
                  AND {_truthy_sql_filter("dc.is_active")}
                  AND {_truthy_sql_filter("cdm.is_active")}
                ORDER BY CAST(dc.sort_order AS INTEGER), CAST(cdm.display_sort_in_group AS INTEGER), cdm.chapter_id
                """,
                (math_type,),
            ).fetchall()
        )
    for row in rows:
        raw_to_display[row["raw_chapter_id"]] = row["chapter_id"]
    return rows, raw_to_display


def raw_chapter_type_rows(conn, math_type: str) -> list[dict[str, Any]]:
    question_type_columns = set(table_columns(conn, "question_types"))
    math_filter = "AND qt.math_type = ?" if "math_type" in question_type_columns else ""
    params: list[Any] = [math_type] if "math_type" in question_type_columns else []
    if table_exists(conn, "chapters"):
        return rows_to_dicts(
            conn.execute(
                f"""
                SELECT qt.type_id, qt.chapter_id, COALESCE(NULLIF(c.chapter_name, ''), qt.chapter_id) AS chapter_name
                FROM question_types qt
                LEFT JOIN chapters c ON c.chapter_id = qt.chapter_id
                WHERE COALESCE(qt.chapter_id, '') != ''
                  {math_filter}
                ORDER BY qt.chapter_id, qt.type_id
                """,
                params,
            ).fetchall()
        )
    return rows_to_dicts(
        conn.execute(
            f"""
            SELECT type_id, chapter_id, COALESCE(NULLIF(chapter_name, ''), chapter_id) AS chapter_name
            FROM question_types qt
            WHERE COALESCE(chapter_id, '') != ''
              {math_filter}
            ORDER BY chapter_id, type_id
            """,
            params,
        ).fetchall()
    )


def mastery_scope_type_ids(
    conn,
    profile: dict[str, Any],
    progress: dict[str, Any] | None,
    plan_stage: str,
    plan_variant_id: str,
) -> set[str] | None:
    sequence = actual_assignable_type_sequence(conn, profile, plan_stage, plan_variant_id)
    if not sequence:
        return None

    start_type_id = str((progress or {}).get("mastery_scope_start_type_id") or "").strip()
    if start_type_id:
        start_type = question_type_for_plan_variant(conn, start_type_id, plan_variant_id, profile.get("level"))
        start_index = sequence_start_index_for_type(sequence, start_type)
    else:
        start_index = 0

    return {
        str(row.get("type_id") or "").strip()
        for row in sequence[start_index:]
        if str(row.get("type_id") or "").strip()
    }


def chapter_mastery_for_curve(
    conn,
    user_id: str,
    profile: dict[str, Any],
    progress: dict[str, Any] | None,
    stat_date: str,
) -> list[dict[str, Any]]:
    math_type = profile.get("math_type") or "数学二"
    plan_stage = (progress or {}).get("plan_stage") or profile.get("current_stage") or "强化"
    plan_variant_id = str(profile.get("plan_variant_id") or "").strip()
    scoped_type_ids = mastery_scope_type_ids(conn, profile, progress, plan_stage, plan_variant_id)
    display_rows, raw_to_display = display_chapter_rows(conn, math_type)
    chapter_stats: dict[str, dict[str, Any]] = {}

    for row in display_rows:
        chapter_id = row["chapter_id"]
        chapter_stats.setdefault(
            chapter_id,
            {
                "chapterId": chapter_id,
                "name": row["chapter_name"],
                "sortOrder": as_int(row.get("sort_order")),
                "submittedQuestionCount": 0,
                "wrongQuestionCount": 0,
                "correctQuestionCount": 0,
                "totalQuestionCount": 0,
                "_eligibleQuestionIds": set(),
            },
        )

    raw_rows = raw_chapter_type_rows(conn, math_type)
    type_to_chapter: dict[str, str] = {}
    for row in raw_rows:
        type_id = str(row.get("type_id") or "").strip()
        if scoped_type_ids is not None and type_id not in scoped_type_ids:
            continue
        raw_chapter_id = row["chapter_id"]
        chapter_id = raw_to_display.get(raw_chapter_id, raw_chapter_id)
        type_to_chapter[type_id] = chapter_id
        chapter_stats.setdefault(
            chapter_id,
            {
                "chapterId": chapter_id,
                "name": row["chapter_name"],
                "sortOrder": len(chapter_stats) + 1,
                "submittedQuestionCount": 0,
                "wrongQuestionCount": 0,
                "correctQuestionCount": 0,
                "totalQuestionCount": 0,
                "_eligibleQuestionIds": set(),
            },
        )

    for row in raw_rows:
        type_id = str(row.get("type_id") or "").strip()
        if scoped_type_ids is not None and type_id not in scoped_type_ids:
            continue
        chapter_id = type_to_chapter.get(type_id)
        if not chapter_id:
            continue
        questions = get_type_questions(conn, type_id, profile, plan_stage)
        target = chapter_stats[chapter_id]
        for question in questions:
            key = (type_id, question["question_id"])
            if key not in target["_eligibleQuestionIds"]:
                target["_eligibleQuestionIds"].add(key)
                target["totalQuestionCount"] += 1

    feedback_rows = rows_to_dicts(
        conn.execute(
            """
            SELECT f.type_id, f.question_id, f.result, f.is_wrong_record, f.unattempted_action, f.submitted_at, f.plan_date
            FROM question_feedbacks f
            WHERE f.user_id = ? AND f.plan_date <= ?
            ORDER BY f.plan_date, f.submitted_at, f.feedback_id
            """,
            (user_id, stat_date),
        ).fetchall()
    )
    latest_by_question: dict[tuple[str, str], dict[str, Any]] = {}
    for row in feedback_rows:
        key = (row.get("type_id") or "", row.get("question_id") or "")
        if key[0] and key[1]:
            latest_by_question[key] = row

    for key, row in latest_by_question.items():
        chapter_id = type_to_chapter.get(key[0])
        if not chapter_id or chapter_id not in chapter_stats:
            continue
        target = chapter_stats[chapter_id]
        if key not in target["_eligibleQuestionIds"]:
            continue
        if row.get("result") == "unattempted" and row.get("unattempted_action") == "defer":
            continue
        target["submittedQuestionCount"] += 1
        is_wrong = (
            str(row.get("is_wrong_record") or "0") == "1"
            or row.get("result") == "wrong"
            or (row.get("result") == "unattempted" and row.get("unattempted_action") == "mark_wrong")
        )
        if row.get("result") == "correct":
            target["correctQuestionCount"] += 1
        if is_wrong:
            target["wrongQuestionCount"] += 1

    mastery: list[dict[str, Any]] = []
    for chapter in sorted(chapter_stats.values(), key=lambda item: (as_int(item.get("sortOrder")), item["chapterId"])):
        total = as_int(chapter.get("totalQuestionCount"))
        correct = as_int(chapter.get("correctQuestionCount"))
        rate = round(correct / total * 100) if total else 0
        if total and correct > 0 and rate == 0:
            rate = 1
        chapter.pop("_eligibleQuestionIds", None)
        chapter.pop("sortOrder", None)
        chapter["rate"] = rate
        chapter["tag"] = "未开始" if chapter["submittedQuestionCount"] == 0 else "良好" if rate >= 75 else "待提升" if rate >= 60 else "薄弱项"
        mastery.append(chapter)
    return mastery


def get_curve(user_id: str, stat_date: str = CURRENT_DATE) -> dict[str, Any]:
    requested_stat_date = str(stat_date or CURRENT_DATE)
    try:
        end_date = date.fromisoformat(stat_date)
    except ValueError as exc:
        raise ValueError("日期格式应为 YYYY-MM-DD") from exc
    start_date = end_date - timedelta(days=13)
    activity_end_date = end_date
    week_start_date = end_date - timedelta(days=end_date.weekday())

    with get_connection() as conn:
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        progress = row_to_dict(
            conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not user or not profile:
            raise ValueError("未找到该学生")
        access = build_student_access_context(user, profile, requested_stat_date)
        user = access["user"]
        profile = access["profile"]
        stat_date = access["effectiveDate"]
        end_date = date.fromisoformat(stat_date)
        start_date = end_date - timedelta(days=13)
        activity_end_date = end_date
        week_start_date = end_date - timedelta(days=end_date.weekday())
        stats = row_to_dict(
            conn.execute(
                "SELECT * FROM student_stats WHERE user_id = ? AND stat_date = ?",
                (user_id, stat_date),
            ).fetchone()
        )
        if not user or not profile:
            raise ValueError("未找到该学生")
        if not stats:
            stats = {
                "overall_accuracy": "0",
                "repeat_wrong_rate": "0",
                "weekly_completion_rate": "0",
            }
        weekly_target = ensure_student_weekly_target(conn, user_id, profile, progress, stat_date)
        weekly_feedback = weekly_feedback_metrics(conn, user_id, week_start_date, end_date)

        activity_rows = rows_to_dicts(
            conn.execute(
                """
                SELECT * FROM student_daily_activity
                WHERE user_id = ? AND activity_date BETWEEN ? AND ?
                ORDER BY activity_date
                """,
                (user_id, start_date.isoformat(), activity_end_date.isoformat()),
            ).fetchall()
        )
        loss_rows = rows_to_dicts(
            conn.execute(
                """
                SELECT * FROM loss_reason_stats
                WHERE user_id = ? AND stat_date = ?
                ORDER BY CAST(sort_order AS INTEGER), reason_code
                """,
                (user_id, stat_date),
            ).fetchall()
        )
        mastery = chapter_mastery_for_curve(conn, user_id, profile, progress, stat_date)

    activity_by_date = {row["activity_date"]: row for row in activity_rows}
    activity = []
    completed_minutes = []
    for offset in range(14):
        current_date = start_date + timedelta(days=offset)
        key = current_date.isoformat()
        row = activity_by_date.get(key, {})
        minutes = as_float(row.get("completed_minutes"))
        completed_question_count = as_int(row.get("completed_questions"))
        completed_minutes.append(minutes)
        studied = (
            str(row.get("is_studied") or ("1" if minutes > 0 or completed_question_count > 0 else "0")) == "1"
            and (minutes > 0 or completed_question_count > 0)
        )
        activity.append(
            {
                "date": current_date.strftime("%m.%d"),
                "rawDate": key,
                "level": activity_level(minutes, row.get("intensity_level")) if minutes > 0 else ("low" if studied else "none"),
                "hours": hours_from_minutes(minutes),
                "minutes": round(minutes),
                "studied": studied,
                "completedQuestions": completed_question_count,
                "submittedTypeCount": as_int(row.get("submitted_type_count")),
            }
        )
    continuous_days = 0
    for row in reversed(activity):
        if not row["studied"]:
            break
        continuous_days += 1

    loss_reasons = [
        {
            "code": row["reason_code"],
            "name": row["reason_name"],
            "rate": pct(row.get("reason_rate")),
            "count": as_int(row.get("wrong_question_count")),
            "trendDelta": pct(row.get("trend_delta")),
        }
        for row in loss_rows
    ]

    total_minutes = sum(completed_minutes)
    max_minutes = max(completed_minutes or [0])
    weekly_completed_questions = weekly_feedback["completed_questions"]
    weekly_planned_questions = as_int(weekly_target.get("target_question_count"))
    weekly_completed_types = weekly_feedback["completed_type_count"]
    weekly_correct_questions = weekly_feedback["correct_questions"]
    weekly_wrong_questions = weekly_feedback["wrong_questions"]
    weekly_completion_rate = (
        min(100, round(weekly_completed_questions / weekly_planned_questions * 100))
        if weekly_planned_questions
        else pct(stats.get("weekly_completion_rate"))
    )
    weekly_accuracy_rate = (
        round(weekly_correct_questions / (weekly_correct_questions + weekly_wrong_questions) * 100)
        if weekly_correct_questions + weekly_wrong_questions
        else pct(stats.get("overall_accuracy"))
    )
    weekly = {
        "completionRate": weekly_completion_rate,
        "accuracyRate": weekly_accuracy_rate,
        "questionCount": weekly_completed_questions,
        "completedTypeCount": weekly_completed_types,
        "plannedQuestionCount": weekly_planned_questions,
    }
    return {
        "student": {
            "userId": user["user_id"],
            "name": user["nickname"],
            "mathType": profile["math_type"],
            "targetScore": as_int(profile["target_score"]),
        },
        "curve": {
            "title": "数学能力成长曲线",
            "targetScore": as_int(profile["target_score"]),
            "overallAccuracy": pct(stats.get("overall_accuracy")),
            "repeatWrongRate": pct(stats.get("repeat_wrong_rate")),
            "weeklyCompletionRate": weekly_completion_rate,
            "weeklyQuestionCount": weekly_completed_questions,
            "weeklyCompletedTypeCount": weekly_completed_types,
            "weeklyAccuracyRate": weekly_accuracy_rate,
            "weekly": weekly,
            "masteryLabel": "章节掌握度",
            "activitySummary": {
                "totalHours": hours_from_minutes(total_minutes),
                "maxHours": hours_from_minutes(max_minutes),
                "continuousDays": continuous_days,
                "studiedDays": sum(1 for row in activity if row["studied"]),
            },
            "activity": activity,
            "mastery": mastery,
            "lossReasons": loss_reasons,
        },
    }


def clamp(value: int, minimum: int, maximum: int) -> int:
    return max(minimum, min(maximum, value))


def study_room_for_profile(conn, profile: dict[str, Any]) -> dict[str, Any] | None:
    ensure_study_room_tables(conn)
    class_id = str(profile.get("class_id") or "").strip()
    if class_id:
        room = row_to_dict(
            conn.execute(
                """
                SELECT * FROM study_rooms
                WHERE class_id = ?
                ORDER BY room_id
                LIMIT 1
                """,
                (class_id,),
            ).fetchone()
        )
        if room:
            return room
    return row_to_dict(
        conn.execute(
            """
            SELECT * FROM study_rooms
            WHERE math_type = ? AND target_exam_year = ?
            ORDER BY room_id
            LIMIT 1
            """,
            (profile.get("math_type") or "", profile.get("target_exam_year") or ""),
        ).fetchone()
    )


def ensure_study_goal(
    conn,
    user_id: str,
    plan_date: str,
    room_id: str,
    task_options: list[dict[str, Any]],
) -> dict[str, Any]:
    goal = row_to_dict(
        conn.execute(
            "SELECT * FROM student_study_goals WHERE user_id = ? AND study_date = ?",
            (user_id, plan_date),
        ).fetchone()
    )
    if goal:
        return goal

    first_task = task_options[0] if task_options else {"typeId": "", "typeName": "今日数学任务"}
    created_at = now_text()
    goal_id = f"SG-{user_id}-{plan_date.replace('-', '')}"
    target_minutes = student_daily_target_minutes(conn, user_id, plan_date)
    conn.execute(
        """
        INSERT INTO student_study_goals (
          goal_id, user_id, study_date, room_id, target_minutes, completed_minutes,
          completion_rate, current_focus_minutes, current_task_type_id, current_task_name,
          focus_status, current_round_started_at, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            goal_id,
            user_id,
            plan_date,
            room_id,
            str(target_minutes),
            "0",
            "0",
            "45",
            first_task["typeId"],
            first_task["typeName"],
            "idle",
            "",
            created_at,
            created_at,
        ),
    )
    conn.commit()
    return row_to_dict(
        conn.execute(
            "SELECT * FROM student_study_goals WHERE user_id = ? AND study_date = ?",
            (user_id, plan_date),
        ).fetchone()
    )


def class_focus_members(
    conn,
    class_id: str,
    plan_date: str,
    limit: int = 4,
) -> list[dict[str, Any]]:
    running_rows = rows_to_dicts(
        conn.execute(
            """
            SELECT
              u.user_id,
              u.nickname,
              u.avatar_url,
              COALESCE(fs.task_name, g.current_task_name, '') AS current_task_name,
              COALESCE(g.completed_minutes, '0') AS completed_minutes
            FROM focus_sessions fs
            JOIN student_profiles p ON p.user_id = fs.user_id
            JOIN users u ON u.user_id = fs.user_id
            LEFT JOIN student_study_goals g
              ON g.user_id = fs.user_id AND g.study_date = fs.plan_date
            WHERE fs.class_id = ?
              AND fs.plan_date = ?
              AND fs.status = 'running'
              AND COALESCE(fs.ended_at, '') = ''
              AND u.role = 'student'
              AND u.account_status = 'active'
              AND fs.started_at = (
                SELECT MAX(fs2.started_at)
                FROM focus_sessions fs2
                WHERE fs2.user_id = fs.user_id
                  AND fs2.plan_date = fs.plan_date
                  AND fs2.status = 'running'
                  AND COALESCE(fs2.ended_at, '') = ''
              )
            ORDER BY fs.started_at DESC, u.nickname
            LIMIT ?
            """,
            (class_id, plan_date, limit),
        ).fetchall()
    )
    members = [
        {
            "userId": row["user_id"],
            "name": classmate_display_name(row["nickname"]),
            "avatarUrl": row["avatar_url"],
            "task": row["current_task_name"] or "今日数学任务",
            "focusedMinutes": as_int(row["completed_minutes"]),
            "focusStatus": "running",
        }
        for row in running_rows
    ]
    if len(members) >= limit:
        return members[:limit]

    running_user_ids = {row["userId"] for row in members}
    completed_rows = rows_to_dicts(
        conn.execute(
            """
            SELECT
              u.user_id,
              u.nickname,
              u.avatar_url,
              COALESCE(g.current_task_name, MAX(fs.task_name), '') AS current_task_name,
              SUM(COALESCE(CAST(NULLIF(fs.counted_minutes, '') AS INTEGER), 0)) AS completed_minutes,
              MAX(fs.ended_at) AS last_ended_at
            FROM focus_sessions fs
            JOIN student_profiles p ON p.user_id = fs.user_id
            JOIN users u ON u.user_id = fs.user_id
            LEFT JOIN student_study_goals g
              ON g.user_id = fs.user_id AND g.study_date = fs.plan_date
            WHERE fs.class_id = ?
              AND fs.plan_date = ?
              AND fs.status = 'completed'
              AND COALESCE(fs.ended_at, '') <> ''
              AND COALESCE(CAST(NULLIF(fs.counted_minutes, '') AS INTEGER), 0) > 0
              AND u.role = 'student'
              AND u.account_status = 'active'
            GROUP BY u.user_id, u.nickname, u.avatar_url, g.current_task_name
            ORDER BY last_ended_at DESC, u.nickname
            LIMIT ?
            """,
            (class_id, plan_date, limit * 2),
        ).fetchall()
    )
    for row in completed_rows:
        if row["user_id"] in running_user_ids:
            continue
        members.append(
            {
                "userId": row["user_id"],
                "name": classmate_display_name(row["nickname"]),
                "avatarUrl": row["avatar_url"],
                "task": row["current_task_name"] or "今日数学任务",
                "focusedMinutes": as_int(row["completed_minutes"]),
                "focusStatus": "completed_today",
            }
        )
        if len(members) >= limit:
            break
    return members


def class_member_count(conn, class_id: str) -> int:
    row = row_to_dict(
        conn.execute(
            """
            SELECT COUNT(*) AS total
            FROM student_profiles p
            JOIN users u ON u.user_id = p.user_id
            WHERE p.class_id = ?
              AND u.role = 'student'
              AND u.account_status = 'active'
            """,
            (class_id,),
        ).fetchone()
    )
    return as_int(row.get("total") if row else 0)


def classmate_display_name(name: Any) -> str:
    compact = str(name or "").strip()
    return f"{compact[0]}同学" if compact else "同学"


def get_study_room(user_id: str, plan_date: str = CURRENT_DATE) -> dict[str, Any]:
    requested_plan_date = str(plan_date or CURRENT_DATE)

    with get_connection() as conn:
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not user or not profile:
            raise ValueError("未找到该学生")
        access = build_student_access_context(user, profile, requested_plan_date)
        user = access["user"]
        profile = access["profile"]
        plan_date = access["effectiveDate"]
        daily_plan = get_daily_plan(user_id, plan_date)
        task_options = [
            {
                "typeId": row["typeId"],
                "typeName": row["typeName"],
                "disabled": len(daily_plan["plan"]["types"]) <= 1,
            }
            for row in daily_plan["plan"]["types"]
        ]
        room = study_room_for_profile(conn, profile)
        if not room:
            raise ValueError("未找到该学生对应自习室")
        goal = ensure_study_goal(conn, user_id, plan_date, room["room_id"], task_options)
        activity = daily_activity_for_date(conn, user_id, plan_date)
        conn.commit()
        member_count = class_member_count(conn, profile["class_id"])
        members = class_focus_members(conn, profile["class_id"], plan_date)

    target_minutes = clamp(as_int(goal.get("target_minutes"), 180), 120, 480)
    completed_minutes = max(0, as_int(activity.get("completed_minutes")))
    completion_rate = min(100, round(completed_minutes / target_minutes * 100)) if target_minutes else 0
    current_task_type_id = goal.get("current_task_type_id") or (task_options[0]["typeId"] if task_options else "")
    current_task = next((row for row in task_options if row["typeId"] == current_task_type_id), None)
    current_task_name = (current_task or {}).get("typeName") or goal.get("current_task_name") or "今日数学任务"

    return {
        "student": {
            "userId": user["user_id"],
            "nickname": user["nickname"],
            "classId": profile["class_id"],
        },
        "studyRoom": {
            "roomId": room["room_id"],
            "roomName": room["room_name"],
            "onlineCount": as_int(room["online_count"]),
            "classMemberCount": member_count,
            "quietMode": room.get("quiet_mode_enabled") == "1",
            "rule": room["room_rule"],
            "durationOptions": [30, 45, 60],
            "customDuration": 45,
            "focusMinutes": clamp(as_int(goal.get("current_focus_minutes"), 45), 10, 180),
            "targetMinutes": target_minutes,
            "targetHours": round(target_minutes / 60, 1),
            "completedMinutes": completed_minutes,
            "completedHours": round(completed_minutes / 60, 1),
            "completionRate": completion_rate,
            "currentTaskTypeId": current_task_type_id,
            "currentTaskName": current_task_name,
            "focusStatus": goal.get("focus_status") or "idle",
            "taskOptions": task_options,
            "members": members,
        },
    }


def task_name_from_options(task_options: list[dict[str, Any]], task_type_id: str) -> str:
    match = next((row for row in task_options if row["typeId"] == task_type_id), None)
    if match:
        return match["typeName"]
    return task_options[0]["typeName"] if task_options else "今日数学任务"


def start_focus_session(
    user_id: str,
    plan_date: str,
    task_type_id: str,
    target_minutes: int,
    daily_target_minutes: int | None = None,
) -> dict[str, Any]:
    with get_connection() as conn:
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not profile:
            raise ValueError("未找到该学生")
        assert_service_active_for_write(conn, profile, plan_date)
    study_room = get_study_room(user_id, plan_date)
    task_options = study_room["studyRoom"]["taskOptions"]
    valid_type_ids = {row["typeId"] for row in task_options}
    if valid_type_ids and task_type_id not in valid_type_ids:
        raise ValueError("当前任务不属于今日计划")
    target_minutes = clamp(as_int(target_minutes, 45), 10, 180)
    daily_target = clamp(as_int(daily_target_minutes, study_room["studyRoom"]["targetMinutes"]), 120, 480)
    task_name = task_name_from_options(task_options, task_type_id)
    started_at = now_text()
    session_id = f"FS-{plan_date.replace('-', '')}-{user_id}-{uuid.uuid4().hex[:8].upper()}"

    with get_connection() as conn:
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not profile:
            raise ValueError("未找到该学生")
        assert_service_active_for_write(conn, profile, plan_date)
        conn.execute(
            """
            INSERT INTO focus_sessions (
              session_id, user_id, class_id, plan_date, task_type_id, task_name,
              target_minutes, started_at, paused_at, resumed_at, ended_at,
              focused_seconds, counted_minutes, status, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                session_id,
                user_id,
                profile["class_id"],
                plan_date,
                task_type_id,
                task_name,
                str(target_minutes),
                started_at,
                "",
                "",
                "",
                "0",
                "0",
                "running",
                started_at,
                started_at,
            ),
        )
        conn.execute(
            """
            UPDATE student_study_goals
            SET target_minutes = ?, current_focus_minutes = ?, current_task_type_id = ?,
                current_task_name = ?, focus_status = 'focusing',
                current_round_started_at = ?, updated_at = ?
            WHERE user_id = ? AND study_date = ?
            """,
            (
                str(daily_target),
                str(target_minutes),
                task_type_id,
                task_name,
                started_at,
                started_at,
                user_id,
                plan_date,
            ),
        )
        conn.commit()

    return {
        "session": {
            "sessionId": session_id,
            "status": "running",
            "targetMinutes": target_minutes,
            "taskTypeId": task_type_id,
            "taskName": task_name,
            "startedAt": started_at,
        },
        "studyRoom": get_study_room(user_id, plan_date)["studyRoom"],
    }


def update_focus_status(session_id: str, status: str) -> dict[str, Any]:
    if status not in {"paused", "running"}:
        raise ValueError("不支持的专注状态")
    timestamp = now_text()
    column = "paused_at" if status == "paused" else "resumed_at"
    with get_connection() as conn:
        session = row_to_dict(
            conn.execute("SELECT * FROM focus_sessions WHERE session_id = ?", (session_id,)).fetchone()
        )
        if not session:
            raise ValueError("未找到专注记录")
        conn.execute(
            f"UPDATE focus_sessions SET status = ?, {column} = ?, updated_at = ? WHERE session_id = ?",
            (status, timestamp, timestamp, session_id),
        )
        conn.commit()
    return {"session": {"sessionId": session_id, "status": status, "updatedAt": timestamp}}


def focus_session_user_id(session_id: str) -> str:
    with get_connection() as conn:
        row = row_to_dict(
            conn.execute("SELECT user_id FROM focus_sessions WHERE session_id = ?", (session_id,)).fetchone()
        )
        if not row:
            raise ValueError("focus session not found")
        return row["user_id"]


def counted_focus_minutes(focused_seconds: int) -> int:
    seconds = max(0, as_int(focused_seconds))
    if seconds < 50:
        return 0
    return (seconds + 10) // 60


def finish_focus_session(session_id: str, focused_seconds: int) -> dict[str, Any]:
    requested_seconds = max(0, as_int(focused_seconds))
    ended_at = now_text()
    with get_connection() as conn:
        session = row_to_dict(
            conn.execute("SELECT * FROM focus_sessions WHERE session_id = ?", (session_id,)).fetchone()
        )
        if not session:
            raise ValueError("未找到专注记录")

        user_id = session["user_id"]
        plan_date = session["plan_date"]
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not profile:
            raise ValueError("未找到该学生")
        assert_service_active_for_write(conn, profile, plan_date)
        target_minutes = clamp(as_int(session.get("target_minutes"), 45), 10, 180)
        target_seconds = target_minutes * 60

        if session.get("ended_at") or session.get("status") == "completed":
            saved_seconds = max(0, as_int(session.get("focused_seconds")))
            saved_minutes = max(0, as_int(session.get("counted_minutes")))
            response_ended_at = session.get("ended_at") or session.get("updated_at") or ended_at
        else:
            saved_seconds = min(requested_seconds, target_seconds)
            saved_minutes = counted_focus_minutes(saved_seconds)
            response_ended_at = ended_at

            activity = daily_activity_for_date(conn, user_id, plan_date)
            previous_completed_minutes = max(0, as_int(activity.get("completed_minutes")))
            completed_minutes = previous_completed_minutes + saved_minutes
            intensity_level = study_minutes_intensity_level(completed_minutes)

            conn.execute(
                """
                UPDATE focus_sessions
                SET status = 'completed', ended_at = ?, focused_seconds = ?,
                    counted_minutes = ?, updated_at = ?
                WHERE session_id = ?
                """,
                (ended_at, str(saved_seconds), str(saved_minutes), ended_at, session_id),
            )
            conn.execute(
                """
                UPDATE student_daily_activity
                SET completed_minutes = ?, is_studied = ?, intensity_level = ?, updated_at = ?
                WHERE user_id = ? AND activity_date = ?
                """,
                (
                    str(completed_minutes),
                    "1" if completed_minutes > 0 else "0",
                    intensity_level,
                    ended_at,
                    user_id,
                    plan_date,
                ),
            )

            goal = row_to_dict(
                conn.execute(
                    "SELECT * FROM student_study_goals WHERE user_id = ? AND study_date = ?",
                    (user_id, plan_date),
                ).fetchone()
            )
            daily_target_minutes = clamp(as_int(goal.get("target_minutes") if goal else 180), 120, 480)
            sync_study_time_projections(
                conn,
                user_id,
                plan_date,
                completed_minutes,
                daily_target_minutes,
                session.get("task_name") or "",
                ended_at,
            )
            conn.commit()

    return {
        "session": {
            "sessionId": session_id,
            "status": "completed",
            "focusedSeconds": saved_seconds,
            "countedMinutes": saved_minutes,
            "endedAt": response_ended_at,
        },
        "studyRoom": get_study_room(user_id, plan_date)["studyRoom"],
        "profile": get_student_profile(user_id, plan_date),
        "curve": get_curve(user_id, plan_date),
        "rankings": get_rankings(user_id, plan_date),
    }


def school_payload(row: dict[str, Any] | None) -> dict[str, Any]:
    if not row:
        return {"schoolId": "", "name": "未选择", "province": "", "tags": []}
    return {
        "schoolId": row.get("school_id") or "",
        "name": row.get("school_name") or "未选择",
        "province": row.get("province") or "",
        "tags": [item for item in str(row.get("school_tags") or "").split(",") if item],
    }


def get_student_profile(user_id: str, stat_date: str = CURRENT_DATE) -> dict[str, Any]:
    requested_stat_date = str(stat_date or CURRENT_DATE)
    with get_connection() as conn:
        access = resolve_student_access_context(conn, user_id, requested_stat_date)
    stat_date = access["effectiveDate"]
    daily_plan = get_daily_plan(user_id, stat_date)
    study_room = get_study_room(user_id, stat_date)
    with get_connection() as conn:
        ensure_student_profile_target_school_columns(conn)
        ensure_student_stats_ranking_columns(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        stats = row_to_dict(
            conn.execute(
                "SELECT * FROM student_stats WHERE user_id = ? AND stat_date = ?",
                (user_id, stat_date),
            ).fetchone()
        ) or {}
        school = None
        if profile and profile.get("target_school_id"):
            school = row_to_dict(
                conn.execute(
                    "SELECT * FROM institutions WHERE school_id = ?",
                    (profile["target_school_id"],),
                ).fetchone()
            )
        if not user or not profile:
            raise ValueError("未找到该学生")
        class_row = row_to_dict(
            conn.execute("SELECT * FROM classes WHERE class_id = ?", (profile.get("class_id") or "",)).fetchone()
        ) or {}
        activity = daily_activity_for_date(conn, user_id, stat_date)
        continuous_days = calculate_continuous_checkin_days(conn, user_id, stat_date)
        conn.commit()

    target_minutes = study_room["studyRoom"]["targetMinutes"]
    completed_minutes = max(0, as_int(activity.get("completed_minutes")))
    return {
        "profile": {
            "userId": user["user_id"],
            "nickname": user["nickname"],
            "avatarUrl": user.get("avatar_url") or "",
            "qqGroup": class_row.get("qq_group_number") or class_row.get("group_name") or "",
            "qq": user.get("QQ") or "",
            "contactQQ": user.get("QQ") or "",
            "classId": profile.get("class_id") or "",
            "className": class_row.get("class_name") or "",
            "mathType": profile["math_type"],
            "examSubject": profile["math_type"],
            "targetScore": as_int(profile["target_score"]),
            "targetExamYear": profile.get("target_exam_year") or "",
            "targetExamDate": profile.get("target_exam_date") or "",
            "daysToExam": days_between(stat_date, profile.get("target_exam_date") or stat_date),
            "continuousDays": continuous_days or as_int(stats.get("continuous_days")),
            "todayStudyHours": hours_from_minutes(completed_minutes),
            "dailyMathGoalHours": hours_from_minutes(target_minutes),
            "primaryBooks": [
                profile_book_display_code(profile.get("primary_book_1")),
                profile_book_display_code(profile.get("primary_book_2")),
            ],
            "backupBooks": [profile_book_display_code(profile.get("backup_book"))],
            "targetSchool": school_payload(school),
            "planTypeCount": len(daily_plan["plan"]["types"]),
        }
    }


def search_institutions(query: str = "", limit: int = 20) -> dict[str, Any]:
    if public_seed.is_public_seed_configured():
        return public_seed.search_public_institutions(query, limit)

    query = str(query or "").strip()
    limit = clamp(as_int(limit, 20), 1, 50)
    with get_connection() as conn:
        if query:
            like = f"%{query}%"
            rows = rows_to_dicts(
                conn.execute(
                    """
                    SELECT * FROM institutions
                    WHERE school_name LIKE ? OR search_keywords LIKE ?
                    ORDER BY CAST(sort_order AS INTEGER), school_name
                    LIMIT ?
                    """,
                    (like, like, limit),
                ).fetchall()
            )
        else:
            rows = rows_to_dicts(
                conn.execute(
                    """
                    SELECT * FROM institutions
                    ORDER BY CAST(sort_order AS INTEGER), school_name
                    LIMIT ?
                    """,
                    (limit,),
                ).fetchall()
            )
    return {"items": [school_payload(row) for row in rows]}


def update_target_school(user_id: str, school_id: str) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_student_profile_target_school_columns(conn)
        school = row_to_dict(conn.execute("SELECT * FROM institutions WHERE school_id = ?", (school_id,)).fetchone())
        if not school:
            raise ValueError("未找到该院校")
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not profile:
            raise ValueError("未找到该学生")
        updated_at = now_text()
        conn.execute(
            """
            UPDATE student_profiles
            SET target_school_id = ?, target_school_name = ?, updated_at = ?
            WHERE user_id = ?
            """,
            (school["school_id"], school["school_name"], updated_at, user_id),
        )
        conn.commit()
    return get_student_profile(user_id, CURRENT_DATE)


def login(
    login_account: str,
    password: str,
    device_id: str = "",
    device_name: str = "",
    ip_address: str = "",
    user_agent: str = "",
    setup_token: str = "",
    totp_code: str = "",
) -> dict[str, Any]:
    with get_connection() as conn:
        user = row_to_dict(
            conn.execute(
                "SELECT * FROM users WHERE login_account = ? OR user_id = ?",
                (login_account, login_account),
            ).fetchone()
        )
        if not user:
            raise ValueError("账号不存在或已停用")

        if not password_matches(user, password):
            raise ValueError("账号或密码错误")

        if user.get("role") == "student" and account_type_for_user(user) == "trial":
            profile = row_to_dict(
                conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user["user_id"],)).fetchone()
            )
            app_today = get_runtime_test_date(conn)
            if user.get("account_status") == "expired_pending_review" or (
                profile and user.get("account_status") == "active" and is_service_expired(profile, app_today)
            ):
                raise ValueError("账号已到期，请联系管理员开通")

        if user["account_status"] != "active":
            raise ValueError("账号不存在或已停用")

        if user.get("role") == "admin":
            ensure_formal_activation_tables(conn)
            state = admin_totp_state(conn, user["user_id"])
            if not state.get("totp_bound_at"):
                if not setup_token:
                    secret = generate_totp_secret()
                    setup_token_plain = generate_setup_token()
                    expires_at = add_minutes_from_now_text(TOTP_SETUP_WINDOW_MINUTES)
                    conn.execute(
                        """
                        UPDATE admin_totp_state
                        SET totp_secret = ?, setup_token_hash = ?, setup_token_expires_at = ?,
                            updated_at = ?
                        WHERE admin_user_id = ?
                        """,
                        (
                            secret,
                            auth_token_hash(setup_token_plain),
                            expires_at,
                            now_text(),
                            user["user_id"],
                        ),
                    )
                    conn.commit()
                    return {
                        "requiresTotpSetup": True,
                        "setupToken": setup_token_plain,
                        "manualSecret": secret,
                        "otpauthUrl": admin_totp_otpauth_url(user["login_account"], secret),
                        "expiresAt": expires_at,
                        "user": {
                            "userId": user["user_id"],
                            "nickname": user["nickname"],
                            "role": "admin",
                        },
                    }
                if state.get("setup_token_hash") != auth_token_hash(setup_token):
                    raise ValueError("TOTP setup expired, please login again")
                if not state.get("setup_token_expires_at") or str(state.get("setup_token_expires_at")) <= now_text():
                    raise ValueError("TOTP setup expired, please login again")
                if not verify_admin_totp_code(state.get("totp_secret", ""), totp_code):
                    return {
                        "requiresTotpSetup": True,
                        "message": "动态码错误，请重新输入。",
                        "setupToken": setup_token,
                        "manualSecret": state.get("totp_secret", ""),
                        "otpauthUrl": admin_totp_otpauth_url(user["login_account"], state.get("totp_secret", "")),
                        "expiresAt": state.get("setup_token_expires_at", ""),
                        "user": {
                            "userId": user["user_id"],
                            "nickname": user["nickname"],
                            "role": "admin",
                        },
                    }
                verified_at = now_text()
                conn.execute(
                    """
                    UPDATE admin_totp_state
                    SET totp_bound_at = ?, setup_token_hash = '', setup_token_expires_at = '',
                        last_verified_at = ?, updated_at = ?
                    WHERE admin_user_id = ?
                    """,
                    (verified_at, verified_at, verified_at, user["user_id"]),
                )
            elif not verify_admin_totp_code(state.get("totp_secret", ""), totp_code):
                return {
                    "requiresTotp": True,
                    "message": "请输入管理员动态码后继续登录。",
                    "user": {
                        "userId": user["user_id"],
                        "nickname": user["nickname"],
                        "role": "admin",
                    },
                }
            else:
                verified_at = now_text()
                conn.execute(
                    """
                    UPDATE admin_totp_state
                    SET last_verified_at = ?, updated_at = ?
                    WHERE admin_user_id = ?
                    """,
                    (verified_at, verified_at, user["user_id"]),
                )

        session = start_auth_session(conn, user, password, device_id, device_name, ip_address, user_agent)
        conn.commit()
        return auth_payload(conn, user, session["token"], session["sessionId"])


def get_daily_plan(user_id: str, plan_date: str = CURRENT_DATE) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_question_source_display_columns(conn)
        ensure_daily_plan_sequence_columns(conn)
        ensure_student_progress_tracking_columns(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        progress = row_to_dict(
            conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not user or not profile:
            raise ValueError("未找到该学生今日计划")

        if is_service_expired(profile, plan_date):
            return {
                "student": {
                    "userId": user["user_id"],
                    "nickname": user["nickname"],
                    "mathType": profile["math_type"],
                    "targetScore": as_int(profile["target_score"]),
                    "paidUntil": profile["paid_until"],
                },
                "plan": {
                    "planId": "",
                    "planDate": plan_date,
                    "planStatus": "expired",
                    "slogan": "服务已到期，请联系老师续费后继续查看今日计划。",
                    "examDate": profile.get("target_exam_date") or "",
                    "daysToExam": 0,
                    "types": [],
                    "feedbacks": [],
                },
            }
        if progress:
            ensure_student_weekly_target(conn, user_id, profile, progress, plan_date)
        question_type_intro_select = question_type_intro_select_clause(conn)
        plan_rows_sql = f"""
            SELECT
              dp.*,
              {question_type_intro_select},
              q.book_display_label,
              q.source_chapter_label,
              q.source_section_label,
              q.source_subsection_label,
              q.source_question_label,
              q.question_note
            FROM daily_plans dp
            LEFT JOIN question_types qt ON qt.type_id = dp.type_id
            LEFT JOIN questions q ON q.question_id = dp.question_id
            WHERE dp.user_id = ? AND dp.plan_date = ?
            ORDER BY CAST(dp.type_index AS INTEGER), CAST(dp.question_order AS INTEGER)
        """
        rows = rows_to_dicts(conn.execute(plan_rows_sql, (user_id, plan_date)).fetchall())
        if not rows:
            if not progress:
                raise ValueError("未找到该学生学习进度")
            rows = generate_daily_plan(conn, user, profile, progress, plan_date)
            rows = rows_to_dicts(conn.execute(plan_rows_sql, (user_id, plan_date)).fetchall())
        if not rows:
            raise ValueError("未找到该学生今日计划")

        feedback_rows = rows_to_dicts(
            conn.execute(
                """
                SELECT
                  f.*,
                  p.original_position_text AS label,
                  p.question_order AS question_order,
                  q.question_note AS question_note
                FROM question_feedbacks f
                LEFT JOIN daily_plans p
                  ON p.user_id = f.user_id
                 AND p.plan_date = f.plan_date
                 AND p.type_id = f.type_id
                 AND p.question_id = f.question_id
                LEFT JOIN questions q ON q.question_id = f.question_id
                WHERE f.user_id = ? AND f.plan_date = ?
                ORDER BY f.submitted_at, CAST(p.question_order AS INTEGER), f.question_id
                """,
                (user_id, plan_date),
            ).fetchall()
        )
        chapter_summary = latest_chapter_summary_for_plan(conn, user_id, plan_date)
        extra_rows_by_type = extra_practice_rows_by_type(conn, profile, rows)

    types: list[dict[str, Any]] = []
    for type_index in sorted({row["type_index"] for row in rows}, key=lambda value: as_int(value)):
        type_rows = [row for row in rows if row["type_index"] == type_index]
        first = type_rows[0]
        sources: dict[str, list[str]] = {}
        source_group_map: dict[str, dict[str, Any]] = {}
        for row in type_rows:
            sources.setdefault(row["source"], []).append(row["original_position_text"])
            display = question_source_display(row)
            source_label = display["bookDisplayLabel"]
            group = source_group_map.setdefault(
                source_label,
                {
                    "sourceName": source_label.strip("《》"),
                    "sourceLabel": source_label,
                    "entries": [],
                },
            )
            group["entries"].append(
                {
                    "chapterLabel": display["sourceChapterLabel"],
                    "phaseLabel": display["sourceSectionLabel"],
                    "questionTypeLabel": display["sourceSubsectionLabel"],
                    "questionTypeShort": display["sourceSubsectionLabel"].replace("题", ""),
                    "title": " · ".join(
                        item
                        for item in (
                            display["sourceChapterLabel"],
                            display["sourceSectionLabel"],
                            display["sourceSubsectionLabel"],
                        )
                        if item
                    ),
                    "continuationTitle": " · ".join(
                        item
                        for item in (
                            display["sourceChapterLabel"],
                            display["sourceSectionLabel"],
                            display["sourceSubsectionLabel"],
                        )
                        if item
                    ),
                    "numbers": [display["sourceQuestionLabel"]],
                }
            )
        intro_payload = question_type_intro_payload(first)
        sequence_label = str(first.get("display_type_label") or "").strip()
        if sequence_label:
            intro_payload["cardTypeLabel"] = sequence_label
        types.append(
            {
                "typeIndex": as_int(type_index),
                "typeId": first["type_id"],
                "typeName": first["type_name"],
                "chapterId": first.get("chapter_id") or "",
                "studentTypeSortInChapter": as_int(first.get("student_type_sort_in_chapter")),
                "displayTypeLabel": sequence_label,
                "isChapterBoundaryAfter": truthy_text(first.get("is_chapter_boundary_after")),
                **intro_payload,
                "feedbackOpenStatus": first["feedback_open_status"],
                "sources": [
                    {"source": source, "items": items}
                    for source, items in sources.items()
                ],
                "sourceGroups": list(source_group_map.values()),
                "questions": [
                    {
                        "questionId": row["question_id"],
                        "label": row["original_position_text"],
                        "order": as_int(row["question_order"]),
                        "questionNote": row.get("question_note") or "",
                        **question_source_display(row),
                    }
                    for row in type_rows
                ],
            }
        )

    plan_payload = {
        "planId": rows[0]["plan_id"],
        "planDate": rows[0]["plan_date"],
        "planStatus": rows[0]["plan_status"],
        "slogan": rows[0]["slogan"],
        "examDate": rows[0]["exam_date"],
        "daysToExam": as_int(rows[0]["days_to_exam"]),
        "types": types,
        "extraPracticeTypes": build_extra_practice_types(types, extra_rows_by_type),
        "feedbacks": feedback_rows,
        "completedMessage": completion_message() if rows[0]["plan_status"] == "completed" else "",
    }
    if chapter_summary:
        plan_payload["chapterSummary"] = chapter_summary

    return {
        "student": {
            "userId": user["user_id"],
            "nickname": user["nickname"],
            "mathType": profile["math_type"],
            "targetScore": as_int(profile["target_score"]),
        },
        "plan": plan_payload,
    }


def advance_test_day(user_id: str) -> dict[str, Any]:
    with get_connection() as conn:
        ensure_test_user_dates_table(conn)
        ensure_student_progress_tracking_columns(conn)
        user = row_to_dict(conn.execute("SELECT * FROM users WHERE user_id = ?", (user_id,)).fetchone())
        profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        progress = row_to_dict(
            conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user_id,)).fetchone()
        )
        if not user or not profile or not progress:
            raise ValueError("未找到该学生测试进度")

        return advance_global_test_day(user_id)

        current_test_date = get_current_test_date(conn, user_id)
        try:
            next_test_date = (date.fromisoformat(current_test_date) + timedelta(days=1)).isoformat()
        except ValueError as exc:
            raise ValueError("测试日期格式应为 YYYY-MM-DD") from exc

        updated_at = now_text()
        conn.execute(
            """
            INSERT INTO test_user_dates (user_id, current_test_date, created_at, updated_at)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(user_id) DO UPDATE SET
              current_test_date = excluded.current_test_date,
              updated_at = excluded.updated_at
            """,
            (user_id, next_test_date, updated_at, updated_at),
        )
        conn.commit()

    daily_plan = get_daily_plan(user_id, next_test_date)
    with get_connection() as conn:
        ensure_test_day_supporting_rows(conn, user_id, next_test_date, daily_plan)
        conn.commit()

    return {
        "currentTestDate": next_test_date,
        "dailyPlan": get_daily_plan(user_id, next_test_date),
    }


def advance_global_test_day(default_user_id: str = "10001") -> dict[str, Any]:
    default_user_id = str(default_user_id or "10001").strip() or "10001"
    with get_connection() as conn:
        next_test_date = advance_shared_test_clock(conn)
        conn.commit()

    daily_plan = get_daily_plan(default_user_id, next_test_date)
    with get_connection() as conn:
        ensure_test_day_supporting_rows(conn, default_user_id, next_test_date, daily_plan)
        conn.commit()

    return {
        "currentTestDate": next_test_date,
        "dailyPlan": get_daily_plan(default_user_id, next_test_date),
    }


def submit_feedback(
    user_id: str,
    plan_date: str,
    type_id: str,
    answers: list[dict[str, Any]],
    unattempted_action: str | None = None,
) -> dict[str, Any]:
    if not answers:
        raise ValueError("answers 不能为空")

    with get_connection() as conn:
        ensure_question_feedbacks_reason_columns(conn)
        ensure_student_progress_tracking_columns(conn)
        ensure_student_stats_ranking_columns(conn)
        plan_rows = rows_to_dicts(
            conn.execute(
                """
                SELECT * FROM daily_plans
                WHERE user_id = ? AND plan_date = ? AND type_id = ?
                ORDER BY CAST(question_order AS INTEGER)
                """,
                (user_id, plan_date, type_id),
            ).fetchall()
        )
        if not plan_rows:
            raise ValueError("未找到要提交的题型计划")
        if plan_rows[0]["feedback_open_status"] != "open":
            raise ValueError("该题型当前不可提交反馈")

        progress_row = row_to_dict(
            conn.execute("SELECT * FROM student_progress WHERE user_id = ?", (user_id,)).fetchone()
        ) or {}
        profile_row = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        ) or {}
        assert_service_active_for_write(conn, profile_row, plan_date)
        plan_variant_id = str(profile_row.get("plan_variant_id") or "").strip()

        expected_ids = {row["question_id"] for row in plan_rows}
        answer_by_id = {item["questionId"]: item["result"] for item in answers}
        if set(answer_by_id) != expected_ids:
            raise ValueError("提交题目必须与当前开放题型完全一致")

        valid_results = {"correct", "wrong", "unattempted"}
        if any(result not in valid_results for result in answer_by_id.values()):
            raise ValueError("result 只能是 correct、wrong、unattempted")

        action_by_id = {
            item["questionId"]: str(item.get("unattemptedAction") or unattempted_action or "").strip()
            for item in answers
        }
        has_unattempted = any(result == "unattempted" for result in answer_by_id.values())
        unattempted_action = str(unattempted_action or "").strip()
        if has_unattempted and unattempted_action not in {"mark_wrong", "defer"}:
            raise ValueError("unattemptedAction must be mark_wrong or defer when answers contain unattempted")
        if not has_unattempted and unattempted_action:
            raise ValueError("unattemptedAction is only allowed when answers contain unattempted")

        previous_rates = previous_loss_reason_rates(conn, user_id, plan_date)
        reason_by_id = {
            item["questionId"]: str(item.get("wrongReason") or "").strip()
            for item in answers
        }
        for question_id, result in answer_by_id.items():
            reason_code = reason_by_id.get(question_id, "")
            row_action = action_by_id.get(question_id, "")
            counts_as_wrong = result == "wrong" or (result == "unattempted" and row_action == "mark_wrong")
            if counts_as_wrong and reason_code not in WRONG_REASON_LABELS:
                raise ValueError("wrongReason must be one of method_error, calculation_error, knowledge_gap, condition_missed")
            if not counts_as_wrong and reason_code:
                raise ValueError("wrongReason is only allowed when result is wrong")

        conn.execute(
            "DELETE FROM question_feedbacks WHERE user_id = ? AND plan_date = ? AND type_id = ?",
            (user_id, plan_date, type_id),
        )

        submitted_at = now_text()
        for index, row in enumerate(plan_rows, start=1):
            result = answer_by_id[row["question_id"]]
            row_unattempted_action = action_by_id.get(row["question_id"], "") if result == "unattempted" else ""
            is_wrong_result = result == "wrong" or (result == "unattempted" and row_unattempted_action == "mark_wrong")
            wrong_reason_code = reason_by_id.get(row["question_id"], "") if is_wrong_result else ""
            wrong_reason_name = WRONG_REASON_LABELS.get(wrong_reason_code, "")
            is_wrong = "1" if is_wrong_result else "0"
            is_blocker = "1" if result == "unattempted" and row_unattempted_action == "defer" else "0"
            conn.execute(
                """
                INSERT INTO question_feedbacks (
                  feedback_id, plan_id, user_id, plan_date, type_id, question_id,
                  result, wrong_reason_code, wrong_reason_name,
                  is_wrong_record, is_unfinished_blocker, submitted_at,
                  created_at, updated_at, unattempted_action
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    f"F-{plan_date.replace('-', '')}-{user_id}-{type_id}-{index}",
                    row["plan_id"],
                    user_id,
                    plan_date,
                    type_id,
                    row["question_id"],
                    result,
                    wrong_reason_code,
                    wrong_reason_name,
                    is_wrong,
                    is_blocker,
                    submitted_at,
                    submitted_at,
                    submitted_at,
                    row_unattempted_action,
                ),
            )

        completed_results = [
            result
            for question_id, result in answer_by_id.items()
            if not (result == "unattempted" and action_by_id.get(question_id, "") == "defer")
        ]
        correct_count = sum(1 for result in completed_results if result == "correct")
        accuracy = correct_count / len(completed_results) if completed_results else 0
        recalculate_daily_activity_from_feedback(conn, user_id, plan_date, submitted_at)
        recalculate_type_mastery_stats(conn, user_id, plan_date, submitted_at)
        recalculate_loss_reason_stats(conn, user_id, plan_date, previous_rates, submitted_at)

        conn.execute(
            """
            UPDATE daily_plans
            SET feedback_open_status = 'submitted', updated_at = ?
            WHERE user_id = ? AND plan_date = ? AND type_id = ?
            """,
            (submitted_at, user_id, plan_date, type_id),
        )

        progress_type_id = type_id
        progress_status = "active"
        current_type_status = "active"
        last_completed_type_id = ""
        unfinished_type_id = ""
        unfinished_question_ids = "[]"
        carryover_type_id = ""
        carryover_question_ids = "[]"
        existing_carryover_type_id = str(progress_row.get("carryover_type_id") or "").strip()
        existing_carryover_ids = parse_json_id_list(progress_row.get("carryover_question_ids"))
        next_type_id = ""

        if has_unattempted and unattempted_action == "defer":
            new_status = "blocked_by_unfinished"
            current_type_status = "unfinished"
            unfinished_type_id = type_id
            deferred_question_ids = [
                question_id for question_id, result in answer_by_id.items() if result == "unattempted"
            ]
            if existing_carryover_type_id == type_id:
                for question_id in existing_carryover_ids:
                    if question_id not in deferred_question_ids:
                        deferred_question_ids.append(question_id)
            unfinished_question_ids = json.dumps(
                deferred_question_ids,
                ensure_ascii=False,
            )
        else:
            last_completed_type_id = type_id
            next_row = row_to_dict(
                conn.execute(
                    """
                    SELECT * FROM daily_plans
                    WHERE user_id = ? AND plan_date = ? AND feedback_open_status = 'locked'
                    ORDER BY CAST(type_index AS INTEGER)
                    LIMIT 1
                    """,
                    (user_id, plan_date),
                ).fetchone()
            )
            if next_row:
                next_type_id = next_row["type_id"]
                conn.execute(
                    """
                    UPDATE daily_plans
                    SET feedback_open_status = 'open', updated_at = ?
                    WHERE user_id = ? AND plan_date = ? AND type_id = ?
                    """,
                    (submitted_at, user_id, plan_date, next_row["type_id"]),
                )
                new_status = "in_progress"
                progress_type_id = next_row["type_id"]
                if existing_carryover_type_id and existing_carryover_ids:
                    carryover_type_id = existing_carryover_type_id
                    carryover_question_ids = json.dumps(existing_carryover_ids, ensure_ascii=False)
            else:
                new_status = "completed"
                if (
                    existing_carryover_type_id == type_id
                    and existing_carryover_ids
                ):
                    progress_type_id = type_id
                    current_type_status = "carryover"
                    carryover_type_id = type_id
                    carryover_question_ids = json.dumps(existing_carryover_ids, ensure_ascii=False)
                else:
                    next_global_type = next_actual_assignable_type(
                        conn,
                        type_id,
                        profile_row,
                        progress_row.get("plan_stage") or profile_row.get("current_stage") or "",
                        plan_variant_id,
                    )
                    if next_global_type:
                        progress_type_id = next_global_type["type_id"]
                    else:
                        progress_status = "completed"

        conn.execute(
            """
            UPDATE daily_plans
            SET plan_status = ?, updated_at = ?
            WHERE user_id = ? AND plan_date = ?
            """,
            (new_status, submitted_at, user_id, plan_date),
        )
        if not (has_unattempted and unattempted_action == "defer"):
            maybe_upsert_chapter_completion_after_feedback(
                conn,
                user_id,
                plan_date,
                type_id,
                submitted_at,
            )

        progress_type = question_type_for_plan_variant(conn, progress_type_id, plan_variant_id, profile_row.get("level"))
        conn.execute(
            """
            UPDATE student_progress
            SET current_type_id = ?, current_type_sort_code = ?, progress_status = ?,
                current_type_status = ?,
                last_completed_type_id = COALESCE(NULLIF(?, ''), last_completed_type_id),
                unfinished_type_id = ?, unfinished_question_ids = ?,
                carryover_type_id = ?, carryover_question_ids = ?, updated_at = ?
            WHERE user_id = ?
            """,
            (
                progress_type_id,
                progress_type.get("sort_code") if progress_type else "",
                progress_status,
                current_type_status,
                last_completed_type_id,
                unfinished_type_id,
                unfinished_question_ids,
                carryover_type_id,
                carryover_question_ids,
                submitted_at,
                user_id,
            ),
        )

        plan_progress = row_to_dict(
            conn.execute(
                """
                SELECT
                  COUNT(DISTINCT type_id) AS total_types,
                  COUNT(DISTINCT CASE WHEN feedback_open_status = 'submitted' THEN type_id END) AS submitted_types
                FROM daily_plans
                WHERE user_id = ? AND plan_date = ?
                """,
                (user_id, plan_date),
            ).fetchone()
        )
        total_types = max(as_int(plan_progress.get("total_types")), 1)
        submitted_types = as_int(plan_progress.get("submitted_types"))
        today_completion_rate = submitted_types / total_types

        student_profile = row_to_dict(
            conn.execute("SELECT * FROM student_profiles WHERE user_id = ?", (user_id,)).fetchone()
        )
        next_checkin_order = row_to_dict(
            conn.execute(
                """
                SELECT COALESCE(MAX(CAST(s.today_checkin_order AS INTEGER)), 0) + 1 AS next_order
                FROM student_stats s
                JOIN student_profiles p ON p.user_id = s.user_id
                WHERE s.stat_date = ? AND p.class_id = ?
                """,
                (plan_date, student_profile.get("class_id") if student_profile else ""),
            ).fetchone()
        )

        conn.execute(
            """
            UPDATE student_stats
            SET today_completion_rate = ?,
                today_checkin_order = CASE
                  WHEN ? = 'completed' AND COALESCE(today_checkin_order, '') IN ('', '0') THEN ?
                  ELSE today_checkin_order
                END,
                today_checkin_at = CASE
                  WHEN ? = 'completed' AND COALESCE(today_checkin_at, '') = '' THEN ?
                  ELSE today_checkin_at
                END,
                overall_accuracy = ?,
                updated_at = ?
            WHERE user_id = ? AND stat_date = ?
            """,
            (
                f"{today_completion_rate:.4f}",
                new_status,
                str(next_checkin_order.get("next_order", 1)),
                new_status,
                submitted_at,
                f"{accuracy:.4f}",
                submitted_at,
                user_id,
                plan_date,
            ),
        )
        sync_student_continuous_days(conn, user_id, plan_date)
        conn.commit()

    return {
        "planStatus": new_status,
        "submittedTypeId": type_id,
        "accuracy": round(accuracy, 4),
        "hasUnattempted": has_unattempted,
        "unattemptedAction": unattempted_action or None,
        "dailyPlan": get_daily_plan(user_id, plan_date),
    }
