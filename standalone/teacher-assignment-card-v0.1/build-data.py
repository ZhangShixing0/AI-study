from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "data" / "production-seed" / "math2_public_seed_v1.xlsx"
OUTPUT = Path(__file__).resolve().parent / "data.js"


def read_sheet(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    return [{headers[index]: value for index, value in enumerate(row)} for row in rows]


def active(value: Any) -> bool:
    return str(value or "").strip() in {"1", "true", "True", "TRUE"}


def as_number(value: Any, default: float = 0) -> float:
    try:
        return float(str(value))
    except (TypeError, ValueError):
        return default


def main() -> None:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    chapters = [
        {
            "chapterId": row["chapter_id"],
            "chapterName": row["chapter_name"],
            "subject": row["subject"],
            "sortOrder": index + 1,
        }
        for index, row in enumerate(read_sheet(workbook, "chapters"))
        if active(row.get("is_active")) and row.get("math_type") == "数学二"
    ]
    chapter_order = {item["chapterId"]: item["sortOrder"] for item in chapters}

    order_rows = [
        row
        for row in read_sheet(workbook, "question_type_orders")
        if active(row.get("is_active")) and row.get("plan_variant_id") == "math2_mixed"
    ]
    order_by_type = {row["type_id"]: row for row in order_rows}

    question_types = [
        {
            "typeId": row["type_id"],
            "chapterId": row["chapter_id"],
            "typeName": row["type_name"],
            "cardTypeLabel": row.get("card_type_label") or "",
            "cardTitle": row.get("card_title") or row["type_name"],
            "importanceTag": row.get("importance_tag") or "",
            "questionKind": row.get("question_kind") or "",
            "introTitle": row.get("intro_title") or "",
            "introBody": row.get("intro_body") or "",
            "sortCode": str(order_by_type.get(row["type_id"], {}).get("sort_code") or ""),
            "duration130Days": as_number(order_by_type.get(row["type_id"], {}).get("duration_130_days")),
            "duration100Days": as_number(order_by_type.get(row["type_id"], {}).get("duration_100_days")),
            "duration80Days": as_number(order_by_type.get(row["type_id"], {}).get("duration_80_days")),
            "chapterSort": chapter_order.get(row["chapter_id"], 999),
        }
        for row in read_sheet(workbook, "question_types")
        if active(row.get("is_active")) and row.get("math_type") == "数学二"
    ]
    question_types.sort(key=lambda item: (item["chapterSort"], item["sortCode"], item["typeId"]))

    book_rows = [
        row
        for row in read_sheet(workbook, "exercise_books")
        if active(row.get("is_active")) and row.get("math_type") == "数学二"
    ]
    books = [
        {
            "bookId": row["book_id"],
            "bookName": row["book_name"],
            "legacyCode": str(row["book_name"]).replace("题", ""),
            "bookType": row.get("book_type") or "",
            "sortOrder": as_number(row.get("sort_order"), 999),
        }
        for row in book_rows
    ]
    books.sort(key=lambda item: (item["sortOrder"], item["bookId"]))
    book_names = {row["book_id"]: row["book_name"] for row in book_rows}
    book_sort = {row["book_id"]: as_number(row.get("sort_order"), 999) for row in book_rows}

    valid_type_ids = {item["typeId"] for item in question_types}
    questions = [
        {
            "questionId": row["question_id"],
            "typeId": row["type_id"],
            "bookId": row["book_id"],
            "bookName": book_names.get(row["book_id"], row["book_id"]),
            "bookSort": book_sort.get(row["book_id"], 999),
            "bookDisplayLabel": row.get("book_display_label") or row["book_id"],
            "applicableLevels": row.get("applicable_levels") or "",
            "originalPositionText": row.get("original_position_text") or "",
            "questionOrder": as_number(row.get("question_order")),
            "sourceChapterLabel": row.get("source_chapter_label") or "",
            "sourceSectionLabel": row.get("source_section_label") or "",
            "sourceSubsectionLabel": row.get("source_subsection_label") or "",
            "sourceQuestionLabel": row.get("source_question_label") or "",
            "questionNote": row.get("question_note") or "",
        }
        for row in read_sheet(workbook, "questions")
        if active(row.get("is_active")) and row.get("type_id") in valid_type_ids
    ]
    questions.sort(key=lambda item: (item["typeId"], item["bookSort"], item["questionOrder"], item["questionId"]))

    data = {
        "version": "math2-teacher-assignment-card-v0.2",
        "source": "data/production-seed/math2_public_seed_v1.xlsx",
        "chapters": chapters,
        "books": books,
        "questionTypes": question_types,
        "questions": questions,
    }
    OUTPUT.write_text(
        "window.TEACHER_ASSIGNMENT_DATA = "
        + json.dumps(data, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUTPUT}")
    print(f"books={len(books)} questionTypes={len(question_types)} questions={len(questions)}")


if __name__ == "__main__":
    main()
